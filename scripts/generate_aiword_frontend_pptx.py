#!/usr/bin/env python3
"""
生成「AI Word 前台」项目介绍 PPT（覆盖当前全部前台功能）。
aicheckword 仅作为后台 API 能力介绍，不配截图。

python scripts/generate_aiword_frontend_pptx.py
默认输出：docs/presentations/aiword_frontend_overview.pptx
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any, Sequence


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[1]


def _load_pptx():
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches, Pt

        return {
            "Presentation": Presentation,
            "RGBColor": RGBColor,
            "MSO_SHAPE": MSO_SHAPE,
            "Inches": Inches,
            "Pt": Pt,
        }
    except ImportError:
        return None


NAVY = BLUE = TEAL = INK = MUTED = WHITE = PAPER = None
FONT = "Microsoft YaHei"


def _init_colors(RGBColor: Any) -> None:
    global NAVY, BLUE, TEAL, INK, MUTED, WHITE, PAPER
    NAVY = RGBColor(0x1B, 0x3A, 0x5F)
    BLUE = RGBColor(0x2E, 0x75, 0xB6)
    TEAL = RGBColor(0x1F, 0x7A, 0x6B)
    INK = RGBColor(0x2B, 0x2B, 0x2B)
    MUTED = RGBColor(0x5B, 0x66, 0x73)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    PAPER = RGBColor(0xF4, 0xF7, 0xFB)


def _light_blue():
    from pptx.dml.color import RGBColor

    return RGBColor(0xC5, 0xD8, 0xEF)


def _set_run(run: Any, text: str, size: int, color: Any, bold: bool = False) -> None:
    run.text = text
    run.font.name = FONT
    run.font.size = size
    run.font.color.rgb = color
    run.font.bold = bold


def _add_rect(slide: Any, MSO_SHAPE: Any, left, top, width, height, fill) -> Any:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    return shape


def _textbox(slide: Any, left, top, width, height) -> Any:
    return slide.shapes.add_textbox(left, top, width, height)


def _write_para(tf: Any, Pt: Any, lines: Sequence[str], size: int = 16, color: Any = None) -> None:
    color = color or INK
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(6)
        run = p.add_run()
        _set_run(run, line, Pt(size), color, False)


def _add_footer(slide: Any, Inches: Any, Pt: Any, page: int, total: int) -> None:
    tf = _textbox(slide, Inches(0.5), Inches(7.15), Inches(10.6), Inches(0.28)).text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "AI Word 前台  ·  项目介绍案例（功能以当前版本为准）", Pt(10), MUTED)
    tf2 = _textbox(slide, Inches(11.4), Inches(7.15), Inches(1.4), Inches(0.28)).text_frame
    p2 = tf2.paragraphs[0]
    run2 = p2.add_run()
    _set_run(run2, f"{page} / {total}", Pt(10), MUTED)


def _new_blank(prs: Any) -> Any:
    return prs.slides.add_slide(prs.slide_layouts[6])


def _decorate(slide: Any, MSO_SHAPE: Any, Inches: Any) -> None:
    _add_rect(slide, MSO_SHAPE, Inches(0), Inches(0), Inches(13.333), Inches(7.5), PAPER)
    _add_rect(slide, MSO_SHAPE, Inches(0), Inches(0), Inches(13.333), Inches(0.12), NAVY)
    _add_rect(slide, MSO_SHAPE, Inches(0), Inches(7.08), Inches(13.333), Inches(0.42), WHITE)


def _title(slide: Any, Inches: Any, Pt: Any, text: str) -> None:
    tf = _textbox(slide, Inches(0.5), Inches(0.26), Inches(12.3), Inches(0.5)).text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, text, Pt(24), NAVY, True)


def add_cover(prs: Any, libs: dict) -> None:
    Inches, Pt, MSO_SHAPE = libs["Inches"], libs["Pt"], libs["MSO_SHAPE"]
    slide = _new_blank(prs)
    _add_rect(slide, MSO_SHAPE, Inches(0), Inches(0), Inches(13.333), Inches(7.5), NAVY)
    _add_rect(slide, MSO_SHAPE, Inches(0), Inches(0), Inches(0.22), Inches(7.5), BLUE)
    tf = _textbox(slide, Inches(0.8), Inches(1.75), Inches(11.5), Inches(0.5)).text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "项目介绍案例", Pt(18), _light_blue())
    tf = _textbox(slide, Inches(0.8), Inches(2.3), Inches(11.8), Inches(1.4)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "AI Word 注册文档协同前台", Pt(36), WHITE, True)
    tf = _textbox(slide, Inches(0.8), Inches(3.9), Inches(11.6), Inches(1.2)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(
        run,
        "公司总览 · 任务协同 · 文档工具 · 文控与考试 · 统计催办",
        Pt(18),
        WHITE,
    )
    tf = _textbox(slide, Inches(0.8), Inches(5.5), Inches(11.6), Inches(0.7)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "覆盖当前全部前台功能  |  智能能力由后台 API 提供（本册不展示后台界面）", Pt(14), _light_blue())


def add_toc(prs: Any, libs: dict, page: int, total: int) -> None:
    Inches, Pt, MSO_SHAPE = libs["Inches"], libs["Pt"], libs["MSO_SHAPE"]
    slide = _new_blank(prs)
    _decorate(slide, MSO_SHAPE, Inches)
    _title(slide, Inches, Pt, "目录")
    items = [
        ("01", "角色、登录与可见范围"),
        ("02", "公司总览：项目、知识训练、发补、文档工具"),
        ("03", "任务管理 / 我的任务 / 统计催办"),
        ("04", "文档工具：初稿、审核、审核后修改、翻译"),
        ("05", "体系与注册：文控、版本任务、项目知识库、文献、考试"),
        ("06", "系统管理、权限开关，以及后台 API 能力"),
    ]
    for i, (num, text) in enumerate(items):
        top = 1.05 + i * 0.85
        _add_rect(slide, MSO_SHAPE, Inches(0.55), Inches(top), Inches(12.2), Inches(0.72), WHITE)
        _add_rect(slide, MSO_SHAPE, Inches(0.55), Inches(top), Inches(0.12), Inches(0.72), BLUE if i % 2 == 0 else TEAL)
        tf = _textbox(slide, Inches(0.85), Inches(top + 0.16), Inches(1.0), Inches(0.42)).text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        _set_run(run, num, Pt(18), BLUE, True)
        tf = _textbox(slide, Inches(2.0), Inches(top + 0.18), Inches(10.4), Inches(0.42)).text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        _set_run(run, text, Pt(16), INK)
    _add_footer(slide, Inches, Pt, page, total)


def add_bullets(
    prs: Any,
    libs: dict,
    title: str,
    bullets: Sequence[str],
    page: int,
    total: int,
    subtitle: str = "",
    size: int = 16,
) -> None:
    Inches, Pt, MSO_SHAPE = libs["Inches"], libs["Pt"], libs["MSO_SHAPE"]
    slide = _new_blank(prs)
    _decorate(slide, MSO_SHAPE, Inches)
    _title(slide, Inches, Pt, title)
    top = 0.95
    if subtitle:
        tf = _textbox(slide, Inches(0.55), Inches(0.78), Inches(12.2), Inches(0.38)).text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        _set_run(run, subtitle, Pt(13), MUTED)
        top = 1.15
    box = _textbox(slide, Inches(0.55), Inches(top), Inches(12.2), Inches(5.7))
    _write_para(box.text_frame, Pt, [f"•  {x}" for x in bullets], size=size)
    _add_footer(slide, Inches, Pt, page, total)


def add_two_col(
    prs: Any,
    libs: dict,
    title: str,
    left_title: str,
    left_items: Sequence[str],
    right_title: str,
    right_items: Sequence[str],
    page: int,
    total: int,
) -> None:
    Inches, Pt, MSO_SHAPE = libs["Inches"], libs["Pt"], libs["MSO_SHAPE"]
    slide = _new_blank(prs)
    _decorate(slide, MSO_SHAPE, Inches)
    _title(slide, Inches, Pt, title)
    _add_rect(slide, MSO_SHAPE, Inches(0.5), Inches(1.0), Inches(5.95), Inches(5.75), WHITE)
    _add_rect(slide, MSO_SHAPE, Inches(6.85), Inches(1.0), Inches(5.95), Inches(5.75), WHITE)
    _add_rect(slide, MSO_SHAPE, Inches(0.5), Inches(1.0), Inches(0.12), Inches(5.75), BLUE)
    _add_rect(slide, MSO_SHAPE, Inches(6.85), Inches(1.0), Inches(0.12), Inches(5.75), TEAL)
    for x, cap, items in ((0.75, left_title, left_items), (7.1, right_title, right_items)):
        tf = _textbox(slide, Inches(x), Inches(1.15), Inches(5.4), Inches(0.4)).text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        _set_run(run, cap, Pt(15), NAVY, True)
        box = _textbox(slide, Inches(x), Inches(1.6), Inches(5.45), Inches(4.95))
        _write_para(box.text_frame, Pt, [f"•  {i}" for i in items], size=13)
    _add_footer(slide, Inches, Pt, page, total)


def add_shot(
    prs: Any,
    libs: dict,
    title: str,
    image_path: Path,
    caption: str,
    bullets: Sequence[str],
    page: int,
    total: int,
) -> None:
    Inches, Pt, MSO_SHAPE = libs["Inches"], libs["Pt"], libs["MSO_SHAPE"]
    slide = _new_blank(prs)
    _decorate(slide, MSO_SHAPE, Inches)
    _title(slide, Inches, Pt, title)
    pic_left, pic_top = Inches(0.4), Inches(0.85)
    pic_w, pic_h = Inches(8.4), Inches(5.95)
    if image_path.exists():
        # 只定宽度，保持实拍比例，避免把当前界面压扁
        slide.shapes.add_picture(str(image_path), pic_left, pic_top, width=pic_w)
    else:
        _add_rect(slide, MSO_SHAPE, pic_left, pic_top, pic_w, pic_h, WHITE)
    _add_rect(slide, MSO_SHAPE, Inches(8.95), Inches(0.85), Inches(3.9), Inches(5.95), WHITE)
    tf = _textbox(slide, Inches(9.1), Inches(1.0), Inches(3.6), Inches(0.7)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, caption, Pt(14), BLUE, True)
    box = _textbox(slide, Inches(9.1), Inches(1.75), Inches(3.6), Inches(4.8))
    _write_para(box.text_frame, Pt, [f"•  {b}" for b in bullets], size=13)
    _add_footer(slide, Inches, Pt, page, total)


def add_end(prs: Any, libs: dict) -> None:
    Inches, Pt, MSO_SHAPE = libs["Inches"], libs["Pt"], libs["MSO_SHAPE"]
    slide = _new_blank(prs)
    _add_rect(slide, MSO_SHAPE, Inches(0), Inches(0), Inches(13.333), Inches(7.5), NAVY)
    tf = _textbox(slide, Inches(0.8), Inches(2.5), Inches(11.7), Inches(1.0)).text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "谢谢", Pt(44), WHITE, True)
    tf = _textbox(slide, Inches(0.8), Inches(3.7), Inches(11.7), Inches(1.3)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "功能按系统开关与账号权限启用。问题可通过右下角「反馈」提交。", Pt(16), _light_blue())


def build_presentation(root: Path) -> Any:
    libs = _load_pptx()
    if libs is None:
        raise RuntimeError("python-pptx 未安装")
    _init_colors(libs["RGBColor"])
    Presentation, Inches = libs["Presentation"], libs["Inches"]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    shot = root / "docs" / "presentations" / "screenshots" / "frontend"
    total = 25

    add_cover(prs, libs)
    add_toc(prs, libs, 2, total)

    add_bullets(
        prs,
        libs,
        "前台定位",
        [
            "AI Word 是注册资料团队的业务前台：登录后按角色协同编写、审核、翻译、文控与考试。",
            "智能生成/审核/翻译/训练由后台 API 完成；前台负责选公司、选项目、传文件、看进度、下发任务。",
            "三种角色首页不同：公司管理员 → 公司总览；项目管理员 → 任务管理与统计；普通用户 → 我的任务。",
            "顶部显示当前公司 / 项目组；项目管理员与普通用户须绑定所属项目组。",
            "功能可按系统开关 + 账号权限开放（文献、文控、知识训练、签字打印等默认可单独关闭）。",
        ],
        3,
        total,
        subtitle="面向医疗器械软件注册资料的内部协同平台",
    )

    add_shot(
        prs,
        libs,
        "登录入口",
        shot / "fe-login.png",
        "统一账号登录",
        [
            "用户名 + 密码。",
            "登录后进入该角色默认首页。",
            "忘记密码联系管理员。",
            "超级管理员另需系统管理访问密码。",
            "三种角色，三条工作线。",
        ],
        4,
        total,
    )

    add_shot(
        prs,
        libs,
        "公司总览 · 工具与知识",
        shot / "fe-company.png",
        "公司管理员首页",
        [
            "文档工具：初稿 / 审核 / 修改 / 翻译（手动上传）。",
            "注册工具：文献检索。",
            "知识训练：法规 / 程序 / 项目案例 / 词条。",
            "发补记录：登记、分组、导入日志。",
            "不从任务列表带入文件。",
        ],
        5,
        total,
    )

    add_shot(
        prs,
        libs,
        "公司总览 · 项目管理",
        shot / "fe-company-projects.png",
        "登记与同步",
        [
            "登记项目：名称、注册国家、注册类别。",
            "分配给项目组，供任务侧同步。",
            "一对多关联任务项目。",
            "批量编辑、移出总览（任务数据保留）。",
            "列表按账号「注册国家可见范围」过滤。",
        ],
        6,
        total,
    )

    add_shot(
        prs,
        libs,
        "任务管理 · 工具入口",
        shot / "fe-upload.png",
        "项目管理员首页",
        [
            "文档：初稿、审核、审核后修改、翻译。",
            "注册：文献检索。",
            "体系：文控中心、项目知识库协同。",
            "项目：版本任务清单生成。",
            "考试训练中心（老师端）。",
        ],
        7,
        total,
    )

    add_shot(
        prs,
        libs,
        "任务管理 · 任务列表",
        shot / "fe-upload-tasks.png",
        "派发与批量操作",
        [
            "可从公司总览同步项目。",
            "录入文件名、任务类型、编写人、截止日。",
            "分组、筛选、列设置、拖动排序。",
            "批量审核：单文档 / 一致性 / 追溯。",
            "批量去签字、去打印（需开启）。",
        ],
        8,
        total,
    )

    add_shot(
        prs,
        libs,
        "我的任务",
        shot / "fe-generate.png",
        "普通用户首页",
        [
            "仅本人任务（或观察员只读）。",
            "标记完成状态；事项型须先填执行备注。",
            "可填文档链接；支持上传/替换模板。",
            "分组、筛选、列设置、拖动排序。",
            "操作列可跳转初稿（自动带入本行信息）。",
        ],
        9,
        total,
    )

    add_shot(
        prs,
        libs,
        "统计与催办",
        shot / "fe-dashboard.png",
        "进度看板",
        [
            "整体完成率、按项目/人员/交叉统计。",
            "模块级联催办与下次自动通知时间。",
            "手动催办：按项目、人员、交叉、单条。",
            "项目组钉钉 Webhook（不使用加签）。",
            "考试中心统计入口（若开启）。",
        ],
        10,
        total,
    )

    add_shot(
        prs,
        libs,
        "初稿生成",
        shot / "fe-draft.png",
        "按案例模板出稿",
        [
            "个人 LLM Key 可加密保存、测试。",
            "必选：所属公司、模板项目案例。",
            "勾选待生成文件、语言、策略。",
            "可选 Base 文档做就地修改。",
            "提交后下载 ZIP；可从任务行带入项目信息。",
        ],
        11,
        total,
    )

    add_shot(
        prs,
        libs,
        "文档审核",
        shot / "fe-audit.png",
        "三种审核模式",
        [
            "单文档 / 多文档一致性 / 跨文档追溯。",
            "必选专属项目；可自动匹配过往案例。",
            "注册国家/类别/组成/形态、文档语言。",
            "支持多文件与 zip/tar；进度条可跟任务。",
            "下载报告 ZIP；可纠正入库、生成待办。",
        ],
        12,
        total,
    )

    add_shot(
        prs,
        libs,
        "审核后修改",
        shot / "fe-audit-modify.png",
        "按报告就地改稿",
        [
            "只落实报告中「立即修改」点。",
            "默认 Word 修订标记、不用案例模板。",
            "先预览修改清单再提交。",
            "1 个目标文件 + 1 个 Base。",
            "可手动上传 report.json，或从任务带入。",
        ],
        13,
        total,
    )

    add_shot(
        prs,
        libs,
        "文档翻译",
        shot / "fe-translate.png",
        "多格式翻译",
        [
            "支持 docx / txt / xlsx / zip。",
            "单次最多 5 文件 + 1 目标语言。",
            "可从我的任务带入或手动上传。",
            "压缩包自动解压内部可译文件。",
            "完成后下载译文。",
        ],
        14,
        total,
    )

    add_shot(
        prs,
        libs,
        "文控中心",
        shot / "fe-document-control.png",
        "受控台账",
        [
            "编号、名称、分类、版本。",
            "编号规则（Scheme）维护。",
            "Excel 导入模板（空表 / 含示例）。",
            "导入操作日志可追溯。",
            "定稿后可由项目知识库同步入台账。",
        ],
        15,
        total,
    )

    add_shot(
        prs,
        libs,
        "版本任务清单生成",
        shot / "fe-version-task.png",
        "按版本下发任务",
        [
            "四步：项目与产品 → 版本链路 → 预览 → 确认。",
            "版本号 X.Y.Z.B：X/Y 变更、Z 缺陷、B 生产发布。",
            "可从文控同步文件编号。",
            "预览可增删改、导出 Word/Excel。",
            "确认后下发到任务列表。",
        ],
        16,
        total,
    )

    add_shot(
        prs,
        libs,
        "项目知识库协同",
        shot / "fe-project-kb.png",
        "任务文件入库",
        [
            "统计：已同步 / 失败 / 待同步 / 待入库。",
            "入库后供初稿、审核、修改、翻译共用。",
            "可重试失败、补历史训练。",
            "「同步到文控」是另一条定稿链路。",
            "与公司知识训练相互独立。",
        ],
        17,
        total,
    )

    add_shot(
        prs,
        libs,
        "文献检索",
        shot / "fe-literature.png",
        "注册文献整理",
        [
            "检索式 + 年份 + 每源条数。",
            "来源：PubMed、Google Scholar。",
            "Embase / Cochrane 可导入 RIS/CSV。",
            "每次检索单独成批次，刷新可恢复。",
            "导出 Clinical Literature Search Result。",
        ],
        18,
        total,
    )

    add_shot(
        prs,
        libs,
        "考试训练中心",
        shot / "fe-exam-center.png",
        "老师 / 学生 / 统计",
        [
            "老师端：AI 录题、组卷、复审发布、下发。",
            "体考轨道 × 日常 / 新标 / 项目案例。",
            "学生端：来一套练习、接收考试、作答。",
            "统计端：次数与通过情况。",
            "左侧条件会被录题、组卷、练习继承。",
        ],
        19,
        total,
    )

    add_two_col(
        prs,
        libs,
        "签字打印 · 多租户 · 问题反馈",
        "去签字 / 去打印",
        [
            "任务列表可批量跳转签字、打印",
            "对接签字打印服务，不在前台内嵌排版",
            "默认按账号禁止，需单独开放",
        ],
        "范围、反馈与多公司",
        [
            "多租户：按公司隔离知识库与业务数据",
            "顶部作用域：当前公司、项目组",
            "右下角「反馈」：模块、优先级、描述、截图",
            "我的反馈可跟进度；超管在系统管理处理工单",
        ],
        20,
        total,
    )

    add_shot(
        prs,
        libs,
        "系统管理（超级管理员）",
        shot / "fe-admin.png",
        "字典 / 账号 / 配置",
        [
            "字典：注册国家、项目组、任务类型、状态。",
            "公司管理：维护公司与知识库对应关系。",
            "账号：角色、项目组/公司、逐项功能权限。",
            "系统配置：功能开关、集成、通知时间。",
            "钉钉催办、作用域诊断、问题反馈处理。",
        ],
        21,
        total,
    )

    add_bullets(
        prs,
        libs,
        "后台 API 能力（不展示后台界面）",
        [
            "知识库：按公司 collection 训练/检索；项目案例元数据与自动匹配；审核点生成（随库规模放大）。",
            "文档智能：初稿生成、单文档审核、多文档一致性、跨文档追溯、审核后就地修改、文档翻译。",
            "质量闭环：审核点纠正入库、项目专属资料、词条与内部管控分类（内部管控不参与审核点/初稿/审核）。",
            "考试：组卷、AI 录题、复审发布、练习提交；项目案例命题与案例库对齐。",
            "前台通过统一集成接口调用；普通用户只看到业务结果，不接触后台地址与调试信息。",
        ],
        22,
        total,
        subtitle="智能计算在后台完成，前台只负责任务、文件与进度",
        size=15,
    )

    add_two_col(
        prs,
        libs,
        "功能覆盖一览（当前版本）",
        "协同与体系",
        [
            "登录 / 三角色 / 观察员",
            "公司总览、任务管理、我的任务、统计催办",
            "文控台账、编号规则、版本任务下发",
            "项目知识库协同、文献检索",
            "考试中心三端、去签字/去打印",
            "钉钉催办、问题反馈、系统管理",
        ],
        "文档智能（前台入口）",
        [
            "知识训练与审核点入库",
            "项目案例中英维度与语言",
            "发补记录",
            "初稿、三种审核、纠正入库、生成待办",
            "审核后修改（修订标记）、文档翻译",
            "个人 LLM Key、自动匹配案例",
        ],
        23,
        total,
    )

    add_bullets(
        prs,
        libs,
        "建议演示路径（约 12 分钟）",
        [
            "1. 登录，说明三种角色首页与顶部公司/项目组。",
            "2. 公司总览：登记项目、打开文档工具；如已开知识训练，展示项目案例字段。",
            "3. 任务管理：同步项目 → 派一条文件型任务 → 指出批量审核/文控/版本任务入口。",
            "4. 初稿：选公司与模板案例；审核：三种模式 + 自动匹配案例。",
            "5. 统计催办与（可选）考试中心老师端，结束时说明功能按开关授权。",
        ],
        24,
        total,
        size=16,
    )

    add_end(prs, libs)
    return prs


def main() -> int:
    parser = argparse.ArgumentParser(description="生成 AI Word 前台项目介绍 PPT")
    parser.add_argument("--out", type=Path, default=None, help="输出 .pptx 路径")
    args = parser.parse_args()
    root = _repo_root()
    out = args.out or (root / "docs" / "presentations" / "aiword_frontend_overview.pptx")
    if _load_pptx() is None:
        print("缺少 python-pptx。请执行：pip install python-pptx", file=sys.stderr)
        return 1
    out.parent.mkdir(parents=True, exist_ok=True)
    prs = build_presentation(root)
    prs.save(str(out))
    print(f"已生成: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
