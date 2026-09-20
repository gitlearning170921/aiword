#!/usr/bin/env python3
"""
生成「医疗器械注册文档智能协同平台」项目案例介绍 PPT。

依赖：pip install python-pptx
默认输出：docs/presentations/system_overview.pptx

python scripts/generate_system_overview_pptx.py
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any, List, Optional, Sequence


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[1]


def _load_pptx():
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Emu, Inches, Pt

        return {
            "Presentation": Presentation,
            "RGBColor": RGBColor,
            "MSO_SHAPE": MSO_SHAPE,
            "PP_ALIGN": PP_ALIGN,
            "Emu": Emu,
            "Inches": Inches,
            "Pt": Pt,
        }
    except ImportError:
        return None


NAVY = None
BLUE = None
TEAL = None
ORANGE = None
INK = None
MUTED = None
WHITE = None
PAPER = None
FONT = "Microsoft YaHei"


def _init_colors(RGBColor: Any) -> None:
    global NAVY, BLUE, TEAL, ORANGE, INK, MUTED, WHITE, PAPER
    NAVY = RGBColor(0x1B, 0x3A, 0x5F)
    BLUE = RGBColor(0x2E, 0x75, 0xB6)
    TEAL = RGBColor(0x1F, 0x7A, 0x6B)
    ORANGE = RGBColor(0xC4, 0x59, 0x11)
    INK = RGBColor(0x2B, 0x2B, 0x2B)
    MUTED = RGBColor(0x5B, 0x66, 0x73)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    PAPER = RGBColor(0xF4, 0xF7, 0xFB)


def _set_run(run: Any, text: str, size: int, color: Any, bold: bool = False) -> None:
    run.text = text
    run.font.name = FONT
    run.font.size = size
    run.font.color.rgb = color
    run.font.bold = bold


def _add_rect(slide: Any, MSO_SHAPE: Any, left, top, width, height, fill) -> Any:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    return shape


def _textbox(slide: Any, left, top, width, height) -> Any:
    return slide.shapes.add_textbox(left, top, width, height)


def _write_para(tf: Any, Pt: Any, lines: Sequence[str], size: int = 16, color: Any = None, bold_first: bool = False) -> None:
    color = color or INK
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        run = p.add_run()
        _set_run(run, line, Pt(size), color, bold=(bold_first and i == 0))


def _add_footer(slide: Any, Inches: Any, Pt: Any, page: int, total: int) -> None:
    tf = _textbox(slide, Inches(0.5), Inches(7.15), Inches(10.5), Inches(0.28)).text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "AI Word 注册文档智能协同平台  ·  内部项目案例介绍", Pt(10), MUTED)
    tf2 = _textbox(slide, Inches(11.4), Inches(7.15), Inches(1.4), Inches(0.28)).text_frame
    p2 = tf2.paragraphs[0]
    run2 = p2.add_run()
    _set_run(run2, f"{page} / {total}", Pt(10), MUTED)


def _new_blank(prs: Any) -> Any:
    return prs.slides.add_slide(prs.slide_layouts[6])


def _decorate_content_slide(slide: Any, MSO_SHAPE: Any, Inches: Any) -> None:
    _add_rect(slide, MSO_SHAPE, Inches(0), Inches(0), Inches(13.333), Inches(7.5), PAPER)
    _add_rect(slide, MSO_SHAPE, Inches(0), Inches(0), Inches(13.333), Inches(0.12), NAVY)
    _add_rect(slide, MSO_SHAPE, Inches(0), Inches(7.08), Inches(13.333), Inches(0.42), WHITE)


def add_cover(prs: Any, libs: dict, page: int, total: int) -> None:
    Inches, Pt, MSO_SHAPE = libs["Inches"], libs["Pt"], libs["MSO_SHAPE"]
    slide = _new_blank(prs)
    _add_rect(slide, MSO_SHAPE, Inches(0), Inches(0), Inches(13.333), Inches(7.5), NAVY)
    _add_rect(slide, MSO_SHAPE, Inches(0), Inches(0), Inches(0.22), Inches(7.5), BLUE)
    tf = _textbox(slide, Inches(0.8), Inches(1.7), Inches(11.5), Inches(1.0)).text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "项目案例介绍", Pt(18), RGB_LIGHT_BLUE())
    tf = _textbox(slide, Inches(0.8), Inches(2.3), Inches(11.8), Inches(1.6)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "医疗器械注册文档智能协同平台", Pt(36), WHITE, True)
    tf = _textbox(slide, Inches(0.8), Inches(4.1), Inches(11.5), Inches(1.1)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "AI Word 任务协同  ·  知识库训练  ·  智能初稿 / 审核 / 翻译  ·  文控与考试", Pt(18), WHITE)
    tf = _textbox(slide, Inches(0.8), Inches(5.6), Inches(11.5), Inches(0.6)).text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "基于当前已上线功能整理  |  含系统实拍界面", Pt(14), RGB_LIGHT_BLUE())


def RGB_LIGHT_BLUE():
    from pptx.dml.color import RGBColor

    return RGBColor(0xC5, 0xD8, 0xEF)


def add_bullets_slide(
    prs: Any,
    libs: dict,
    title: str,
    bullets: Sequence[str],
    page: int,
    total: int,
    subtitle: str = "",
) -> None:
    Inches, Pt, MSO_SHAPE = libs["Inches"], libs["Pt"], libs["MSO_SHAPE"]
    slide = _new_blank(prs)
    _decorate_content_slide(slide, MSO_SHAPE, Inches)
    tf = _textbox(slide, Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.55)).text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, title, Pt(26), NAVY, True)
    top = 1.0
    if subtitle:
        tf = _textbox(slide, Inches(0.55), Inches(0.82), Inches(12.2), Inches(0.4)).text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        _set_run(run, subtitle, Pt(14), MUTED)
        top = 1.25
    box = _textbox(slide, Inches(0.6), Inches(top), Inches(12.1), Inches(5.6))
    _write_para(box.text_frame, Pt, [f"•  {x}" for x in bullets], size=17)
    _add_footer(slide, Inches, Pt, page, total)


def add_two_col_slide(
    prs: Any,
    libs: dict,
    title: str,
    left_title: str,
    left_items: Sequence[str],
    right_title: str,
    right_items: Sequence[str],
    page: int,
    total: int,
) -> None:
    Inches, Pt, MSO_SHAPE = libs["Inches"], libs["Pt"], libs["MSO_SHAPE"]
    slide = _new_blank(prs)
    _decorate_content_slide(slide, MSO_SHAPE, Inches)
    tf = _textbox(slide, Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.5)).text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, title, Pt(26), NAVY, True)

    _add_rect(slide, MSO_SHAPE, Inches(0.5), Inches(1.05), Inches(5.95), Inches(5.7), WHITE)
    _add_rect(slide, MSO_SHAPE, Inches(6.85), Inches(1.05), Inches(5.95), Inches(5.7), WHITE)
    _add_rect(slide, MSO_SHAPE, Inches(0.5), Inches(1.05), Inches(0.12), Inches(5.7), BLUE)
    _add_rect(slide, MSO_SHAPE, Inches(6.85), Inches(1.05), Inches(0.12), Inches(5.7), TEAL)

    for x, cap, items in (
        (0.75, left_title, left_items),
        (7.1, right_title, right_items),
    ):
        tf = _textbox(slide, Inches(x), Inches(1.2), Inches(5.4), Inches(0.4)).text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        _set_run(run, cap, Pt(16), NAVY, True)
        box = _textbox(slide, Inches(x), Inches(1.7), Inches(5.4), Inches(4.8))
        _write_para(box.text_frame, Pt, [f"•  {i}" for i in items], size=14)
    _add_footer(slide, Inches, Pt, page, total)


def add_screenshot_slide(
    prs: Any,
    libs: dict,
    title: str,
    image_path: Path,
    caption: str,
    bullets: Sequence[str],
    page: int,
    total: int,
) -> None:
    Inches, Pt, MSO_SHAPE = libs["Inches"], libs["Pt"], libs["MSO_SHAPE"]
    slide = _new_blank(prs)
    _decorate_content_slide(slide, MSO_SHAPE, Inches)
    tf = _textbox(slide, Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.45)).text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, title, Pt(22), NAVY, True)

    # 左侧截图
    pic_left, pic_top = Inches(0.45), Inches(0.85)
    pic_w, pic_h = Inches(8.35), Inches(5.95)
    if image_path.exists():
        slide.shapes.add_picture(str(image_path), pic_left, pic_top, pic_w, pic_h)
    else:
        _add_rect(slide, MSO_SHAPE, pic_left, pic_top, pic_w, pic_h, WHITE)
        tf = _textbox(slide, pic_left, Inches(3.4), pic_w, Inches(0.4)).text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        _set_run(run, f"（截图缺失：{image_path.name}）", Pt(12), MUTED)

    # 右侧说明
    _add_rect(slide, MSO_SHAPE, Inches(8.95), Inches(0.85), Inches(3.9), Inches(5.95), WHITE)
    tf = _textbox(slide, Inches(9.1), Inches(1.0), Inches(3.6), Inches(0.7)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, caption, Pt(14), BLUE, True)
    box = _textbox(slide, Inches(9.1), Inches(1.75), Inches(3.6), Inches(4.8))
    _write_para(box.text_frame, Pt, [f"•  {b}" for b in bullets], size=13)
    _add_footer(slide, Inches, Pt, page, total)


def add_toc_slide(prs: Any, libs: dict, page: int, total: int) -> None:
    items = [
        ("01", "背景与定位：注册资料编写、审核、协同的一体化需求"),
        ("02", "系统架构：AI Word 前台 + 审核引擎 + 文控 / 签字打印"),
        ("03", "核心能力：知识训练、项目案例、初稿、审核、翻译"),
        ("04", "业务协同：公司总览、任务分配、统计催办、文控与考试"),
        ("05", "界面实拍：训练 / 审核点 / 项目维度 / 文档审核 / 登录"),
        ("06", "价值闭环：经验复用、质量留痕、可扩展 API"),
    ]
    Inches, Pt, MSO_SHAPE = libs["Inches"], libs["Pt"], libs["MSO_SHAPE"]
    slide = _new_blank(prs)
    _decorate_content_slide(slide, MSO_SHAPE, Inches)
    tf = _textbox(slide, Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.5)).text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "目录", Pt(26), NAVY, True)
    for i, (num, text) in enumerate(items):
        top = 1.05 + i * 0.85
        _add_rect(slide, MSO_SHAPE, Inches(0.55), Inches(top), Inches(12.2), Inches(0.72), WHITE)
        _add_rect(slide, MSO_SHAPE, Inches(0.55), Inches(top), Inches(0.12), Inches(0.72), BLUE if i % 2 == 0 else TEAL)
        tf = _textbox(slide, Inches(0.85), Inches(top + 0.16), Inches(1.0), Inches(0.42)).text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        _set_run(run, num, Pt(18), BLUE, True)
        tf = _textbox(slide, Inches(2.0), Inches(top + 0.18), Inches(10.4), Inches(0.42)).text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        _set_run(run, text, Pt(16), INK)
    _add_footer(slide, Inches, Pt, page, total)


def add_end_slide(prs: Any, libs: dict) -> None:
    Inches, Pt, MSO_SHAPE = libs["Inches"], libs["Pt"], libs["MSO_SHAPE"]
    slide = _new_blank(prs)
    _add_rect(slide, MSO_SHAPE, Inches(0), Inches(0), Inches(13.333), Inches(7.5), NAVY)
    tf = _textbox(slide, Inches(0.8), Inches(2.5), Inches(11.7), Inches(1.0)).text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "谢谢", Pt(44), WHITE, True)
    tf = _textbox(slide, Inches(0.8), Inches(3.7), Inches(11.7), Inches(1.2)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "欢迎试用与反馈。功能持续迭代中，界面以当前环境为准。", Pt(18), RGB_LIGHT_BLUE())


def build_presentation(root: Path) -> Any:
    libs = _load_pptx()
    if libs is None:
        raise RuntimeError("python-pptx 未安装")
    _init_colors(libs["RGBColor"])
    Presentation, Inches = libs["Presentation"], libs["Inches"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    shot = root / "docs" / "presentations" / "screenshots"
    img = root / "docs" / "images"

    # Pre-count: cover + toc + content + screenshots + end
    # We'll build sequentially and number at the end by rewriting footers is hard,
    # so fix total in advance.
    total = 18

    add_cover(prs, libs, 1, total)
    add_toc_slide(prs, libs, 2, total)

    add_bullets_slide(
        prs,
        libs,
        "背景与痛点",
        [
            "医疗器械软件注册资料种类多、版本密、中英文并存，人工核对成本高。",
            "法规、程序、过往项目经验分散在文件夹中，难以按产品/国家复用。",
            "编写、审核、修改、翻译、文控编号常拆在多个工具里，进度难以统一跟踪。",
            "审核意见难沉淀：同类误报会反复出现，缺少「纠正入库」闭环。",
            "需要一套可落地的协同平台：任务可分配、知识可训练、文档可生成可审核。",
        ],
        3,
        total,
        subtitle="面向有源医疗器械软件 / 独立软件注册资料场景",
    )

    add_two_col_slide(
        prs,
        libs,
        "系统定位",
        "AI Word（业务前台）",
        [
            "公司总览、任务分配、我的任务、统计催办",
            "文档工具：初稿、审核、审核后修改、翻译",
            "体系工具：文控中心、项目知识库、版本任务清单",
            "考试训练中心、文献检索、去签字/去打印",
            "按角色授权（公司管理员 / 项目管理员 / 普通用户）",
        ],
        "审核引擎（aicheckword）",
        [
            "法规 / 程序 / 项目案例知识库（RAG 向量检索）",
            "审核点自动生成，随知识库规模放大覆盖面",
            "单文档、多文档一致性、跨文档可追溯性审核",
            "过往项目案例自动匹配（中英文维度）",
            "开放 REST API，供 AI Word 与外部系统集成",
        ],
        4,
        total,
    )

    add_bullets_slide(
        prs,
        libs,
        "当前功能全景",
        [
            "知识训练：法规、程序、项目案例、词条、内部管控；支持 Word/PDF/Excel/压缩包，.doc 可走 WPS。",
            "项目案例：中英名称/产品/国家、注册类别/组成/形态、适用范围、文档语言；审核时优先产品名+适用范围匹配。",
            "审核点：按知识库规模动态生成；法规/程序审核点全语言通用；项目案例相关审核点可选中/英/中英。",
            "文档能力：初稿生成（选模板案例）、三种审核模式、纠正入库、审核后修改、文档翻译。",
            "协同能力：任务模板与钉钉催办、文控台账与编号、版本任务清单、考试中心（含项目案例命题）。",
        ],
        5,
        total,
    )

    add_screenshot_slide(
        prs,
        libs,
        "界面实拍 · 登录入口",
        shot / "01-login.png",
        "AI Word 统一登录",
        [
            "账号 + 密码登录，按角色进入默认首页。",
            "公司管理员 → 公司项目总览。",
            "项目管理员 → 任务管理 / 统计。",
            "普通用户 → 我的任务。",
            "地址示例：本机 http://localhost:5000",
        ],
        6,
        total,
    )

    add_screenshot_slide(
        prs,
        libs,
        "界面实拍 · 法规/案例训练",
        shot / "03-step1-train.png",
        "第一步：知识入库",
        [
            "按公司切换独立知识库。",
            "文件分类：法规、程序、项目案例等。",
            "支持拖拽上传与压缩包自动解压。",
            "也可从服务器目录批量训练。",
            "当前库已有上千文件、数万向量块。",
        ],
        7,
        total,
    )

    add_screenshot_slide(
        prs,
        libs,
        "界面实拍 · 生成审核点",
        shot / "04-step1-checklist.png",
        "知识越多，审核点越全",
        [
            "基于已训练知识库自动生成清单。",
            "可按注册国家扩展法规检索（如 CE→MDR）。",
            "项目案例相关审核点可选语言。",
            "规模提示：超大库预计 ≥500 个审核点。",
            "可上传已有基础审核点再优化。",
        ],
        8,
        total,
    )

    add_screenshot_slide(
        prs,
        libs,
        "界面实拍 · 审核点管理",
        shot / "06-step2-checkpoints.png",
        "第二步：确认后入库",
        [
            "查看、勾选、批量训练审核点清单。",
            "内置「有源软件二类」通用审核点。",
            "支持导入 JSON 审核点。",
            "trained / draft 状态可区分是否已入库。",
            "项目与专属资料与当前审核项目绑定。",
        ],
        9,
        total,
    )

    add_screenshot_slide(
        prs,
        libs,
        "界面实拍 · 项目专属维度",
        shot / "07-step2-projects.png",
        "中英字段统一维护",
        [
            "项目名称 / 产品名称 / 注册国家均支持中英。",
            "注册类别、组成、项目形态可配置。",
            "审核时按项目维度过滤适用审核点。",
            "与「过往项目案例」区分：案例是经验复用，项目是当前任务。",
            "型号等字段可写入一致性核对。",
        ],
        10,
        total,
    )

    add_screenshot_slide(
        prs,
        libs,
        "界面实拍 · 文档审核",
        shot / "05-step3-review.png",
        "第三步：按知识库审核",
        [
            "通用审核 或 按项目审核。",
            "自动匹配过往项目案例（12 个案例可用）。",
            "产品名称 + 适用范围优先，否则按维度。",
            "支持批量文件与压缩包。",
            "另有多文档一致性、跨文档追溯。",
        ],
        11,
        total,
    )

    add_two_col_slide(
        prs,
        libs,
        "项目案例闭环（经验复用）",
        "训练侧",
        [
            "分类选「项目案例文件」",
            "新建或选择已有案例",
            "填写中英名称、产品、国家",
            "注册类别 / 组成 / 形态 / 适用范围",
            "文档语言：不指定 / 中 / 英 / 中英",
            "入库到通用知识库并绑定 case_id",
        ],
        "使用侧",
        [
            "审核：自动匹配后注入案例上下文",
            "初稿：选择模板项目案例生成格式与法规部分",
            "考试：考试类型可选「项目案例」命题",
            "语言：法规/程序审核全语言通用",
            "案例相关审核可指定中英文规范",
            "纠正入库：审核误报可回写知识库",
        ],
        12,
        total,
    )

    add_screenshot_slide(
        prs,
        libs,
        "业务协同 · 公司项目总览（结构）",
        img / "manual-company-projects.png",
        "公司管理员首页",
        [
            "登记项目：国家、类别、项目组。",
            "一对多关联任务项目。",
            "文档工具：初稿 / 审核 / 翻译（手动上传）。",
            "知识库训练与发补记录（可开关）。",
            "示意图摘自角色操作手册。",
        ],
        13,
        total,
    )

    add_two_col_slide(
        prs,
        libs,
        "任务、统计与体系工具",
        "任务协同",
        [
            "页面1：模板上传、负责人、截止日、批量编辑",
            "页面2：我的任务、完成状态、执行备注",
            "页面3：完成率、按项目/人员统计、催办",
            "钉钉：任务分配、逾期催告、周期统计",
            "三种角色首页不同，权限可按账号细化",
        ],
        "体系与质量",
        [
            "文控中心：受控台账、编号规则、Excel 导入",
            "版本任务清单：按版本生成编写/变更任务",
            "项目知识库：任务文件同步供初稿/审核共用",
            "考试训练中心：日常 / 新标 / 项目案例命题",
            "文献检索、去签字、去打印（按开关启用）",
        ],
        14,
        total,
    )

    add_screenshot_slide(
        prs,
        libs,
        "开放能力 · 审核引擎 API",
        shot / "02-api-docs.png",
        "可被前台与外部调用",
        [
            "知识检索、训练、审核、初稿、翻译。",
            "Quiz：组卷、录题、项目案例列表。",
            "AI Word 通过集成接口代理，不把内部地址暴露给普通用户。",
            "Swagger：http://localhost:8000/docs",
            "多公司 Header：X-Aiword-Company-Id",
        ],
        15,
        total,
    )

    add_bullets_slide(
        prs,
        libs,
        "价值与落地要点",
        [
            "把「过往项目经验」从文件夹变成可匹配的项目案例知识，而不是为每个新项目重新训练。",
            "审核、初稿、考试共用同一套案例与审核点，避免各模块口径不一致。",
            "任务协同把编写进度、催办、文控编号接到同一平台，减少邮件和表格对账。",
            "纠正入库让误报变成组织资产；操作记录与历史报告可追溯。",
            "功能按开关与角色授权，便于按公司逐步启用（知识训练、文控、考试等）。",
        ],
        16,
        total,
        subtitle="当前为内部试用版本，界面与开关以实际部署为准",
    )

    add_bullets_slide(
        prs,
        libs,
        "建议演示路径（10 分钟）",
        [
            "1. 登录 AI Word，展示角色首页差异（总览 / 任务 / 我的任务）。",
            "2. 打开知识训练：说明项目案例元数据（中英、国家、适用范围）。",
            "3. 生成审核点：展示知识库规模与预计审核点数。",
            "4. 文档审核：勾选自动匹配案例，上传 1～2 份文档看报告结构。",
            "5. 如已开启：文控中心编号、版本任务清单、考试中心项目案例命题。",
        ],
        17,
        total,
    )

    add_end_slide(prs, libs)
    return prs


def main() -> int:
    parser = argparse.ArgumentParser(description="生成系统项目案例介绍 PPT")
    parser.add_argument(
        "--out",
        type=Path,
        default=None,
        help="输出 .pptx 路径",
    )
    args = parser.parse_args()
    root = _repo_root()
    out = args.out if args.out is not None else root / "docs" / "presentations" / "system_overview.pptx"

    if _load_pptx() is None:
        print("缺少依赖 python-pptx。请执行：pip install python-pptx", file=sys.stderr)
        return 1

    out.parent.mkdir(parents=True, exist_ok=True)
    prs = build_presentation(root)
    prs.save(str(out))
    print(f"已生成: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
