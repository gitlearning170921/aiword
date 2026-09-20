"""导出当前版本任务预览清单为 Excel。"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_PREVIEW_EXCEL_HEADERS: list[tuple[str, str, int]] = [
    ("seq", "序号", 8),
    ("recordStatus", "状态", 10),
    ("applied", "下发", 12),
    ("isSystemRecord", "体系记录", 10),
    ("changeKind", "变更", 10),
    ("fileName", "文件名", 28),
    ("taskType", "任务类型", 14),
    ("targetVersion", "目标版本", 14),
    ("author", "责任人", 14),
    ("dueDate", "完成日期", 14),
    ("documentDisplayDate", "文档日期", 14),
    ("belongingModule", "模块", 12),
    ("documentNumber", "文件编号", 20),
    ("fileVersion", "文件版本号", 14),
    ("explanation", "说明", 28),
    ("notes", "备注", 18),
    ("displayedAuthor", "编", 12),
    ("reviewer", "审", 12),
    ("approver", "批", 12),
    ("archiveFrequency", "归档频率", 14),
    ("triggeredBy", "触发位", 12),
    ("changeReason", "变更原因", 22),
    ("chapter", "章节分类", 16),
]

_RECORD_STATUS_LABELS = {"adopt": "选用", "discard": "弃用", "pending": "待定"}
_CHANGE_KIND_LABELS = {"add": "新增", "update": "已修改", "delete": "已删除"}
_APPLIED_LABELS = {"applied": "已下发", "partial": "部分下发", "none": "未下发"}
_TRIGGER_LABELS = {"X": "X位", "Y": "Y位", "Z": "Z位", "B": "B位"}


def _preview_cell_text(item: dict[str, Any], key: str, seq: int) -> str:
    if key == "seq":
        return str(seq)
    if key == "recordStatus":
        raw = str(item.get("recordStatus") or "").strip().lower()
        return _RECORD_STATUS_LABELS.get(raw) or (str(item.get("recordStatus") or "").strip() or "选用")
    if key == "applied":
        raw = str(item.get("applied") or "").strip().lower()
        return _APPLIED_LABELS.get(raw, "")
    if key == "isSystemRecord":
        raw = item.get("isSystemRecord")
        if raw in {True, 1, "1", "true", "yes", "是"}:
            return "是"
        return "否"
    if key == "changeKind":
        raw = str(item.get("changeKind") or "").strip().lower()
        return _CHANGE_KIND_LABELS.get(raw, "")
    if key == "triggeredBy":
        raw = item.get("triggeredBy")
        bits = raw if isinstance(raw, list) else str(raw or "").split(",")
        labels = []
        for bit in bits:
            text = str(bit or "").strip()
            if not text:
                continue
            labels.append(_TRIGGER_LABELS.get(text.upper(), text))
        return "、".join(labels)
    if key == "targetVersion":
        return str(item.get("targetVersion") or item.get("registrationVersion") or "").strip()
    if key == "chapter":
        return str(item.get("chapter") or item.get("processBranchLabel") or "").strip()
    return str(item.get(key) or "").strip()


def _identity_key(file_name: Any, task_type: Any, author: Any, target_version: Any) -> str:
    ver = str(target_version or "").strip() or "未指定"
    return "|".join(
        [
            str(file_name or "").strip().lower(),
            str(task_type or "").strip().lower(),
            (str(author or "").strip() or "待分配").lower(),
            ver.lower(),
        ]
    )


def _split_authors(raw: Any) -> list[str]:
    text = str(raw or "").strip()
    if not text:
        return ["待分配"]
    names: list[str] = []
    seen: set[str] = set()
    for part in text.split(","):
        name = str(part or "").strip()
        if not name:
            continue
        key = name.casefold()
        if key in seen:
            continue
        seen.add(key)
        names.append(name)
    return names or ["待分配"]


def _canonical_task_type(value: Any) -> str:
    text = str(value or "").strip()
    if text in {"", "版本变更任务", "归档文件", "变更控制流程", "缺陷管理流程", "生产发布流程"}:
        return "初稿待编写"
    return text


def collect_issued_key_set(issued_items: list[Any] | None) -> set[str]:
    keys: set[str] = set()
    for row in issued_items or []:
        if not isinstance(row, dict):
            continue
        key = str(row.get("key") or "").strip()
        if key:
            keys.add(key)
        for extra in row.get("keys") or []:
            text = str(extra or "").strip()
            if text:
                keys.add(text)
        file_name = str(row.get("fileName") or "").strip()
        author = str(row.get("author") or "").strip() or "待分配"
        ver = str(row.get("targetVersion") or "").strip()
        raw_type = str(row.get("taskType") or "").strip()
        for task_type in {raw_type, _canonical_task_type(raw_type)}:
            if file_name:
                keys.add(_identity_key(file_name, task_type, author, ver))
    return keys


def stamp_preview_applied_status(items: list[Any], issued_items: list[Any] | None) -> list[dict[str, Any]]:
    issued_keys = collect_issued_key_set(issued_items)
    out: list[dict[str, Any]] = []
    for raw in items or []:
        if not isinstance(raw, dict):
            continue
        item = dict(raw)
        if not issued_keys:
            item["applied"] = "none"
            out.append(item)
            continue
        authors = _split_authors(item.get("author"))
        file_name = str(item.get("fileName") or "").strip()
        raw_type = str(item.get("taskType") or "").strip()
        ver = str(item.get("targetVersion") or item.get("registrationVersion") or "").strip()
        types = {raw_type, _canonical_task_type(raw_type)}
        hit = 0
        for author in authors:
            if any(
                _identity_key(file_name, task_type, author, ver) in issued_keys
                for task_type in types
                if file_name
            ):
                hit += 1
        if authors and hit >= len(authors):
            item["applied"] = "applied"
        elif hit:
            item["applied"] = "partial"
        else:
            item["applied"] = "none"
        out.append(item)
    return out


def export_version_task_preview_excel(
    *,
    items: list[Any],
    product_name: str = "",
    from_version: str = "",
    to_version: str = "",
) -> tuple[bytes, str, str]:
    rows = [row for row in (items or []) if isinstance(row, dict)]
    if not rows:
        raise ValueError("当前版本任务清单没有可导出的记录，请先生成预览")

    wb = Workbook()
    ws = wb.active
    ws.title = "版本任务清单"
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D6EAF8")
    wrap = Alignment(wrap_text=True, vertical="center")
    for col_idx, (_key, title, width) in enumerate(_PREVIEW_EXCEL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    for seq, item in enumerate(rows, start=1):
        for col_idx, (key, _title, _width) in enumerate(_PREVIEW_EXCEL_HEADERS, start=1):
            cell = ws.cell(row=seq + 1, column=col_idx, value=_preview_cell_text(item, key, seq))
            cell.alignment = wrap
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_PREVIEW_EXCEL_HEADERS))}{len(rows) + 1}"

    bits = [str(product_name or "").strip(), str(from_version or "").strip(), str(to_version or "").strip()]
    bits = [x for x in bits if x]
    label = "_".join(bits) if bits else datetime.now().strftime("%Y%m%d")
    filename = f"版本任务清单_{label}.xlsx"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), filename, "version-task-list.xlsx"
