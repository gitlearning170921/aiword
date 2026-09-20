"""按版本任务清单更新并导出主文档清单 / 技术文件清单 / 中文目录 Word 附件。"""
from __future__ import annotations

import io
import re
import zipfile
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

from docx import Document
from docx.oxml.ns import qn

from .version_task_generator import (
    _filename_match_keys,
    _title_fuzzy_score,
    normalize_record_status,
    parse_version,
)

_PACKAGED_DIR = Path(__file__).resolve().parent / "list_templates"
_OFFICE_DIR = Path(r"g:\互联网产品部\质量体系\AIWORD\任务清单")

LIST_KINDS: dict[str, dict[str, str]] = {
    "master": {
        "source_name": "QR-QP4.2.3-01 医疗器械主文档清单（呼吸护理管理系统）.docx",
        "download_prefix": "QR-QP4.2.3-01 医疗器械主文档清单",
        "value_field": "fileVersion",
        "empty_value": "/",
        "ascii_name": "QR-QP4.2.3-01-master-document-list.docx",
        "layout": "notes",
        "label": "主文档清单",
    },
    "tech": {
        "source_name": "QR-QP4.2.4-07 技术文件清单（2025.10）.docx",
        "download_prefix": "QR-QP4.2.4-07 技术文件清单",
        "value_field": "documentDisplayDate",
        "empty_value": "",
        "ascii_name": "QR-QP4.2.4-07-technical-file-list.docx",
        "layout": "notes",
        "label": "技术文件清单",
    },
    "catalog": {
        "source_name": "中文目录.docx",
        "download_prefix": "中文目录",
        "value_field": "fileVersion",
        "empty_value": "",
        "ascii_name": "chinese-catalog.docx",
        "layout": "catalog",
        "label": "中文目录",
    },
}

_KIND_ORDER = ("master", "tech", "catalog")
_FUZZY_MATCH_MIN = 0.82
_DOC_REV_RE = re.compile(r"^([A-Za-z]+)[/／\s]*(\d+)$")


def resolve_list_template_path(kind: str) -> Path:
    spec = LIST_KINDS.get(kind)
    if not spec:
        raise ValueError("不支持的清单类型")
    name = spec["source_name"]
    office = _OFFICE_DIR / name
    if office.is_file():
        return office
    packaged = _PACKAGED_DIR / name
    if packaged.is_file():
        return packaged
    raise FileNotFoundError(f"未找到清单模板：{name}")


def _norm_code(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).casefold()


def _display_name(value: Any) -> str:
    text = str(value or "").strip()
    text = re.sub(r"\.(docx?|pdf|xlsx?|pptx?)$", "", text, flags=re.I).strip()
    return text


def _name_keys(value: Any) -> list[str]:
    return _filename_match_keys(_display_name(value))


def _is_live_preview_row(row: Any) -> bool:
    if not isinstance(row, dict):
        return False
    if str(row.get("changeKind") or "").strip().lower() == "delete":
        return False
    if row.get("hideInPreview"):
        return False
    if normalize_record_status(row.get("recordStatus")) == "discard":
        return False
    return bool(_display_name(row.get("fileName")))


def _target_version_sort_key(item: dict[str, Any]) -> tuple:
    raw = str(item.get("targetVersion") or item.get("registrationVersion") or "").strip()
    if not raw or raw == "未指定":
        return (0, 0, 0, 0, 0)
    try:
        parsed = parse_version(raw)
        return (1, parsed.x, parsed.y, parsed.z, parsed.b)
    except Exception:
        return (0, 0, 0, 0, 0)


def _file_version_sort_key(raw: Any) -> tuple:
    """文件版本号排序：A/3、A3、03 都视为 3；空值最低。"""
    text = re.sub(r"\s+", "", str(raw or "").strip())
    if not text:
        return (0, 0, 0)
    match = _DOC_REV_RE.match(text)
    if match:
        return (1, ord(match.group(1).upper()), int(match.group(2)))
    match = re.match(r"^(\d+)$", text)
    if match:
        return (1, ord("A"), int(match.group(1)))
    return (0, 0, 0)


def _has_letter_revision(raw: Any) -> bool:
    text = re.sub(r"\s+", "", str(raw or "").strip())
    return bool(_DOC_REV_RE.match(text))


def _date_sort_key(raw: Any) -> tuple:
    text = str(raw or "").strip()
    if not text:
        return (0, 0, 0, 0)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            dt = datetime.strptime(text[:10], fmt)
            return (1, dt.year, dt.month, dt.day)
        except ValueError:
            continue
    return (0, 0, 0, 0)


def _merge_latest_item(prev: dict[str, Any], rec: dict[str, Any]) -> dict[str, Any]:
    if _target_version_sort_key(rec) >= _target_version_sort_key(prev):
        base, other = dict(rec), prev
    else:
        base, other = dict(prev), rec
    other_fv = str(other.get("fileVersion") or "").strip()
    base_fv = str(base.get("fileVersion") or "").strip()
    other_key = _file_version_sort_key(other_fv)
    base_key = _file_version_sort_key(base_fv)
    if other_key > base_key or (
        other_key == base_key and other_fv and _has_letter_revision(other_fv) and not _has_letter_revision(base_fv)
    ):
        base["fileVersion"] = other_fv
    if _date_sort_key(other.get("documentDisplayDate")) > _date_sort_key(base.get("documentDisplayDate")):
        base["documentDisplayDate"] = str(other.get("documentDisplayDate") or "").strip()
    if not str(base.get("documentNumber") or "").strip() and str(other.get("documentNumber") or "").strip():
        base["documentNumber"] = str(other.get("documentNumber") or "").strip()
    if not str(base.get("fileName") or "").strip() and str(other.get("fileName") or "").strip():
        base["fileName"] = _display_name(other.get("fileName"))
    return base


def _item_identity(item: dict[str, Any]) -> str:
    code = _norm_code(item.get("documentNumber"))
    if code:
        return f"no:{code}"
    return f"name:{_display_name(item.get('fileName')).casefold()}"


def collect_unique_latest_items(items: Optional[list[Any]]) -> list[dict[str, Any]]:
    latest: dict[str, dict[str, Any]] = {}
    for raw in items or []:
        if not _is_live_preview_row(raw):
            continue
        rec = {
            "fileName": _display_name(raw.get("fileName")),
            "documentNumber": str(raw.get("documentNumber") or "").strip(),
            "fileVersion": str(raw.get("fileVersion") or "").strip(),
            "documentDisplayDate": str(raw.get("documentDisplayDate") or "").strip(),
            "targetVersion": str(raw.get("targetVersion") or raw.get("registrationVersion") or "").strip(),
            "notes": str(raw.get("notes") or "").strip(),
            "explanation": str(raw.get("explanation") or "").strip(),
        }
        key = _item_identity(rec)
        prev = latest.get(key)
        latest[key] = rec if prev is None else _merge_latest_item(prev, rec)
    out = list(latest.values())
    out.sort(key=lambda row: (_display_name(row.get("fileName")).casefold(), str(row.get("documentNumber") or "")))
    return out


def _format_list_date(raw: Any) -> str:
    text = str(raw or "").strip()
    if not text:
        return ""
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            dt = datetime.strptime(text[:10], fmt)
            return f"{dt.year}/{dt.month}/{dt.day}"
        except ValueError:
            continue
    return text


def _list_value(item: dict[str, Any], field: str) -> str:
    if field == "documentDisplayDate":
        return _format_list_date(item.get("documentDisplayDate"))
    return str(item.get(field) or "").strip()


def _format_master_version(raw: Any) -> str:
    text = re.sub(r"\s+", "", str(raw or "").strip())
    match = _DOC_REV_RE.match(text)
    if match:
        return f"{match.group(1).upper()}/{int(match.group(2))}"
    match = re.match(r"^(\d+)$", text)
    if match:
        return f"A/{int(match.group(1))}"
    return str(raw or "").strip()


def _replace_cell_text(cell, text: str) -> None:
    """整格重写，避免同一单元格多个 w:t 残留（A/ + 1 改成 A/3 变成 A/31）。"""
    cell.text = str(text or "")


def _parse_existing_seq(raw: Any) -> int:
    match = re.match(r"^\s*(\d+)", str(raw or ""))
    return int(match.group(1)) if match else 0


def _cell_numpr_ppr(cell):
    for para in cell.paragraphs:
        p_pr = para._p.find(qn("w:pPr"))
        if p_pr is None:
            continue
        if p_pr.find(qn("w:numPr")) is not None:
            return p_pr
    return None


def _apply_auto_seq(cell, sample_cell) -> bool:
    src_ppr = _cell_numpr_ppr(sample_cell) if sample_cell is not None else None
    if src_ppr is None:
        return False
    cell.text = ""
    dest = cell.paragraphs[0]
    dest_ppr = dest._p.find(qn("w:pPr"))
    if dest_ppr is not None:
        dest._p.remove(dest_ppr)
    dest._p.insert(0, deepcopy(src_ppr))
    return True


def _row_texts(row) -> list[str]:
    return [str(cell.text or "").strip() for cell in row.cells]


def _is_section_row(row) -> bool:
    texts = _row_texts(row)
    nonempty = [t for t in texts if t]
    if not nonempty:
        return False
    unique_tc = len({id(cell._tc) for cell in row.cells})
    if unique_tc == 1:
        return True
    return len(set(nonempty)) == 1 and len(nonempty) == len(texts)


def _is_header_row(row) -> bool:
    blob = "".join(re.sub(r"\s+", "", t) for t in _row_texts(row))
    return "文件编号" in blob and "文件名称" in blob


def _is_empty_row(row) -> bool:
    return not any(_row_texts(row))


def _match_score(item: dict[str, Any], code: str, name: str) -> float:
    item_code = _norm_code(item.get("documentNumber"))
    row_code = _norm_code(code)
    if item_code and row_code and item_code == row_code:
        return 2.0
    item_name = _display_name(item.get("fileName"))
    row_name = _display_name(name)
    if item_name and row_name and item_name.casefold() == row_name.casefold():
        return 1.5
    item_keys = set(_name_keys(item_name))
    row_keys = set(_name_keys(row_name))
    if item_keys and row_keys and item_keys & row_keys:
        return 1.2
    score = _title_fuzzy_score(item_name, row_name)
    return float(score or 0.0)


def _pick_item(items: list[dict[str, Any]], used: set[int], code: str, name: str) -> Optional[dict[str, Any]]:
    best_idx = None
    best_score = 0.0
    for idx, item in enumerate(items):
        if idx in used:
            continue
        score = _match_score(item, code, name)
        if score > best_score:
            best_score = score
            best_idx = idx
    if best_idx is None or best_score < _FUZZY_MATCH_MIN:
        return None
    used.add(best_idx)
    return items[best_idx]


def _find_item(items: list[dict[str, Any]], code: str, name: str) -> Optional[dict[str, Any]]:
    best = None
    best_score = 0.0
    for item in items:
        score = _match_score(item, code, name)
        if score > best_score:
            best_score = score
            best = item
    if best is None or best_score < _FUZZY_MATCH_MIN:
        return None
    return best


def _note_text(item: dict[str, Any], fallback: str) -> str:
    text = str(item.get("notes") or "").strip() or str(item.get("explanation") or "").strip()
    return text or fallback


def _catalog_should_replace_name(row_name: str, item_name: str) -> bool:
    row_name = _display_name(row_name)
    item_name = _display_name(item_name)
    if not item_name:
        return False
    if not row_name:
        return True
    if row_name.casefold() == item_name.casefold():
        return False
    if row_name.startswith(item_name):
        extra = row_name[len(item_name) :].lstrip()
        if extra.startswith("（") or extra.startswith("("):
            return False
    return True


def _format_value_for_spec(item: dict[str, Any], spec: dict[str, str]) -> str:
    new_value = _list_value(item, spec["value_field"])
    if new_value and spec["value_field"] == "fileVersion":
        return _format_master_version(new_value)
    return new_value


def _update_notes_table(table, catalog: list[dict[str, Any]], spec: dict[str, str]) -> None:
    used: set[int] = set()
    empty_value = spec["empty_value"]
    max_seq = 0
    data_rows = 0
    sample_seq_cell = None

    for idx, row in enumerate(table.rows):
        if idx == 0 or _is_section_row(row) or _is_header_row(row) or _is_empty_row(row):
            continue
        cells = row.cells
        if len(cells) < 5:
            continue
        code = cells[1].text
        name = cells[2].text
        if not str(code or "").strip() and not str(name or "").strip():
            continue
        data_rows += 1
        max_seq = max(max_seq, _parse_existing_seq(cells[0].text), data_rows)
        if sample_seq_cell is None and _cell_numpr_ppr(cells[0]) is not None:
            sample_seq_cell = cells[0]
        item = _pick_item(catalog, used, code, name)
        if item is None:
            continue
        if item.get("documentNumber"):
            _replace_cell_text(cells[1], item["documentNumber"])
        if item.get("fileName"):
            _replace_cell_text(cells[2], item["fileName"])
        new_value = _format_value_for_spec(item, spec)
        if new_value:
            _replace_cell_text(cells[3], new_value)
        _replace_cell_text(cells[4], _note_text(item, empty_value or "/"))

    leftover = [
        item
        for i, item in enumerate(catalog)
        if i not in used and str(item.get("documentNumber") or "").strip()
    ]
    for item in leftover:
        row = table.add_row()
        cells = row.cells
        if not _apply_auto_seq(cells[0], sample_seq_cell):
            max_seq += 1
            _replace_cell_text(cells[0], str(max_seq))
        new_value = _format_value_for_spec(item, spec)
        values = [
            item.get("documentNumber") or "/",
            item.get("fileName") or "",
            new_value or empty_value or "/",
            _note_text(item, empty_value or "/"),
        ]
        for cell, text in zip(cells[1:], values):
            _replace_cell_text(cell, text)


def _update_catalog_doc(doc, catalog: list[dict[str, Any]], spec: dict[str, str]) -> None:
    matched_ids: set[str] = set()
    last_table = None
    for table in doc.tables:
        if not table.rows or len(table.rows[0].cells) < 4:
            continue
        last_table = table
        for idx, row in enumerate(table.rows):
            if idx == 0 or _is_section_row(row) or _is_header_row(row) or _is_empty_row(row):
                continue
            cells = row.cells
            if len(cells) < 4:
                continue
            code = cells[1].text
            name = cells[2].text
            if not str(code or "").strip() and not str(name or "").strip():
                continue
            item = _find_item(catalog, code, name)
            if item is None:
                continue
            matched_ids.add(_item_identity(item))
            if item.get("documentNumber"):
                _replace_cell_text(cells[1], item["documentNumber"])
            if item.get("fileName") and _catalog_should_replace_name(name, item["fileName"]):
                _replace_cell_text(cells[2], item["fileName"])
            new_value = _format_value_for_spec(item, spec)
            if new_value:
                _replace_cell_text(cells[3], new_value)

    leftover = [
        item
        for item in catalog
        if _item_identity(item) not in matched_ids and str(item.get("documentNumber") or "").strip()
    ]
    if not leftover or last_table is None:
        return
    max_seq = 0
    for idx, row in enumerate(last_table.rows):
        if idx == 0 or _is_section_row(row) or _is_header_row(row) or _is_empty_row(row):
            continue
        if len(row.cells) < 4:
            continue
        max_seq = max(max_seq, _parse_existing_seq(row.cells[0].text))
    for item in leftover:
        max_seq += 1
        row = last_table.add_row()
        new_value = _format_value_for_spec(item, spec)
        values = [
            str(max_seq),
            item.get("documentNumber") or "/",
            item.get("fileName") or "",
            new_value or "/",
        ]
        for cell, text in zip(row.cells, values):
            _replace_cell_text(cell, text)


def normalize_export_kinds(raw: Any) -> list[str]:
    if isinstance(raw, str):
        values = [raw]
    elif isinstance(raw, list):
        values = raw
    else:
        values = []
    seen: set[str] = set()
    out: list[str] = []
    for item in values:
        kind = str(item or "").strip().lower()
        if kind not in LIST_KINDS or kind in seen:
            continue
        seen.add(kind)
        out.append(kind)
    out.sort(key=lambda kind: _KIND_ORDER.index(kind) if kind in _KIND_ORDER else 99)
    return out


def export_version_task_list_docx(
    *,
    kind: str,
    items: list[Any],
    product_name: str = "",
) -> tuple[bytes, str, str]:
    spec = LIST_KINDS.get(kind)
    if not spec:
        raise ValueError("不支持的清单类型，请选择主文档清单、技术文件清单或中文目录")
    catalog = collect_unique_latest_items(items)
    if not catalog:
        raise ValueError("当前版本任务清单没有可导出的记录，请先生成预览")

    path = resolve_list_template_path(kind)
    doc = Document(str(path))
    if not doc.tables:
        raise ValueError("清单模板没有表格，无法更新")
    if spec.get("layout") == "catalog":
        _update_catalog_doc(doc, catalog, spec)
    else:
        _update_notes_table(doc.tables[0], catalog, spec)

    label = str(product_name or "").strip()
    prefix = spec["download_prefix"]
    filename = f"{prefix}（{label}）.docx" if label else f"{prefix}.docx"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), filename, spec["ascii_name"]


def _zip_writestr(zf: zipfile.ZipFile, name: str, data: bytes) -> None:
    info = zipfile.ZipInfo(filename=name)
    info.compress_type = zipfile.ZIP_DEFLATED
    info.flag_bits |= 0x800
    zf.writestr(info, data)


def export_version_task_lists(
    *,
    kinds: Any,
    items: list[Any],
    product_name: str = "",
) -> tuple[bytes, str, str, str]:
    selected = normalize_export_kinds(kinds)
    if not selected:
        raise ValueError("请选择要导出的清单：主文档清单、技术文件清单或中文目录")
    files: list[tuple[str, str, bytes]] = []
    for kind in selected:
        raw, filename, ascii_name = export_version_task_list_docx(
            kind=kind,
            items=items,
            product_name=product_name,
        )
        files.append((filename, ascii_name, raw))
    if len(files) == 1:
        filename, ascii_name, raw = files[0]
        return raw, filename, ascii_name, "docx"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for filename, _ascii_name, raw in files:
            _zip_writestr(zf, filename, raw)
    label = str(product_name or "").strip()
    zip_name = f"体系文件清单（{label}）.zip" if label else "体系文件清单.zip"
    return buf.getvalue(), zip_name, "system-file-lists.zip", "zip"
