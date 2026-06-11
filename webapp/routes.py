# -*- coding: utf-8 -*-
from __future__ import annotations

import csv
import html
import io
import json
import os
import re
import secrets
import hashlib
import uuid
import socket
import time as pytime
from datetime import date, datetime, time
from functools import wraps
from pathlib import Path
from typing import Any, Optional

from sqlalchemy import or_
from urllib.error import HTTPError, URLError
from urllib.parse import quote as _urlquote
from urllib.parse import urlencode
from urllib.request import Request, urlopen
import requests

from flask import (
    Blueprint,
    current_app,
    make_response,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    send_from_directory,
    session,
    url_for,
)
from werkzeug.utils import secure_filename

from . import db
from .doc_service import (
    download_template_from_url,
    extract_placeholders,
    extract_placeholders_from_bytes,
    generate_document,
)
from .task_template_archive import resolve_task_template_from_saved_path
from .models import (
    GenerateRecord, GenerationSummary, NoteAttachmentFile, UploadRecord, User,
    TaskTypeConfig, CompletionStatusConfig, AuditStatusConfig, NotifyTemplateConfig, AppConfig,
    Project, ProjectTeam, ModuleCascadeReminder, now_local,
    ExamBankIngestJob,
    ExamSetReviewJob,
    ExamCenterActivity,
    ExamCenterAssignment,
    ExamCenterActivityDetail,
    ExamCenterAssignmentExtra,
    ExamCenterAssignmentAudience,
    ExamAttempt,
    ExamAttemptItem,
    ExamGradingJob,
)
from . import dingtalk_service
from .dingtalk_callback_crypto import (
    DingTalkCallbackCrypto,
    DingTalkCallbackCryptoError,
    build_text_reply_json,
    parse_callback_query_args,
)

bp = Blueprint("pages", __name__)
_chatbot_group_last_reply_at: dict[str, float] = {}


def _dingtalk_webhook_str() -> str:
    from .app_settings import get_setting
    return (get_setting("DINGTALK_WEBHOOK") or "").strip()


def _dingtalk_secret_opt() -> Optional[str]:
    from .app_settings import get_setting
    s = (get_setting("DINGTALK_SECRET") or "").strip()
    return s or None


def _chatbot_dingtalk_webhook_str() -> str:
    """体系记录机器人发回复用 Webhook（与催办 DINGTALK_WEBHOOK 独立；未配时兼容旧单键部署）。"""
    from .app_settings import get_setting

    w = (get_setting("CHATBOT_DINGTALK_WEBHOOK") or "").strip()
    if w:
        return w
    return (get_setting("DINGTALK_WEBHOOK") or "").strip()


def _chatbot_dingtalk_secret_opt() -> Optional[str]:
    from .app_settings import get_setting

    s = (get_setting("CHATBOT_DINGTALK_SECRET") or "").strip()
    if s:
        return s or None
    s = (get_setting("DINGTALK_SECRET") or "").strip()
    return s or None


def _dingtalk_callback_crypto_optional() -> Optional[DingTalkCallbackCrypto]:
    from .app_settings import get_setting

    token = (get_setting("DINGTALK_CALLBACK_TOKEN", default="") or "").strip()
    aes_key = (get_setting("DINGTALK_CALLBACK_AES_KEY", default="") or "").strip()
    owner = (
        get_setting("DINGTALK_CALLBACK_OWNER_KEY", default="")
        or get_setting("DINGTALK_APP_KEY", default="")
        or ""
    ).strip()
    if not token or not aes_key or not owner:
        return None
    try:
        return DingTalkCallbackCrypto(token, aes_key, owner)
    except DingTalkCallbackCryptoError:
        return None


def _chatbot_process_incoming_payload(payload: dict[str, Any]) -> tuple[Optional[dict[str, Any]], Optional[str]]:
    """处理一条已解密的钉钉消息体；返回 (结果 dict, 错误 message)。"""
    if not _chatbot_enabled():
        return {"ignored": "chatbot disabled"}, None

    group_id = _chatbot_extract_group_id(payload)
    groups = _chatbot_enabled_groups()
    if groups and group_id and group_id not in groups:
        return {"ignored": "group not enabled"}, None

    text = _chatbot_extract_text(payload)
    if not text:
        return {"ignored": "empty text"}, None
    if not _chatbot_should_trigger(text):
        return {"ignored": "trigger not matched"}, None
    if _chatbot_is_cooldown(group_id):
        return {"ignored": "cooldown"}, None

    message_id = (
        str(payload.get("msgId") or payload.get("messageId") or payload.get("eventId") or "").strip()
        or uuid.uuid4().hex
    )
    trigger_type = "at_bot" if "@" in text else "keyword"
    recent_messages: list[str] = []
    hist = payload.get("recentMessages")
    if isinstance(hist, list):
        recent_messages = [str(x).strip() for x in hist if str(x).strip()][:6]

    reply_data, err = _chatbot_call_aicheckword(
        query=text,
        group_id=group_id,
        message_id=message_id,
        trigger_type=trigger_type,
        recent_messages=recent_messages,
        provider=_chatbot_llm_provider_from_settings(),
    )
    if err:
        return None, err

    need_human = bool(reply_data.get("need_human"))
    answer = (reply_data.get("answer_summary") or reply_data.get("answer") or "").strip()
    confidence = float(reply_data.get("confidence") or 0.0)
    conf_threshold = _chatbot_confidence_threshold()
    if need_human or confidence < conf_threshold or not answer:
        answer = "这个问题我先帮你转人工确认，稍后给你准确答复。"

    return (
        {
            "success": True,
            "answer": answer,
            "need_human": need_human,
            "confidence": confidence,
            "session_webhook": str(payload.get("sessionWebhook") or "").strip(),
            "reply_data": reply_data,
        },
        None,
    )


def _resolve_dingtalk_for_team(team_id: str | None) -> tuple[str, Optional[str], str]:
    from .dingtalk_team import resolve_dingtalk_credentials

    webhook, secret, source = resolve_dingtalk_credentials(team_id)
    return webhook, secret, source


def _resolve_team_id_by_project_name(project_name: str | None) -> str | None:
    from .dingtalk_team import resolve_team_id_by_project_name

    return resolve_team_id_by_project_name(project_name)


def _chatbot_enabled() -> bool:
    from .app_settings import get_setting
    raw = (get_setting("CHATBOT_ENABLE", default="") or "").strip().lower()
    return raw in ("1", "true", "yes", "on", "y")


def _chatbot_keywords() -> list[str]:
    from .app_settings import get_setting
    raw = (get_setting("DINGTALK_TRIGGER_KEYWORDS", default="") or "").strip()
    if not raw:
        return []
    return [x.strip().lower() for x in raw.split(",") if x.strip()]


def _chatbot_enabled_groups() -> set[str]:
    from .app_settings import get_setting
    raw = (get_setting("CHATBOT_ENABLED_GROUPS", default="") or "").strip()
    if not raw:
        return set()
    return {x.strip() for x in raw.split(",") if x.strip()}


def _chatbot_cooldown_seconds() -> int:
    from .app_settings import get_setting
    raw = (get_setting("CHATBOT_REPLY_COOLDOWN_SECONDS", default="10") or "10").strip()
    try:
        val = int(raw)
    except ValueError:
        val = 10
    return max(1, min(300, val))


def _chatbot_confidence_threshold() -> float:
    from .app_settings import get_setting
    raw = (get_setting("CHATBOT_CONFIDENCE_THRESHOLD", default="0.65") or "0.65").strip()
    try:
        val = float(raw)
    except ValueError:
        val = 0.65
    return max(0.0, min(1.0, val))


def _chatbot_api_base() -> str:
    """聊天回复 API 根地址：未单独配置时与考试训练中心（QUIZ_API_BASE_URL）复用同一 aicheckword 实例。"""
    from .app_settings import get_setting
    raw = (
        get_setting("AICHECKWORD_CHAT_API_BASE", default="")
        or get_setting("QUIZ_API_BASE_URL", default="")
        or get_setting("AICHECKWORD_DRAFT_API_BASE", default="")
        or str(current_app.config.get("QUIZ_API_BASE_URL") or "")
    ).strip()
    return raw.rstrip("/")


def _chatbot_api_key() -> str:
    from .app_settings import get_setting
    return (get_setting("AICHECKWORD_CHAT_API_KEY", default="") or "").strip()


def _chatbot_api_timeout_seconds() -> int:
    """聊天接口含两次 LLM + 向量检索，读超时独立于 QUIZ_API（后者常被压到 20～30s）。"""
    from .app_settings import get_setting

    raw = (get_setting("AICHECKWORD_CHAT_TIMEOUT_SECONDS", default="") or "").strip()
    if raw:
        try:
            val = int(raw)
        except ValueError:
            val = 120
    else:
        val = max(120, _quiz_api_timeout_seconds())
    if val < 30:
        return 30
    if val > 600:
        return 600
    return val


def _chatbot_requests_timeout() -> tuple[int, int]:
    from ._integration_common import integration_requests_timeout

    return integration_requests_timeout(read_seconds=_chatbot_api_timeout_seconds())


_AIWORD_CHATBOT_LLM_PROVIDERS: tuple[str, ...] = ("deepseek", "tongyi", "ollama", "openai", "lingyi")


def _chatbot_llm_provider_from_settings() -> str:
    from .app_settings import get_setting

    raw = (get_setting("CHATBOT_LLM_PROVIDER", default="") or "").strip().lower()
    if raw in _AIWORD_CHATBOT_LLM_PROVIDERS:
        return raw
    return "deepseek"


def _chatbot_normalize_provider(requested: Optional[str]) -> tuple[str, Optional[str]]:
    p = (requested or "").strip().lower()
    if not p:
        p = _chatbot_llm_provider_from_settings()
    if p in _AIWORD_CHATBOT_LLM_PROVIDERS:
        return p, None
    fb = _chatbot_llm_provider_from_settings()
    return fb, f"不支持的 provider={p}，已改用 {fb}"


def _chatbot_client_headers(provider: str) -> dict[str, str]:
    """页面2 个人 LLM 凭据透传；personalKeysOnly 时禁止回落系统 Key。"""
    from .draft_generation_routes import (
        _client_llm_headers,
        _draft_personal_key_headers,
        _draft_personal_keys_enforced,
        _normalize_requested_provider,
    )

    prov = (provider or "").strip().lower()
    headers: dict[str, str] = {}
    if prov:
        headers["X-Client-Llm-Provider"] = prov
    uid = session.get("user_id")
    if not uid:
        return headers
    try:
        norm = _normalize_requested_provider(prov) or prov
        if norm in ("deepseek", "tongyi", "cursor"):
            enforced = _draft_personal_keys_enforced()
            if enforced:
                headers.update(_draft_personal_key_headers() or {})
            personal = _client_llm_headers(str(uid), provider=norm) or {}
            if enforced or personal.get("X-Client-Llm-Api-Key"):
                headers.update(personal)
            headers["X-Client-Llm-Provider"] = norm
    except Exception:
        pass
    return headers


def _chatbot_extract_text(payload: dict[str, Any]) -> str:
    candidates = [
        payload.get("text", {}).get("content") if isinstance(payload.get("text"), dict) else None,
        payload.get("content"),
        payload.get("msg"),
        payload.get("msgContent"),
        payload.get("conversationText"),
    ]
    for c in candidates:
        if isinstance(c, str) and c.strip():
            return c.strip()
    return ""


def _chatbot_extract_group_id(payload: dict[str, Any]) -> str:
    for k in ("conversationId", "openConversationId", "chatbotConversationId", "groupId"):
        v = payload.get(k)
        if isinstance(v, str) and v.strip():
            return v.strip()
    return ""


def _chatbot_should_trigger(text: str) -> bool:
    t = (text or "").strip().lower()
    if not t:
        return False
    if "@" in t:
        return True
    for kw in _chatbot_keywords():
        if kw and kw in t:
            return True
    return False


def _chatbot_is_cooldown(group_id: str) -> bool:
    gid = (group_id or "").strip() or "__default__"
    now = pytime.time()
    last = _chatbot_group_last_reply_at.get(gid) or 0.0
    if now - last < _chatbot_cooldown_seconds():
        return True
    _chatbot_group_last_reply_at[gid] = now
    return False


def _chatbot_call_aicheckword(
    *,
    query: str,
    group_id: str,
    message_id: str,
    trigger_type: str,
    recent_messages: list[str],
    provider: Optional[str] = None,
) -> tuple[Optional[dict[str, Any]], Optional[str]]:
    base = _chatbot_api_base()
    if not base:
        return None, "未配置 aicheckword 地址：请填写 QUIZ_API_BASE_URL 或 AICHECKWORD_CHAT_API_BASE"
    eff_provider, _prov_note = _chatbot_normalize_provider(provider)
    url = f"{base}/api/chat/reply/generate"
    headers = {"Accept": "application/json", "Content-Type": "application/json; charset=utf-8"}
    token = _chatbot_api_key()
    if token:
        headers["Authorization"] = f"Bearer {token}"
    headers.update(_chatbot_client_headers(eff_provider))
    payload = {
        "query": query,
        "group_id": group_id,
        "message_id": message_id,
        "trigger_type": trigger_type,
        "current_provider": eff_provider,
        "context": {"recent_messages": recent_messages},
        "options": {
            "domain": "system_record_writing",
            "knowledge_category": "program",
            "top_k": 6,
            "min_similarity": 0.55,
            "max_reply_chars": 320,
            "max_detail_chars": 2400,
        },
    }
    try:
        r = requests.post(url, headers=headers, json=payload, timeout=_chatbot_requests_timeout())
    except requests.Timeout as e:
        sec = _chatbot_api_timeout_seconds()
        return (
            None,
            f"调用 aicheckword 超时（已等待约 {sec} 秒）：{e}。"
            f"可在系统配置调大 AICHECKWORD_CHAT_TIMEOUT_SECONDS（建议 120～300），"
            f"并确认 aicheckword 已启动且 LLM 可连通。",
        )
    except requests.RequestException as e:
        return None, f"调用 aicheckword 失败: {e}"
    if r.status_code != 200:
        txt = (r.text or "")[:300]
        return None, f"aicheckword HTTP {r.status_code}: {txt}"
    try:
        data = r.json()
    except Exception:
        return None, "aicheckword 返回非 JSON"
    if not isinstance(data, dict):
        return None, "aicheckword 返回格式异常"
    return data, None


def _quiz_api_base_url() -> str:
    from .app_settings import get_setting
    from ._integration_common import resolve_integration_api_base

    raw = get_setting(
        "QUIZ_API_BASE_URL",
        default=str(current_app.config.get("QUIZ_API_BASE_URL") or ""),
    )
    return resolve_integration_api_base((raw or "").strip())


def _quiz_api_timeout_seconds() -> int:
    from .app_settings import get_setting

    raw = (
        get_setting(
            "QUIZ_API_TIMEOUT_SECONDS",
            default=str(current_app.config.get("QUIZ_API_TIMEOUT_SECONDS") or "20"),
        )
        or "20"
    ).strip()
    try:
        val = int(raw)
    except ValueError:
        val = 20
    if val < 3:
        return 3
    # 与 _quiz_api_call 上限对齐：LLM 类接口（如法规更新提示）可能需要数分钟
    if val > 600:
        return 600
    return val


def _stats_upstream_quick_timeout_seconds() -> int:
    """带本地兜底（quiz/stats/*）的 GET：限制上游阻塞，默认最多 8s，仍可被较大 QUIZ_API_TIMEOUT_SECONDS 压低。"""
    return max(3, min(8, _quiz_api_timeout_seconds()))


def _quiz_attempt_answers_quick_timeout_seconds() -> int:
    """quiz/attempts/*/answers：题目列表略大，允许略长于 stats 兜底，但仍限制整段阻塞。"""
    return max(5, min(15, _quiz_api_timeout_seconds()))


def _items_from_activity_detail_snapshot(det: ExamCenterActivityDetail | None) -> list[Any]:
    """从提交时落库的 upstream_payload 尝试恢复题目行，供上游不可用时降级展示。"""
    if not det or not det.upstream_payload or not isinstance(det.upstream_payload, dict):
        return []
    pl = det.upstream_payload
    for key in ("items", "questions", "details", "attempt_items"):
        v = pl.get(key)
        if isinstance(v, list) and len(v) > 0:
            return v
    inner = pl.get("data") if isinstance(pl.get("data"), dict) else None
    if isinstance(inner, dict):
        for key in ("items", "questions", "details"):
            v = inner.get(key)
            if isinstance(v, list) and len(v) > 0:
                return v
    return []


def _exam_org_scoping_enabled() -> bool:
    from .exam_scope import exam_org_scoping_enabled

    return exam_org_scoping_enabled()


def _resolve_exam_organization_id(
    *,
    explicit_org_id: str | None = None,
    assignment_id: str | None = None,
    attempt_id: str | None = None,
    activity_id: str | None = None,
    project_id: str | None = None,
) -> str:
    """解析考试中心上游请求应携带的 organization_id。"""
    from .tenant_context import resolve_organization_context

    if not _exam_org_scoping_enabled():
        return ""

    oid = str(explicit_org_id or "").strip()
    if oid:
        return oid

    aid = str(assignment_id or "").strip()
    if aid:
        row = ExamCenterAssignment.query.filter_by(assignment_id=aid).first()
        oid = str(getattr(row, "organization_id", "") or "").strip() if row else ""
        if oid:
            return oid

    atid = str(attempt_id or "").strip()
    if atid:
        att = ExamAttempt.query.filter_by(attempt_id=atid).first()
        if att:
            oid = str(getattr(att, "organization_id", "") or "").strip()
            if oid:
                return oid
            aid2 = str(getattr(att, "assignment_id", "") or "").strip()
            if aid2:
                row = ExamCenterAssignment.query.filter_by(assignment_id=aid2).first()
                oid = str(getattr(row, "organization_id", "") or "").strip() if row else ""
                if oid:
                    return oid

    actid = str(activity_id or "").strip()
    if actid:
        act = ExamCenterActivity.query.filter_by(id=actid).first()
        if act:
            oid = str(getattr(act, "organization_id", "") or "").strip()
            if oid:
                return oid
            aid3 = str(getattr(act, "assignment_id", "") or "").strip()
            if aid3:
                row = ExamCenterAssignment.query.filter_by(assignment_id=aid3).first()
                oid = str(getattr(row, "organization_id", "") or "").strip() if row else ""
                if oid:
                    return oid

    pid = str(project_id or "").strip()
    if pid:
        proj = Project.query.get(pid)
        oid = str(getattr(proj, "organization_id", "") or "").strip() if proj else ""
        if oid:
            return oid

    from .authz import is_normal_user, is_page13_super_admin, is_project_admin

    if is_page13_super_admin() or is_normal_user() or is_project_admin():
        from .exam_scope import resolve_active_organization_id

        exam_oid = resolve_active_organization_id(write_session=False)
        if exam_oid:
            return exam_oid

    oid, _ = resolve_organization_context()
    return str(oid or "").strip()


def _current_exam_scope_organization_id() -> str:
    """当前会话在考试中心的数据作用域 organization_id（未启用公司体系时返回空串）。"""
    from .exam_scope import exam_org_scoping_enabled, resolve_active_organization_id

    if not exam_org_scoping_enabled():
        return ""
    return resolve_active_organization_id()


def _exam_org_id_sql_clause(org_id: str, column):
    """公司过滤；organization_id 为空的历史记录在当前公司下仍可见（升级期兼容）。"""
    oid = str(org_id or "").strip()
    if not oid:
        return None
    from sqlalchemy import or_

    return or_(column == oid, column.is_(None))


def _ensure_exam_scope_data_repaired() -> None:
    """首次查询前幂等补齐考试历史默认公司/项目组（避免启动后仍查不到旧数据）。"""
    try:
        from flask import g, has_request_context

        if has_request_context():
            if getattr(g, "_exam_scope_data_repaired", False):
                return
            g._exam_scope_data_repaired = True
    except Exception:
        pass
    try:
        from .historical_migration import repair_exam_scope_defaults

        repair_exam_scope_defaults()
    except Exception:
        try:
            from . import db

            db.session.rollback()
        except Exception:
            pass


def _exam_team_scoped_user_ids() -> frozenset[str] | None:
    from .authz import exam_team_scoped_user_ids

    return exam_team_scoped_user_ids()


def _exam_assignment_ids_for_team_scope(
    uids: frozenset[str],
    *,
    org_id: str = "",
) -> set[str]:
    """项目组可见的考试任务 id：受众含组内成员，或组内成员已有活动记录。"""
    if not uids:
        return set()
    aud_aids = {
        str(r.assignment_id).strip()
        for r in ExamCenterAssignmentAudience.query.filter(
            ExamCenterAssignmentAudience.user_id.in_(list(uids))
        ).all()
        if str(r.assignment_id).strip()
    }
    act_q = ExamCenterActivity.query.filter(
        ExamCenterActivity.user_id.in_(list(_expand_exam_activity_user_keys(uids))),
        ExamCenterActivity.assignment_id.isnot(None),
        ExamCenterActivity.assignment_id != "",
    )
    org_clause = _exam_org_id_sql_clause(org_id, ExamCenterActivity.organization_id)
    if org_clause is not None:
        act_q = act_q.filter(org_clause)
    act_aids = {
        str(r.assignment_id).strip()
        for r in act_q.all()
        if str(r.assignment_id).strip()
    }
    return aud_aids | act_aids


def _expand_exam_activity_user_keys(user_keys: frozenset[str] | set[str] | list[str]) -> set[str]:
    from .exam_display_labels import exam_activity_user_id_match_keys

    out: set[str] = set()
    for k in user_keys or []:
        s = str(k or "").strip()
        if not s:
            continue
        out.update(exam_activity_user_id_match_keys(s))
    return out


def _scope_exam_activity_query(q):
    """人员相关：当前公司 + 项目组（学员 user_id）过滤 ExamCenterActivity。"""
    from sqlalchemy import false as sql_false

    _ensure_exam_scope_data_repaired()
    org_id = _current_exam_scope_organization_id()
    org_clause = _exam_org_id_sql_clause(org_id, ExamCenterActivity.organization_id)
    if org_clause is not None:
        q = q.filter(org_clause)
    uids = _exam_team_scoped_user_ids()
    if uids is not None:
        if not uids:
            q = q.filter(sql_false())
        else:
            expanded = _expand_exam_activity_user_keys(uids)
            q = q.filter(ExamCenterActivity.user_id.in_(list(expanded)))
    return q


def _active_exam_team_id_for_observer() -> str:
    """只读列表展示项目组名时，与当前 session 选中的考试中心项目组一致。"""
    from .exam_scope import resolve_active_exam_filter_team_id

    tid = resolve_active_exam_filter_team_id()
    return tid if tid else ""


def _scope_exam_assignment_query(q):
    """下发考试任务：仅按当前公司（organization_id）隔离，不按项目组。"""
    org_id = _current_exam_scope_organization_id()
    org_clause = _exam_org_id_sql_clause(org_id, ExamCenterAssignment.organization_id)
    if org_clause is not None:
        q = q.filter(org_clause)
    return q


def _exam_activity_in_staff_scope(row: ExamCenterActivity | None) -> bool:
    if row is None:
        return False
    from .authz import exam_row_organization_id_matches, user_in_exam_team_scope

    org_id = _current_exam_scope_organization_id()
    if not exam_row_organization_id_matches(getattr(row, "organization_id", None), org_id):
        return False
    return user_in_exam_team_scope(str(getattr(row, "user_id", "") or ""))


def _exam_activity_deletable_by_staff(row: ExamCenterActivity | None) -> bool:
    """删除权限：超管可删当前公司内任意记录；项目管理员仅可删所属项目组学员记录。"""
    if row is None:
        return False
    from .authz import exam_row_organization_id_matches, is_page13_super_admin, is_project_admin
    from .exam_scope import activity_user_belongs_to_teams, project_admin_team_ids_for_org

    org_id = _current_exam_scope_organization_id()
    if not exam_row_organization_id_matches(getattr(row, "organization_id", None), org_id):
        return False
    if is_page13_super_admin():
        return True
    if not is_project_admin():
        return False
    pa_teams = project_admin_team_ids_for_org(org_id)
    if not pa_teams:
        return False
    return activity_user_belongs_to_teams(str(getattr(row, "user_id", "") or ""), pa_teams)


def _exam_assignment_in_staff_scope(row: ExamCenterAssignment | None) -> bool:
    """下发任务：仅校验公司作用域（与项目组无关）。"""
    if row is None:
        return False
    from .authz import exam_row_organization_id_matches

    org_id = _current_exam_scope_organization_id()
    return exam_row_organization_id_matches(getattr(row, "organization_id", None), org_id)


def _validate_exam_audience_user_ids(audience_ids: list[str]) -> str | None:
    """项目管理员下发考试时，受众须为所属项目组内账号。"""
    from .authz import user_in_exam_team_scope

    uids = _exam_team_scoped_user_ids()
    if uids is None:
        return None
    bad = [x for x in audience_ids if not user_in_exam_team_scope(x)]
    if bad:
        return "考试对象须为所属项目组内人员"
    return None


def _project_teams_for_organization(org_id: str) -> list[ProjectTeam]:
    """当前公司下启用且已绑定该公司的项目组（考试中心下拉用，不含测试组）。"""
    from .team_organizations import teams_for_organization

    return teams_for_organization(org_id, active_only=True)


def _resolve_default_exam_team_id_for_org(org_id: str) -> str:
    """默认项目组：互联网产品部（优先匹配当前公司绑定）。"""
    from .exam_scope import default_exam_team_id_for_org

    return default_exam_team_id_for_org(org_id)


def _apply_super_admin_active_exam_team(org_id: str) -> str:
    """超管：校验 session 中项目组；首次进入默认互联网产品部；可选「全部项目组」。"""
    from .authz import is_page13_super_admin

    if not is_page13_super_admin():
        return ""
    if session.get("exam_team_scope_all"):
        session.pop("active_exam_team_id", None)
        return ""
    allowed = {t["id"] for t in _exam_teams_payload_for_scope(org_id)}
    active = str(session.get("active_exam_team_id") or "").strip()
    if active and active in allowed:
        return active
    if active and active not in allowed:
        session.pop("active_exam_team_id", None)
    default_id = _resolve_default_exam_team_id_for_org(org_id)
    if default_id and default_id in allowed:
        session["active_exam_team_id"] = default_id
        session.pop("exam_team_scope_all", None)
        return default_id
    session.pop("active_exam_team_id", None)
    session.pop("exam_team_scope_all", None)
    return ""


def _exam_teams_payload_for_scope(org_id: str) -> list[dict[str, str]]:
    """考试中心作用域内可选项目组（超管按当前公司；项目管理员仅所属组）。"""
    from .authz import is_page13_super_admin, is_project_admin, user_team_ids
    from .team_organizations import organization_ids_for_team

    oid = str(org_id or "").strip()
    if is_page13_super_admin():
        scoped = _project_teams_for_organization(oid)
    elif is_project_admin():
        allow = [str(x).strip() for x in user_team_ids() if str(x).strip()]
        scoped = []
        seen: set[str] = set()
        for tid in allow:
            if oid:
                linked = organization_ids_for_team(tid)
                if linked and oid not in linked:
                    continue
            t = ProjectTeam.query.get(tid)
            if not t or not bool(getattr(t, "is_active", True)):
                continue
            sid = str(t.id or "").strip()
            if not sid or sid in seen:
                continue
            seen.add(sid)
            scoped.append(t)
        scoped.sort(key=lambda x: (int(getattr(x, "sort_order", 0) or 0), str(x.name or "")))
    else:
        scoped = []
    return [{"id": str(t.id), "name": str(t.name or t.id)} for t in scoped]


def _sync_project_admin_active_exam_team(org_id: str) -> str:
    """项目管理员：校验/默认当前公司下的所属项目组。"""
    from .authz import is_page13_super_admin, is_project_admin

    if is_page13_super_admin() or not is_project_admin():
        return ""
    session.pop("exam_team_scope_all", None)
    allowed = {t["id"] for t in _exam_teams_payload_for_scope(org_id)}
    if not allowed:
        session.pop("active_exam_team_id", None)
        return ""
    active = str(session.get("active_exam_team_id") or "").strip()
    if active and active in allowed:
        return active
    if active and active not in allowed:
        session.pop("active_exam_team_id", None)
    teams = _exam_teams_payload_for_scope(org_id)
    if teams:
        pick = str(teams[0].get("id") or "").strip()
        if pick:
            session["active_exam_team_id"] = pick
            return pick
    return ""


def _sync_super_admin_active_exam_team(org_id: str) -> str:
    """超管切换公司后，剔除非法 active_exam_team_id，并回退默认组。"""
    from .authz import is_page13_super_admin

    if not is_page13_super_admin():
        return ""
    active = str(session.get("active_exam_team_id") or "").strip()
    allowed = {t["id"] for t in _exam_teams_payload_for_scope(org_id)}
    if active and active not in allowed:
        session.pop("active_exam_team_id", None)
    return _apply_super_admin_active_exam_team(org_id)


def _quiz_api_headers(*, organization_id: str | None = None) -> dict[str, str]:
    from .app_settings import get_setting

    headers = {
        "Accept": "application/json",
        "Content-Type": "application/json; charset=utf-8",
    }
    bearer = (get_setting("QUIZ_API_BEARER_TOKEN") or "").strip()
    secret = (get_setting("QUIZ_API_SECRET") or "").strip()
    if bearer:
        headers["Authorization"] = f"Bearer {bearer}"
    if secret:
        headers["X-Integration-Secret"] = secret
    oid = str(organization_id or "").strip()
    if oid:
        headers["X-Aiword-Company-Id"] = oid
    return headers


def _quiz_api_call(
    upstream_path: str,
    method: str = "GET",
    payload: Optional[dict[str, Any]] = None,
    query: Optional[dict[str, Any]] = None,
    timeout_seconds: Optional[int] = None,
    organization_id: Optional[str] = None,
) -> tuple[int, dict[str, Any]]:
    base_url = _quiz_api_base_url()
    trace_id = uuid.uuid4().hex
    req_method = (method or "GET").upper()
    request_url = ""
    if not base_url:
        return 503, {
            "code": "QUIZ_API_NOT_CONFIGURED",
            "message": "未配置考试训练中心后端地址，请在页面4 · 系统与钉钉「系统配置」中设置 QUIZ_API_BASE_URL",
            "data": None,
            "trace_id": trace_id,
            "request": {"url": request_url, "method": req_method, "upstreamPath": upstream_path},
        }

    url = f"{base_url}/{upstream_path.lstrip('/')}"
    if query:
        q = {k: v for k, v in query.items() if v is not None and str(v).strip() != ""}
        if q:
            url = f"{url}?{urlencode(q)}"
    request_url = url

    body_bytes = None
    if payload is not None and req_method in {"POST", "PUT", "PATCH", "DELETE"}:
        body_bytes = json.dumps(payload, ensure_ascii=False).encode("utf-8")

    p = payload if isinstance(payload, dict) else {}
    qd = query if isinstance(query, dict) else {}
    oid = _resolve_exam_organization_id(
        explicit_org_id=(
            str(organization_id or "").strip()
            or str(p.get("organization_id") or p.get("organizationId") or qd.get("organization_id") or qd.get("organizationId") or "")
        ),
        assignment_id=(
            str(p.get("assignment_id") or p.get("assignmentId") or qd.get("assignment_id") or qd.get("assignmentId") or "")
        ),
        attempt_id=(
            str(
                p.get("attempt_id")
                or p.get("attemptId")
                or p.get("session_id")
                or qd.get("attempt_id")
                or qd.get("attemptId")
                or ""
            )
        ),
        activity_id=(
            str(p.get("activity_id") or p.get("activityId") or qd.get("activity_id") or qd.get("activityId") or "")
        ),
        project_id=(
            str(p.get("project_id") or p.get("projectId") or qd.get("project_id") or qd.get("projectId") or "")
        ),
    )

    req = Request(
        url=url,
        data=body_bytes,
        headers=_quiz_api_headers(organization_id=oid),
        method=req_method,
    )

    try:
        _timeout = _quiz_api_timeout_seconds() if timeout_seconds is None else int(timeout_seconds)
        if _timeout < 1:
            _timeout = 1
        if _timeout > 600:
            _timeout = 600
        with urlopen(req, timeout=_timeout) as resp:
            raw = resp.read().decode("utf-8", errors="ignore")
            try:
                data = json.loads(raw) if raw else {}
            except json.JSONDecodeError:
                data = {"raw": raw}
            return 200, {
                "code": 0,
                "message": "ok",
                "data": data,
                "trace_id": trace_id,
                "request": {"url": request_url, "method": req_method, "upstreamPath": upstream_path},
                "organization_id": oid or None,
            }
    except HTTPError as e:
        raw = e.read().decode("utf-8", errors="ignore") if hasattr(e, "read") else ""
        try:
            upstream = json.loads(raw) if raw else {}
        except json.JSONDecodeError:
            upstream = {"raw": raw}
        msg = (
            upstream.get("message")
            or upstream.get("detail")
            or f"考试训练中心请求失败（HTTP {e.code}）"
        )
        return e.code, {
            "code": "QUIZ_API_UPSTREAM_ERROR",
            "message": str(msg),
            "data": upstream,
            "trace_id": trace_id,
            "request": {"url": request_url, "method": req_method, "upstreamPath": upstream_path},
        }
    except URLError as e:
        return 503, {
            "code": "QUIZ_API_NETWORK_ERROR",
            "message": f"无法连接考试训练中心后端：{e.reason}",
            "data": None,
            "trace_id": trace_id,
            "request": {"url": request_url, "method": req_method, "upstreamPath": upstream_path},
        }
    except socket.timeout:
        return 504, {
            "code": "QUIZ_API_TIMEOUT",
            "message": (
                "考试训练中心后端响应超时。"
                "请检查 QUIZ_API_BASE_URL 是否正确、aicheckword 服务是否在运行，"
                "或在系统配置中调大 QUIZ_API_TIMEOUT_SECONDS（如 180；法规提示等 LLM 接口上限 600）。"
            ),
            "data": {"timeoutSeconds": _timeout},
            "trace_id": trace_id,
            "request": {"url": request_url, "method": req_method, "upstreamPath": upstream_path},
        }
    except Exception as e:
        return 500, {
            "code": "QUIZ_API_UNKNOWN_ERROR",
            "message": f"考试训练中心请求异常：{e}",
            "data": None,
            "trace_id": trace_id,
            "request": {"url": request_url, "method": req_method, "upstreamPath": upstream_path},
        }


def _quiz_try_paths(
    paths: list[str],
    *,
    method: str,
    payload: Optional[dict[str, Any]] = None,
    query: Optional[dict[str, Any]] = None,
    organization_id: Optional[str] = None,
) -> tuple[int, dict[str, Any], list[str]]:
    """依次尝试多个上游路径（新旧路由不一致时兼容）。"""
    tried: list[str] = []
    last_status = 502
    last_payload: dict[str, Any] = {
        "code": "QUIZ_API_UNKNOWN_ERROR",
        "message": "未尝试任何上游请求",
        "data": None,
        "trace_id": uuid.uuid4().hex,
    }
    for p in paths:
        pp = (p or "").strip().lstrip("/")
        if not pp:
            continue
        tried.append(pp)
        st, pl = _quiz_api_call(
            pp,
            method=method,
            payload=payload,
            query=query,
            organization_id=organization_id,
        )
        last_status, last_payload = int(st), pl if isinstance(pl, dict) else {"data": pl}
        if 200 <= last_status < 300:
            return last_status, last_payload, tried
        if last_status in (404, 405):
            continue
        return last_status, last_payload, tried
    return last_status, last_payload, tried


def _unwrap_quiz_api_success_data(pl: Any) -> Any:
    """_quiz_api_call 成功时整包为 {code:0, data: 上游JSON}，取出上游根对象。"""
    if not isinstance(pl, dict):
        return {}
    if pl.get("code") == 0 and "data" in pl:
        return pl.get("data")
    return pl


def _exam_activity_upstream_root_ok(root: dict[str, Any]) -> bool:
    """上游 JSON 根若含 code 字段，则仅当为成功码时视为业务成功。"""
    if "code" not in root:
        return True
    return root.get("code") in (0, "0", None, "")


def _json_sanitize_for_db(obj: Any, *, max_bytes: int = 4 * 1024 * 1024) -> dict[str, Any]:
    """将嵌套结构转为可被 db.JSON/MySQL 稳定序列化的 plain dict（避免 Decimal 等类型导致 flush 静默失败）。"""
    if not isinstance(obj, dict):
        obj = {}
    try:
        raw = json.dumps(obj, ensure_ascii=False, default=str)
    except Exception:
        return {"aiword_upstream_payload_sanitize_error": True, "hint": "json.dumps_failed"}
    if len(raw.encode("utf-8")) > max_bytes:
        return {
            "aiword_upstream_payload_truncated": True,
            "max_bytes": max_bytes,
            "preview": raw[:8192],
        }
    try:
        out = json.loads(raw)
        return out if isinstance(out, dict) else {"aiword_upstream_payload_wrap": True, "value": out}
    except Exception:
        return {"aiword_upstream_payload_sanitize_error": True, "hint": "json.loads_failed"}


def _scan_grading_pending_failed(obj: Any, depth: int = 0) -> tuple[bool, bool]:
    """在嵌套字典中查找 grading_status：pending → 阅卷中；failed → 阅卷失败。"""
    if depth > 16 or obj is None:
        return False, False
    if isinstance(obj, list):
        p = f = False
        for x in obj:
            if isinstance(x, (dict, list)):
                cp, cf = _scan_grading_pending_failed(x, depth + 1)
                p = p or cp
                f = f or cf
        return p, f
    if not isinstance(obj, dict):
        return False, False
    g = str(obj.get("grading_status") or "").strip().casefold()
    pend = g == "pending"
    failed = g == "failed"
    for _k, v in obj.items():
        if isinstance(v, (dict, list)):
            cp, cf = _scan_grading_pending_failed(v, depth + 1)
            pend = pend or cp
            failed = failed or cf
        if pend and failed:
            break
    return pend, failed


def _extract_result_metrics(root: Any) -> dict[str, Any]:
    """尽量从上游结果中提取分数/对错/薄弱项建议（字段名兼容多形态）。"""
    if not isinstance(root, dict):
        return {}
    inner = root.get("data") if isinstance(root.get("data"), dict) else {}

    def _pick(*vals):
        for v in vals:
            if v is not None and str(v).strip() != "":
                return v
        return None

    score = _pick(root.get("score"), root.get("total_score"), inner.get("score"), inner.get("total_score"))
    total = _pick(root.get("max_score"), root.get("full_score"), inner.get("max_score"), inner.get("full_score"))
    correct = _pick(root.get("correct_count"), root.get("correct"), inner.get("correct_count"), inner.get("correct"))
    wrong = _pick(root.get("wrong_count"), root.get("wrong"), inner.get("wrong_count"), inner.get("wrong"))
    weakness = _pick(root.get("weakness"), root.get("weak_points"), inner.get("weakness"), inner.get("weak_points"))
    reco = _pick(
        root.get("next_step"), root.get("recommendation"), root.get("suggestion"),
        inner.get("next_step"), inner.get("recommendation"), inner.get("suggestion")
    )
    gp, gf = _scan_grading_pending_failed(root)
    out: dict[str, Any] = {
        "score": None,
        "total_score": None,
        "correct_count": None,
        "wrong_count": None,
        "grading_pending": gp,
        "grading_failed": gf,
    }
    try:
        out["score"] = float(score) if score is not None else None
    except Exception:
        out["score"] = None
    try:
        out["total_score"] = float(total) if total is not None else None
    except Exception:
        out["total_score"] = None
    try:
        out["correct_count"] = int(correct) if correct is not None else None
    except Exception:
        out["correct_count"] = None
    try:
        out["wrong_count"] = int(wrong) if wrong is not None else None
    except Exception:
        out["wrong_count"] = None
    if gp:
        out["score"] = None
        out["total_score"] = None
        out["correct_count"] = None
        out["wrong_count"] = None
    out["weakness"] = str(weakness)[:1000] if weakness is not None and str(weakness).strip() else None
    out["recommendation"] = str(reco)[:1000] if reco is not None and str(reco).strip() else None
    return out


def _exam_pass_score() -> float:
    from .app_settings import get_setting

    raw = str(get_setting("EXAM_PASS_SCORE") or "80").strip()
    try:
        v = float(raw)
    except Exception:
        v = 80.0
    if v < 0:
        v = 0.0
    if v > 1000:
        v = 1000.0
    return v


def _norm_answer_plain(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (dict, list)):
        try:
            return json.dumps(v, ensure_ascii=False, sort_keys=True)
        except Exception:
            return str(v)
    return str(v).strip()


def _activity_item_options(row: dict[str, Any]) -> list[Any]:
    raw = row.get("options") or row.get("choices")
    if isinstance(raw, str) and raw.strip():
        try:
            p = json.loads(raw)
            return list(p) if isinstance(p, list) else []
        except Exception:
            return []
    return list(raw) if isinstance(raw, list) else []


def _letter_choice_index_activity(raw: Any) -> int | None:
    if raw is None or isinstance(raw, (dict, list)):
        return None
    s = str(raw).strip().upper()
    if len(s) != 1 or s < "A" or s > "Z":
        return None
    return ord(s) - ord("A")


def _resolve_letter_to_option_aiword(value: Any, options: list[Any]) -> Any:
    if not options:
        return value
    ix = _letter_choice_index_activity(value)
    if ix is None or ix >= len(options):
        return value
    return options[ix]


def _true_false_to_bool_aiword(v: Any) -> bool:
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return float(v) != 0.0
    if isinstance(v, str):
        t = v.strip().lower()
        if t in ("false", "0", "no", "n", "f", "wrong", "错误", "错", "否", "不正确", "不对"):
            return False
        if t in ("true", "1", "yes", "y", "t", "正确", "对", "是", "√"):
            return True
        return False
    return False


def _tf_maybe_literal_aiword(v: Any) -> bool:
    if isinstance(v, bool):
        return True
    if not isinstance(v, str):
        return False
    s = v.strip()
    sl = s.lower()
    if sl in ("true", "false", "1", "0"):
        return True
    return s in ("正确", "错误", "对", "错", "是", "否")


def _norm_single_choice_key_aiword(value: Any, options: list[Any]) -> str:
    v = _resolve_letter_to_option_aiword(value, options)
    if v is None:
        return ""
    return str(v).strip().lower()


def _norm_multiple_choice_set_aiword(raw: Any, options: list[Any]) -> set[str]:
    if raw is None:
        return set()
    items = raw if isinstance(raw, list) else [raw]
    out: set[str] = set()
    for x in items:
        v = _resolve_letter_to_option_aiword(x, options)
        k = str(v).strip().lower() if v is not None else ""
        if k:
            out.add(k)
    return out


def _question_type_from_activity_row(row: dict[str, Any]) -> str:
    return str(row.get("question_type") or row.get("type") or "").strip().lower()


def _objective_answers_equivalent_aiword(row: dict[str, Any], ua: Any, ca: Any) -> bool:
    """与 aicheckword quiz 判分维度对齐：选项字母先解析为选项原文，判断题再按真假比较。"""
    opts = _activity_item_options(row)
    qt = _question_type_from_activity_row(row)
    if qt == "true_false":
        aa = _resolve_letter_to_option_aiword(ca, opts)
        uu = _resolve_letter_to_option_aiword(ua, opts)
        return _true_false_to_bool_aiword(aa) == _true_false_to_bool_aiword(uu)
    if qt == "single_choice":
        return _norm_single_choice_key_aiword(ua, opts) == _norm_single_choice_key_aiword(ca, opts)
    if qt == "multiple_choice":
        s_ca = _norm_multiple_choice_set_aiword(ca, opts)
        s_ua = _norm_multiple_choice_set_aiword(ua, opts)
        return bool(s_ca) and s_ca == s_ua
    if opts:
        if _norm_single_choice_key_aiword(ua, opts) == _norm_single_choice_key_aiword(ca, opts):
            return True
        u_r = _resolve_letter_to_option_aiword(ua, opts)
        c_r = _resolve_letter_to_option_aiword(ca, opts)
        if isinstance(c_r, bool) or _tf_maybe_literal_aiword(c_r) or _tf_maybe_literal_aiword(u_r):
            return _true_false_to_bool_aiword(c_r) == _true_false_to_bool_aiword(u_r)
    return _norm_answer_plain(ua) == _norm_answer_plain(ca)


def _exam_activity_history_result_text(mode: str, metrics: dict[str, Any], raw_msg: Any) -> str:
    """列表「结果」列：用分数推导通过/不通过，避免出现英文 ok。"""
    try:
        if metrics.get("grading_pending"):
            return "阅卷中"
        if metrics.get("grading_failed"):
            return "阅卷失败"
        sc = metrics.get("score")
        if sc is not None:
            thr = float(_exam_pass_score())
            fv = float(sc)
            return "通过" if fv >= thr - 1e-9 else "不通过"
    except Exception:
        pass
    s = str(raw_msg or "").strip()
    low = s.casefold()
    if not s or low in ("ok", "success"):
        return "已提交"
    return s[:500]


def _first_nested_item_list(up: Any, depth: int = 6) -> list[dict[str, Any]]:
    """从上游结果树中挑出首个「题目明细」数组（dict list）。"""
    if depth <= 0 or not isinstance(up, dict):
        return []
    for key in ("items", "questions", "attempt_items", "details"):
        raw = up.get(key)
        if isinstance(raw, list) and raw and isinstance(raw[0], dict):
            return [x for x in raw if isinstance(x, dict)]
    inn = up.get("data") if isinstance(up.get("data"), dict) else None
    if isinstance(inn, dict):
        nested = _first_nested_item_list(inn, depth - 1)
        if nested:
            return nested
    return []


def _qid_from_any_row(it: dict[str, Any]) -> str:
    return str(it.get("question_id") or it.get("questionId") or it.get("id") or "").strip()


def _upstream_row_user_answer(base: dict[str, Any]) -> Any:
    for k in (
        "user_answer",
        "selected_answer",
        "userAnswer",
        "student_answer",
        "response",
        "selected",
        "response_text",
    ):
        if k in base and base.get(k) not in (None, ""):
            return base.get(k)
    return None


def _standard_answer_any_row(it: dict[str, Any]) -> Any:
    return (
        it.get("correct_answer")
        if it.get("correct_answer") is not None
        else it.get("answer")
        if it.get("answer") is not None
        else (
            it.get("standard_answer")
            if it.get("standard_answer") is not None
            else it.get("solution")
        )
    )


def _merge_upstream_snapshot_with_submitted_answers(up_root: dict[str, Any], body: dict[str, Any]) -> dict[str, Any]:
    """
    合并请求体中学生作答与题干快照（submission_questions_snapshot），写入顶层 items，
    供详情弹窗降级为本地快照时仍能展示 student_answer（user_answer）。
    """
    out = dict(up_root or {})
    raw_answers = body.get("answers") or []
    ans_rows: list[dict[str, Any]] = []
    if isinstance(raw_answers, dict):
        ans_rows = [{"question_id": str(k), "answer": v} for k, v in raw_answers.items() if str(k).strip()]
    elif isinstance(raw_answers, list):
        ans_rows = [x for x in raw_answers if isinstance(x, dict)]
    by_qid: dict[str, Any] = {}
    for r in ans_rows:
        q = str(r.get("question_id") or "").strip()
        if not q:
            continue
        if r.get("answer") is not None:
            by_qid[q] = r.get("answer")
        elif r.get("user_answer") is not None:
            by_qid[q] = r.get("user_answer")
    stem_map: dict[str, str] = {}
    snaps = body.get("submission_questions_snapshot") or body.get("questions_snapshot") or []
    if isinstance(snaps, list):
        for s in snaps:
            if isinstance(s, dict):
                qi = _qid_from_any_row(s)
                if qi:
                    stem_map[qi] = str(s.get("stem") or s.get("title") or "").strip()

    upstream_items = _first_nested_item_list(out)

    def _finalize_row(base: dict[str, Any], qid: str) -> dict[str, Any]:
        row = dict(base)
        ua_v = _upstream_row_user_answer(row)
        if ua_v is None:
            ua_v = by_qid.get(qid)
        row["user_answer"] = ua_v
        stem_f = stem_map.get(qid) or ""
        if stem_f and not str(row.get("stem") or "").strip():
            row["stem"] = stem_f
        ca_v = row.get("correct_answer")
        if ca_v is None:
            ca_v = _standard_answer_any_row(row)
            if ca_v is not None:
                row.setdefault("answer", ca_v)

        needs_compare = ua_v is not None and ca_v is not None
        ic_any = row.get("is_correct")
        if ic_any is None:
            ic_any = row.get("isCorrect")
        if ic_any in (0, 1, True, False):
            row["is_correct"] = bool(ic_any)
        elif needs_compare:
            row["is_correct"] = _objective_answers_equivalent_aiword(row, ua_v, ca_v)
        else:
            row["is_correct"] = None
        return row

    merged: list[dict[str, Any]] = []

    if upstream_items:
        by_up: dict[str, dict[str, Any]] = {}
        ordered_ids: list[str] = []
        for ix, x in enumerate(upstream_items):
            qid = _qid_from_any_row(x) or ("uq-" + str(ix))
            if qid not in by_up:
                ordered_ids.append(qid)
                by_up[qid] = dict(x)
        for qid in by_qid:
            if qid not in by_up:
                ordered_ids.append(qid)

        ids = ordered_ids if ordered_ids else sorted(by_qid.keys())

        if not ids:
            ids = [
                (_qid_from_any_row(upstream_items[idx]) or ("uq-" + str(idx)))
                for idx in range(len(upstream_items))
            ]

        seen: set[str] = set()

        def _consume(qid: str) -> dict[str, Any]:
            blk = dict(by_up[qid]) if qid in by_up else {"question_id": qid}
            return _finalize_row(blk, qid)

        for qid in ids:
            qid_key = qid.strip()
            if not qid_key or qid_key in seen:
                continue
            merged.append(_consume(qid_key))
            seen.add(qid_key)

    if not merged and by_qid:
        for qid, ua in sorted(by_qid.items()):
            merged.append(_finalize_row({"question_id": qid}, qid))

    if merged:
        out["items"] = merged
        out["attempt_items"] = merged
        out["attempt_snapshot_aiword_v1"] = True
        if isinstance(out.get("data"), dict):
            od = dict(out["data"])
            od["items"] = merged
            out["data"] = od

    elif ans_rows:
        mini = [
            {
                "question_id": str(a.get("question_id") or ""),
                "stem": stem_map.get(str(a.get("question_id") or ""), ""),
                "user_answer": a.get("answer"),
                "answer": None,
                "is_correct": None,
            }
            for a in ans_rows
            if isinstance(a, dict)
        ]
        out["attempt_items"] = mini
        out["items"] = mini
        out["attempt_snapshot_aiword_v1"] = True

    return out


def _persist_exam_center_activity(
    *,
    user_id: str,
    organization_id: str | None,
    username: str | None,
    display_name: str | None,
    mode: str,
    exam_track: str | None,
    exam_category: str | None = None,
    set_id: str | None,
    assignment_id: str | None,
    assignment_label: str | None,
    attempt_id: str | None,
    upstream_http_status: int,
    upstream_trace_id: str | None,
    result_summary: str | None,
    upstream_result_payload: dict[str, Any] | None = None,
    created_at=None,
    commit: bool = True,
) -> str | None:
    """写入 exam_center_activities（及明细）；返回 activity id。"""
    uid = str(user_id or "").strip()
    if not uid:
        return None
    try:
        row = ExamCenterActivity(
            organization_id=(str(organization_id or "").strip() or None),
            user_id=uid,
            username=str(username) if username else None,
            display_name=str(display_name) if display_name else None,
            mode=str(mode or "").strip()[:16] or "unknown",
            exam_track=(exam_track or "").strip() or None,
            exam_category=_normalize_exam_category(exam_category or "daily"),
            set_id=(set_id or "").strip() or None,
            assignment_id=(assignment_id or "").strip() or None,
            assignment_label=(assignment_label or "").strip() or None,
            attempt_id=(attempt_id or "").strip() or None,
            upstream_http_status=int(upstream_http_status),
            upstream_trace_id=(upstream_trace_id or "").strip() or None,
            result_summary=(result_summary or "").strip()[:500] or None,
        )
        if created_at is not None:
            row.created_at = created_at
        db.session.add(row)
        db.session.flush()

        if upstream_result_payload and isinstance(upstream_result_payload, dict):
            try:
                safe_pl = _json_sanitize_for_db(upstream_result_payload)
                m = _extract_result_metrics(safe_pl)
                detail = ExamCenterActivityDetail(
                    activity_id=row.id,
                    mode=row.mode,
                    score=m.get("score"),
                    total_score=m.get("total_score"),
                    correct_count=m.get("correct_count"),
                    wrong_count=m.get("wrong_count"),
                    weakness=m.get("weakness"),
                    recommendation=m.get("recommendation"),
                    upstream_payload=safe_pl,
                )
                with db.session.begin_nested():
                    db.session.add(detail)
            except Exception as detail_exc:
                try:
                    current_app.logger.warning(
                        "exam_center_activity_detail_skipped uid=%s mode=%s activity_id=%s: %s",
                        uid,
                        mode,
                        row.id,
                        detail_exc,
                        exc_info=True,
                    )
                except Exception:
                    pass

        if commit:
            db.session.commit()
        return str(row.id)
    except Exception as exc:
        db.session.rollback()
        try:
            current_app.logger.exception(
                "exam_center_activity_persist_failed uid=%s mode=%s attempt_id=%s: %s",
                uid,
                mode,
                (attempt_id or "").strip(),
                exc,
            )
        except Exception:
            pass
        return None


def _backfill_exam_center_activities_from_attempts() -> int:
    """历史 exam_attempts 补写 exam_center_activities（旧版 submit-local 未落活动表）。"""
    from .models import ExamAttempt, User
    from .tenant_context import default_organization

    dorg = default_organization()
    default_oid = str(getattr(dorg, "id", "") or "").strip()
    created = 0
    for att in ExamAttempt.query.order_by(ExamAttempt.created_at.asc()).all():
        attempt_key = str(getattr(att, "attempt_id", None) or "").strip()
        if not attempt_key:
            continue
        if ExamCenterActivity.query.filter_by(attempt_id=attempt_key).first():
            continue
        uid = str(getattr(att, "user_id", None) or "").strip()
        if not uid:
            continue
        u = User.query.filter_by(id=uid).first()
        assign_row = ExamCenterAssignment.query.filter_by(
            assignment_id=str(getattr(att, "assignment_id", None) or "").strip()
        ).first()
        set_id_log = str(getattr(assign_row, "set_id", None) or "").strip() or None
        assignment_label_log = str(getattr(assign_row, "title", None) or "").strip() or None
        org_id = str(getattr(att, "organization_id", None) or "").strip() or default_oid or None
        state_l = str(getattr(att, "state", None) or "").strip().lower()
        gs = "pending" if state_l == "grading" else ("graded" if state_l == "graded" else state_l or "submitted")
        items_snap = _local_exam_attempt_items_payload(attempt_key, att=att) or []
        merged_log: dict[str, Any] = {
            "code": 0,
            "message": "backfill-from-attempt",
            "grading_status": gs,
            "score": getattr(att, "score", None),
            "total_score": getattr(att, "total_score", None),
            "correct_count": getattr(att, "correct_count", None),
            "wrong_count": getattr(att, "wrong_count", None),
            "data": {
                "score": getattr(att, "score", None),
                "total_score": getattr(att, "total_score", None),
                "correct_count": getattr(att, "correct_count", None),
                "wrong_count": getattr(att, "wrong_count", None),
                "grading_status": gs,
                "attempt_id": attempt_key,
                "assignment_id": getattr(att, "assignment_id", None),
                "set_id": set_id_log,
            },
            "attempt_items": items_snap,
            "aiword_backfill_from_attempt_v1": True,
        }
        mx = _extract_result_metrics(merged_log)
        summ = _exam_activity_history_result_text("exam", mx, merged_log.get("message"))
        ts = getattr(att, "submitted_at", None) or getattr(att, "started_at", None) or getattr(att, "created_at", None)
        act_id = _persist_exam_center_activity(
            user_id=uid,
            organization_id=org_id,
            username=str(getattr(u, "username", None) or "") or None,
            display_name=str(getattr(u, "display_name", None) or getattr(u, "username", None) or "") or None,
            mode="exam",
            exam_track=str(getattr(att, "exam_track", None) or "").strip() or None,
            exam_category=str(getattr(att, "exam_category", None) or "").strip() or "daily",
            set_id=set_id_log,
            assignment_id=str(getattr(att, "assignment_id", None) or "").strip() or None,
            assignment_label=assignment_label_log,
            attempt_id=attempt_key,
            upstream_http_status=200,
            upstream_trace_id=None,
            result_summary=summ,
            upstream_result_payload=merged_log,
            created_at=ts,
            commit=False,
        )
        if act_id:
            created += 1
    if created:
        db.session.commit()
    return created


def _log_student_exam_center_activity(
    *,
    mode: str,
    exam_track: str | None,
    exam_category: str | None = None,
    set_id: str | None,
    assignment_id: str | None,
    assignment_label: str | None,
    attempt_id: str | None,
    upstream_http_status: int,
    upstream_trace_id: str | None,
    result_summary: str | None,
    upstream_result_payload: dict[str, Any] | None = None,
) -> None:
    from .exam_scope import organization_id_for_exam_write

    uid = str(session.get("user_id") or "").strip()
    if not uid:
        try:
            current_app.logger.warning("exam_center_activity_skip_no_uid_in_session mode=%s", mode)
        except Exception:
            pass
        return
    s_username = str(session.get("username") or "").strip()
    s_display_name = str(session.get("display_name") or "").strip()
    u = User.query.filter_by(id=uid).first()
    username = s_username or (u.username if u else "")
    display_name = s_display_name or ((u.display_name or u.username) if u else "")
    org_id = (
        organization_id_for_exam_write(
            _resolve_exam_organization_id(
                assignment_id=(assignment_id or "").strip(),
                attempt_id=(attempt_id or "").strip(),
            )
        )
        or None
    )
    _persist_exam_center_activity(
        user_id=uid,
        organization_id=org_id,
        username=str(username) if username else None,
        display_name=str(display_name) if display_name else None,
        mode=mode,
        exam_track=exam_track,
        exam_category=exam_category,
        set_id=set_id,
        assignment_id=assignment_id,
        assignment_label=assignment_label,
        attempt_id=attempt_id,
        upstream_http_status=upstream_http_status,
        upstream_trace_id=upstream_trace_id,
        result_summary=result_summary,
        upstream_result_payload=upstream_result_payload,
        commit=True,
    )


def _extract_assignments_from_quiz_root(root: Any) -> list[dict[str, Any]]:
    if not isinstance(root, dict):
        return []
    inner = root.get("data") if isinstance(root.get("data"), dict) else None
    cands: list[Any] = [
        root.get("assignments"),
        root.get("items"),
        inner.get("assignments") if inner else None,
        inner.get("items") if inner else None,
    ]
    for c in cands:
        if isinstance(c, list) and c:
            return [x for x in c if isinstance(x, dict)]
    for c in cands:
        if isinstance(c, list):
            return [x for x in c if isinstance(x, dict)]
    return []


_ASSIGNMENT_DUE_KEYS: tuple[str, ...] = (
    "due_at",
    "dueAt",
    "due_date",
    "dueDate",
    "deadline",
    "complete_before",
    "completeBefore",
)


def _parse_assignment_due_from_request(req: dict[str, Any]) -> tuple[Optional[datetime], bool]:
    """从老师下发请求解析截止完成时间。(due_at, 是否应更新镜像字段)；未出现任何 due 相关键则第二项为 False。"""
    if not isinstance(req, dict):
        return None, False
    specified = False
    raw: Any = None
    for k in _ASSIGNMENT_DUE_KEYS:
        if k not in req:
            continue
        specified = True
        v = req.get(k)
        if v is not None and str(v).strip() != "":
            raw = v
            break
    if not specified:
        return None, False
    if raw is None:
        return None, True
    s = str(raw).strip()
    if not s:
        return None, True
    if re.match(r"^\d{4}-\d{2}-\d{2}$", s) and "T" not in s and " " not in s:
        try:
            d0 = date.fromisoformat(s)
            return datetime.combine(d0, time(23, 59, 59)), True
        except ValueError:
            return None, True
    s_norm = s.replace("Z", "+00:00")
    cand = s_norm.replace(" ", "T", 1)
    try:
        dt = datetime.fromisoformat(cand)
        if dt.tzinfo:
            dt = dt.replace(tzinfo=None)
        return dt, True
    except ValueError:
        return None, True


def _deadline_completion_from_exam_rows(
    rows: list[Any], due_end: Optional[datetime]
) -> dict[str, Any]:
    """按 assignment 下每名用户首次 mode=exam 的 created_at 与截止时刻统计按时/逾期。"""
    first: dict[str, datetime] = {}
    for r in rows:
        if str(getattr(r, "mode", None) or "").lower() != "exam":
            continue
        uid = str(getattr(r, "user_id", None) or "").strip()
        ca = getattr(r, "created_at", None)
        if not uid or not ca:
            continue
        prev = first.get(uid)
        if prev is None or ca < prev:
            first[uid] = ca
    out: dict[str, Any] = {
        "due_at": due_end.isoformat() if due_end else None,
        "exam_students_submitted": len(first),
        "on_time_count": None,
        "late_count": None,
    }
    if not due_end:
        out["note"] = "未配置截止时间，无法区分按时/逾期。"
        return out
    on_time = 0
    late = 0
    for _uid, ts in first.items():
        if ts <= due_end:
            on_time += 1
        else:
            late += 1
    out["on_time_count"] = on_time
    out["late_count"] = late
    out["note"] = (
        "按每名学生在任务下首次考试提交时间与截止时刻比对（本地 activity.created_at）；"
        "截止时间后仍可开考，晚于截止的提交计入逾期。"
    )
    return out


def _normalize_student_assignment_row(x: dict[str, Any]) -> dict[str, Any]:
    aid = str(x.get("assignment_id") or x.get("id") or x.get("assignmentId") or "").strip()
    title = str(x.get("title") or x.get("name") or x.get("label") or aid).strip() or aid
    row: dict[str, Any] = {"id": aid, "name": title, "label": title}
    diff = str(x.get("difficulty") or x.get("difficulty_level") or x.get("difficultyLevel") or "").strip().lower()
    if diff in ("easy", "medium", "hard"):
        row["difficulty"] = diff
    for dk in ("due_at", "dueAt", "due_date", "dueDate"):
        dv = x.get(dk)
        if dv is not None and str(dv).strip():
            row["due_at"] = str(dv).strip()
            break
    for sk in ("set_id", "setId", "quiz_set_id", "quizSetId"):
        sv = x.get(sk)
        if sv is not None and str(sv).strip():
            svs = str(sv).strip()
            row["set_id"] = svs
            row["setId"] = svs
            break
    return row


def _difficulty_from_quiz_get_set_payload(pl: dict[str, Any]) -> str:
    """从 quiz/sets/{id} 网关响应中解析套题难度（与 aicheckword set_config.difficulty 一致）。"""
    up = _unwrap_quiz_api_success_data(pl)
    if not isinstance(up, dict):
        return ""
    root = up.get("data") if up.get("ok") is True and isinstance(up.get("data"), dict) else up
    if not isinstance(root, dict):
        return ""
    cfg = root.get("set_config") or root.get("setConfig") or {}
    if isinstance(cfg, str):
        try:
            cfg = json.loads(cfg)
        except Exception:
            cfg = {}
    if not isinstance(cfg, dict):
        return ""
    d = str(cfg.get("difficulty") or "").strip().lower()
    return d if d in ("easy", "medium", "hard") else ""


def _extract_assignment_from_quiz_payload(payload: dict[str, Any]) -> tuple[str, str, str | None]:
    """从 quiz/assignments 返回中提取 assignment_id/title/set_id。"""
    root = _unwrap_quiz_api_success_data(payload)
    if not isinstance(root, dict):
        root = {}
    inner = root.get("data") if isinstance(root.get("data"), dict) else {}
    aid = str(
        root.get("assignment_id")
        or root.get("assignmentId")
        or root.get("id")
        or inner.get("assignment_id")
        or inner.get("assignmentId")
        or inner.get("id")
        or ""
    ).strip()
    title = str(
        root.get("title")
        or root.get("name")
        or root.get("label")
        or inner.get("title")
        or inner.get("name")
        or inner.get("label")
        or aid
    ).strip() or aid
    set_id = str(root.get("set_id") or root.get("setId") or inner.get("set_id") or inner.get("setId") or "").strip() or None
    return aid, title, set_id


def _expand_question_count_aliases(body: dict[str, Any]) -> dict[str, Any]:
    """aicheckword 各版本字段名不一致时，把题量同步写入多种常见键名。"""
    out = dict(body or {})
    raw = (
        out.get("question_count")
        or out.get("questionCount")
        or out.get("count")
        or out.get("size")
        or out.get("num_questions")
        or out.get("numQuestions")
        or out.get("n")
        or out.get("target_count")
        or out.get("targetCount")
    )
    try:
        n = int(raw) if raw is not None and str(raw).strip() != "" else None
    except (TypeError, ValueError):
        n = None
    if n is None:
        return out
    n = max(1, min(n, 500))
    for key in (
        "question_count",
        "questionCount",
        "count",
        "size",
        "num_questions",
        "numQuestions",
        "n",
        "target_count",
        "targetCount",
    ):
        out[key] = n
    return out


def _normalize_exam_category(v: Any) -> str:
    s = str(v or "").strip().lower()
    if s in ("new_standard", "new_standard_release", "newstd", "新标", "新标发布"):
        return "new_standard"
    if s in (
        "project_case",
        "projectcase",
        "case_audit",
        "caseaudit",
        "项目案例",
        "案例考试",
        "项目案例考试",
    ):
        return "project_case"
    return "daily"


def _expand_project_case_id_aliases(body: dict[str, Any]) -> dict[str, Any]:
    """project_case_id / projectCaseId → 同步为 int，便于上游 Pydantic 校验。"""
    out = dict(body or {})
    raw = out.get("project_case_id")
    if raw is None:
        raw = out.get("projectCaseId")
    if raw is None or str(raw).strip() == "":
        return out
    try:
        pid = int(raw)
    except (TypeError, ValueError):
        out["project_case_id"] = raw
        out["projectCaseId"] = raw
        return out
    if pid <= 0:
        return out
    out["project_case_id"] = pid
    out["projectCaseId"] = pid
    return out


def _expand_exam_category_aliases(body: dict[str, Any]) -> dict[str, Any]:
    """考试类型：daily=日常考试；new_standard=新标发布（与体考类型 exam_track 正交）。历史数据迁移后统一为 daily。"""
    out = dict(body or {})
    raw = (
        out.get("exam_category")
        or out.get("examCategory")
        or out.get("exam_kind")
        or out.get("examKind")
        or ""
    )
    cat = _normalize_exam_category(raw)
    for key in ("exam_category", "examCategory", "exam_kind", "examKind"):
        out[key] = cat
    return out


def _expand_quiz_request_body(body: dict[str, Any] | None) -> dict[str, Any]:
    return _expand_project_case_id_aliases(
        _expand_exam_category_aliases(_expand_exam_track_and_difficulty_aliases(_expand_question_count_aliases(body or {})))
    )


def _expand_exam_track_and_difficulty_aliases(body: dict[str, Any]) -> dict[str, Any]:
    """体考类型、难度与页面设置对齐：写入多种常见键名，避免上游只认 camelCase 或其它字段。"""
    out = dict(body or {})
    track_raw = (
        out.get("exam_track")
        or out.get("examTrack")
        or out.get("track")
        or out.get("exam_type")
        or out.get("examType")
        or out.get("exam_track_code")
        or ""
    )
    track = str(track_raw).strip() or "cn"
    for key in (
        "exam_track",
        "examTrack",
        "track",
        "exam_type",
        "examType",
    ):
        out[key] = track

    diff_raw = (
        out.get("difficulty")
        or out.get("difficulty_level")
        or out.get("difficultyLevel")
        or out.get("level")
        or out.get("diffLevel")
    )
    if diff_raw is None or str(diff_raw).strip() == "":
        return out
    diff = str(diff_raw).strip()
    for key in ("difficulty", "difficulty_level", "difficultyLevel", "level", "diffLevel"):
        out[key] = diff
    return out


_EXAM_TRACK_REQUIREMENT_PROFILES: dict[str, dict[str, Any]] = {
    "cn": {
        "title": "国内体考",
        "scopes": ["GMP", "核查指南", "42061法规"],
        "min_total_questions": 200,
        "topic_groups": [
            {"topic": "GMP", "keywords": ["gmp", "药品生产质量管理规范", "生产质量管理规范"]},
            {"topic": "核查指南", "keywords": ["核查指南", "核查要点", "检查指南"]},
            {
                "topic": "42061法规",
                "keywords": [
                    "42061",
                    "yy/t 42061",
                    "yy/t42061",
                    "t 42061",
                    "yyt42061",
                ],
            },
        ],
    },
    "iso13485": {
        "title": "13485体考",
        "scopes": ["ISO13485", "MDR"],
        "min_total_questions": 180,
        "topic_groups": [
            {"topic": "ISO13485", "keywords": ["iso13485", "iso 13485", "13485"]},
            {"topic": "MDR", "keywords": ["mdr", "eu 2017/745", "2017/745"]},
        ],
    },
    "mdsap": {
        "title": "MDSAP体考",
        "scopes": ["MDSAP", "5国法规"],
        "min_total_questions": 220,
        "topic_groups": [
            {"topic": "MDSAP", "keywords": ["mdsap", "审计方法", "audit model"]},
            {"topic": "美国法规", "keywords": ["fda", "21 cfr", "usa", "美国"]},
            {"topic": "加拿大法规", "keywords": ["canada", "sor/98-282", "cmdr", "加拿大"]},
            {"topic": "巴西法规", "keywords": ["anvisa", "brazil", "巴西"]},
            {"topic": "日本法规", "keywords": ["pmda", "mhlw", "japan", "日本"]},
            {"topic": "澳大利亚法规", "keywords": ["tga", "australia", "澳大利亚"]},
        ],
    },
}
_EXAM_BATCH_INGEST_SIZE = 50
_EXAM_REQUIREMENT_BASELINE_KEY = "EXAM_BANK_REQUIREMENT_BASELINES_JSON"


def _exam_ingest_target_count() -> int:
    from .app_settings import get_setting

    raw = str(get_setting("EXAM_INGEST_TARGET_COUNT") or "").strip()
    try:
        n = int(raw)
    except Exception:
        n = _EXAM_BATCH_INGEST_SIZE
    return max(1, min(n, 200))


def _parse_csv_four_weights(raw: str, default: tuple[float, float, float, float]) -> list[float]:
    s = (raw or "").replace("，", ",").strip()
    if not s:
        return list(default)
    parts = [p.strip() for p in s.split(",") if p.strip() != ""]
    if len(parts) != 4:
        return list(default)
    out: list[float] = []
    for p in parts:
        try:
            out.append(float(p))
        except Exception:
            return list(default)
    if any(x < 0 for x in out):
        return list(default)
    tot = sum(out)
    if tot <= 0:
        return list(default)
    return [x / tot for x in out]


def _exam_ingest_knowledge_weights() -> list[float]:
    from .app_settings import get_setting

    raw = str(get_setting("EXAM_INGEST_KNOWLEDGE_WEIGHTS") or "").strip()
    return _parse_csv_four_weights(raw, (0.3, 0.3, 0.2, 0.2))


def _exam_ingest_question_type_weights() -> list[float]:
    from .app_settings import get_setting

    raw = str(get_setting("EXAM_INGEST_QUESTION_TYPE_WEIGHTS") or "").strip()
    return _parse_csv_four_weights(raw, (0.3, 0.1, 0.1, 0.5))


def _exam_ingest_max_similar_frac() -> float:
    from .app_settings import get_setting

    raw = str(get_setting("EXAM_INGEST_MAX_SIMILAR_FRAC") or "").strip()
    try:
        x = float(raw)
    except Exception:
        x = 0.1
    return max(0.0, min(x, 0.5))


def _exam_track_policy_version_key(track: str) -> str:
    return f"EXAM_TRACK_REG_POLICY_VERSION_{_normalize_exam_track(track).upper()}"


def _get_exam_track_policy_version(track: str) -> str:
    """人工配置法规版本（兜底/强制覆盖）。"""
    row = AppConfig.query.filter_by(config_key=_exam_track_policy_version_key(track)).first()
    return str(row.config_value if row and row.config_value is not None else "").strip()


def _set_exam_track_policy_version(track: str, version_value: str) -> None:
    key = _exam_track_policy_version_key(track)
    row = AppConfig.query.filter_by(config_key=key).first()
    v = str(version_value or "").strip()
    if row:
        row.config_value = v
    else:
        db.session.add(AppConfig(config_key=key, config_value=v))


def _extract_policy_version_tokens_from_text(text: str) -> list[str]:
    s = str(text or "")
    if not s:
        return []
    pats = [
        r"ISO\s*13485[:：]?\s*\d{4}",
        r"(?:EU\s*)?20\d{2}\s*/\s*\d{3,4}",
        r"MDSAP(?:\s*Audit\s*Model)?\s*(?:v|V)?\d+(?:\.\d+)?",
        r"21\s*CFR(?:\s*Part)?\s*\d+",
        r"SOR\s*/\s*\d{2,4}\s*-\s*\d+",
        r"YY/T\s*\d{3,6}(?:\.\d+)?(?:-\d{4})?",
        r"GB\s*/\s*T\s*\d{3,6}(?:\.\d+)?(?:-\d{4})?",
        r"(?:GMP|核查指南|MDR|FDA|ANVISA|PMDA|TGA)[^,\n;，；]{0,18}(?:20\d{2}|v\d+(?:\.\d+)?)",
    ]
    out: list[str] = []
    for p in pats:
        try:
            for m in re.findall(p, s, flags=re.IGNORECASE):
                t = str(m).strip()
                if t and t not in out:
                    out.append(t)
        except Exception:
            continue
    return out


def _get_exam_track_policy_version_auto(track_raw: Any) -> dict[str, Any]:
    """自动识别法规版本，返回版本+置信度+证据。"""
    track = _normalize_exam_track(track_raw)
    empty = {"version": "", "confidence": "none", "evidence": None}
    track_keys_map: dict[str, list[str]] = {
        "cn": ["gmp", "核查", "核查指南", "检查指南", "生产质量管理规范"],
        "iso13485": ["iso13485", "iso 13485", "13485", "mdr", "2017/745", "eu 2017"],
        "mdsap": ["mdsap", "fda", "21 cfr", "anvisa", "pmda", "tga", "sor/", "canada"],
    }
    keys = track_keys_map.get(track, [])
    rows = UploadRecord.query.order_by(UploadRecord.updated_at.desc()).limit(3000).all()
    for r in rows:
        txt = " ".join(
            [
                str(getattr(r, "task_type", None) or ""),
                str(getattr(r, "file_name", None) or ""),
                str(getattr(r, "project_name", None) or ""),
                str(getattr(r, "notes", None) or ""),
                str(getattr(r, "project_notes", None) or ""),
            ]
        )
        txt_l = txt.lower()
        if keys and not any(k in txt_l for k in keys):
            continue
        ev_base = {
            "upload_id": str(getattr(r, "id", "") or ""),
            "file_name": str(getattr(r, "file_name", "") or ""),
            "updated_at": getattr(r, "updated_at", None).isoformat() if getattr(r, "updated_at", None) else None,
        }
        fv = str(getattr(r, "file_version", None) or "").strip()
        rv = str(getattr(r, "registration_version", None) or "").strip()
        if fv:
            return {
                "version": fv,
                "confidence": "high",
                "evidence": {**ev_base, "matched_by": "file_version"},
            }
        if rv:
            return {
                "version": rv,
                "confidence": "high",
                "evidence": {**ev_base, "matched_by": "registration_version"},
            }
        toks = _extract_policy_version_tokens_from_text(txt)
        if toks:
            return {
                "version": toks[0],
                "confidence": "medium",
                "evidence": {**ev_base, "matched_by": "text_regex", "matched_token": toks[0]},
            }
    return empty


def _resolve_exam_track_policy_version(track_raw: Any) -> dict[str, Any]:
    """返回生效版本、来源、置信度与证据。"""
    track = _normalize_exam_track(track_raw)
    auto_info = _get_exam_track_policy_version_auto(track)
    auto_v = str(auto_info.get("version") or "").strip() if isinstance(auto_info, dict) else ""
    fallback_v = _get_exam_track_policy_version(track)
    if auto_v:
        return {
            "effective_policy_version": auto_v,
            "policy_version_source": "auto",
            "auto_policy_version": auto_v,
            "fallback_policy_version": fallback_v,
            "policy_version_confidence": str(auto_info.get("confidence") or "medium"),
            "policy_version_evidence": auto_info.get("evidence"),
        }
    if fallback_v:
        return {
            "effective_policy_version": fallback_v,
            "policy_version_source": "fallback",
            "auto_policy_version": "",
            "fallback_policy_version": fallback_v,
            "policy_version_confidence": "fallback",
            "policy_version_evidence": None,
        }
    return {
        "effective_policy_version": "",
        "policy_version_source": "none",
        "auto_policy_version": "",
        "fallback_policy_version": "",
        "policy_version_confidence": "none",
        "policy_version_evidence": None,
    }


def _normalize_exam_track(v: Any) -> str:
    s = str(v or "").strip().lower()
    if s in _EXAM_TRACK_REQUIREMENT_PROFILES:
        return s
    if s in {"iso", "13485", "iso_13485"}:
        return "iso13485"
    if s in {"cn", "china"}:
        return "cn"
    if s in {"mdsap", "m_dsap"}:
        return "mdsap"
    return "cn"


def _app_config_get_json(config_key: str, fallback: Any) -> Any:
    row = AppConfig.query.filter_by(config_key=config_key).first()
    raw = str(row.config_value).strip() if row and row.config_value is not None else ""
    if not raw:
        return fallback
    try:
        return json.loads(raw)
    except Exception:
        return fallback


def _app_config_set_json(config_key: str, value: Any) -> None:
    row = AppConfig.query.filter_by(config_key=config_key).first()
    payload = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if row:
        row.config_value = payload
    else:
        db.session.add(AppConfig(config_key=config_key, config_value=payload))


def _extract_total_from_bank_payload(payload: dict[str, Any]) -> Optional[int]:
    if not isinstance(payload, dict):
        return None
    root = _unwrap_quiz_api_success_data(payload)
    if not isinstance(root, dict):
        root = {}
    inner = root.get("data") if isinstance(root.get("data"), dict) else {}
    cands = [
        root.get("total"),
        root.get("count"),
        root.get("matched_total"),
        inner.get("total"),
        inner.get("count"),
        inner.get("matched_total"),
    ]
    for c in cands:
        try:
            if c is None or str(c).strip() == "":
                continue
            n = int(c)
            if n >= 0:
                return n
        except Exception:
            continue
    items = _extract_bank_items_from_list_root(root)
    if isinstance(items, list):
        return len(items)
    return None


def _query_bank_total_for_track(track: str, keyword: str | None = None) -> int:
    q: dict[str, Any] = {
        "limit": "1",
        "offset": "0",
        "exam_track": track,
        "examTrack": track,
        "track": track,
    }
    if keyword:
        q["q"] = keyword
    st, pl, _ = _quiz_try_paths(
        ["quiz/bank/questions", "quiz/questions", "quiz/bank/question-list"],
        method="GET",
        query=q,
    )
    if not (200 <= int(st) < 300 and isinstance(pl, dict)):
        return 0
    n = _extract_total_from_bank_payload(pl)
    return max(0, int(n or 0))


def _collect_exam_knowledge_markers(track: str) -> dict[str, Optional[datetime]]:
    track = _normalize_exam_track(track)
    track_keys_map: dict[str, list[str]] = {
        "cn": ["gmp", "核查", "核查指南", "检查指南", "生产质量管理规范"],
        "iso13485": ["iso13485", "13485", "mdr", "2017/745", "欧盟"],
        "mdsap": ["mdsap", "fda", "anvisa", "pmda", "tga", "canada", "5国", "五国"],
    }
    track_keys = track_keys_map.get(track, [])
    regs_global = ["法规", "regulation", "guideline", "指南", "标准", "standard"]
    proc_keys = ["程序", "procedure", "sop", "作业指导", "管理规程"]
    case_keys = ["案例", "case", "项目案例", "project case"]
    marker: dict[str, Optional[datetime]] = {
        "regulations_at": None,
        "procedures_at": None,
        "project_case_at": None,
    }
    rows = UploadRecord.query.order_by(UploadRecord.updated_at.desc()).limit(3000).all()
    for r in rows:
        upd = getattr(r, "updated_at", None)
        if not upd:
            continue
        txt = " ".join(
            [
                str(getattr(r, "task_type", None) or ""),
                str(getattr(r, "file_name", None) or ""),
                str(getattr(r, "project_name", None) or ""),
                str(getattr(r, "notes", None) or ""),
                str(getattr(r, "project_notes", None) or ""),
            ]
        ).lower()
        if any(k in txt for k in (track_keys + regs_global)):
            if marker["regulations_at"] is None or upd > marker["regulations_at"]:
                marker["regulations_at"] = upd
        if any(k in txt for k in proc_keys):
            if marker["procedures_at"] is None or upd > marker["procedures_at"]:
                marker["procedures_at"] = upd
        if any(k in txt for k in case_keys):
            if marker["project_case_at"] is None or upd > marker["project_case_at"]:
                marker["project_case_at"] = upd
    return marker


def _build_exam_bank_requirement_status(track_raw: Any) -> dict[str, Any]:
    track = _normalize_exam_track(track_raw)
    prof = _EXAM_TRACK_REQUIREMENT_PROFILES.get(track) or _EXAM_TRACK_REQUIREMENT_PROFILES["cn"]
    total = _query_bank_total_for_track(track, keyword=None)
    topic_checks: list[dict[str, Any]] = []
    missing_topics: list[str] = []
    for g in prof.get("topic_groups") or []:
        kws = [str(k).strip() for k in (g.get("keywords") or []) if str(k).strip()]
        # 每个主题只取第 1 关键词做轻量检索，避免单次检查请求过多。
        k0 = kws[0] if kws else ""
        hit = _query_bank_total_for_track(track, keyword=k0) if k0 else 0
        ok = hit > 0
        topic_checks.append({"topic": g.get("topic") or k0, "keyword": k0, "hit_count": hit, "is_met": ok})
        if not ok:
            missing_topics.append(str(g.get("topic") or k0))

    min_total = int(prof.get("min_total_questions") or 0)
    missing_cnt = max(0, min_total - total)
    baseline_all = _app_config_get_json(_EXAM_REQUIREMENT_BASELINE_KEY, {})
    base = baseline_all.get(track) if isinstance(baseline_all, dict) else None
    if not isinstance(base, dict):
        base = {}

    marker_now = _collect_exam_knowledge_markers(track)
    pv_info = _resolve_exam_track_policy_version(track)
    current_policy_ver = str(pv_info.get("effective_policy_version") or "").strip()
    policy_source = str(pv_info.get("policy_version_source") or "none")
    auto_policy_ver = str(pv_info.get("auto_policy_version") or "").strip()
    fallback_policy_ver = str(pv_info.get("fallback_policy_version") or "").strip()
    policy_conf = str(pv_info.get("policy_version_confidence") or "none")
    policy_evidence = pv_info.get("policy_version_evidence")
    base_policy_ver = str(base.get("policy_version") or "").strip()

    def _parse_iso(v: Any) -> Optional[datetime]:
        s = str(v or "").strip()
        if not s:
            return None
        try:
            return datetime.fromisoformat(s.replace("Z", "+00:00")).replace(tzinfo=None)
        except Exception:
            return None

    changed_flags: dict[str, bool] = {}
    for k in ("regulations_at", "procedures_at", "project_case_at"):
        cur = marker_now.get(k)
        old = _parse_iso(base.get(k))
        changed_flags[k] = bool(cur and (old is None or cur > old))
    changed_flags["policy_version"] = bool(current_policy_ver and current_policy_ver != base_policy_ver)

    reasons: list[str] = []
    if missing_cnt > 0:
        reasons.append(f"当前题量 {total}，低于达标题量 {min_total}。")
    if missing_topics:
        reasons.append("主题覆盖不足：" + "、".join(missing_topics) + "。")
    if not base:
        reasons.append("尚未建立“达标基线”，请补全后点击“设为当前达标基线”。")
    if changed_flags.get("policy_version"):
        reasons.append("法规版本标识已变化，需补充新发布/升级要求对应题目。")
    if changed_flags.get("regulations_at"):
        reasons.append("检测到法规资料有新增/修改，需补充题库。")
    if changed_flags.get("procedures_at"):
        reasons.append("检测到程序文件有新增/修改，建议补全题库。")
    if changed_flags.get("project_case_at"):
        reasons.append("检测到项目案例知识有新增/修改，建议补全题库。")
    is_satisfied = len(reasons) == 0
    return {
        "track": track,
        "track_label": prof.get("title") or track,
        "scopes": prof.get("scopes") or [],
        "required_min_total": min_total,
        "bank_total": total,
        "missing_total_count": missing_cnt,
        "topic_checks": topic_checks,
        "is_satisfied": is_satisfied,
        "reasons": reasons,
        "next_batch_target_count": _EXAM_BATCH_INGEST_SIZE,
        "can_continue_ingest": True,
        "knowledge_markers": {
            "current": {k: (v.isoformat() if v else None) for k, v in marker_now.items()},
            "baseline": {
                "regulations_at": base.get("regulations_at"),
                "procedures_at": base.get("procedures_at"),
                "project_case_at": base.get("project_case_at"),
                "policy_version": base.get("policy_version"),
                "checked_at": base.get("checked_at"),
            },
            "changed_flags": changed_flags,
            "current_policy_version": current_policy_ver,
            "policy_version_source": policy_source,
            "policy_version_confidence": policy_conf,
            "policy_version_evidence": policy_evidence,
            "auto_policy_version": auto_policy_ver,
            "fallback_policy_version": fallback_policy_ver,
        },
        "recommendations": {
            "teacher_generate_count": 20 if track == "cn" else (25 if track == "iso13485" else 30),
            "teacher_generate_difficulty": "easy" if track == "cn" else ("medium" if track == "iso13485" else "hard"),
            "student_practice_count": 20 if track == "cn" else (25 if track == "iso13485" else 30),
            "student_practice_difficulty": "easy" if track == "cn" else ("medium" if track == "iso13485" else "hard"),
            "ingest_batch_size": _EXAM_BATCH_INGEST_SIZE,
        },
    }


def _normalize_answers_list_shape(out: dict[str, Any]) -> dict[str, Any]:
    """answers 统一为 list[{question_id, answer, user_answer}]；上游 aicheckword落库字段为 user_answer。"""
    answers = out.get("answers")
    if isinstance(answers, dict):
        out["answers"] = [
            {"question_id": str(k), "answer": v, "user_answer": v}
            for k, v in answers.items()
            if str(k).strip() != ""
        ]
    elif isinstance(answers, list):
        normed: list[dict[str, Any]] = []
        for it in answers:
            if not isinstance(it, dict):
                continue
            row = dict(it)
            ua = row.get("user_answer")
            if ua is None:
                ua = row.get("answer")
            if ua is None:
                ua = row.get("value") or row.get("selected") or row.get("response")
            if ua is not None:
                row["user_answer"] = ua
            normed.append(row)
        out["answers"] = normed
    return out


def _coerce_attempt_id_int_str(v: Any) -> str:
    """将任意 attempt 标识归一成可被 int 解析的字符串。"""
    raw = str(v or "").strip()
    if not raw:
        return ""
    m = re.search(r"\d+", raw)
    if m:
        return m.group(0)
    # 纯字符串/UUID 等无数字时，退化为稳定正整数（避免上游 int_parsing）
    h = hashlib.md5(raw.encode("utf-8")).hexdigest()[:8]
    try:
        n = int(h, 16)
    except Exception:
        n = 1
    if n <= 0:
        n = 1
    return str(n)


def _normalize_practice_submit_upstream_body(body: dict[str, Any]) -> dict[str, Any]:
    """练习提交：会话 id / set_id 常见别名一并带上，避免上游只认其中一种字段名。"""
    raw = body if isinstance(body, dict) else {}
    out = dict(raw)
    # 上游有些版本要求 attemptId 为整数；空串会触发 422，故空值直接移除
    if str(out.get("attempt_id") or "").strip() == "":
        out.pop("attempt_id", None)
    if str(out.get("attemptId") or "").strip() == "":
        out.pop("attemptId", None)
    sid = str(
        out.get("practice_session_id")
        or out.get("session_id")
        or out.get("practiceSessionId")
        or out.get("sessionId")
        or ""
    ).strip()
    if sid:
        out["practice_session_id"] = sid
        out["session_id"] = sid
        out["practiceSessionId"] = sid
        out["sessionId"] = sid
    aid_raw = str(
        out.get("attempt_id")
        or out.get("attemptId")
        or out.get("practice_attempt_id")
        or out.get("practiceAttemptId")
        or out.get("practice_id")
        or out.get("practiceId")
        or out.get("id")
        or sid
        or ""
    ).strip()
    aid = _coerce_attempt_id_int_str(aid_raw)
    if aid:
        out["attempt_id"] = aid
        out["attemptId"] = aid
    # 空串 set_id 会导致上游 int_parsing，先移除再尝试从嵌套对象补齐
    if str(out.get("set_id") or "").strip() == "":
        out.pop("set_id", None)
    if str(out.get("setId") or "").strip() == "":
        out.pop("setId", None)
    set_id = str(out.get("set_id") or out.get("setId") or _extract_set_id_from_any(out) or "").strip()
    if set_id:
        out["set_id"] = set_id
        out["setId"] = set_id
    return _normalize_answers_list_shape(out)


def _normalize_exam_submit_upstream_body(body: dict[str, Any], attempt_id: str) -> dict[str, Any]:
    raw = body if isinstance(body, dict) else {}
    out = dict(raw)
    aid = _coerce_attempt_id_int_str(attempt_id)
    if aid:
        out["attempt_id"] = aid
        out["attemptId"] = aid
    return _normalize_answers_list_shape(out)


def _extract_bank_items_from_list_root(root: Any) -> list[dict[str, Any]]:
    """从题库列表上游 JSON 根对象中取出题目 dict 列表。"""
    if not isinstance(root, dict):
        return []
    inner = root.get("data") if isinstance(root.get("data"), dict) else None
    cands: list[Any] = [
        root.get("items"),
        root.get("questions"),
        inner.get("items") if inner else None,
        inner.get("questions") if inner else None,
    ]
    for c in cands:
        if isinstance(c, list) and c:
            return [x for x in c if isinstance(x, dict)]
    for c in cands:
        if isinstance(c, list):
            return [x for x in c if isinstance(x, dict)]
    return []


def _question_id_from_bank_item(it: dict[str, Any]) -> str:
    return str(it.get("question_id") or it.get("questionId") or it.get("id") or "").strip()


def _find_set_item_dicts(root: Any, depth: int = 0) -> list[dict[str, Any]]:
    """从 GET quiz/sets/{id} 上游根对象中取出套内题目/条目 dict 列表。"""
    if depth > 10:
        return []
    cand_keys = ("items", "questions", "question_items", "questionItems", "entries")
    if isinstance(root, list):
        return [x for x in root if isinstance(x, dict)]
    if not isinstance(root, dict):
        return []
    for hop in ("load_set", "set", "quiz_set", "quizSet", "payload", "result"):
        inner = root.get(hop)
        if isinstance(inner, dict):
            got = _find_set_item_dicts(inner, depth + 1)
            if got:
                return got
    for k in cand_keys:
        v = root.get(k)
        if isinstance(v, list) and v and isinstance(v[0], dict):
            return [x for x in v if isinstance(x, dict)]
    for k in cand_keys:
        v = root.get(k)
        if isinstance(v, list):
            out = [x for x in v if isinstance(x, dict)]
            if out:
                return out
    inner = root.get("data")
    if isinstance(inner, dict):
        return _find_set_item_dicts(inner, depth + 1)
    return []


def _question_id_from_set_item(it: dict[str, Any]) -> str:
    for key in (
        "question_id",
        "questionId",
        "id",
        "bank_question_id",
        "bankQuestionId",
        "question_bank_id",
        "questionBankId",
        "ref_question_id",
        "refQuestionId",
    ):
        v = it.get(key)
        if v is not None and str(v).strip():
            return str(v).strip()
    q = it.get("question")
    if isinstance(q, dict):
        for key in ("id", "question_id", "questionId", "bank_question_id", "bankQuestionId"):
            v = q.get(key)
            if v is not None and str(v).strip():
                return str(v).strip()
    return ""


def _ordered_question_ids_from_set_upstream_root(root: Any) -> list[str]:
    out: list[str] = []
    for it in _find_set_item_dicts(root):
        qid = _question_id_from_set_item(it)
        if qid:
            out.append(qid)
    return out


def _bank_row_matches_set_id(it: dict[str, Any], set_sid: str) -> bool:
    want = str(set_sid).strip()
    if not want:
        return False
    blobs: list[dict[str, Any]] = [it]
    q = it.get("question")
    if isinstance(q, dict):
        blobs.append(q)
    for mk in ("meta", "metadata"):
        m = it.get(mk)
        if isinstance(m, dict):
            blobs.append(m)
    id_keys = (
        "set_id",
        "setId",
        "quiz_set_id",
        "quizSetId",
        "parent_set_id",
        "parentSetId",
        "exam_set_id",
        "examSetId",
    )
    list_keys = ("set_ids", "setIds", "quiz_set_ids", "quizSetIds")
    for cur in blobs:
        if not isinstance(cur, dict):
            continue
        for k in id_keys:
            v = cur.get(k)
            if v is None or v == "":
                continue
            if str(v).strip() == want:
                return True
        for lk in list_keys:
            v = cur.get(lk)
            if isinstance(v, list) and want in {str(x).strip() for x in v}:
                return True
    return False


# ---------- 辅助函数 ----------

def _save_file(file_storage, target_dir: Path) -> tuple[str, str]:
    from .upload_filename import preserved_secure_filename

    filename = preserved_secure_filename(file_storage.filename or "")
    generated_name = f"{now_local().strftime('%Y%m%d%H%M%S%f')}_{filename}"
    file_path = target_dir / generated_name
    file_storage.save(file_path)
    return generated_name, str(file_path)


def _project_priority_label(priority: int | None) -> str:
    p = int(priority or Project.PRIORITY_MEDIUM)
    if p >= Project.PRIORITY_HIGH:
        return "高"
    if p <= Project.PRIORITY_LOW:
        return "低"
    return "中"


def _project_status_label(status: str | None) -> str:
    s = (status or "").strip().lower()
    return "已结束" if s == Project.STATUS_ENDED else "进行中"


def _sort_upload_records_by_project_priority(
    records: list[UploadRecord], proj_meta: dict[str, dict[str, Any]]
) -> list[UploadRecord]:
    """按项目优先级降序（数值越大越靠前），同优先级按 sort_order 升序、创建时间升序。"""

    def _created_ts(r: UploadRecord) -> float:
        ct = getattr(r, "created_at", None)
        if not ct:
            return 0.0
        try:
            return float(ct.timestamp())
        except (OSError, OverflowError, ValueError):
            return 0.0

    def _key(r: UploadRecord) -> tuple:
        pr = int((proj_meta.get(r.project_name) or {}).get("priority") or Project.PRIORITY_MEDIUM)
        so = int(r.sort_order or 0)
        return (-pr, so, _created_ts(r))

    return sorted(records, key=_key)


def _project_display_label_from_fields(
    name: str | None,
    registered_country: str | None,
    registered_category: str | None,
) -> str:
    n = (name or "").strip()
    c = (registered_country or "").strip()
    cat = (registered_category or "").strip()
    if not c and not cat:
        return n
    return f"{n}（{c or '—'} / {cat or '—'}）"


def _project_display_label(p: Project) -> str:
    return _project_display_label_from_fields(
        p.name,
        getattr(p, "registered_country", None),
        getattr(p, "registered_category", None),
    )


def _project_registered_product_name_hints(project_ids: list[str]) -> dict[str, str]:
    """从上传记录中取各 project_id 下一条非空「注册产品名称」，仅用于下拉展示（不参与 projectKey 匹配）。"""
    ids = [str(x).strip() for x in (project_ids or []) if str(x).strip()]
    if not ids:
        return {}
    out: dict[str, str] = {}
    rows = (
        UploadRecord.query.filter(
            UploadRecord.project_id.in_(ids),
            UploadRecord.registered_product_name.isnot(None),
        )
        .order_by(UploadRecord.updated_at.desc())
        .all()
    )
    for u in rows:
        pid = (getattr(u, "project_id", None) or "").strip()
        if not pid or pid in out:
            continue
        pn = (getattr(u, "registered_product_name", None) or "").strip()
        if pn:
            out[pid] = pn
    return out


def _filter_nullable_eq(q, col, value):
    """value=None 用 IS NULL；否则用 =。避免 MySQL 生成 `IS 'xxx'` 导致 500。"""
    return q.filter(col.is_(None)) if value is None else q.filter(col == value)


def _backfill_project_ids() -> None:
    """将历史 upload_records/module_cascade_reminders/generation_summary 的 project_id 回填（按展示键匹配）。"""
    from .models import UploadRecord, ModuleCascadeReminder, GenerationSummary
    rows = Project.query.order_by(Project.updated_at.asc(), Project.id.asc()).all()
    # 同一展示键可能有重复项目：选择最早的作为“主项目ID”
    key_to_pid: dict[str, str] = {}
    for p in rows:
        k = _project_display_label(p)
        if k and k not in key_to_pid:
            key_to_pid[k] = p.id

    for k, pid in key_to_pid.items():
        UploadRecord.query.filter(UploadRecord.project_id.is_(None), UploadRecord.project_name == k).update(
            {"project_id": pid}
        )
        ModuleCascadeReminder.query.filter(ModuleCascadeReminder.project_id.is_(None), ModuleCascadeReminder.project_name == k).update(
            {"project_id": pid}
        )
        try:
            GenerationSummary.query.filter(GenerationSummary.project_id.is_(None), GenerationSummary.project_name == k).update(
                {"project_id": pid}
            )
        except Exception:
            pass
    db.session.commit()


def _resolve_assigned_team_id_for_project_autocreate() -> str | None:
    """任务录入自动补项目行时：仅使用当前账号已绑定的项目组，不回退默认组。"""
    from .authz import rbac_enforced, user_team_ids

    if not rbac_enforced():
        return None
    tids = [str(x).strip() for x in user_team_ids() if str(x).strip()]
    return tids[0] if tids else None


def _realign_project_team_for_creator(project: Project | None) -> None:
    """创建人再次录入时，仅当项目尚未绑定项目组时补写其所属组（不覆盖已有归属）。"""
    from .authz import rbac_enforced, user_team_ids

    if project is None or not rbac_enforced():
        return
    tid = str(getattr(project, "assigned_team_id", "") or "").strip()
    if tid:
        return
    uid = str(session.get("user_id") or "").strip()
    if not uid:
        return
    creator = str(getattr(project, "created_by_user_id", "") or "").strip()
    if creator and creator != uid:
        return
    tids = [str(x).strip() for x in user_team_ids() if str(x).strip()]
    if not tids:
        return
    project.assigned_team_id = tids[0]
    if not creator:
        project.created_by_user_id = uid
    db.session.add(project)
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
    from .authz import _invalidate_project_lookup_maps

    _invalidate_project_lookup_maps()


def _ensure_project_row(project_name: str) -> Project | None:
    # 兼容旧数据：project_name 可能是 base name，也可能是展示键(label)
    from .authz import _invalidate_project_lookup_maps, _project_lookup_maps

    label = (project_name or "").strip()
    if not label:
        return None

    _by_id, by_label, by_name = _project_lookup_maps()
    existing = by_label.get(label) or by_name.get(label)
    if existing is not None:
        _realign_project_team_for_creator(existing)
        return existing

    team_id = _resolve_assigned_team_id_for_project_autocreate()
    row = Project(
        name=label,
        priority=Project.PRIORITY_MEDIUM,
        status=Project.STATUS_ACTIVE,
        registered_country=None,
        registered_category=None,
        assigned_team_id=team_id,
        created_by_user_id=str(session.get("user_id") or "").strip() or None,
    )
    db.session.add(row)
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        _invalidate_project_lookup_maps()
        _by_id, by_label, by_name = _project_lookup_maps()
        return by_label.get(label) or by_name.get(label)
    _invalidate_project_lookup_maps()
    return row


def _project_meta_map(auto_create_from_uploads: bool = False) -> dict[str, dict[str, Any]]:
    """
    返回项目元信息映射：name -> {priority, status}。
    若 auto_create_from_uploads=True，会把 upload_records 中出现但 projects 表不存在的项目补齐（默认中/进行中）。
    """
    if auto_create_from_uploads:
        labels = (
            db.session.query(UploadRecord.project_name)
            .filter(UploadRecord.project_name.isnot(None), UploadRecord.project_name != "")
            .distinct()
            .all()
        )
        for (lab,) in labels:
            _ensure_project_row(lab)

    rows = Project.query.order_by(Project.priority.desc(), Project.name.asc()).all()
    out: dict[str, dict[str, Any]] = {}
    for r in rows:
        if not (r.name or "").strip():
            continue
        key = _project_display_label(r)
        out[key] = {
            "priority": int(r.priority or Project.PRIORITY_MEDIUM),
            "status": r.status or Project.STATUS_ACTIVE,
            "baseName": r.name,
            "registeredCountry": getattr(r, "registered_country", None),
            "registeredCategory": getattr(r, "registered_category", None),
        }
    return out


def _ended_project_names() -> set[str]:
    rows = Project.query.filter(Project.status == Project.STATUS_ENDED).all()
    return {_project_display_label(r) for r in rows if (r.name or "").strip()}


def _normalize_doc_link(line: str) -> str:
    """截断 https 之前的内容，返回从 http(s) 开始的有效可打开地址。"""
    line = (line or "").strip()
    if not line:
        return line
    lower = line.lower()
    for prefix in ("https://", "http://"):
        idx = lower.find(prefix)
        if idx != -1:
            return line[idx:]
    return line


def _normalize_template_links(value: str) -> str:
    """多行文档地址，每行截断 https 之前的信息后合并。"""
    if not value or not value.strip():
        return value.strip() or ""
    lines = [_normalize_doc_link(ln) for ln in value.split("\n") if ln.strip()]
    return "\n".join(ln for ln in lines if ln)


def _is_valid_doc_link(value: str) -> bool:
    """校验文档链接是否合理：每行经归一化后需以 http:// 或 https:// 开头。"""
    if not value or not value.strip():
        return True
    for line in value.strip().split("\n"):
        ln = (line or "").strip()
        if not ln:
            continue
        normalized = _normalize_doc_link(ln)
        lower = normalized.lower()
        if not (lower.startswith("http://") or lower.startswith("https://")):
            return False
    return True


AUDIT_REJECT_PENDING_STATUS = "审核不通过待修改"


def _maybe_bump_audit_reject_count(
    upload: UploadRecord,
    *,
    previous_completion_status: Optional[str],
    previous_audit_status: Optional[str],
    target_audit_status: Optional[str] = None,
    target_completion_status: Optional[str] = None,
) -> None:
    """仅在「原已完成」且新设为审核不通过待修改时 +1；重复保存或其它编辑不计数。"""
    prev_completed = (previous_completion_status or "").strip()
    if not prev_completed:
        return
    if target_audit_status is not None:
        new_audit = (target_audit_status or "").strip()
        old_audit = (previous_audit_status or "").strip()
        if new_audit == AUDIT_REJECT_PENDING_STATUS and old_audit != AUDIT_REJECT_PENDING_STATUS:
            upload.audit_reject_count = (getattr(upload, "audit_reject_count", 0) or 0) + 1
        return
    if target_completion_status is not None:
        new_cs = (target_completion_status or "").strip()
        if new_cs == AUDIT_REJECT_PENDING_STATUS and prev_completed != AUDIT_REJECT_PENDING_STATUS:
            upload.audit_reject_count = (getattr(upload, "audit_reject_count", 0) or 0) + 1


def _prepare_summary(upload: UploadRecord) -> GenerationSummary:
    if upload.summary:
        return upload.summary
    summary = GenerationSummary(
        upload=upload,
        project_name=upload.project_name,
        file_name=upload.file_name,
        author=upload.author,
    )
    db.session.add(summary)
    return summary


def _safe_ftp_task_filename(name: str) -> str:
    """FTP 远端文件名：保留中文等 Unicode，仅去除路径分隔符与非法文件名字符。"""
    base = os.path.basename(name or "document")
    base = base.strip() or "document"
    base = re.sub(r'[\x00-\x1f\\/:*?"<>|]+', "_", base)
    base = re.sub(r"\s+", " ", base).strip()
    return (base or "document")[:200]


def _normalize_handoff_display_filename(name: str) -> str:
    """修复 UTF-8 文件名在部分链路中被误读为 Latin-1 的乱码。"""
    s = (name or "").strip().replace("\\", "/")
    s = os.path.basename(s) or s
    if not s:
        return "document.docx"
    if re.search(r"[\u4e00-\u9fff]", s):
        return s[:512]
    if not any(ord(c) > 127 for c in s):
        return s[:512]
    for enc in ("latin-1", "cp1252"):
        try:
            repaired = s.encode(enc).decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            continue
        if repaired and repaired != s and re.search(r"[\u4e00-\u9fff]", repaired):
            return repaired[:512]
    return s[:512]


def _is_internal_handoff_cache_basename(name: str) -> bool:
    """本机/FTP 缓存临时名，不可作为用户可见文件名。"""
    base = os.path.basename((name or "").replace("\\", "/"))
    low = base.lower()
    if low.startswith(("_ftptpl_", "_dbtpl_", "link_")):
        return True
    if re.match(
        r"^_ftptpl_[0-9a-f\-]{36}(?:_\d{14})?\.(docx|doc|xlsx|xls|pdf)$",
        low,
    ):
        return True
    return False


def _upload_handoff_display_filename(upload: UploadRecord) -> str:
    """交接展示名：优先 original_file_name，跳过内部缓存文件名。"""
    for cand in (
        (getattr(upload, "original_file_name", None) or "").strip(),
        (upload.file_name or "").strip(),
    ):
        if not cand:
            continue
        base = os.path.basename(cand.replace("\\", "/"))
        if _is_internal_handoff_cache_basename(base):
            continue
        return _normalize_handoff_display_filename(base)
    return "document.docx"


def _resolve_handoff_display_name(upload: UploadRecord, fallback: str = "") -> str:
    fb = (fallback or "").strip()
    if fb and not _is_internal_handoff_cache_basename(fb):
        return _normalize_handoff_display_filename(fb)
    return _upload_handoff_display_filename(upload)


def _unlink_task_template_cache_files(upload_id: str) -> None:
    """删除本机为某条上传记录缓存的模板副本。"""
    uploads_dir = Path(current_app.config["UPLOAD_FOLDER"])
    for suf in (".docx", ".doc"):
        (uploads_dir / f"_ftptpl_{upload_id}{suf}").unlink(missing_ok=True)
    (uploads_dir / f"_dbtpl_{upload_id}.docx").unlink(missing_ok=True)


def _upload_record_has_task_file(r: UploadRecord) -> bool:
    """是否有任务模板文件（BLOB、本机路径或 FTP）。"""
    if r.template_file_blob or r.storage_path:
        return True
    return bool((getattr(r, "ftp_path", None) or "").strip())


def _upload_record_visible_to_page2_user(rec: UploadRecord) -> bool:
    """与页面2「我的任务」/初稿带入一致；分级管理员只增权。"""
    from .authz import upload_record_visible_to_user

    return upload_record_visible_to_user(rec)


def _can_access_upload_template(upload: UploadRecord) -> bool:
    """页面1（访问密码超级管理员）或页面2（登录且有权）可下载任务模板。"""
    from .authz import is_page13_super_admin

    if not _page13_password_configured():
        return True
    if is_page13_super_admin():
        return True
    if session.get("user_id"):
        return _upload_record_visible_to_page2_user(upload)
    return False


def _upload_template_download_filename(upload: UploadRecord) -> str:
    name = (upload.original_file_name or upload.file_name or "template").strip() or "template"
    low = name.lower()
    if not (low.endswith(".doc") or low.endswith(".docx")):
        name = f"{name}.docx"
    return name


def _apply_task_template_ftp_after_flush(row: UploadRecord, original_display_name: str | None) -> None:
    """在 row.id 已落库后尝试将 template_file_blob 上传 FTP；成功则清空 BLOB。"""
    blob = row.template_file_blob
    if not blob:
        current_app.logger.debug(
            "FTP 任务模板跳过 upload_id=%s：无 template_file_blob（可能仅用链接或未上传文件）",
            getattr(row, "id", None),
        )
        return
    from .ftp_store import try_upload_bytes

    safe = _safe_ftp_task_filename(original_display_name or row.file_name or "document")
    rel = f"task_assignments/{row.id}/{safe}"
    pth, err = try_upload_bytes(blob, rel)
    if pth:
        current_app.logger.info(
            "任务模板 FTP 已写入 upload_id=%s rel=%s ftp_path=%s",
            row.id,
            rel,
            pth,
        )
        row.ftp_path = pth
        row.ftp_last_error = None
        row.template_file_blob = None
    elif err:
        current_app.logger.warning(
            "任务模板 FTP 上传失败 upload_id=%s rel=%s（已写入 ftp_last_error 供前端展示）: %s",
            row.id,
            rel,
            err,
        )
        row.ftp_last_error = err[:512] if len(err) > 512 else err
    else:
        current_app.logger.info(
            "任务模板未上传 FTP upload_id=%s rel=%s：未配置 FTP_HOST 等；模板仍保存在数据库 BLOB",
            row.id,
            rel,
        )


def _get_template_path_for_upload(upload: UploadRecord, link_index: int = 0) -> str:
    """返回上传记录对应的模板本地路径：库内模板先落盘，否则 FTP 缓存，否则 storage_path，链接则下载到 uploads/。"""
    if upload.template_file_blob:
        uploads_dir = Path(current_app.config["UPLOAD_FOLDER"])
        mat = uploads_dir / f"_dbtpl_{upload.id}.docx"
        if not mat.exists() or mat.stat().st_size != len(upload.template_file_blob):
            mat.write_bytes(upload.template_file_blob)
        return str(mat)
    fp = (getattr(upload, "ftp_path", None) or "").strip()
    if fp:
        from .ftp_store import download_bytes

        uploads_dir = Path(current_app.config["UPLOAD_FOLDER"])
        fn = ((upload.original_file_name or upload.file_name or "document") or "document").lower()
        ext = ".docx"
        if fn.endswith(".doc") and not fn.endswith(".docx"):
            ext = ".doc"
        elif fn.endswith(".docx"):
            ext = ".docx"
        mat = uploads_dir / f"_ftptpl_{upload.id}{ext}"
        data = download_bytes(fp)
        mat.write_bytes(data)
        return str(mat)
    if upload.storage_path and Path(upload.storage_path).exists():
        return upload.storage_path
    links = upload.get_template_links_list()
    if links and link_index < len(links):
        uploads_dir = Path(current_app.config["UPLOAD_FOLDER"])
        save_path = uploads_dir / f"link_{upload.id}_{link_index}.docx"
        if not save_path.exists():
            download_template_from_url(links[link_index], str(save_path))
        return str(save_path)
    raise ValueError("上传记录未关联有效模板（文件或链接）")


def _resolve_handoff_doc_for_print_sign(
    upload_id: str, mode: str = "sign"
) -> tuple[Optional[bytes], str, Optional[str], Optional[str]]:
    """为 aiprintword 交接解析任务文档。
    返回 (bytes, filename, error, reuse_ftp_path)：
    - 签字模式优先复用 upload.ftp_path（.docx/.xlsx），避免在签字端重复上传 FTP；
    - 其余场景优先最新成功生成结果，否则回退到模板。"""
    uid = (upload_id or "").strip()
    reuse_fp: Optional[str] = None
    mode_k = (mode or "sign").strip().lower()
    if mode_k not in ("sign", "print"):
        mode_k = "sign"
    if not uid:
        return None, "", "缺少 upload_id", None
    upload = UploadRecord.query.get(uid)
    if not upload:
        return None, "", "任务不存在", None
    # 去签字/去打印：优先复用 aiword 任务模板 FTP，避免重复上传且保留中文展示名。
    if mode_k in ("sign", "print"):
        fp_tpl = (getattr(upload, "ftp_path", None) or "").strip()
        if fp_tpl:
            display = _upload_handoff_display_filename(upload)
            ext0 = Path(display).suffix.lower()
            if not ext0:
                ext0 = ".docx"
                display = Path(display).stem + ext0
            if ext0 in (".docx", ".xlsx"):
                return None, display, None, fp_tpl
    rec = (
        GenerateRecord.query.filter_by(upload_id=uid, success=True)
        .order_by(GenerateRecord.created_at.desc())
        .first()
    )
    if rec:
        blob: Optional[bytes] = rec.output_file_blob
        if not blob and rec.output_path and Path(rec.output_path).is_file():
            try:
                blob = Path(rec.output_path).read_bytes()
            except OSError:
                blob = None
        if blob:
            name = (rec.output_file_name or "").strip()
            if not name or _is_internal_handoff_cache_basename(name):
                name = _upload_handoff_display_filename(upload)
                if not name.endswith((".docx", ".xlsx")):
                    name = Path(name).stem + ".docx"
            else:
                name = _normalize_handoff_display_filename(name)
            return blob, name, None, None
    try:
        path = _get_template_path_for_upload(upload, 0)
    except Exception as e:
        return None, "", f"无已生成文档且无法加载模板：{e}", None
    p = Path(path)
    if not p.is_file():
        return None, "", "模板文件不可用", None
    ext = p.suffix.lower()
    if ext not in (".doc", ".docx", ".xls", ".xlsx", ".pdf"):
        return None, "", f"当前任务文件类型不支持签字/批量打印（{ext or '无扩展名'}）", None
    display = _resolve_handoff_display_name(upload, (upload.file_name or "").strip() or p.name)
    if not Path(display).suffix:
        display = Path(display).stem + ext
    elif ext in (".docx", ".xlsx") and not display.lower().endswith(ext):
        display = Path(display).stem + ext
    fp_tpl = (getattr(upload, "ftp_path", None) or "").strip()
    if fp_tpl and ext in (".docx", ".xlsx"):
        # 与 aiprintword 同一套 FTP：交接仅登记路径，避免服务端再读本地/再传一遍字节
        return None, display, None, fp_tpl
    try:
        raw = p.read_bytes()
    except OSError as e:
        return None, "", str(e), None
    return raw, display, None, None


def _aiprintword_go_redirect(mode: str):
    """mode: sign | print"""
    from .app_settings import get_setting

    uid = (request.args.get("upload_id") or "").strip()
    raw, fname, err, reuse_ftp = _resolve_handoff_doc_for_print_sign(uid, mode=mode)
    fname = _normalize_handoff_display_filename(fname or "document.docx")
    reuse_ok = (reuse_ftp or "").strip()
    if err or (raw is None and not reuse_ok):
        emsg = html.escape(err or "无法读取文档", quote=True)
        htm = (
            "<!DOCTYPE html><html lang=zh-CN><head><meta charset=utf-8><title>跳转失败</title></head><body>"
            f"<p>{emsg}</p><p><a href=\"javascript:history.back()\">返回</a></p></body></html>"
        )
        return make_response(htm, 400)

    base = (get_setting("AIPRINTWORD_BASE_URL") or "").strip().rstrip("/")
    secret = (get_setting("AIPRINTWORD_HANDOFF_SECRET") or "").strip()
    if not base or not secret:
        page_html = (
            "<!DOCTYPE html><html lang=zh-CN><head><meta charset=utf-8><title>未配置</title></head><body>"
            "<p>请在系统配置中填写 AIPRINTWORD_BASE_URL 与 AIPRINTWORD_HANDOFF_SECRET（并与 aiprintword 的 AIWORD_HANDOFF_SECRET 保持一致）。"
            "</p></body></html>"
        )
        return make_response(page_html, 503)

    import requests as _req

    upload_row = UploadRecord.query.get(uid)
    handoff_ctx = _build_aiprintword_handoff_context(upload_row)

    post_data: dict[str, str] = {"purpose": mode, "filename": fname}
    if handoff_ctx:
        post_data["handoff_context"] = json.dumps(handoff_ctx, ensure_ascii=False)

    url = f"{base}/api/handoff"
    try:
        reuse_fp = (reuse_ftp or "").strip()
        if reuse_fp:
            post_data["reuse_ftp_path"] = reuse_fp
            r = _req.post(
                url,
                headers={"X-Aiword-Handoff-Secret": secret},
                data=post_data,
                timeout=120,
            )
        else:
            r = _req.post(
                url,
                headers={"X-Aiword-Handoff-Secret": secret},
                files={"file": (fname, raw)},
                data=post_data,
                timeout=120,
            )
    except _req.RequestException as e:
        page_html = (
            "<!DOCTYPE html><html lang=zh-CN><head><meta charset=utf-8><title>连接失败</title></head><body>"
            f"<p>无法连接 aiprintword（{e}）。请检查 AIPRINTWORD_BASE_URL 与网络。</p></body></html>"
        )
        return make_response(page_html, 502)
    if r.status_code != 200:
        snippet = (r.text or "")[:500].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        page_html = (
            "<!DOCTYPE html><html lang=zh-CN><head><meta charset=utf-8><title>交接失败</title></head><body>"
            f"<p>aiprintword 返回 HTTP {r.status_code}：{snippet}</p></body></html>"
        )
        return make_response(page_html, 502)
    try:
        payload = r.json()
    except Exception:
        page_html = "<!DOCTYPE html><html><body><p>aiprintword 返回非 JSON</p></body></html>"
        return make_response(page_html, 502)
    if not payload.get("ok"):
        msg = html.escape(str(payload.get("error") or "unknown"), quote=True)
        page_html = f"<!DOCTYPE html><html><body><p>交接失败：{msg}</p></body></html>"
        return make_response(page_html, 400)
    token = (payload.get("token") or "").strip()
    if not token:
        return make_response("<html><body><p>未返回 token</p></body></html>", 502)
    if mode == "sign":
        target = f"{base}/sign?from=aiword&handoff_token={_urlquote(token)}"
    else:
        target = f"{base}/?from=aiword&handoff_token={_urlquote(token)}"
    return redirect(target, code=302)


def _build_aiprintword_handoff_context(upload_row: UploadRecord | None) -> dict[str, str]:
    handoff_ctx: dict[str, str] = {}
    if not upload_row:
        return handoff_ctx
    ed = (getattr(upload_row, "displayed_author", None) or "").strip()
    auth = (upload_row.author or "").strip()
    if ed:
        handoff_ctx["editor"] = ed
    elif auth:
        handoff_ctx["editor"] = auth
    if auth:
        handoff_ctx["writer"] = auth
    rv = (getattr(upload_row, "reviewer", None) or "").strip()
    if rv:
        handoff_ctx["reviewer"] = rv
    ap = (getattr(upload_row, "approver", None) or "").strip()
    if ap:
        handoff_ctx["approver"] = ap
    dd = getattr(upload_row, "document_display_date", None)
    if dd is not None:
        try:
            handoff_ctx["doc_date"] = dd.isoformat()
        except Exception:
            pass
    # 阶段：优先任务类型，其次所属模块（供签字页按阶段分组）
    phase = (getattr(upload_row, "task_type", None) or "").strip()
    if not phase:
        phase = (getattr(upload_row, "belonging_module", None) or "").strip()
    if phase:
        handoff_ctx["phase"] = phase
    reg_country = (getattr(upload_row, "country", None) or "").strip()
    if not reg_country:
        try:
            pid_for_country = (getattr(upload_row, "project_id", None) or "").strip()
            proj_row = None
            if pid_for_country:
                proj_row = Project.query.get(pid_for_country)
            if proj_row is None:
                pname_for_country = (getattr(upload_row, "project_name", None) or "").strip()
                if pname_for_country:
                    proj_row = Project.query.filter_by(name=pname_for_country).first()
            if proj_row is not None:
                reg_country = (getattr(proj_row, "registered_country", None) or "").strip()
        except Exception:
            pass
    if reg_country:
        handoff_ctx["country"] = reg_country
    pid = (getattr(upload_row, "project_id", None) or "").strip()
    if pid:
        handoff_ctx["project_id"] = pid
    pname = (getattr(upload_row, "project_name", None) or "").strip()
    if pname:
        handoff_ctx["project_name"] = pname
    pcode = (getattr(upload_row, "project_code", None) or "").strip()
    if pcode:
        handoff_ctx["project_code"] = pcode
    return handoff_ctx


def _aiprintword_batch_handoff_redirect(mode: str, upload_ids: list[str]) -> tuple[dict[str, Any], int]:
    from .app_settings import get_setting
    import requests as _req

    mode_k = (mode or "").strip().lower()
    if mode_k not in ("sign", "print"):
        mode_k = "sign"
    ids = []
    seen: set[str] = set()
    for x in upload_ids or []:
        s = (x or "").strip()
        if not s or s in seen:
            continue
        seen.add(s)
        ids.append(s)
    if len(ids) < 2:
        return {"ok": False, "error": "请至少选择 2 条任务"}, 400

    rows: list[UploadRecord] = []
    projects: set[str] = set()
    for uid in ids:
        row = UploadRecord.query.get(uid)
        if not row:
            return {"ok": False, "error": f"任务不存在：{uid}"}, 404
        rows.append(row)
        projects.add((row.project_name or "").strip())
    if len(projects) > 1:
        return {"ok": False, "error": "请只勾选同一项目的任务"}, 400

    base = (get_setting("AIPRINTWORD_BASE_URL") or "").strip().rstrip("/")
    secret = (get_setting("AIPRINTWORD_HANDOFF_SECRET") or "").strip()
    if not base or not secret:
        return {
            "ok": False,
            "error": "请在系统配置中填写 AIPRINTWORD_BASE_URL 与 AIPRINTWORD_HANDOFF_SECRET",
        }, 503

    url = f"{base}/api/handoff/batch"
    failures: list[dict[str, str]] = []
    manifest: list[dict[str, Any]] = []
    files_payload: dict[str, tuple[str, bytes]] = {}

    for idx, row in enumerate(rows):
        raw, fname, err, reuse_ftp = _resolve_handoff_doc_for_print_sign(row.id, mode=mode_k)
        reuse_ok = (reuse_ftp or "").strip()
        if err or (raw is None and not reuse_ok):
            failures.append(
                {
                    "upload_id": row.id,
                    "file_name": row.file_name or "",
                    "error": err or "无法读取文档",
                }
            )
            continue
        fname = _normalize_handoff_display_filename(fname or "document.docx")
        item: dict[str, Any] = {"purpose": mode_k, "filename": fname}
        handoff_ctx = _build_aiprintword_handoff_context(row)
        if handoff_ctx:
            item["handoff_context"] = handoff_ctx
        if reuse_ok:
            item["reuse_ftp_path"] = reuse_ok
        else:
            ff = f"file_{idx}"
            item["file_field"] = ff
            files_payload[ff] = (fname, raw or b"")
        manifest.append(item)

    if not manifest:
        return {"ok": False, "error": "批量交接失败", "failures": failures}, 502

    post_data: dict[str, str] = {
        "purpose": mode_k,
        "manifest": json.dumps(manifest, ensure_ascii=False),
    }
    try:
        r = _req.post(
            url,
            headers={"X-Aiword-Handoff-Secret": secret},
            files=files_payload or None,
            data=post_data,
            timeout=180,
        )
    except _req.RequestException as e:
        return {"ok": False, "error": f"连接 aiprintword 失败：{e}", "failures": failures}, 502
    if r.status_code != 200:
        return {"ok": False, "error": f"aiprintword 返回 HTTP {r.status_code}", "failures": failures}, 502
    try:
        payload = r.json()
    except Exception:
        return {"ok": False, "error": "aiprintword 返回非 JSON", "failures": failures}, 502
    if not payload.get("ok"):
        return {"ok": False, "error": str(payload.get("error") or "unknown"), "failures": failures}, 502

    batch_token = (payload.get("batch_token") or "").strip()
    if not batch_token:
        return {"ok": False, "error": "未返回 batch_token", "failures": failures}, 502
    api_failures = payload.get("failures")
    if isinstance(api_failures, list):
        for x in api_failures:
            if isinstance(x, dict):
                failures.append(
                    {
                        "upload_id": str(x.get("upload_id") or ""),
                        "file_name": str(x.get("filename") or ""),
                        "error": str(x.get("error") or "unknown"),
                    }
                )

    if mode_k == "sign":
        redirect_url = f"{base}/sign?from=aiword&handoff_batch_token={_urlquote(batch_token)}"
    else:
        redirect_url = f"{base}/?from=aiword&handoff_batch_token={_urlquote(batch_token)}"
    success_count = int(payload.get("success_count") or 0)
    return {
        "ok": True,
        "redirect_url": redirect_url,
        "success_count": success_count,
        "failure_count": len(failures),
        "failures": failures,
    }, 200


def _build_option_tree(records: list[UploadRecord]) -> list[dict[str, Any]]:
    projects: dict[str, dict[str, Any]] = {}
    for record in records:
        proj = projects.setdefault(
            record.project_name,
            {"projectName": record.project_name, "files": {}},
        )
        file_entry = proj["files"].setdefault(
            record.file_name,
            {"fileName": record.file_name, "authors": []},
        )
        file_entry["authors"].append(
            {
                "author": record.author,
                "uploadId": record.id,
                "hasLinks": bool(record.template_links),
                "taskStatus": record.task_status,
                "quickCompleted": record.quick_completed,
            }
        )
    formatted = []
    for project in projects.values():
        files = list(project["files"].values())
        project["files"] = files
        formatted.append(project)
    return formatted


def _summary_payload():
    """
    统计逻辑：
    - completion_status 有值 => 已完成
    - completion_status 为空 => 未完成
    - 项目+人员统计融合各完成状态数量
    """
    proj_meta = _project_meta_map(auto_create_from_uploads=True)
    ended = {n for n, m in proj_meta.items() if (m.get("status") or "").strip().lower() == Project.STATUS_ENDED}
    q = UploadRecord.query
    if ended:
        q = q.filter(~UploadRecord.project_name.in_(list(ended)))
    uploads = _sort_upload_records_by_project_priority(
        q.order_by(UploadRecord.sort_order.asc(), UploadRecord.created_at.asc()).all(),
        proj_meta,
    )
    from .authz import filter_upload_records_in_scope, is_page13_super_admin

    if not is_page13_super_admin():
        uploads = filter_upload_records_in_scope(uploads)
    total_files = len(uploads)
    
    def _rate(done: int, total: int) -> float:
        return round(done / total, 4) if total else 0.0
    
    completed_files = sum(1 for u in uploads if u.completion_status)
    
    by_project: dict[str, dict[str, Any]] = {}
    by_author: dict[str, dict[str, Any]] = {}
    by_project_author: dict[str, dict[str, Any]] = {}

    for u in uploads:
        proj_key = u.project_name
        auth_key = u.author
        proj_auth_key = f"{u.project_name}__{u.author}"
        is_completed = bool(u.completion_status)
        status = u.completion_status or "未完成"

        for bucket, key in (
            (by_project, proj_key),
            (by_author, auth_key),
            (by_project_author, proj_auth_key),
        ):
            stats = bucket.setdefault(key, {
                "total": 0, "completed": 0, "pending": 0,
                "byStatus": {}, "pendingAuthors": set(),
                "auditRejectCount": 0,
            })
            stats["total"] += 1
            if is_completed:
                stats["completed"] += 1
            else:
                stats["pending"] += 1
                stats["pendingAuthors"].add(u.author)
            stats["byStatus"][status] = stats["byStatus"].get(status, 0) + 1
            stats["auditRejectCount"] = stats.get("auditRejectCount", 0) + (getattr(u, "audit_reject_count", None) or 0)

    def _format_with_status(bucket: dict[str, dict[str, Any]], label_join: str = "", include_project_author_keys: bool = False):
        formatted = []
        for key, stats in bucket.items():
            label = key
            if label_join and "__" in key:
                parts = key.split("__")
                label = label_join.join(parts)
            by_status_list = [
                {"status": s, "count": c}
                for s, c in sorted(stats["byStatus"].items(), key=lambda x: (x[0] != "未完成", -x[1]))
            ]
            item = {
                "label": label,
                "completed": stats["completed"],
                "pending": stats["pending"],
                "total": stats["total"],
                "rate": _rate(stats["completed"], stats["total"]),
                "byStatus": by_status_list,
                "pendingAuthors": list(stats["pendingAuthors"]),
                "auditRejectCount": stats.get("auditRejectCount", 0),
            }
            if not include_project_author_keys:
                # byProject / byAuthor
                if bucket is by_project:
                    m = proj_meta.get(key) or {}
                    item["projectPriority"] = int(m.get("priority") or Project.PRIORITY_MEDIUM)
                    item["projectStatus"] = (m.get("status") or Project.STATUS_ACTIVE)
            if include_project_author_keys and "__" in key:
                p, a = key.split("__", 1)
                item["projectName"] = p
                item["author"] = a
                m = proj_meta.get(p) or {}
                item["projectPriority"] = int(m.get("priority") or Project.PRIORITY_MEDIUM)
                item["projectStatus"] = (m.get("status") or Project.STATUS_ACTIVE)
            formatted.append(item)
        return formatted

    detail_rows = [
        {
            "seq": idx + 1,
            "uploadId": u.id,
            "projectName": u.project_name,
            "fileName": u.file_name,
            "taskType": u.task_type,
            "author": u.author,
            "completionStatus": u.completion_status,
            "isCompleted": bool(u.completion_status),
            "dueDate": u.due_date.strftime("%Y-%m-%d") if u.due_date else None,
            "sortOrder": u.sort_order,
            "businessSide": u.business_side,
            "product": u.product,
            "country": u.country,
            "projectCode": getattr(u, "project_code", None),
            "projectPriority": int((proj_meta.get(u.project_name) or {}).get("priority") or Project.PRIORITY_MEDIUM),
            "fileVersion": getattr(u, "file_version", None),
            "documentDisplayDate": (lambda d: d.strftime("%Y-%m-%d") if d else None)(getattr(u, "document_display_date", None)),
            "reviewer": getattr(u, "reviewer", None),
            "approver": getattr(u, "approver", None),
            "belongingModule": getattr(u, "belonging_module", None),
            "displayedAuthor": getattr(u, "displayed_author", None),
            "docLink": (u.get_template_links_list() or [None])[0] or None,
            "createdAt": u.created_at.isoformat() if u.created_at else None,
            "notes": u.notes,
            "projectNotes": getattr(u, "project_notes", None),
            "executionNotes": u.execution_notes,
            "registeredProductName": getattr(u, "registered_product_name", None),
            "model": getattr(u, "model", None),
            "registrationVersion": getattr(u, "registration_version", None),
        }
        for idx, u in enumerate(uploads)
    ]

    return {
        "overall": {
            "completed": completed_files,
            "pending": total_files - completed_files,
            "total": total_files,
            "rate": _rate(completed_files, total_files),
        },
        "byProject": sorted(
            _format_with_status(by_project),
            key=lambda x: (-int(x.get("projectPriority") or Project.PRIORITY_MEDIUM), str(x.get("label") or "")),
        ),
        "byAuthor": _format_with_status(by_author),
        "byProjectAuthor": sorted(
            _format_with_status(by_project_author, label_join=" / ", include_project_author_keys=True),
            key=lambda x: (-int(x.get("projectPriority") or Project.PRIORITY_MEDIUM), str(x.get("label") or "")),
        ),
        "detail": detail_rows,
    }


# ---------- 登录验证装饰器 ----------

def login_required(f):
    """页面2 等：须登录；页面4 访问密码视同超级管理员；公司管理员仅页面0。"""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        from .authz import (
            _company_admin_blocked_response,
            block_until_super_admin_or_user_id,
            is_company_admin,
            is_page13_super_admin,
        )

        if is_page13_super_admin():
            return f(*args, **kwargs)
        blocked = block_until_super_admin_or_user_id()
        if blocked is not None:
            return blocked
        if is_company_admin():
            return _company_admin_blocked_response()
        return f(*args, **kwargs)
    return decorated_function


def _page13_or_login_required(f):
    """访问密码超级管理员，或任意已登录账号（含页面2 普通账号；公司管理员除外）。"""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        from .authz import (
            _company_admin_blocked_response,
            block_until_super_admin_or_user_id,
            is_company_admin,
            is_page13_super_admin,
        )

        if is_page13_super_admin():
            return f(*args, **kwargs)
        if session.get("user_id"):
            if is_company_admin():
                return _company_admin_blocked_response()
            blocked = block_until_super_admin_or_user_id()
            if blocked is not None:
                return blocked
            return f(*args, **kwargs)
        blocked = block_until_super_admin_or_user_id()
        if blocked is not None:
            return blocked
        return f(*args, **kwargs)
    return decorated_function


def _page13_password_configured() -> bool:
    from .authz import page13_password_configured

    return page13_password_configured()


def super_admin_required(f):
    """页面4：仅访问密码超级管理员（字典/账号/配置/系统与钉钉维护）。"""
    from .authz import super_admin_required as _authz_super_admin_required

    return _authz_super_admin_required(f)


def page4_access_required(f):
    """页面4 HTML：须访问密码验证。"""
    from .authz import page4_access_required as _authz_page4_access_required

    return _authz_page4_access_required(f)


def page13_access_required(f):
    """页面1、页面3：默认走账号登录（项目管理员 + 项目组）；超级管理员仅须访问密码。"""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        from .authz import (
            _company_admin_blocked_response,
            block_until_super_admin_or_user_id,
            company_registry_enabled,
            is_company_admin,
            is_page13_super_admin,
            is_project_admin,
        )

        if is_page13_super_admin():
            return f(*args, **kwargs)
        if session.get("user_id") and is_company_admin():
            return _company_admin_blocked_response()
        if session.get("user_id") and is_project_admin():
            blocked = block_until_super_admin_or_user_id()
            if blocked is not None:
                return blocked
            return f(*args, **kwargs)
        blocked = block_until_super_admin_or_user_id()
        if blocked is not None:
            return blocked
        if company_registry_enabled():
            if request.is_json or request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.path.startswith("/api/"):
                return jsonify({"message": "页面1/3 仅项目管理员可访问"}), 403
            return (
                render_template(
                    "error.html",
                    title="无访问权限",
                    message="页面1/3 仅项目管理员可访问，请使用项目管理员账号登录。",
                    back_url=url_for("pages.login_page"),
                    back_label="重新登录",
                    hide_main_nav=True,
                    gate_page=True,
                ),
                403,
            )
        return f(*args, **kwargs)
    return decorated_function


# ---------- 页面路由 ----------

@bp.route("/favicon.ico")
def favicon():
    """避免浏览器请求 favicon 时 404，返回空响应。"""
    return "", 204


@bp.route("/")
def index():
    # 两套用户体系：
    # - 页面4 访问密码（超级管理员）：进入页面1/3 或考试中心老师/统计端
    # - user_id（学生端/页面2）：进入页面2
    if session.get("user_id"):
        from .authz import role_home_url, user_access_binding_block_response

        blocked = user_access_binding_block_response()
        if blocked is not None:
            return blocked
        return redirect(role_home_url())
    if session.get("page13_authenticated"):
        from .authz import role_home_url

        return redirect(role_home_url())
    return redirect(url_for("pages.login_page"))


@bp.route("/upload")
@page13_access_required
def upload_page():
    return render_template("upload.html")


@bp.route("/go/sign")
@page13_access_required
def go_aiprintword_sign():
    """页面1：交接当前任务文档到 aiprintword 在线签字页。"""
    return _aiprintword_go_redirect("sign")


@bp.route("/go/print")
@page13_access_required
def go_aiprintword_print():
    """页面1：交接当前任务文档到 aiprintword 批量打印页。"""
    return _aiprintword_go_redirect("print")


@bp.route("/api/go/batch-sign", methods=["POST"])
@page13_access_required
def api_go_aiprintword_batch_sign():
    data = request.get_json(silent=True) or {}
    ids = data.get("upload_ids")
    if not isinstance(ids, list):
        return jsonify({"ok": False, "error": "upload_ids 必须为数组"}), 400
    payload, status = _aiprintword_batch_handoff_redirect("sign", [str(x or "").strip() for x in ids])
    return jsonify(payload), status


@bp.route("/api/go/batch-print", methods=["POST"])
@page13_access_required
def api_go_aiprintword_batch_print():
    data = request.get_json(silent=True) or {}
    ids = data.get("upload_ids")
    if not isinstance(ids, list):
        return jsonify({"ok": False, "error": "upload_ids 必须为数组"}), 400
    payload, status = _aiprintword_batch_handoff_redirect("print", [str(x or "").strip() for x in ids])
    return jsonify(payload), status


@bp.route("/login")
def login_page():
    if request.args:
        return redirect(url_for("pages.login_page"))
    if session.get("page13_authenticated"):
        from .authz import role_home_url

        return redirect(role_home_url())
    if session.get("user_id"):
        from .authz import role_home_url, validate_user_access_binding
        from .models import User

        user = User.query.get(session.get("user_id"))
        ok, message, _ = validate_user_access_binding(user)
        if not ok:
            session.clear()
            return render_template(
                "login.html",
                page13_password_configured=_page13_password_configured(),
                login_error=message,
            )
        return redirect(role_home_url())
    return render_template(
        "login.html",
        page13_password_configured=_page13_password_configured(),
    )


@bp.route("/generate")
@login_required
def generate_page():
    return render_template("generate.html")


@bp.route("/dashboard")
@page13_access_required
def dashboard_page():
    return render_template("dashboard.html")


def _exam_center_display_user() -> str:
    if session.get("user_id") is not None:
        name = (session.get("display_name") or session.get("username") or "").strip()
        if name:
            return name
        return str(session.get("user_id"))
    if session.get("page13_authenticated"):
        return "超级管理员（页面4）"
    return ""


@bp.route("/exam-center")
def exam_center_page():
    from .authz import is_company_admin, is_exam_center_staff, is_page13_super_admin, is_project_admin

    if session.get("user_id") and is_company_admin() and not is_page13_super_admin():
        return (
            render_template(
                "error.html",
                title="无访问权限",
                message="公司管理员仅可访问页面0（公司总览），无法进入考试训练中心。",
                back_url=url_for("company.company_registry_page"),
                back_label="返回页面0",
                hide_main_nav=True,
            ),
            403,
        )

    role_arg = request.args.get("role")
    staff = is_exam_center_staff()
    if role_arg is None or str(role_arg).strip() == "":
        if staff:
            role = "teacher"
        else:
            role = "student"
    else:
        role = str(role_arg).strip().lower()
    if role not in {"teacher", "student", "analytics"}:
        role = "student"

    allowed_roles = []
    if session.get("user_id"):
        if is_project_admin():
            allowed_roles.extend(["teacher", "student", "analytics"])
        else:
            allowed_roles.append("student")
    if session.get("page13_authenticated"):
        for r in ("teacher", "student", "analytics"):
            if r not in allowed_roles:
                allowed_roles.append(r)

    if role in {"teacher", "analytics"}:
        if not staff:
            from .authz import (
                block_until_super_admin_or_user_id,
                page13_password_configured,
                super_admin_password_gate_response,
            )

            if page13_password_configured():
                return super_admin_password_gate_response(
                    gate_description="请输入访问密码以进入考试训练中心老师/统计端（超级管理员可见全部公司与项目组数据）。",
                )
            blocked = block_until_super_admin_or_user_id()
            if blocked is not None:
                return blocked
    elif role == "student":
        if not is_page13_super_admin():
            from .authz import block_until_super_admin_or_user_id

            blocked = block_until_super_admin_or_user_id()
            if blocked is not None:
                return blocked
    else:
        from .authz import block_until_super_admin_or_user_id

        blocked = block_until_super_admin_or_user_id()
        if blocked is not None:
            return blocked

    session["exam_center_active_role"] = role

    return render_template(
        "exam_center.html",
        exam_role=role,
        exam_allowed_roles=allowed_roles,
        exam_display_user=_exam_center_display_user(),
        hide_main_nav=(role == "student"),
    )


def _exam_allowed_organization_ids_for_scope() -> list[str]:
    from .exam_scope import allowed_organization_ids

    return allowed_organization_ids()


def _exam_scope_organizations_payload() -> list[dict]:
    from .exam_scope import allowed_organization_ids
    from .team_organizations import organizations_payload_for_ids

    return organizations_payload_for_ids(allowed_organization_ids())


def _exam_student_assigned_teams_payload() -> list[dict[str, str]]:
    from .authz import user_team_ids
    from .models import ProjectTeam

    ids = user_team_ids()
    if not ids:
        return []
    rows = (
        ProjectTeam.query.filter(ProjectTeam.id.in_(ids), ProjectTeam.is_active.is_(True))
        .order_by(ProjectTeam.sort_order.asc(), ProjectTeam.name.asc())
        .all()
    )
    by_id = {str(t.id or "").strip(): t for t in rows}
    out: list[dict[str, str]] = []
    for tid in ids:
        t = by_id.get(str(tid or "").strip())
        if not t:
            continue
        out.append({"id": str(t.id or "").strip(), "name": str(t.name or t.id or "").strip()})
    return out


def _exam_student_scope_context_payload() -> dict:
    """学生端：项目组只读；公司来自所属项目组关联公司（不含直绑公司 membership）。"""
    from .authz import is_page13_super_admin
    from .team_organizations import organizations_payload_for_ids
    from .tenant_context import collection_for_organization, user_allowed_organization_ids

    org_ids = user_allowed_organization_ids()
    orgs = organizations_payload_for_ids(org_ids)
    assigned = _exam_student_assigned_teams_payload()
    if not orgs:
        return {
            "organizations": [],
            "activeOrganizationId": None,
            "activeKnowledgeCollection": "regulations",
            "teams": assigned,
            "assignedTeams": assigned,
            "activeTeamId": assigned[0]["id"] if assigned else None,
            "scopeAllTeams": False,
            "defaultTeamId": assigned[0]["id"] if assigned else None,
            "canSwitchTeam": False,
            "canSwitchOrganization": False,
            "page13SuperAdmin": is_page13_super_admin(),
            "message": "当前账号未分配所属公司，请联系管理员在页面4配置",
        }
    from .exam_scope import resolve_active_organization_id

    active = resolve_active_organization_id()
    coll = collection_for_organization(active) if active else "regulations"
    return {
        "organizations": orgs,
        "activeOrganizationId": active or None,
        "activeKnowledgeCollection": coll,
        "teams": assigned,
        "assignedTeams": assigned,
        "activeTeamId": assigned[0]["id"] if assigned else None,
        "scopeAllTeams": False,
        "defaultTeamId": assigned[0]["id"] if assigned else None,
        "canSwitchTeam": False,
        "canSwitchOrganization": len(orgs) > 1,
        "page13SuperAdmin": is_page13_super_admin(),
    }


def _exam_project_admin_organization_ids() -> list[str]:
    from .exam_scope import allowed_organization_ids

    return allowed_organization_ids()


def _exam_project_admin_scope_context_payload() -> dict:
    """项目管理员：所属项目组只读；公司来自项目组关联（一人仅一组）。"""
    from .authz import is_page13_super_admin, user_team_ids
    from .team_organizations import organizations_payload_for_ids
    from .tenant_context import collection_for_organization

    assigned = _exam_student_assigned_teams_payload()
    if not user_team_ids():
        return {
            "organizations": [],
            "activeOrganizationId": None,
            "activeKnowledgeCollection": "regulations",
            "teams": [],
            "assignedTeams": [],
            "activeTeamId": None,
            "scopeAllTeams": False,
            "defaultTeamId": None,
            "canSwitchTeam": False,
            "canSwitchOrganization": False,
            "page13SuperAdmin": False,
            "isProjectAdmin": True,
            "message": "当前账号未分配项目组，请联系管理员在页面4配置",
        }
    org_ids = _exam_project_admin_organization_ids()
    orgs = organizations_payload_for_ids(org_ids)
    if not orgs:
        return {
            "organizations": [],
            "activeOrganizationId": None,
            "activeKnowledgeCollection": "regulations",
            "teams": assigned,
            "assignedTeams": assigned,
            "activeTeamId": assigned[0]["id"] if assigned else None,
            "scopeAllTeams": False,
            "defaultTeamId": assigned[0]["id"] if assigned else None,
            "canSwitchTeam": False,
            "canSwitchOrganization": False,
            "page13SuperAdmin": is_page13_super_admin(),
            "isProjectAdmin": True,
            "message": "所属项目组未关联任何公司，请超级管理员在页面4维护项目组关联公司",
        }
    from .exam_scope import resolve_active_organization_id

    active = resolve_active_organization_id()
    org_id = active
    active_team = _sync_project_admin_active_exam_team(org_id)
    teams = _exam_teams_payload_for_scope(org_id)
    can_switch_team = False
    coll = collection_for_organization(active) if active else "regulations"
    return {
        "organizations": orgs,
        "activeOrganizationId": active or None,
        "activeKnowledgeCollection": coll,
        "teams": teams,
        "assignedTeams": assigned,
        "activeTeamId": active_team or None,
        "scopeAllTeams": False,
        "defaultTeamId": teams[0]["id"] if teams else None,
        "canSwitchTeam": can_switch_team,
        "canSwitchOrganization": len(orgs) > 1,
        "page13SuperAdmin": is_page13_super_admin(),
        "isProjectAdmin": True,
    }


@bp.get("/api/scope/context")
def api_scope_context():
    """全站作用域条：角色、公司/项目组、本页过滤说明与空数据提示。"""
    from .authz import block_until_super_admin_or_user_id, is_page13_super_admin
    from .scope_context import infer_page_key, scope_context_payload

    if not is_page13_super_admin() and not session.get("user_id"):
        blocked = block_until_super_admin_or_user_id()
        if blocked is not None:
            return blocked
    page = str(request.args.get("page") or request.args.get("pageKey") or "").strip()
    if not page:
        page = infer_page_key()
    return jsonify(scope_context_payload(page_key=page))


@bp.get("/api/scope/diagnostics")
@page13_access_required
def api_scope_diagnostics():
    """超级管理员：session/绑定/有效过滤诊断。"""
    from .scope_context import scope_diagnostics_payload

    payload = scope_diagnostics_payload()
    if not payload.get("ok"):
        return jsonify(payload), 403
    return jsonify(payload)


@bp.get("/api/exam-center/scope-context")
@_page13_or_login_required
def api_exam_scope_context():
    """考试中心：当前公司与项目组作用域（超管可切换全部主体，数据按所选过滤）。"""
    from .authz import is_normal_user, is_page13_super_admin, is_project_admin
    from .tenant_context import integration_org_context_payload

    if not (
        is_page13_super_admin()
        or session.get("user_id")
    ):
        return jsonify({"message": "请先登录或使用页面4 访问密码"}), 401
    if is_normal_user():
        return jsonify(_exam_student_scope_context_payload())
    if is_project_admin() and not is_page13_super_admin():
        return jsonify(_exam_project_admin_scope_context_payload())
    org_payload = integration_org_context_payload()
    if not org_payload.get("organizations") and org_payload.get("message"):
        org_payload.setdefault("teams", [])
        org_payload.setdefault("canSwitchTeam", False)
        org_payload.setdefault("canSwitchOrganization", False)
        org_payload.setdefault("page13SuperAdmin", is_page13_super_admin())
        return jsonify(org_payload)
    from .exam_scope import resolve_active_organization_id

    org_id = resolve_active_organization_id()
    active_team = _sync_super_admin_active_exam_team(org_id) if is_page13_super_admin() else ""
    teams = _exam_teams_payload_for_scope(org_id) if is_page13_super_admin() else []
    orgs = org_payload.get("organizations") or []
    return jsonify(
        {
            "organizations": orgs,
            "activeOrganizationId": org_payload.get("activeOrganizationId") or org_id or None,
            "activeKnowledgeCollection": org_payload.get("activeKnowledgeCollection"),
            "teams": teams,
            "activeTeamId": active_team or None,
            "scopeAllTeams": bool(session.get("exam_team_scope_all")),
            "defaultTeamId": _resolve_default_exam_team_id_for_org(org_id) or None,
            "canSwitchTeam": is_page13_super_admin(),
            "canSwitchOrganization": is_page13_super_admin() or len(orgs) > 1,
            "page13SuperAdmin": is_page13_super_admin(),
        }
    )


@bp.post("/api/exam-center/scope-context/organization")
@page13_access_required
def api_exam_set_active_organization():
    """切换考试中心当前公司（写入 session.active_organization_id）。"""
    from .authz import is_page13_super_admin
    from .tenant_context import collection_for_organization

    data = request.get_json(force=True) or {}
    target = str(data.get("organizationId") or data.get("organization_id") or "").strip()
    allowed = {str(x.get("id") or "").strip() for x in _exam_scope_organizations_payload()}
    if not target:
        return jsonify({"message": "缺少 organizationId"}), 400
    if target not in allowed:
        return jsonify({"message": "无权切换到该公司"}), 403
    session["active_organization_id"] = target
    if is_page13_super_admin():
        session.pop("active_exam_team_id", None)
        session.pop("exam_team_scope_all", None)
        _apply_super_admin_active_exam_team(target)
    else:
        from .authz import is_project_admin

        if is_project_admin():
            _sync_project_admin_active_exam_team(target)
    return jsonify(
        {
            "ok": True,
            "activeOrganizationId": target,
            "activeKnowledgeCollection": collection_for_organization(target),
        }
    )


@bp.post("/api/exam-center/scope-context/team")
@page13_access_required
def api_exam_set_active_team():
    """切换考试中心当前项目组（超管可选全部；项目管理员仅限所属组）。"""
    from .authz import is_page13_super_admin, is_project_admin

    if not (is_page13_super_admin() or is_project_admin()):
        return jsonify({"message": "无权切换项目组"}), 403
    data = request.get_json(force=True) or {}
    team_id = str(data.get("teamId") or data.get("team_id") or "").strip()
    org_id = _current_exam_scope_organization_id()
    allowed = {t["id"] for t in _exam_teams_payload_for_scope(org_id)}
    if is_project_admin() and not is_page13_super_admin():
        if not team_id:
            return jsonify({"message": "项目管理员须选择具体项目组"}), 400
        if team_id not in allowed:
            return jsonify({"message": "无权选择该项目组"}), 403
        session["active_exam_team_id"] = team_id
        session.pop("exam_team_scope_all", None)
        return jsonify({"ok": True, "activeTeamId": team_id, "scopeAllTeams": False})
    if team_id:
        if team_id not in allowed:
            return jsonify({"message": "无权选择该项目组"}), 403
        session["active_exam_team_id"] = team_id
        session.pop("exam_team_scope_all", None)
    else:
        session.pop("active_exam_team_id", None)
        session["exam_team_scope_all"] = True
    return jsonify({"ok": True, "activeTeamId": team_id or None, "scopeAllTeams": not bool(team_id)})


# ---------- 认证 API ----------

@bp.post("/api/login")
def api_login():
    data = request.get_json(force=True) or {}
    username = (data.get("username") or "").strip()
    password = data.get("password")
    if password is None:
        password = ""
    else:
        # 与创建账号时 strip 保持一致，避免复制密码带首尾空格导致误判
        password = str(password).strip()
    if not username or not password:
        return jsonify({"message": "用户名和密码不能为空"}), 400
    user = User.query.filter_by(username=username).first()
    if not user or not user.check_password(password):
        return jsonify({"message": "用户名或密码错误，请核对账号与密码后重试"}), 401
    from .authz import company_registry_enabled, role_home_url, sync_user_session, validate_user_access_binding

    sync_user_session(user)
    role = session.get("admin_role") or "none"

    ok, message, extra = validate_user_access_binding(user)
    if not ok:
        session.clear()
        return jsonify({"message": message, **extra}), 403

    home = role_home_url()
    return jsonify({
        "message": "登录成功",
        "homeUrl": home,
        "redirectUrl": home,
        "user": {
            "id": user.id,
            "username": user.username,
            "displayName": user.display_name,
            "adminRole": role,
            "teamIds": list(session.get("team_ids") or []),
            "canAccessCompanyRegistry": bool(session.get("can_access_company_registry")),
            "registeredCountries": list(session.get("country_scopes") or []),
        },
        "companyRegistryEnabled": company_registry_enabled(),
    })


@bp.post("/api/logout")
def api_logout():
    session.clear()
    return jsonify({"message": "已退出登录"})


@bp.get("/api/me")
def api_me():
    from .authz import is_page13_super_admin, role_home_url

    if not session.get("user_id"):
        if is_page13_super_admin():
            from .app_settings import effective_feature_flags_for_request
            from .authz import company_registry_enabled

            return jsonify({
                "loggedIn": False,
                "page13SuperAdmin": True,
                "page2ViewMode": "super_admin_readonly",
                "examStudentViewMode": "super_admin_readonly",
                "homeUrl": role_home_url(),
                "featureAdminViewer": True,
                "featureFlags": effective_feature_flags_for_request(),
                "companyRegistryEnabled": company_registry_enabled(),
            })
        return jsonify({"loggedIn": False})
    from .app_settings import effective_feature_flags_for_request

    from .authz import (
        company_registry_enabled,
        current_admin_role,
        is_company_registry_user,
        is_page13_super_admin,
        is_project_admin,
        role_home_url,
        user_country_scopes,
        validate_user_access_binding,
    )
    from .models import User
    from .observer_view import exam_student_view_mode, page2_view_mode

    scopes = user_country_scopes()
    user = User.query.get(session.get("user_id"))
    binding_ok, binding_message, binding_flags = validate_user_access_binding(user)
    return jsonify({
        "loggedIn": True,
        "homeUrl": role_home_url(),
        "accessBindingOk": binding_ok,
        "accessBindingMessage": binding_message or None,
        **binding_flags,
        "page13SuperAdmin": is_page13_super_admin(),
        "page2ViewMode": page2_view_mode(),
        "examStudentViewMode": exam_student_view_mode(),
        "readOnlyObserver": page2_view_mode() != "normal",
        "user": {
            "id": session.get("user_id"),
            "username": session.get("username"),
            "displayName": session.get("display_name"),
            "adminRole": current_admin_role(),
            "teamIds": list(session.get("team_ids") or []),
            "canAccessCompanyRegistry": is_company_registry_user(),
            "isProjectAdmin": is_project_admin(),
            "registeredCountries": list(scopes or []),
        },
        "featureAdminViewer": bool(session.get("page13_authenticated")),
        "featureFlags": effective_feature_flags_for_request(),
        "companyRegistryEnabled": company_registry_enabled(),
    })


@bp.get("/api/page13-auth-state")
def api_page13_auth_state():
    """获取页面1/3 是否需密码及本次验证用的 nonce（不校验身份，任何人可调）。"""
    required = _page13_password_configured()
    authenticated = bool(session.get("page13_authenticated"))
    if not required:
        return jsonify({"required": False, "authenticated": True})
    if authenticated:
        return jsonify({"required": True, "authenticated": True})
    nonce = secrets.token_hex(16)
    session["page13_nonce"] = nonce
    return jsonify({"required": True, "authenticated": False, "nonce": nonce})


@bp.post("/api/page13-auth")
def api_page13_auth():
    """提交 hash(nonce+password)，校验通过后设置 session，密码不明文传输。"""
    if not _page13_password_configured():
        return jsonify({"success": True, "message": "未配置访问密码"})
    try:
        data = request.get_json(force=True, silent=True) or {}
    except Exception:
        data = {}
    client_hash = (data.get("hash") or "").strip()
    if not client_hash:
        return jsonify({"message": "缺少校验参数"}), 400
    nonce = session.get("page13_nonce")
    if not nonce:
        return jsonify({"message": "请先获取验证码（刷新页面后重试）"}), 400
    from .app_settings import get_setting
    raw = get_setting("PAGE13_ACCESS_PASSWORD", default=str(current_app.config.get("PAGE13_ACCESS_PASSWORD") or ""))
    password = (str(raw).replace("\ufeff", "").strip() if raw else "")
    expected = hashlib.sha256((nonce + password).encode("utf-8")).hexdigest()
    session.pop("page13_nonce", None)
    if not secrets.compare_digest(expected, client_hash):
        return jsonify({"message": "访问密码错误"}), 401
    session["page13_authenticated"] = True
    from .authz import role_home_url

    return jsonify({
        "success": True,
        "message": "验证成功",
        "homeUrl": role_home_url(),
        "page13SuperAdmin": True,
    })


# ---------- 考试训练中心 API（aiword 代理） ----------

def _json_payload() -> dict[str, Any]:
    data = request.get_json(silent=True)
    if isinstance(data, dict):
        return data
    return {}


def _guess_job_id_from_payload(p: dict[str, Any]) -> str:
    """从代理层整包中提取上游 ingest 的 job_id。上游常见形态：p['data'] = {ok, data: {job_id, set_id}}。"""
    if not isinstance(p, dict):
        return ""
    data = p.get("data")
    if isinstance(data, dict):
        for k in ("job_id", "jobId", "id", "jobID"):
            v = data.get(k)
            if v:
                return str(v).strip()
        inner = data.get("data")
        if isinstance(inner, dict):
            for k in ("job_id", "jobId", "id", "jobID"):
                v = inner.get(k)
                if v:
                    return str(v).strip()
    for k in ("job_id", "jobId", "id", "jobID"):
        v = p.get(k)
        if v:
            return str(v).strip()
    return ""


def _normalize_job_status_from_upstream_data(upstream_data: Any) -> str:
    if not isinstance(upstream_data, dict):
        return "unknown"
    raw = (
        upstream_data.get("status")
        or upstream_data.get("state")
        or upstream_data.get("job_status")
        or upstream_data.get("jobStatus")
        or ""
    )
    if not raw:
        inner = upstream_data.get("data")
        if isinstance(inner, dict):
            raw = (
                inner.get("status")
                or inner.get("state")
                or inner.get("job_status")
                or inner.get("jobStatus")
                or ""
            )
        job = upstream_data.get("job")
        if not raw and isinstance(job, dict):
            raw = job.get("status") or job.get("state") or ""
    s = str(raw).strip().lower()
    if s in {"done", "success", "completed"}:
        return "done"
    if s in {"failed", "error"}:
        return "failed"
    if s in {"running", "processing", "in_progress", "progress"}:
        return "running"
    if s in {"pending", "queued", "created"}:
        return "pending"
    return s or "unknown"


def _extract_set_id_from_dict(obj: Any) -> str:
    """从任意 dict 中尽力提取上游 set_id（兼容 setId / quiz_sets.id 等常见形态）。"""
    if not isinstance(obj, dict):
        return ""
    for k in ("set_id", "setId", "quiz_set_id", "quizSetId"):
        v = obj.get(k)
        if v:
            s = str(v).strip()
            if s:
                return s
    inner = obj.get("set")
    if isinstance(inner, dict):
        for k in ("id", "set_id", "setId"):
            v = inner.get(k)
            if v:
                s = str(v).strip()
                if s:
                    return s
    return ""


def _extract_set_id_from_any(obj: Any) -> str:
    """递归浅层扫描：优先顶层字段，其次常见嵌套 data/job/result。"""
    sid = _extract_set_id_from_dict(obj) if isinstance(obj, dict) else ""
    if sid:
        return sid
    if not isinstance(obj, dict):
        return ""
    for child_key in ("data", "job", "result", "payload"):
        child = obj.get(child_key)
        if isinstance(child, dict):
            sid = _extract_set_id_from_dict(child)
            if sid:
                return sid
        if isinstance(child, list) and child and isinstance(child[0], dict):
            sid = _extract_set_id_from_dict(child[0])
            if sid:
                return sid
    return ""


def _extract_set_id_from_quiz_proxy_payload(payload: Any) -> str:
    """
    aiword 代理返回结构通常是 {code,message,data,request,trace_id}，
    其中 data 可能是 dict 或 list；set_id 可能在 data 内或更深。
    """
    if not isinstance(payload, dict):
        return ""
    sid = _extract_set_id_from_any(payload.get("data"))
    if sid:
        return sid
    return _extract_set_id_from_any(payload)


@bp.post("/api/exam-center/teacher/sets/generate")
@page13_access_required
def api_exam_teacher_generate_set():
    body = _expand_quiz_request_body(_json_payload())
    status, payload = _quiz_api_call("quiz/sets/generate", method="POST", payload=body)
    return jsonify(payload), status


@bp.get("/api/exam-center/teacher/project-cases")
@page13_access_required
def api_exam_teacher_project_cases():
    """已训练项目案例列表（透传 aicheckword quiz/tools/project-cases）。"""
    coll = (request.args.get("collection") or "regulations").strip() or "regulations"
    status, payload = _quiz_api_call("quiz/tools/project-cases", method="GET", query={"collection": coll})
    return jsonify(payload), status


@bp.get("/api/exam-center/student/project-cases")
@login_required
def api_exam_student_project_cases():
    """学生端练习：已训练项目案例列表（同上，不要求 page13 老师权限）。"""
    coll = (request.args.get("collection") or "regulations").strip() or "regulations"
    status, payload = _quiz_api_call("quiz/tools/project-cases", method="GET", query={"collection": coll})
    return jsonify(payload), status


@bp.post("/api/exam-center/teacher/bank/ingest-by-ai")
@page13_access_required
def api_exam_teacher_ingest_bank():
    req_payload = _expand_quiz_request_body(_json_payload())
    org_id = _resolve_exam_organization_id(
        explicit_org_id=str(
            req_payload.get("organization_id") or req_payload.get("organizationId") or ""
        ).strip(),
        project_id=str(req_payload.get("project_id") or req_payload.get("projectId") or "").strip(),
    )
    tc = _exam_ingest_target_count()
    wk = _exam_ingest_knowledge_weights()
    wt = _exam_ingest_question_type_weights()
    msf = _exam_ingest_max_similar_frac()
    req_payload["target_count"] = tc
    req_payload["targetCount"] = tc
    req_payload["question_count"] = tc
    req_payload["questionCount"] = tc
    req_payload["ingest_knowledge_weights"] = wk
    req_payload["ingest_question_type_weights"] = wt
    req_payload["max_similar_frac"] = msf
    status, payload = _quiz_api_call(
        "quiz/bank/ingest-by-ai",
        method="POST",
        payload=req_payload,
        organization_id=org_id,
    )

    # 成功拿到 job_id 后：落库生成任务记录（用于后续从记录查询轮询状态与结果）
    if 200 <= int(status) < 300 and isinstance(payload, dict):
        upstream_job_id = _guess_job_id_from_payload(payload)
        if upstream_job_id:
            created_by = (session.get("display_name") or session.get("username") or "").strip() or None
            exam_track = (req_payload.get("exam_track") or req_payload.get("examTrack") or "").strip() or None
            exam_category = _normalize_exam_category(
                req_payload.get("exam_category") or req_payload.get("examCategory") or "daily"
            )
            try:
                target_count = int(req_payload.get("target_count") or req_payload.get("targetCount") or 0) or None
            except Exception:
                target_count = None
            review_mode = (req_payload.get("review_mode") or req_payload.get("reviewMode") or "").strip() or None
            row_q = ExamBankIngestJob.query.filter_by(upstream_job_id=upstream_job_id)
            if org_id:
                row_q = row_q.filter(ExamBankIngestJob.organization_id == org_id)
            row = row_q.first()
            if not row:
                row = ExamBankIngestJob(
                    organization_id=org_id or None,
                    upstream_job_id=upstream_job_id,
                    exam_track=exam_track,
                    exam_category=exam_category,
                    target_count=target_count,
                    review_mode=review_mode,
                    status="pending",
                    created_by=created_by,
                )
            row.organization_id = org_id or row.organization_id
            row.exam_category = exam_category
            row.last_upstream_http_status = int(status)
            row.last_upstream_data = payload.get("data") if isinstance(payload.get("data"), dict) else payload.get("data")
            row.last_upstream_request_url = ((payload.get("request") or {}).get("url") if isinstance(payload.get("request"), dict) else None) or None
            row.last_upstream_trace_id = (payload.get("trace_id") or None)
            row.last_message = (payload.get("message") or None)
            set_id = _extract_set_id_from_quiz_proxy_payload(payload)
            if set_id:
                row.upstream_set_id = set_id
            db.session.add(row)
            try:
                db.session.commit()
            except Exception as e:
                db.session.rollback()
                if isinstance(payload, dict):
                    payload["job_record_error"] = f"aiword 落库失败（不影响上游任务）：{e}"
            else:
                if isinstance(payload, dict):
                    payload["job_record"] = {
                        "id": row.id,
                        "upstream_job_id": row.upstream_job_id,
                        "upstream_set_id": row.upstream_set_id,
                        "status": row.status,
                    }

    return jsonify(payload), status


@bp.get("/api/exam-center/teacher/bank/requirements-check")
@page13_access_required
def api_exam_teacher_bank_requirements_check():
    track = _normalize_exam_track(
        request.args.get("exam_track")
        or request.args.get("examTrack")
        or request.args.get("track")
        or "cn"
    )
    data = _build_exam_bank_requirement_status(track)
    msg = "当前题库已满足体考要求。仍可继续录题以增强覆盖。" if data.get("is_satisfied") else "当前题库尚未满足体考要求。"
    return jsonify({"code": 0, "message": msg, "data": data, "trace_id": uuid.uuid4().hex}), 200


@bp.get("/api/exam-center/teacher/system-settings")
@page13_access_required
def api_exam_teacher_system_settings_get():
    """老师端「考试与录题配置」：与页面4 系统配置同源读写 app_configs，仅返回考试相关键列表。"""
    from .app_settings import (
        EXAM_CENTER_TEACHER_SETTINGS_KEYS,
        SYSTEM_CONFIG_KEYS,
        persist_config_json_into_empty_db_keys,
        sync_authoritative_sources_into_db,
        system_settings_for_api_get,
    )

    app = current_app._get_current_object()
    project_root = Path(app.root_path).resolve().parent
    sync_authoritative_sources_into_db(project_root, app)
    persist_config_json_into_empty_db_keys(project_root, app)
    meta_by_key = {k: (k, lbl, sens) for k, lbl, sens in SYSTEM_CONFIG_KEYS}
    keys_meta: list[dict[str, Any]] = []
    for key in EXAM_CENTER_TEACHER_SETTINGS_KEYS:
        row = meta_by_key.get(key)
        if not row:
            continue
        kk, lbl, sens = row
        keys_meta.append({"key": kk, "label": lbl, "sensitive": sens})
    return jsonify({"settings": system_settings_for_api_get(app, project_root), "keys": keys_meta})


@bp.post("/api/exam-center/teacher/regulatory-updates-hint")
@page13_access_required
def api_exam_teacher_regulatory_updates_hint():
    """「新标发布」备考：按体考类型请求 aicheckword 归纳需关注的法规/标准/指南更新方向（模型输出，非官方清单）。"""
    body = _expand_quiz_request_body(_json_payload())
    # 法规提示走 LLM，耗时常超过默认 20s/120s；在配置值基础上至少给 3 分钟，且不超过代理硬上限 600s
    hint_to = max(int(_quiz_api_timeout_seconds()), 180)
    st, pl = _quiz_api_call(
        "quiz/tools/regulatory-updates-hint",
        method="POST",
        payload=body,
        timeout_seconds=min(hint_to, 600),
    )
    return jsonify(pl), st


@bp.get("/api/exam-center/teacher/bank/policy-version")
@page13_access_required
def api_exam_teacher_bank_policy_version_get():
    track = _normalize_exam_track(
        request.args.get("exam_track")
        or request.args.get("examTrack")
        or request.args.get("track")
        or "cn"
    )
    pv_info = _resolve_exam_track_policy_version(track)
    return jsonify(
        {
            "code": 0,
            "message": "ok",
            "data": {
                "track": track,
                "policy_version": pv_info.get("fallback_policy_version"),
                "auto_policy_version": pv_info.get("auto_policy_version"),
                "effective_policy_version": pv_info.get("effective_policy_version"),
                "policy_version_source": pv_info.get("policy_version_source"),
                "policy_version_confidence": pv_info.get("policy_version_confidence"),
                "policy_version_evidence": pv_info.get("policy_version_evidence"),
            },
            "trace_id": uuid.uuid4().hex,
        }
    ), 200


@bp.put("/api/exam-center/teacher/bank/policy-version")
@page13_access_required
def api_exam_teacher_bank_policy_version_put():
    body = _json_payload()
    track = _normalize_exam_track(
        body.get("exam_track")
        or body.get("examTrack")
        or body.get("track")
        or request.args.get("exam_track")
        or "cn"
    )
    version_value = str(body.get("policy_version") or body.get("policyVersion") or "").strip()
    _set_exam_track_policy_version(track, version_value)
    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({"code": "DB_ERROR", "message": f"保存法规版本失败：{e}", "data": None, "trace_id": uuid.uuid4().hex}), 500
    return jsonify(
        {
            "code": 0,
            "message": "法规版本标识（兜底）已保存。",
            "data": {"track": track, "policy_version": version_value},
            "trace_id": uuid.uuid4().hex,
        }
    ), 200


@bp.post("/api/exam-center/teacher/bank/requirements-baseline")
@page13_access_required
def api_exam_teacher_bank_requirements_baseline():
    body = _json_payload()
    track = _normalize_exam_track(
        body.get("exam_track")
        or body.get("examTrack")
        or body.get("track")
        or request.args.get("exam_track")
        or "cn"
    )
    snap = _build_exam_bank_requirement_status(track)
    km = snap.get("knowledge_markers") if isinstance(snap, dict) else {}
    current = km.get("current") if isinstance(km, dict) else {}
    policy_v = km.get("current_policy_version") if isinstance(km, dict) else ""
    all_base = _app_config_get_json(_EXAM_REQUIREMENT_BASELINE_KEY, {})
    if not isinstance(all_base, dict):
        all_base = {}
    all_base[track] = {
        "regulations_at": (current.get("regulations_at") if isinstance(current, dict) else None),
        "procedures_at": (current.get("procedures_at") if isinstance(current, dict) else None),
        "project_case_at": (current.get("project_case_at") if isinstance(current, dict) else None),
        "policy_version": str(policy_v or "").strip(),
        "checked_at": datetime.now().isoformat(timespec="seconds"),
        "checked_by": (session.get("display_name") or session.get("username") or "").strip() or None,
    }
    _app_config_set_json(_EXAM_REQUIREMENT_BASELINE_KEY, all_base)
    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({"code": "DB_ERROR", "message": f"写入基线失败：{e}", "data": None, "trace_id": uuid.uuid4().hex}), 500
    latest = _build_exam_bank_requirement_status(track)
    return jsonify({"code": 0, "message": "已设置当前达标基线。", "data": latest, "trace_id": uuid.uuid4().hex}), 200


@bp.get("/api/exam-center/teacher/bank/ingest-jobs/<job_id>")
@page13_access_required
def api_exam_teacher_ingest_job(job_id: str):
    job_id = (job_id or "").strip()
    if not job_id:
        return (
            jsonify(
                {
                    "code": "BAD_REQUEST",
                    "message": "缺少 job_id",
                    "data": None,
                    "trace_id": uuid.uuid4().hex,
                    "request": {"url": "", "method": "GET", "upstreamPath": ""},
                }
            ),
            400,
        )
    # 1) 先查本地任务记录（如果不存在，也允许继续向上游查询）
    row = ExamBankIngestJob.query.filter_by(upstream_job_id=job_id).first()
    scope_org_id = _current_exam_scope_organization_id()
    if row is not None and scope_org_id and str(getattr(row, "organization_id", "") or "").strip() != scope_org_id:
        return jsonify({"code": "NOT_FOUND", "message": "本地无该任务记录", "data": None, "trace_id": uuid.uuid4().hex}), 404

    # 2) 默认刷新上游状态（refresh=0 可只看本地快照）
    refresh = (request.args.get("refresh") or "1").strip()
    if refresh not in {"0", "false", "no"}:
        req_org_id = (
            str(getattr(row, "organization_id", "") or "").strip()
            or scope_org_id
            or _resolve_exam_organization_id()
        )
        status, payload = _quiz_api_call(
            f"quiz/bank/ingest-jobs/{job_id}",
            method="GET",
            query={k: v for k, v in request.args.to_dict().items() if k != "refresh"},
            organization_id=req_org_id,
        )
        # 回写本地快照
        if row is None:
            row = ExamBankIngestJob(
                upstream_job_id=job_id,
                status="unknown",
                organization_id=req_org_id or None,
            )
        row.organization_id = req_org_id or row.organization_id
        row.last_upstream_http_status = int(status)
        if isinstance(payload, dict):
            row.last_message = (payload.get("message") or None)
            row.last_upstream_trace_id = (payload.get("trace_id") or None)
            req_meta = payload.get("request") if isinstance(payload.get("request"), dict) else {}
            row.last_upstream_request_url = (req_meta.get("url") or None)
            row.last_upstream_data = payload.get("data")
            row.status = _normalize_job_status_from_upstream_data(payload.get("data"))
            set_id = _extract_set_id_from_quiz_proxy_payload(payload)
            if set_id:
                row.upstream_set_id = set_id
        db.session.add(row)
        try:
            db.session.commit()
        except Exception as e:
            db.session.rollback()
            if isinstance(payload, dict):
                payload["job_record_error"] = f"aiword 落库失败（不影响上游查询）：{e}"
        else:
            if isinstance(payload, dict):
                payload["job_record"] = {
                    "id": row.id,
                    "upstream_job_id": row.upstream_job_id,
                    "upstream_set_id": row.upstream_set_id,
                    "status": row.status,
                }
        return jsonify(payload), status

    # 仅返回本地记录快照
    if row is None:
        return jsonify({"code": "NOT_FOUND", "message": "本地无该任务记录", "data": None, "trace_id": uuid.uuid4().hex}), 404
    return jsonify(
        {
            "code": 0,
            "message": "ok",
            "data": {
                "job_record": {
                    "id": row.id,
                    "upstream_job_id": row.upstream_job_id,
                    "upstream_set_id": getattr(row, "upstream_set_id", None),
                    "status": row.status,
                    "exam_track": row.exam_track,
                    "target_count": row.target_count,
                    "review_mode": row.review_mode,
                    "created_by": row.created_by,
                    "created_at": row.created_at.isoformat() if row.created_at else None,
                    "updated_at": row.updated_at.isoformat() if row.updated_at else None,
                    "last_message": row.last_message,
                    "last_upstream_http_status": row.last_upstream_http_status,
                    "last_upstream_request_url": row.last_upstream_request_url,
                    "last_upstream_trace_id": row.last_upstream_trace_id,
                    "last_upstream_data": row.last_upstream_data,
                }
            },
            "trace_id": uuid.uuid4().hex,
        }
    )


@bp.get("/api/exam-center/teacher/bank/ingest-jobs")
@page13_access_required
def api_exam_teacher_ingest_jobs_list():
    limit_raw = (request.args.get("limit") or "20").strip()
    try:
        limit = int(limit_raw)
    except ValueError:
        limit = 20
    if limit < 1:
        limit = 1
    if limit > 200:
        limit = 200

    scope_org_id = _current_exam_scope_organization_id()
    q_rows = ExamBankIngestJob.query
    if scope_org_id:
        q_rows = q_rows.filter(ExamBankIngestJob.organization_id == scope_org_id)
    rows = q_rows.order_by(ExamBankIngestJob.created_at.desc()).limit(limit).all()
    return jsonify(
        {
            "code": 0,
            "message": "ok",
            "data": {
                "jobs": [
                    {
                        "id": r.id,
                        "upstream_job_id": r.upstream_job_id,
                        "upstream_set_id": getattr(r, "upstream_set_id", None),
                        "status": r.status,
                        "exam_track": r.exam_track,
                        "exam_category": getattr(r, "exam_category", None) or "daily",
                        "target_count": r.target_count,
                        "review_mode": r.review_mode,
                        "created_by": r.created_by,
                        "created_at": r.created_at.isoformat() if r.created_at else None,
                        "updated_at": r.updated_at.isoformat() if r.updated_at else None,
                        "last_message": r.last_message,
                        "last_upstream_http_status": r.last_upstream_http_status,
                        "last_upstream_request_url": r.last_upstream_request_url,
                        "last_upstream_trace_id": r.last_upstream_trace_id,
                    }
                    for r in rows
                ]
            },
            "trace_id": uuid.uuid4().hex,
        }
    )


@bp.get("/api/exam-center/teacher/sets")
@page13_access_required
def api_exam_teacher_sets_list():
    """老师端：套题列表（上游）+ 本地 ingest 关联。"""
    status, payload = _quiz_api_call("quiz/sets", method="GET", query=request.args.to_dict())
    scope_org_id = _current_exam_scope_organization_id()
    q_rows = ExamBankIngestJob.query
    if scope_org_id:
        q_rows = q_rows.filter(ExamBankIngestJob.organization_id == scope_org_id)
    rows = q_rows.order_by(ExamBankIngestJob.created_at.desc()).limit(200).all()
    by_set: dict[str, Any] = {}
    without_set: list[dict[str, Any]] = []
    for r in rows:
        sid = (getattr(r, "upstream_set_id", None) or "").strip()
        jr = {
            "id": r.id,
            "upstream_job_id": r.upstream_job_id,
            "upstream_set_id": getattr(r, "upstream_set_id", None),
            "status": r.status,
            "exam_category": getattr(r, "exam_category", None) or "daily",
            "created_at": r.created_at.isoformat() if r.created_at else None,
            "updated_at": r.updated_at.isoformat() if r.updated_at else None,
        }
        if sid:
            if sid not in by_set:
                by_set[sid] = jr
        else:
            without_set.append(jr)
    if isinstance(payload, dict):
        data = payload.get("data")
        if isinstance(data, dict):
            data.setdefault("aiword", {})
            if isinstance(data.get("aiword"), dict):
                data["aiword"]["ingest_jobs_by_set_id"] = by_set
                data["aiword"]["ingest_jobs_without_set_id"] = without_set[:50]
    return jsonify(payload), status


@bp.get("/api/exam-center/teacher/sets/<set_id>")
@page13_access_required
def api_exam_teacher_set_detail(set_id: str):
    """老师端：套题详情（上游透传）。"""
    sid = (set_id or "").strip()
    if not sid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 set_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    status, payload = _quiz_api_call(f"quiz/sets/{sid}", method="GET", query=request.args.to_dict())
    return jsonify(payload), status


@bp.delete("/api/exam-center/teacher/sets/<set_id>")
@page13_access_required
def api_exam_teacher_delete_set(set_id: str):
    """老师端：删除套题（上游透传）。"""
    sid = (set_id or "").strip()
    if not sid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 set_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    status, payload = _quiz_api_call(f"quiz/sets/{sid}", method="DELETE", payload=_json_payload())
    return jsonify(payload), status


@bp.get("/api/exam-center/teacher/bank/questions")
@page13_access_required
def api_exam_teacher_bank_questions_list():
    """
    老师端：题库题目列表。
    无 set_id：GET quiz/bank/questions 透传（多路径尝试）。
    有 set_id：先 GET quiz/sets/{id} 取题序并与题库分页扫描求交（上游常忽略 set_id 查询参数）。
    """
    args = dict(request.args.items())
    set_sid = str(
        args.get("set_id")
        or args.get("setId")
        or args.get("bank_set_id")
        or args.get("bankSetId")
        or ""
    ).strip()
    paths = [
        "quiz/bank/questions",
        "quiz/questions",
        "quiz/bank/question-list",
    ]
    if not set_sid:
        status, payload, tried = _quiz_try_paths(paths, method="GET", query=args)
        if isinstance(payload, dict) and not (200 <= int(status) < 300):
            payload.setdefault("aiword_upstream_tried_paths", tried)
            payload.setdefault(
                "message",
                (payload.get("message") or "上游题库列表接口不可用")
                + "（已尝试路径：" + ", ".join(tried) + "）",
            )
        return jsonify(payload), status

    _strip_bank = frozenset({"set_id", "setId", "bank_set_id", "bankSetId"})
    bank_args = {k: v for k, v in args.items() if k not in _strip_bank}
    trace_id = uuid.uuid4().hex

    def _int_arg(name: str, default: int, lo: int, hi: int) -> int:
        try:
            v = int(str(bank_args.get(name) or default))
        except (TypeError, ValueError):
            v = default
        return max(lo, min(hi, v))

    user_off = _int_arg("offset", 0, 0, 500000)
    user_lim = _int_arg("limit", 50, 1, 200)
    scan_lim = 200
    max_scan = 12000

    def _pull_bank_page(off: int) -> tuple[list[dict[str, Any]], int, dict[str, Any]]:
        q = dict(bank_args)
        q["limit"] = str(scan_lim)
        q["offset"] = str(off)
        st_b, pl_b, _tried = _quiz_try_paths(paths, method="GET", query=q)
        if not (200 <= int(st_b) < 300):
            return [], int(st_b), pl_b if isinstance(pl_b, dict) else {"data": pl_b}
        root_b = _unwrap_quiz_api_success_data(pl_b)
        if not isinstance(root_b, dict):
            root_b = {}
        items_page = _extract_bank_items_from_list_root(root_b)
        return items_page, 200, pl_b

    st_set, pl_set = _quiz_api_call(f"quiz/sets/{set_sid}", method="GET", query={})
    root_set = _unwrap_quiz_api_success_data(pl_set) if isinstance(pl_set, dict) else {}
    qids_ordered: list[str] = []
    if 200 <= int(st_set) < 300:
        qids_ordered = _ordered_question_ids_from_set_upstream_root(root_set)

    collected_in_order: list[dict[str, Any]] = []

    if qids_ordered:
        wanted = set(qids_ordered)
        by_id: dict[str, dict[str, Any]] = {}
        scan_off = 0
        while len(by_id) < len(wanted) and scan_off < max_scan:
            items_page, st_b, pl_err = _pull_bank_page(scan_off)
            if st_b != 200:
                return jsonify(pl_err), st_b
            for it in items_page:
                bid = _question_id_from_bank_item(it)
                if bid in wanted and bid not in by_id:
                    by_id[bid] = it
            if len(items_page) < scan_lim:
                break
            scan_off += scan_lim
        collected_in_order = [by_id[q] for q in qids_ordered if q in by_id]
    else:
        scan_off = 0
        while scan_off < max_scan:
            items_page, st_b, pl_err = _pull_bank_page(scan_off)
            if st_b != 200:
                if scan_off == 0:
                    return jsonify(pl_err), st_b
                break
            for it in items_page:
                if _bank_row_matches_set_id(it, set_sid):
                    collected_in_order.append(it)
            if len(items_page) < scan_lim:
                break
            scan_off += scan_lim

    total = len(collected_in_order)
    sliced = collected_in_order[user_off : user_off + user_lim]
    note = (
        "aiword：已按套题合并过滤。"
        + ("题序来自 GET quiz/sets。" if qids_ordered else "套题明细无题序，已按题目上的 set_id 类字段扫描题库。")
    )
    out_payload: dict[str, Any] = {
        "code": 0,
        "message": "ok",
        "data": {
            "items": sliced,
            "total": total,
            "limit": user_lim,
            "offset": user_off,
            "meta": {
                "aiword_bank_set_filter": True,
                "aiword_set_id": set_sid,
                "aiword_note": note,
                "aiword_set_detail_http": int(st_set),
                "aiword_set_ordered_ids_count": len(qids_ordered),
            },
        },
        "trace_id": trace_id,
        "request": {"url": "", "method": "GET", "upstreamPath": "quiz/bank/questions + aiword set merge"},
    }
    return jsonify(out_payload), 200


@bp.patch("/api/exam-center/teacher/bank/questions/<question_id>")
@page13_access_required
def api_exam_teacher_bank_question_update(question_id: str):
    qid = (question_id or "").strip()
    if not qid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 question_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    status, payload = _quiz_api_call(
        f"quiz/bank/questions/{qid}",
        method="PATCH",
        payload=_json_payload(),
        query=request.args.to_dict(),
    )
    return jsonify(payload), status


@bp.delete("/api/exam-center/teacher/bank/questions/<question_id>")
@page13_access_required
def api_exam_teacher_bank_question_delete(question_id: str):
    qid = (question_id or "").strip()
    if not qid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 question_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    status, payload = _quiz_api_call(
        f"quiz/bank/questions/{qid}",
        method="DELETE",
        payload=_json_payload(),
        query=request.args.to_dict(),
    )
    return jsonify(payload), status


def _exam_stats_options_local() -> dict[str, Any]:
    from .exam_display_labels import exam_user_filter_options, human_assignment_label, normalize_user_key

    org_id = _current_exam_scope_organization_id()
    student_keys: set[str] = set()
    try:
        q_students = _scope_exam_activity_query(
            db.session.query(ExamCenterActivity.user_id)
            .filter(ExamCenterActivity.user_id.isnot(None))
            .filter(ExamCenterActivity.user_id != "")
        )
        for row in q_students.distinct().all():
            uid = row[0] if isinstance(row, (tuple, list)) else row
            nk = normalize_user_key(str(uid or ""))
            if nk:
                student_keys.add(nk)
    except Exception:
        student_keys = set()
    students = exam_user_filter_options(student_keys)
    assignments_map: dict[str, dict[str, Any]] = {}
    try:
        q_assign = _scope_exam_activity_query(
            db.session.query(ExamCenterActivity.assignment_id, ExamCenterActivity.assignment_label)
            .filter(ExamCenterActivity.assignment_id.isnot(None))
            .filter(ExamCenterActivity.assignment_id != "")
        )
        rows = q_assign.distinct().limit(500).all()
        for aid, alab in rows:
            k = str(aid or "").strip()
            if not k:
                continue
            lab = human_assignment_label(k, activity_label=str(alab or "").strip())
            assignments_map[k] = {"id": k, "name": lab, "label": lab, "set_id": ""}
        q_local_assign = _scope_exam_assignment_query(
            ExamCenterAssignment.query.filter(
                or_(
                    ExamCenterAssignment.status.is_(None),
                    ~ExamCenterAssignment.status.in_(("inactive", "cancelled", "archived", "deleted")),
                )
            )
        )
        local_rows = q_local_assign.limit(500).all()
        for r in local_rows:
            k = str(r.assignment_id or "").strip()
            if not k:
                continue
            lab = human_assignment_label(
                k,
                title=str(getattr(r, "title", None) or "").strip(),
                activity_label=assignments_map.get(k, {}).get("name"),
            )
            sid = str(getattr(r, "set_id", None) or "").strip()
            assignments_map[k] = {"id": k, "name": lab, "label": lab, "set_id": sid}
        all_aids = [x for x in assignments_map.keys() if x]
        if all_aids:
            q_rows = _scope_exam_assignment_query(
                ExamCenterAssignment.query.filter(ExamCenterAssignment.assignment_id.in_(all_aids))
            )
            for ar in q_rows.all():
                kk = str(ar.assignment_id or "").strip()
                if kk in assignments_map:
                    s2 = str(getattr(ar, "set_id", None) or "").strip()
                    if s2:
                        assignments_map[kk]["set_id"] = s2
    except Exception:
        assignments_map = {}
    assignments = sorted(assignments_map.values(), key=lambda x: str(x.get("name") or ""))
    return {"students": students, "assignments": assignments}


def _local_focus_flags_for_student(practice_count: int, wrong_total: int, exam_started: int, exam_submitted: int) -> list[str]:
    flags: list[str] = []
    if practice_count < 3:
        flags.append("练习次数少")
    if wrong_total >= 10:
        flags.append("错题多")
    if exam_started > exam_submitted:
        flags.append("未完成考试")
    return flags


def _local_stats_for_student(user_id: str) -> dict[str, Any]:
    uid = str(user_id or "").strip()
    if not uid:
        return {}
    org_id = _current_exam_scope_organization_id()
    q_rows = _scope_exam_activity_query(ExamCenterActivity.query.filter_by(user_id=uid))
    rows = q_rows.order_by(ExamCenterActivity.created_at.desc()).all()
    if not rows:
        return {
            "student_id": uid,
            "student_name": uid,
            "practice_count": 0,
            "exam_submitted_count": 0,
            "exam_started_count": 0,
            "wrong_total": 0,
            "pass_count": 0,
            "fail_count": 0,
            "pass_rate_percent": None,
            "focus_flags": ["练习次数少"],
        }
    ids = [str(r.id) for r in rows if getattr(r, "id", None)]
    det_map: dict[str, ExamCenterActivityDetail] = {}
    if ids:
        ds = ExamCenterActivityDetail.query.filter(ExamCenterActivityDetail.activity_id.in_(ids)).all()
        det_map = {str(d.activity_id): d for d in ds}
    pass_score = _exam_pass_score()
    practice_count = 0
    exam_submitted = 0
    exam_started = 0
    wrong_total = 0
    pass_count = 0
    fail_count = 0
    student_name = ""
    for r in rows:
        if not student_name:
            student_name = str(r.display_name or r.username or r.user_id or "").strip()
        m = str(r.mode or "").strip().lower()
        if m == "practice":
            practice_count += 1
        if m == "exam":
            exam_submitted += 1
        if r.assignment_id:
            exam_started += 1
        d = det_map.get(str(r.id))
        if d and d.wrong_count is not None:
            wrong_total += int(d.wrong_count)
        if d and d.score is not None and m == "exam":
            if float(d.score) >= pass_score:
                pass_count += 1
            else:
                fail_count += 1
    if exam_started < exam_submitted:
        exam_started = exam_submitted
    graded = pass_count + fail_count
    pass_rate = round((pass_count * 100.0 / graded), 2) if graded > 0 else None
    return {
        "student_id": uid,
        "student_name": student_name or uid,
        "practice_count": practice_count,
        "exam_submitted_count": exam_submitted,
        "exam_started_count": exam_started,
        "wrong_total": wrong_total,
        "pass_score": pass_score,
        "pass_count": pass_count,
        "fail_count": fail_count,
        "pass_rate_percent": pass_rate,
        "focus_flags": _local_focus_flags_for_student(practice_count, wrong_total, exam_started, exam_submitted),
    }


def _local_stats_all_students_rows() -> list[dict[str, Any]]:
    """按学生在本地聚合多行指标（与会话下拉中的学生列表同源）。"""
    opts = _exam_stats_options_local()
    studs = opts.get("students") if isinstance(opts.get("students"), list) else []
    rows: list[dict[str, Any]] = []
    for s in studs:
        if not isinstance(s, dict):
            continue
        sid = str(s.get("id") or "").strip()
        if not sid:
            continue
        row = _local_stats_for_student(sid)
        if not row:
            continue
        pc = int(row.get("pass_count") or 0)
        fc = int(row.get("fail_count") or 0)
        row["graded_exam_count"] = pc + fc
        row["total_learning_count"] = int(row.get("practice_count") or 0) + int(row.get("exam_submitted_count") or 0)
        rows.append(row)
    rows.sort(key=lambda x: str(x.get("student_name") or x.get("student_id") or ""))
    return rows


def _local_stats_rows_student_by_mode() -> list[dict[str, Any]]:
    org_id = _current_exam_scope_organization_id()
    """学生 ×（考试/练习）交叉分组：记录数、已判分、通过/不通过/通过率；判分口径与 /stats/mode 一致（有 score 即计分）。"""
    opts = _exam_stats_options_local()
    studs = opts.get("students") if isinstance(opts.get("students"), list) else []
    pass_score = _exam_pass_score()
    out: list[dict[str, Any]] = []
    for s in studs:
        if not isinstance(s, dict):
            continue
        sid = str(s.get("id") or "").strip()
        if not sid:
            continue
        label = str(s.get("label") or s.get("name") or sid).strip() or sid
        rows_q = _scope_exam_activity_query(ExamCenterActivity.query.filter_by(user_id=sid))
        rows_u = rows_q.order_by(ExamCenterActivity.created_at.desc()).all()
        ids = [str(r.id) for r in rows_u if getattr(r, "id", None)]
        det_map: dict[str, ExamCenterActivityDetail] = {}
        if ids:
            ds = ExamCenterActivityDetail.query.filter(ExamCenterActivityDetail.activity_id.in_(ids)).all()
            det_map = {str(d.activity_id): d for d in ds}
        for mode_key, mode_label in (("exam", "考试"), ("practice", "练习")):
            subset = [r for r in rows_u if str(r.mode or "").strip().lower() == mode_key]
            total = len(subset)
            graded = 0
            pc = 0
            fc = 0
            for r in subset:
                d = det_map.get(str(r.id))
                if not d or d.score is None:
                    continue
                graded += 1
                if float(d.score) >= pass_score:
                    pc += 1
                else:
                    fc += 1
            pr = round((pc * 100.0 / graded), 2) if graded > 0 else None
            out.append(
                {
                    "student_id": sid,
                    "student_name": label,
                    "mode": mode_key,
                    "mode_label": mode_label,
                    "total_count": total,
                    "graded_count": graded,
                    "pass_score": pass_score,
                    "pass_count": pc,
                    "fail_count": fc,
                    "pass_rate_percent": pr,
                }
            )
    out.sort(key=lambda x: (str(x.get("student_name") or ""), 0 if x.get("mode") == "exam" else 1))
    return out


def _local_stats_overview_all() -> dict[str, Any]:
    """全库聚合（不按条数截断），与 /stats/mode、按学生聚合口径一致。"""
    pass_score = _exam_pass_score()
    q_practice = _scope_exam_activity_query(ExamCenterActivity.query.filter_by(mode="practice"))
    q_exam = _scope_exam_activity_query(ExamCenterActivity.query.filter_by(mode="exam"))
    practice_count = q_practice.count()
    exam_act = q_exam.all()
    exam_count = len(exam_act)
    ids_exam = [str(r.id) for r in exam_act if getattr(r, "id", None)]
    det_map: dict[str, ExamCenterActivityDetail] = {}
    if ids_exam:
        ds = ExamCenterActivityDetail.query.filter(ExamCenterActivityDetail.activity_id.in_(ids_exam)).all()
        det_map = {str(d.activity_id): d for d in ds}
    pass_count = 0
    fail_count = 0
    for r in exam_act:
        d = det_map.get(str(r.id))
        if not d or d.score is None:
            continue
        if float(d.score) >= pass_score:
            pass_count += 1
        else:
            fail_count += 1
    graded_exam_count = pass_count + fail_count
    pass_rate = round((pass_count * 100.0 / graded_exam_count), 2) if graded_exam_count > 0 else None
    try:
        uq = _scope_exam_activity_query(
            db.session.query(ExamCenterActivity.user_id)
            .filter(ExamCenterActivity.user_id.isnot(None))
            .filter(ExamCenterActivity.user_id != "")
        )
        urows = uq.distinct().all()
        uid_list = [str(u[0]).strip() for u in urows if u and u[0]]
    except Exception:
        uid_list = []
    by_student: dict[str, dict[str, Any]] = {}
    for uid in uid_list:
        by_student[uid] = _local_stats_for_student(uid)
    focus_students = [v for v in by_student.values() if v.get("focus_flags")]
    return {
        "practice_count": practice_count,
        "exam_count": exam_count,
        "graded_exam_count": graded_exam_count,
        "pass_count": pass_count,
        "fail_count": fail_count,
        "pass_score": pass_score,
        "pass_rate_percent": pass_rate,
        "students_total": len(by_student),
        "focus_students": focus_students,
    }


def _exam_stats_recent_activity_local(limit: int) -> list[dict[str, Any]]:
    lim = max(1, min(200, int(limit or 80)))
    out: list[dict[str, Any]] = []
    pass_score = _exam_pass_score()
    org_id = _current_exam_scope_organization_id()
    try:
        rq = _scope_exam_activity_query(ExamCenterActivity.query)
        rows = rq.order_by(ExamCenterActivity.created_at.desc()).limit(lim).all()
        ids = [str(r.id) for r in rows if getattr(r, "id", None)]
        det_map: dict[str, ExamCenterActivityDetail] = {}
        if ids:
            ds = ExamCenterActivityDetail.query.filter(ExamCenterActivityDetail.activity_id.in_(ids)).all()
            det_map = {str(d.activity_id): d for d in ds}
        exam_attempt_ids = [
            str(x.attempt_id).strip()
            for x in rows
            if str(x.mode or "").strip().lower() == "exam" and str(getattr(x, "attempt_id", None) or "").strip()
        ]
        att_by_tid: dict[str, ExamAttempt] = {}
        gj_by_tid: dict[str, ExamGradingJob] = {}
        pending_subj: set[str] = set()
        if exam_attempt_ids:
            att_by_tid = {str(x.attempt_id): x for x in ExamAttempt.query.filter(ExamAttempt.attempt_id.in_(exam_attempt_ids)).all()}
            gj_by_tid = {str(x.attempt_id): x for x in ExamGradingJob.query.filter(ExamGradingJob.attempt_id.in_(exam_attempt_ids)).all()}
            pending_subj = _local_exam_attempt_ids_with_pending_subjective(exam_attempt_ids)
        for a in rows:
            d = det_map.get(str(a.id))
            mode = str(a.mode or "").strip().lower()
            mode_label = "练习" if mode == "practice" else ("考试" if mode == "exam" else mode or "-")
            from .exam_display_labels import human_user_label

            who = human_user_label(
                str(getattr(a, "user_id", "") or ""),
                activity_display=str(getattr(a, "display_name", None) or "").strip() or None,
                activity_username=str(getattr(a, "username", None) or "").strip() or None,
            )
            tgt = (a.assignment_label or a.set_id or a.attempt_id or "-") or "-"
            res = (a.result_summary or "").strip() or "-"
            tid = str(getattr(a, "attempt_id", None) or "").strip()
            att_row = att_by_tid.get(tid) if tid else None
            gj_row = gj_by_tid.get(tid) if tid else None
            regrade_ok = False
            if mode == "exam" and tid and att_row and tid in pending_subj:
                regrade_ok, _ = _local_exam_admin_regrade_allowed(att_row, gj_row)
            rec: dict[str, Any] = {
                "id": a.id,
                "created_at": a.created_at.isoformat() if a.created_at else "",
                "user_id": a.user_id,
                "student_name": who,
                "mode": a.mode,
                "mode_label": mode_label,
                "assignment_id": str(a.assignment_id).strip() if getattr(a, "assignment_id", None) else "",
                "set_id": str(getattr(a, "set_id", None) or "").strip() or None,
                "setId": str(getattr(a, "set_id", None) or "").strip() or None,
                "attempt_id": tid or None,
                "target_label": str(tgt),
                "result": res[:500],
                "score": d.score if d else None,
                "total_score": d.total_score if d else None,
                "correct_count": d.correct_count if d else None,
                "wrong_count": d.wrong_count if d else None,
                "pass_score": pass_score,
                "passed": (d.score is not None and float(d.score) >= pass_score) if d and d.score is not None else None,
                "weakness": d.weakness if d else None,
                "recommendation": d.recommendation if d else None,
                "local_exam_regrade_eligible": bool(regrade_ok),
                "grading_job_status": (gj_row.status if gj_row else None),
            }
            out.append(rec)
    except Exception:
        return []
    return out


@bp.get("/api/exam-center/stats/options")
@page13_access_required
def api_exam_stats_options():
    local = _exam_stats_options_local()
    return jsonify(
        {
            "code": 0,
            "message": "ok（本地统计）",
            "data": local,
            "trace_id": uuid.uuid4().hex,
            "request": {"url": "", "method": "GET", "upstreamPath": "local aggregate"},
        }
    ), 200


@bp.get("/api/exam-center/stats/recent-activity")
@page13_access_required
def api_exam_stats_recent_activity():
    try:
        lim = int(str(request.args.get("limit") or "80"))
    except (TypeError, ValueError):
        lim = 80
    recs = _exam_stats_recent_activity_local(lim)
    # 本地筛选：仅影响“记录列表”，不影响看板
    want_uid = str(request.args.get("student_id") or request.args.get("user_id") or "").strip()
    want_aid = str(request.args.get("assignment_id") or "").strip()
    if want_uid:
        from .exam_display_labels import exam_activity_user_id_match_keys

        match_keys = exam_activity_user_id_match_keys(want_uid)
        if match_keys:
            recs = [x for x in recs if str((x or {}).get("user_id") or "").strip() in match_keys]
        else:
            recs = []
    if want_aid:
        recs = [x for x in recs if str((x or {}).get("assignment_id") or "").strip() == want_aid]
    return jsonify(
        {
            "code": 0,
            "message": "ok（本地统计）",
            "data": {"records": recs},
            "trace_id": uuid.uuid4().hex,
            "request": {"url": "", "method": "GET", "upstreamPath": "local aggregate"},
        }
    ), 200


@bp.post("/api/exam-center/teacher/sets/review-by-ai")
@page13_access_required
def api_exam_teacher_review_set():
    data = _json_payload()
    set_id = (data.pop("set_id", "") or "").strip()
    if not set_id:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 set_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    org_id = _resolve_exam_organization_id(
        explicit_org_id=str(data.get("organization_id") or data.get("organizationId") or "").strip(),
        project_id=str(data.get("project_id") or data.get("projectId") or "").strip(),
    )
    status, payload = _quiz_api_call(
        f"quiz/sets/{set_id}/review-by-ai",
        method="POST",
        payload=data,
        organization_id=org_id,
    )
    # 异步复审：上游立即返回 job_id；本地落库便于轮询与任务列表（与录题 ingest 一致）
    if 200 <= int(status) < 300 and isinstance(payload, dict):
        upstream_job_id = _guess_job_id_from_payload(payload)
        if upstream_job_id:
            created_by = (session.get("display_name") or session.get("username") or "").strip() or None
            row_q = ExamSetReviewJob.query.filter_by(upstream_job_id=upstream_job_id)
            if org_id:
                row_q = row_q.filter(ExamSetReviewJob.organization_id == org_id)
            row = row_q.first()
            if not row:
                row = ExamSetReviewJob(
                    organization_id=org_id or None,
                    upstream_job_id=upstream_job_id,
                    set_id=set_id,
                    status="pending",
                    created_by=created_by,
                )
            row.organization_id = org_id or row.organization_id
            row.last_upstream_http_status = int(status)
            row.last_upstream_data = payload.get("data") if isinstance(payload.get("data"), dict) else payload.get("data")
            row.last_upstream_request_url = ((payload.get("request") or {}).get("url") if isinstance(payload.get("request"), dict) else None) or None
            row.last_upstream_trace_id = (payload.get("trace_id") or None)
            row.last_message = (payload.get("message") or None)
            row.status = _normalize_job_status_from_upstream_data(payload.get("data")) or "running"
            db.session.add(row)
            try:
                db.session.commit()
            except Exception as e:
                db.session.rollback()
                if isinstance(payload, dict):
                    payload["job_record_error"] = f"aiword 落库失败（不影响上游任务）：{e}"
            else:
                if isinstance(payload, dict):
                    payload["job_record"] = {
                        "id": row.id,
                        "upstream_job_id": row.upstream_job_id,
                        "set_id": row.set_id,
                        "status": row.status,
                    }
    return jsonify(payload), status


@bp.get("/api/exam-center/teacher/sets/review-jobs/<job_id>")
@page13_access_required
def api_exam_teacher_review_job(job_id: str):
    job_id = (job_id or "").strip()
    if not job_id:
        return (
            jsonify(
                {
                    "code": "BAD_REQUEST",
                    "message": "缺少 job_id",
                    "data": None,
                    "trace_id": uuid.uuid4().hex,
                    "request": {"url": "", "method": "GET", "upstreamPath": ""},
                }
            ),
            400,
        )
    row = ExamSetReviewJob.query.filter_by(upstream_job_id=job_id).first()
    scope_org_id = _current_exam_scope_organization_id()
    if row is not None and scope_org_id and str(getattr(row, "organization_id", "") or "").strip() != scope_org_id:
        return jsonify({"code": "NOT_FOUND", "message": "本地无该任务记录", "data": None, "trace_id": uuid.uuid4().hex}), 404
    refresh = (request.args.get("refresh") or "1").strip()
    if refresh not in {"0", "false", "no"}:
        req_org_id = (
            str(getattr(row, "organization_id", "") or "").strip()
            or scope_org_id
            or _resolve_exam_organization_id()
        )
        status, payload = _quiz_api_call(
            f"quiz/sets/review-jobs/{job_id}",
            method="GET",
            query={k: v for k, v in request.args.to_dict().items() if k != "refresh"},
            organization_id=req_org_id,
        )
        if row is None:
            row = ExamSetReviewJob(
                upstream_job_id=job_id,
                status="unknown",
                organization_id=req_org_id or None,
            )
        row.organization_id = req_org_id or row.organization_id
        row.last_upstream_http_status = int(status)
        if isinstance(payload, dict):
            row.last_message = (payload.get("message") or None)
            row.last_upstream_trace_id = (payload.get("trace_id") or None)
            req_meta = payload.get("request") if isinstance(payload.get("request"), dict) else {}
            row.last_upstream_request_url = (req_meta.get("url") or None)
            row.last_upstream_data = payload.get("data")
            row.status = _normalize_job_status_from_upstream_data(payload.get("data"))
        db.session.add(row)
        try:
            db.session.commit()
        except Exception as e:
            db.session.rollback()
            if isinstance(payload, dict):
                payload["job_record_error"] = f"aiword 落库失败（不影响上游查询）：{e}"
        else:
            if isinstance(payload, dict):
                payload["job_record"] = {
                    "id": row.id,
                    "upstream_job_id": row.upstream_job_id,
                    "set_id": row.set_id,
                    "status": row.status,
                }
        return jsonify(payload), status

    if row is None:
        return jsonify({"code": "NOT_FOUND", "message": "本地无该任务记录", "data": None, "trace_id": uuid.uuid4().hex}), 404
    return jsonify(
        {
            "code": 0,
            "message": "ok",
            "data": {
                "job_record": {
                    "id": row.id,
                    "upstream_job_id": row.upstream_job_id,
                    "set_id": row.set_id,
                    "status": row.status,
                    "created_by": row.created_by,
                    "created_at": row.created_at.isoformat() if row.created_at else None,
                    "updated_at": row.updated_at.isoformat() if row.updated_at else None,
                    "last_message": row.last_message,
                    "last_upstream_http_status": row.last_upstream_http_status,
                    "last_upstream_request_url": row.last_upstream_request_url,
                    "last_upstream_trace_id": row.last_upstream_trace_id,
                    "last_upstream_data": row.last_upstream_data,
                }
            },
            "trace_id": uuid.uuid4().hex,
        }
    )


@bp.get("/api/exam-center/teacher/sets/review-jobs")
@page13_access_required
def api_exam_teacher_review_jobs_list():
    limit_raw = (request.args.get("limit") or "20").strip()
    try:
        limit = int(limit_raw)
    except ValueError:
        limit = 20
    if limit < 1:
        limit = 1
    if limit > 200:
        limit = 200

    scope_org_id = _current_exam_scope_organization_id()
    q_rows = ExamSetReviewJob.query
    if scope_org_id:
        q_rows = q_rows.filter(ExamSetReviewJob.organization_id == scope_org_id)
    rows = q_rows.order_by(ExamSetReviewJob.created_at.desc()).limit(limit).all()
    return jsonify(
        {
            "code": 0,
            "message": "ok",
            "data": {
                "jobs": [
                    {
                        "id": r.id,
                        "upstream_job_id": r.upstream_job_id,
                        "set_id": r.set_id,
                        "status": r.status,
                        "created_by": r.created_by,
                        "created_at": r.created_at.isoformat() if r.created_at else None,
                        "updated_at": r.updated_at.isoformat() if r.updated_at else None,
                        "last_message": r.last_message,
                        "last_upstream_http_status": r.last_upstream_http_status,
                        "last_upstream_request_url": r.last_upstream_request_url,
                        "last_upstream_trace_id": r.last_upstream_trace_id,
                    }
                    for r in rows
                ]
            },
            "trace_id": uuid.uuid4().hex,
        }
    )


@bp.post("/api/exam-center/teacher/sets/publish")
@page13_access_required
def api_exam_teacher_publish_set():
    data = _json_payload()
    set_id = (data.pop("set_id", "") or "").strip()
    if not set_id:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 set_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    status, payload = _quiz_api_call(
        f"quiz/sets/{set_id}/publish",
        method="POST",
        payload=data,
    )
    return jsonify(payload), status


@bp.post("/api/exam-center/teacher/papers")
@page13_access_required
def api_exam_teacher_create_paper():
    status, payload = _quiz_api_call("quiz/papers", method="POST", payload=_json_payload())
    return jsonify(payload), status


@bp.post("/api/exam-center/teacher/assignments")
@page13_access_required
def api_exam_teacher_create_assignment():
    req = _json_payload()
    from .exam_scope import organization_id_for_exam_write

    org_id = organization_id_for_exam_write(
        str(req.get("organization_id") or req.get("organizationId") or "").strip()
    )
    due_val, due_update = _parse_assignment_due_from_request(req)
    forward = dict(req) if isinstance(req, dict) else {}
    if due_val and "due_at" not in forward and "dueAt" not in forward:
        forward["due_at"] = due_val.isoformat(sep=" ", timespec="seconds")
    # 上游（aicheckword）在部分版本中不提供 quiz/assignments 下发接口；
    # 404/405 时回退为“仅本地下发”（学生端列表已有上游失败→本地镜像兜底）。
    status, payload, tried = _quiz_try_paths(
        [
            "quiz/assignments",  # 某些上游实现
            "quiz/teacher/assignments",
            "quiz/exam-center/assignments",
        ],
        method="POST",
        payload=forward,
        query=None,
        organization_id=org_id,
    )
    if isinstance(payload, dict) and tried:
        payload.setdefault("request", {})
        if isinstance(payload.get("request"), dict):
            payload["request"]["upstreamTried"] = tried

    if int(status) in (404, 405) or (
        isinstance(payload, dict)
        and payload.get("code") in ("QUIZ_API_UPSTREAM_ERROR", "QUIZ_API_NOT_CONFIGURED")
        and int(status) in (404, 405, 503)
    ):
        # 本地下发：生成 assignment_id 并写入镜像表
        aid_local = "loc-" + uuid.uuid4().hex[:24]
        title_local = str(req.get("title") or req.get("name") or req.get("label") or "").strip() or None
        set_id_local = str(req.get("set_id") or req.get("setId") or "").strip() or None
        exam_track_local = str(req.get("exam_track") or req.get("examTrack") or "").strip() or None
        exam_cat_local = _normalize_exam_category(req.get("exam_category") or req.get("examCategory") or "daily")
        diff_req = str(req.get("difficulty") or req.get("difficultyLevel") or "").strip().lower()
        if diff_req not in ("easy", "medium", "hard"):
            diff_req = ""
        row = ExamCenterAssignment(assignment_id=aid_local)
        row.organization_id = org_id or None
        row.title = title_local
        row.set_id = set_id_local
        row.exam_track = exam_track_local
        row.exam_category = exam_cat_local
        row.difficulty = diff_req or None
        row.status = "published"
        row.created_by = (session.get("display_name") or session.get("username") or "").strip() or None
        if due_update:
            row.due_at = due_val
        db.session.add(row)
        try:
            db.session.commit()
            return (
                jsonify(
                    {
                        "code": 0,
                        "message": "ok（上游不支持下发接口，已改为仅本地下发）",
                        "data": {
                            "assignment_id": row.assignment_id,
                            "title": row.title,
                            "set_id": row.set_id,
                            "exam_track": row.exam_track,
                            "exam_category": getattr(row, "exam_category", None) or "daily",
                            "difficulty": row.difficulty,
                            "status": row.status,
                            "due_at": row.due_at.isoformat(timespec="seconds") if getattr(row, "due_at", None) else None,
                            "aiword_local_only": True,
                        },
                        "trace_id": uuid.uuid4().hex,
                    }
                ),
                200,
            )
        except Exception as e:
            db.session.rollback()
            return jsonify({"code": "DB_ERROR", "message": str(e), "data": None, "trace_id": uuid.uuid4().hex}), 500

    if 200 <= int(status) < 300 and isinstance(payload, dict):
        aid, title, set_id = _extract_assignment_from_quiz_payload(payload)
        if aid:
            row = ExamCenterAssignment.query.filter_by(assignment_id=aid).first()
            if not row:
                row = ExamCenterAssignment(assignment_id=aid)
            row.organization_id = org_id or row.organization_id
            row.title = title or row.title
            row.set_id = set_id or row.set_id or str(req.get("set_id") or req.get("setId") or "").strip() or None
            row.exam_track = str(req.get("exam_track") or req.get("examTrack") or "").strip() or row.exam_track
            row.exam_category = _normalize_exam_category(
                req.get("exam_category") or req.get("examCategory") or getattr(row, "exam_category", None) or "daily"
            )
            diff_req = str(req.get("difficulty") or req.get("difficultyLevel") or "").strip().lower()
            if diff_req in ("easy", "medium", "hard"):
                row.difficulty = diff_req
            if due_update:
                row.due_at = due_val
            sid = str(row.set_id or "").strip()
            if sid and (not row.difficulty or row.difficulty not in ("easy", "medium", "hard")):
                st_set, pl_set = _quiz_api_call(
                    f"quiz/sets/{_urlquote(sid, safe='')}",
                    method="GET",
                    query=None,
                    organization_id=(org_id or row.organization_id or ""),
                )
                if 200 <= int(st_set) < 300 and isinstance(pl_set, dict):
                    d2 = _difficulty_from_quiz_get_set_payload(pl_set)
                    if d2:
                        row.difficulty = d2
            row.status = "published"
            row.created_by = (session.get("display_name") or session.get("username") or "").strip() or row.created_by
            db.session.add(row)
            try:
                db.session.commit()
            except Exception:
                db.session.rollback()
            else:
                payload.setdefault("aiword", {})
                if isinstance(payload.get("aiword"), dict):
                    payload["aiword"]["assignment"] = {
                        "assignment_id": row.assignment_id,
                        "title": row.title,
                        "set_id": row.set_id,
                        "exam_track": row.exam_track,
                        "exam_category": getattr(row, "exam_category", None) or "daily",
                    }
    return jsonify(payload), status


@bp.post("/api/exam-center/teacher/assignments/issue")
@page13_access_required
def api_exam_teacher_issue_assignments_modal():
    """老师端弹窗：本地下发考试任务（含受众/截止/目的），学生端仅受众可见。"""
    req = _json_payload()
    from .exam_scope import organization_id_for_exam_write

    org_id = organization_id_for_exam_write(
        str(req.get("organization_id") or req.get("organizationId") or "").strip()
    )
    due_val, due_update = _parse_assignment_due_from_request(req)
    purpose = str(req.get("purpose") or req.get("exam_purpose") or "").strip() or None
    exam_track = str(req.get("exam_track") or req.get("examTrack") or "").strip() or None
    exam_category = _normalize_exam_category(req.get("exam_category") or req.get("examCategory") or "daily")

    aud = req.get("audience_user_ids") or req.get("audienceUserIds") or req.get("user_ids") or []
    audience_ids = [str(x).strip() for x in aud] if isinstance(aud, list) else []
    audience_ids = [x for x in audience_ids if x]

    items = req.get("items")
    set_ids: list[str] = []
    title_by_set: dict[str, str] = {}
    if isinstance(items, list) and items:
        for it in items:
            if not isinstance(it, dict):
                continue
            sid = str(it.get("set_id") or it.get("setId") or "").strip()
            if not sid:
                continue
            set_ids.append(sid)
            ttl = str(it.get("title") or it.get("name") or "").strip()
            if ttl:
                title_by_set[sid] = ttl
    else:
        sid0 = str(req.get("set_id") or req.get("setId") or "").strip()
        if sid0:
            set_ids = [sid0]
            ttl0 = str(req.get("title") or req.get("name") or "").strip()
            if ttl0:
                title_by_set[sid0] = ttl0

    if not set_ids:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 set_id/items", "data": None, "trace_id": uuid.uuid4().hex}), 400
    if not audience_ids:
        return jsonify({"code": "BAD_REQUEST", "message": "请选择考试对象（至少1人）", "data": None, "trace_id": uuid.uuid4().hex}), 400
    aud_err = _validate_exam_audience_user_ids(audience_ids)
    if aud_err:
        return jsonify({"code": "FORBIDDEN", "message": aud_err, "data": None, "trace_id": uuid.uuid4().hex}), 403

    created: list[dict[str, Any]] = []
    try:
        for sid in set_ids:
            aid_local = "loc-" + uuid.uuid4().hex[:24]
            title = (title_by_set.get(sid) or "").strip()
            if title:
                nm = title + " 考试任务"
            else:
                nm = f"{sid} 考试任务"
            row = ExamCenterAssignment(assignment_id=aid_local)
            row.organization_id = org_id or None
            row.title = nm
            row.set_id = sid
            row.exam_track = exam_track
            row.exam_category = exam_category
            row.status = "published"
            row.created_by = (session.get("display_name") or session.get("username") or "").strip() or None
            if due_update:
                row.due_at = due_val
            db.session.add(row)
            db.session.flush()

            if purpose:
                ex = ExamCenterAssignmentExtra.query.filter_by(assignment_id=aid_local).first()
                if not ex:
                    ex = ExamCenterAssignmentExtra(assignment_id=aid_local)
                ex.purpose = purpose
                db.session.add(ex)

            for uid in audience_ids:
                db.session.add(ExamCenterAssignmentAudience(assignment_id=aid_local, user_id=uid))

            created.append(
                {
                    "assignment_id": aid_local,
                    "title": nm,
                    "set_id": sid,
                    "exam_track": exam_track,
                    "exam_category": exam_category,
                    "due_at": due_val.isoformat(timespec="seconds") if (due_update and due_val) else None,
                    "purpose": purpose,
                    "audience_user_ids": audience_ids,
                    "aiword_local_only": True,
                }
            )
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({"code": "DB_ERROR", "message": str(e), "data": None, "trace_id": uuid.uuid4().hex}), 500

    return jsonify({"code": 0, "message": "ok", "data": {"created": created}, "trace_id": uuid.uuid4().hex}), 200


def _exam_try_upstream_modify_assignment_proxy(assignment_id: str, op: str) -> tuple[list[dict[str, Any]], int, dict[str, Any]]:
    """对已下发任务调用上游可变路径；不要求全部成功以便本地仍可收口。"""
    aid = (assignment_id or "").strip()
    org_id = _resolve_exam_organization_id(assignment_id=aid)
    q = _urlquote(aid, safe="")
    attempts: list[dict[str, Any]] = []
    last_status = 599
    last_pl: dict[str, Any] = {}
    ops = []
    op_l = op.lower().strip()
    if op_l == "delete":
        ops = [(f"quiz/assignments/{q}", "DELETE", None)]
    elif op_l == "unpublish":
        ops = [
            # 常见语义：设为无效/归档
            (
                f"quiz/assignments/{q}",
                "PATCH",
                {"status": "inactive", "is_active": False, "cancelled": True},
            ),
            (f"quiz/assignments/{q}/cancel", "POST", {}),
        ]
    elif op_l == "publish":
        ops = [
            (
                f"quiz/assignments/{q}",
                "PATCH",
                {"status": "published", "is_active": True, "cancelled": False},
            ),
        ]
    for path, meth, bod in ops:
        st0, pl0 = _quiz_api_call(
            path.strip("/"),
            method=meth,
            payload=bod if isinstance(bod, dict) else None,
            organization_id=org_id,
        )
        attempts.append({"path": path, "method": meth, "http_status": int(st0)})
        last_status = int(st0)
        last_pl = pl0 if isinstance(pl0, dict) else {}
        if 200 <= int(st0) < 300:
            break
        if int(st0) == 204:
            break
    return attempts, last_status, last_pl


@bp.get("/api/exam-center/teacher/assignable-users")
@page13_access_required
def api_exam_teacher_assignable_users():
    """老师端下发考试：可选人员（按当前公司与项目组作用域过滤）。"""
    from .exam_scope import exam_teacher_assignable_users
    from .user_access import serialize_user_access

    org_id = _current_exam_scope_organization_id()
    users = exam_teacher_assignable_users(org_id=org_id or None)
    return jsonify(
        {
            "users": [
                {
                    "id": u.id,
                    "username": u.username,
                    "displayName": u.display_name,
                    "display_name": u.display_name,
                    **serialize_user_access(u),
                }
                for u in users
            ]
        }
    )


@bp.get("/api/exam-center/teacher/assignments-local")
@page13_access_required
def api_teacher_assignments_local_list():
    """老师端：列出 aiword 本地镜像的考试任务。"""
    try:
        q_rows = _scope_exam_assignment_query(ExamCenterAssignment.query)
        rows = q_rows.order_by(ExamCenterAssignment.created_at.desc()).limit(500).all()
    except Exception as e:
        return jsonify(
            {
                "code": "DB_ERROR",
                "message": f"查询失败：{e}",
                "data": {"rows": []},
                "trace_id": uuid.uuid4().hex,
            }
        ), 500
    out_rows: list[dict[str, Any]] = []
    for r in rows:
        out_rows.append(
            {
                "assignment_id": r.assignment_id,
                "title": (r.title or r.assignment_id or "").strip(),
                "set_id": r.set_id,
                "exam_track": r.exam_track,
                "exam_category": getattr(r, "exam_category", None) or "daily",
                "difficulty": (r.difficulty or "").strip() if getattr(r, "difficulty", None) else "",
                "status": (r.status or "").strip(),
                "due_at": r.due_at.isoformat(timespec="seconds") if getattr(r, "due_at", None) else None,
                "created_at": r.created_at.isoformat() if getattr(r, "created_at", None) else "",
            }
        )
    return jsonify(
        {"code": 0, "message": "ok", "data": {"rows": out_rows}, "trace_id": uuid.uuid4().hex}
    ), 200


@bp.get("/api/exam-center/teacher/assignments/<assignment_id>")
@page13_access_required
def api_teacher_get_assignment(assignment_id: str):
    """老师端：单条考试任务详情（用于编辑已下发任务）。"""
    aid = (assignment_id or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 assignment_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    row = ExamCenterAssignment.query.filter_by(assignment_id=aid).first()
    if not row:
        return jsonify({"code": "NOT_FOUND", "message": "任务不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404
    if not _exam_assignment_in_staff_scope(row):
        return jsonify({"code": "NOT_FOUND", "message": "任务不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404
    purpose = ""
    try:
        ex = ExamCenterAssignmentExtra.query.filter_by(assignment_id=aid).first()
        if ex and ex.purpose:
            purpose = str(ex.purpose).strip()
    except Exception:
        purpose = ""
    audience_ids: list[str] = []
    try:
        for ar in ExamCenterAssignmentAudience.query.filter_by(assignment_id=aid).all():
            uid = str(getattr(ar, "user_id", None) or "").strip()
            if uid:
                audience_ids.append(uid)
    except Exception:
        audience_ids = []
    data = {
        "assignment_id": row.assignment_id,
        "title": (row.title or "").strip(),
        "set_id": (row.set_id or "").strip() if getattr(row, "set_id", None) else "",
        "exam_track": (row.exam_track or "").strip() if getattr(row, "exam_track", None) else "",
        "exam_category": getattr(row, "exam_category", None) or "daily",
        "difficulty": (row.difficulty or "").strip() if getattr(row, "difficulty", None) else "",
        "status": (row.status or "").strip() if getattr(row, "status", None) else "",
        "due_at": row.due_at.isoformat(timespec="seconds") if getattr(row, "due_at", None) else None,
        "purpose": purpose,
        "audience_user_ids": audience_ids,
    }
    return jsonify({"code": 0, "message": "ok", "data": data, "trace_id": uuid.uuid4().hex}), 200


@bp.patch("/api/exam-center/teacher/assignments/<assignment_id>")
@page13_access_required
def api_teacher_patch_assignment(assignment_id: str):
    """更新已下发的本地考试任务镜像（标题/截止/目的/受众/体考维度等）；并尽力同步上游 PATCH（失败不阻断）。"""
    aid = (assignment_id or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 assignment_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    req = _json_payload()
    if not isinstance(req, dict):
        req = {}
    row = ExamCenterAssignment.query.filter_by(assignment_id=aid).first()
    if not row:
        return jsonify({"code": "NOT_FOUND", "message": "任务不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404
    if not _exam_assignment_in_staff_scope(row):
        return jsonify({"code": "NOT_FOUND", "message": "任务不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404

    if "title" in req:
        row.title = str(req.get("title") or "").strip() or None
    elif "name" in req:
        row.title = str(req.get("name") or "").strip() or None

    if "set_id" in req or "setId" in req:
        sid = str((req.get("set_id") if "set_id" in req else req.get("setId")) or "").strip()
        if not sid:
            return jsonify({"code": "BAD_REQUEST", "message": "set_id 不能为空", "data": None, "trace_id": uuid.uuid4().hex}), 400
        row.set_id = sid

    if "exam_track" in req or "examTrack" in req:
        et = str(req.get("exam_track") if "exam_track" in req else req.get("examTrack") or "").strip()
        if et:
            row.exam_track = et
    if "exam_category" in req or "examCategory" in req:
        row.exam_category = _normalize_exam_category(
            req.get("exam_category") if "exam_category" in req else req.get("examCategory")
        )

    if "difficulty" in req or "difficultyLevel" in req:
        diff_req = str(req.get("difficulty") if "difficulty" in req else req.get("difficultyLevel") or "").strip().lower()
        if diff_req in ("easy", "medium", "hard"):
            row.difficulty = diff_req
        elif diff_req == "" and ("difficulty" in req or "difficultyLevel" in req):
            row.difficulty = None

    due_val, due_update = _parse_assignment_due_from_request(req)
    if due_update:
        row.due_at = due_val

    if "purpose" in req:
        pv = str(req.get("purpose") or "").strip()
        ex = ExamCenterAssignmentExtra.query.filter_by(assignment_id=aid).first()
        if pv:
            if not ex:
                ex = ExamCenterAssignmentExtra(assignment_id=aid)
            ex.purpose = pv
            db.session.add(ex)
        else:
            if ex:
                db.session.delete(ex)

    audience_updated = False
    audience_ids: list[str] = []
    if "audience_user_ids" in req or "audienceUserIds" in req:
        raw = req.get("audience_user_ids") if "audience_user_ids" in req else req.get("audienceUserIds")
        audience_ids = [str(x).strip() for x in raw] if isinstance(raw, list) else []
        audience_ids = [x for x in audience_ids if x]
        aud_err = _validate_exam_audience_user_ids(audience_ids)
        if aud_err:
            return jsonify({"code": "FORBIDDEN", "message": aud_err, "data": None, "trace_id": uuid.uuid4().hex}), 403
        ExamCenterAssignmentAudience.query.filter_by(assignment_id=aid).delete()
        for uid in audience_ids:
            db.session.add(ExamCenterAssignmentAudience(assignment_id=aid, user_id=uid))
        audience_updated = True

    try:
        db.session.add(row)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({"code": "DB_ERROR", "message": str(e), "data": None, "trace_id": uuid.uuid4().hex}), 500

    upstream_attempts: list[dict[str, Any]] = []
    last_st = 599
    last_pl: dict[str, Any] = {}
    forward: dict[str, Any] = {}
    if row.title:
        forward["title"] = row.title
    if row.set_id:
        forward["set_id"] = str(row.set_id).strip()
    if row.due_at:
        forward["due_at"] = row.due_at.isoformat(sep=" ", timespec="seconds")
    else:
        forward["due_at"] = None
    if row.exam_track:
        forward["exam_track"] = row.exam_track
    if getattr(row, "exam_category", None):
        forward["exam_category"] = row.exam_category
    if row.difficulty:
        forward["difficulty"] = row.difficulty
    if audience_updated:
        forward["audience_user_ids"] = audience_ids

    q = _urlquote(aid, safe="")
    for path in (f"quiz/assignments/{q}", f"quiz/teacher/assignments/{q}", f"quiz/exam-center/assignments/{q}"):
        st0, pl0 = _quiz_api_call(
            path.strip("/"),
            method="PATCH",
            payload=forward,
            organization_id=str(getattr(row, "organization_id", "") or "").strip(),
        )
        upstream_attempts.append({"path": path, "method": "PATCH", "http_status": int(st0)})
        last_st = int(st0)
        last_pl = pl0 if isinstance(pl0, dict) else {}
        if 200 <= int(st0) < 300 or int(st0) == 204:
            break

    return jsonify(
        {
            "code": 0,
            "message": "已更新本地任务" + ("；上游 PATCH 未成功（可忽略，仅 loc- 任务时正常）" if not (200 <= int(last_st) < 300 or int(last_st) == 204) else ""),
            "data": {
                "assignment_id": aid,
                "upstream_attempts": upstream_attempts,
                "last_upstream": last_pl,
                "last_http_status": last_st,
            },
            "trace_id": uuid.uuid4().hex,
        }
    ), 200


@bp.delete("/api/exam-center/teacher/assignments/<assignment_id>")
@page13_access_required
def api_teacher_delete_local_assignment(assignment_id: str):
    aid = (assignment_id or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 assignment_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    attempts, last_st, last_pl = _exam_try_upstream_modify_assignment_proxy(aid, "delete")
    row = ExamCenterAssignment.query.filter_by(assignment_id=aid).first()
    if row and not _exam_assignment_in_staff_scope(row):
        return jsonify({"code": "NOT_FOUND", "message": "任务不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404
    if row:
        try:
            ExamCenterAssignment.query.filter_by(assignment_id=aid).delete()
            db.session.commit()
        except Exception as e:
            db.session.rollback()
            return jsonify({"code": "DB_ERROR", "message": str(e), "data": None, "trace_id": uuid.uuid4().hex}), 500
    return jsonify(
        {
            "code": 0,
            "message": "已删除本地任务记录" + ("；上游返回值见 data.last_upstream。" if attempts else ""),
            "data": {"assignment_id": aid, "upstream_attempts": attempts, "last_upstream": last_pl, "last_http_status": last_st},
            "trace_id": uuid.uuid4().hex,
        }
    ), 200


@bp.post("/api/exam-center/teacher/assignments/<assignment_id>/unpublish")
@page13_access_required
def api_teacher_unpublish_local_assignment(assignment_id: str):
    aid = (assignment_id or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 assignment_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    attempts, last_st, last_pl = _exam_try_upstream_modify_assignment_proxy(aid, "unpublish")
    row = ExamCenterAssignment.query.filter_by(assignment_id=aid).first()
    if row and not _exam_assignment_in_staff_scope(row):
        return jsonify({"code": "NOT_FOUND", "message": "任务不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404
    if row:
        try:
            row.status = "inactive"
            db.session.add(row)
            db.session.commit()
        except Exception as e:
            db.session.rollback()
            return jsonify({"code": "DB_ERROR", "message": str(e), "data": None, "trace_id": uuid.uuid4().hex}), 500
    return jsonify(
        {
            "code": 0,
            "message": "已标记本地任务下架（inactive）；" + ("上游返回码=" + str(last_st)),
            "data": {"assignment_id": aid, "upstream_attempts": attempts, "last_upstream": last_pl, "last_http_status": last_st},
            "trace_id": uuid.uuid4().hex,
        }
    ), 200


@bp.post("/api/exam-center/teacher/assignments/<assignment_id>/publish")
@page13_access_required
def api_teacher_publish_local_assignment(assignment_id: str):
    """重新上架：本地标记 published，并尽力请求上游恢复（失败不阻断）。"""
    aid = (assignment_id or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 assignment_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    row = ExamCenterAssignment.query.filter_by(assignment_id=aid).first()
    if not row:
        return jsonify({"code": "NOT_FOUND", "message": "本地无该任务记录，无法上架", "data": None, "trace_id": uuid.uuid4().hex}), 404
    if not _exam_assignment_in_staff_scope(row):
        return jsonify({"code": "NOT_FOUND", "message": "本地无该任务记录，无法上架", "data": None, "trace_id": uuid.uuid4().hex}), 404
    attempts, last_st, last_pl = _exam_try_upstream_modify_assignment_proxy(aid, "publish")
    try:
        row.status = "published"
        db.session.add(row)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({"code": "DB_ERROR", "message": str(e), "data": None, "trace_id": uuid.uuid4().hex}), 500
    return jsonify(
        {
            "code": 0,
            "message": "已标记本地任务上架（published）；" + ("上游返回码=" + str(last_st)),
            "data": {"assignment_id": aid, "upstream_attempts": attempts, "last_upstream": last_pl, "last_http_status": last_st},
            "trace_id": uuid.uuid4().hex,
        }
    ), 200


@bp.post("/api/exam-center/student/practice/generate-set")
@login_required
def api_exam_student_generate_practice_set():
    from .observer_view import exam_student_mutation_allowed, observer_mutation_blocked_response

    if not exam_student_mutation_allowed():
        return observer_mutation_blocked_response()
    body = _expand_quiz_request_body(_json_payload())
    uid = str(session.get("user_id") or "").strip()
    if uid:
        body["user_id"] = uid
        body["userId"] = uid
    status, payload = _quiz_api_call("quiz/practice/generate-set", method="POST", payload=body)
    return jsonify(payload), status


@bp.post("/api/exam-center/student/practice/submit")
@login_required
def api_exam_student_submit_practice():
    from .observer_view import exam_student_mutation_allowed, observer_mutation_blocked_response

    if not exam_student_mutation_allowed():
        return observer_mutation_blocked_response()
    body = _normalize_practice_submit_upstream_body(_json_payload())
    uid = str(session.get("user_id") or "").strip()
    if uid:
        body.setdefault("user_id", uid)
        body.setdefault("userId", uid)
    attempt_q = str(body.get("attempt_id") or body.get("attemptId") or "").strip()
    q = {"attempt_id": attempt_q} if attempt_q else None
    status, payload = _quiz_api_call("quiz/practice/submit", method="POST", payload=body, query=q)
    if 200 <= int(status) < 300 and isinstance(payload, dict):
        # HTTP 成功即落库：上游 data 内 code 非成功时过去会跳过写库，导致历史「丢记录」
        up_raw = payload.get("data")
        up_root = up_raw if isinstance(up_raw, dict) else {}
        merged = _merge_upstream_snapshot_with_submitted_answers(up_root, body)
        if isinstance(up_raw, dict) and not _exam_activity_upstream_root_ok(up_raw):
            merged["aiword_upstream_business_uncertain"] = True
        inner_m = merged.get("data") if isinstance(merged.get("data"), dict) else {}
        mx = _extract_result_metrics(merged)
        summ = _exam_activity_history_result_text(
            "practice",
            mx,
            merged.get("message") or payload.get("message"),
        )
        set_id = str(body.get("set_id") or inner_m.get("set_id") or inner_m.get("setId") or "").strip() or None
        etrack = str(body.get("exam_track") or "").strip() or None
        ecat = _normalize_exam_category(body.get("exam_category") or body.get("examCategory") or "daily")
        _log_student_exam_center_activity(
            mode="practice",
            exam_track=etrack,
            exam_category=ecat,
            set_id=set_id,
            assignment_id=None,
            assignment_label=None,
            attempt_id=str(body.get("attempt_id") or body.get("attemptId") or body.get("session_id") or "").strip() or None,
            upstream_http_status=int(status),
            upstream_trace_id=str(payload.get("trace_id") or "").strip() or None,
            result_summary=summ,
            upstream_result_payload=merged,
        )
    return jsonify(payload), status


@bp.get("/api/exam-center/student/quiz/grading-status/<path:attempt_id>")
@login_required
def api_student_quiz_grading_status(attempt_id: str):
    aid = (attempt_id or "").strip().lstrip("/")
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 attempt_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    st, pl = _quiz_api_call(
        f"quiz/attempts/{_urlquote(aid, safe='')}/grading-status",
        method="GET",
        organization_id=_resolve_exam_organization_id(attempt_id=aid),
    )
    return jsonify(pl), st


@bp.post("/api/exam-center/student/quiz/sync-attempt-result")
@login_required
def api_student_sync_quiz_attempt_result():
    body = _json_payload()
    aid = str(body.get("attempt_id") or body.get("attemptId") or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 attempt_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    uid = str(session.get("user_id") or "").strip()
    if not uid:
        return jsonify({"code": "UNAUTHORIZED", "message": "未登录", "data": None, "trace_id": uuid.uuid4().hex}), 401

    row = (
        ExamCenterActivity.query.filter_by(user_id=uid, attempt_id=aid)
        .order_by(ExamCenterActivity.created_at.desc())
        .first()
    )
    if not row:
        return jsonify(
            {"code": "NOT_FOUND", "message": "未找到与该 attempt 匹配的本地记录", "data": {"updated": False}, "trace_id": uuid.uuid4().hex}
        ), 404

    st, pl = _quiz_api_call(
        f"quiz/attempts/{_urlquote(aid, safe='')}/grading-status",
        method="GET",
        organization_id=(
            str(getattr(row, "organization_id", "") or "").strip()
            or _resolve_exam_organization_id(attempt_id=aid)
        ),
    )
    if not (200 <= int(st) < 300) or not isinstance(pl, dict):
        return jsonify(pl), st

    uq = _unwrap_quiz_api_success_data(pl)
    if isinstance(uq, dict) and uq.get("ok") is True and isinstance(uq.get("data"), dict):
        gdata = uq["data"]
    elif isinstance(uq, dict):
        gdata = uq
    else:
        gdata = {}

    ready = bool(gdata.get("ready"))
    if not ready:
        return jsonify(
            {
                "code": 0,
                "message": "ok（仍阅卷中）",
                "data": {"updated": False, "grading": gdata},
                "trace_id": uuid.uuid4().hex,
            }
        ), 200

    metrics = _extract_result_metrics(dict(gdata))
    summ = _exam_activity_history_result_text(str(row.mode or ""), metrics, gdata.get("grading_message"))
    row.result_summary = (summ or "")[:500]

    det = ExamCenterActivityDetail.query.filter_by(activity_id=str(row.id)).first()
    if det:
        try:
            det.score = int(round(float(metrics.get("score") or 0)))
        except Exception:
            det.score = metrics.get("score")
        ts_m = metrics.get("total_score")
        try:
            det.total_score = int(round(float(ts_m))) if ts_m is not None else det.score
        except Exception:
            det.total_score = ts_m if ts_m is not None else det.score
        det.correct_count = metrics.get("correct_count")
        det.wrong_count = metrics.get("wrong_count")
        merged_up = dict(det.upstream_payload or {})
        merged_up["grading_sync"] = gdata
        det.upstream_payload = merged_up
    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({"code": "DB_ERROR", "message": str(e), "data": {"updated": False}, "trace_id": uuid.uuid4().hex}), 500

    return jsonify(
        {
            "code": 0,
            "message": "ok",
            "data": {"updated": True, "grading": gdata, "result_summary": summ},
            "trace_id": uuid.uuid4().hex,
        }
    ), 200


@bp.get("/api/exam-center/student/wrongbook")
@login_required
def api_exam_student_wrongbook():
    q = dict(request.args.to_dict())
    uid = str(session.get("user_id") or "").strip()
    if uid:
        q["user_id"] = uid
        q["userId"] = uid
    status, payload = _quiz_api_call("quiz/wrongbook", method="GET", query=q)
    return jsonify(payload), status


@bp.get("/api/exam-center/student/unpracticed-bank")
@login_required
def api_exam_student_unpracticed_bank():
    q = dict(request.args.to_dict())
    uid = str(session.get("user_id") or "").strip()
    if uid:
        q["user_id"] = uid
        q["userId"] = uid
    status, payload = _quiz_api_call("quiz/student/unpracticed-bank", method="GET", query=q)
    return jsonify(payload), status


@bp.get("/api/exam-center/student/bank/tracks")
@login_required
def api_exam_student_tracks():
    status, payload = _quiz_api_call("quiz/bank/tracks", method="GET", query=request.args.to_dict())
    return jsonify(payload), status


@bp.get("/api/exam-center/student/assignments")
@login_required
def api_exam_student_assignments_list():
    org_id = _current_exam_scope_organization_id()
    paths = [
        "quiz/student/assignments",
        "quiz/me/assignments",
        "quiz/student/exams",
    ]
    st, pl, tried = _quiz_try_paths(
        paths,
        method="GET",
        query=request.args.to_dict(),
        organization_id=org_id or None,
    )
    upstream_http_ok = 200 <= int(st) < 300 and isinstance(pl, dict)
    raw_list: list[dict[str, Any]] = []
    if upstream_http_ok:
        root = _unwrap_quiz_api_success_data(pl)
        raw_list = _extract_assignments_from_quiz_root(root)
    normalized: list[dict[str, Any]] = []
    for x in raw_list:
        if not isinstance(x, dict):
            continue
        row = _normalize_student_assignment_row(x)
        if row["id"]:
            normalized.append(row)
    aids = [str(r.get("id") or "").strip() for r in normalized if str(r.get("id") or "").strip()]
    due_map: dict[str, Optional[datetime]] = {}
    if aids:
        try:
            q_rows = ExamCenterAssignment.query.filter(ExamCenterAssignment.assignment_id.in_(aids))
            org_clause = _exam_org_id_sql_clause(org_id, ExamCenterAssignment.organization_id)
            if org_clause is not None:
                q_rows = q_rows.filter(org_clause)
            for ar in q_rows.all():
                aid_k = str(ar.assignment_id or "").strip()
                if aid_k:
                    due_map[aid_k] = getattr(ar, "due_at", None)
        except Exception:
            due_map = {}
    for r in normalized:
        aid_k = str(r.get("id") or "").strip()
        dt = due_map.get(aid_k)
        if dt:
            r["due_at"] = dt.isoformat(timespec="seconds")
    if aids:
        try:
            q_rows = ExamCenterAssignment.query.filter(ExamCenterAssignment.assignment_id.in_(aids))
            org_clause = _exam_org_id_sql_clause(org_id, ExamCenterAssignment.organization_id)
            if org_clause is not None:
                q_rows = q_rows.filter(org_clause)
            for ar in q_rows.all():
                aid_k = str(ar.assignment_id or "").strip()
                if not aid_k:
                    continue
                sid = str(getattr(ar, "set_id", None) or "").strip()
                if not sid:
                    continue
                for r in normalized:
                    if str(r.get("id") or "").strip() != aid_k:
                        continue
                    if not str(r.get("set_id") or r.get("setId") or "").strip():
                        r["set_id"] = sid
                        r["setId"] = sid
        except Exception:
            pass
    # 上游 200 但 assignments 为空（aicheckword 未实现该列表、或兼容桩返回空）时，仍须合并 aiword 本地下发任务。
    # 仅当「上游已成功解析出非空任务列表」时才不合并本地，避免与上游真实列表混用过期 assignment_id。
    if not normalized:
        try:
            local_q = ExamCenterAssignment.query.filter(
                or_(
                    ExamCenterAssignment.status.is_(None),
                    ~ExamCenterAssignment.status.in_(("inactive", "cancelled", "archived", "deleted")),
                )
            )
            org_clause = _exam_org_id_sql_clause(org_id, ExamCenterAssignment.organization_id)
            if org_clause is not None:
                local_q = local_q.filter(org_clause)
            local_rows = local_q.order_by(ExamCenterAssignment.created_at.desc()).limit(200).all()
            uid_cur = str(session.get("user_id") or "").strip()
            allow_ids: set[str] = set()
            any_audience_ids: set[str] = set()
            try:
                # 若任务配置了受众，则仅允许受众可见；未配置受众的历史任务默认全员可见。
                aids_all = [str(r.assignment_id or "").strip() for r in local_rows if str(r.assignment_id or "").strip()]
                if aids_all:
                    any_audience_ids = {
                        str(x.assignment_id)
                        for x in ExamCenterAssignmentAudience.query.filter(
                            ExamCenterAssignmentAudience.assignment_id.in_(aids_all)
                        ).all()
                    }
                    if uid_cur:
                        allow_ids = {
                            str(x.assignment_id)
                            for x in ExamCenterAssignmentAudience.query.filter_by(user_id=uid_cur).filter(
                                ExamCenterAssignmentAudience.assignment_id.in_(aids_all)
                            ).all()
                        }
            except Exception:
                allow_ids = set()
                any_audience_ids = set()
            for r in local_rows:
                aid = str(r.assignment_id or "").strip()
                if not aid:
                    continue
                if any_audience_ids and aid in any_audience_ids and aid not in allow_ids:
                    continue
                nm = str(r.title or aid).strip() or aid
                diff_l = (getattr(r, "difficulty", None) or "").strip().lower()
                due_loc = getattr(r, "due_at", None)
                sid_loc = str(getattr(r, "set_id", None) or "").strip()
                normalized.append(
                    {
                        "id": aid,
                        "name": nm,
                        "label": nm,
                        **({"set_id": sid_loc, "setId": sid_loc} if sid_loc else {}),
                        "exam_track": (getattr(r, "exam_track", None) or "").strip() or None,
                        "exam_category": getattr(r, "exam_category", None) or "daily",
                        **({"difficulty": diff_l} if diff_l in ("easy", "medium", "hard") else {}),
                        **(
                            {"due_at": due_loc.isoformat(timespec="seconds")}
                            if due_loc
                            else {}
                        ),
                    }
                )
        except Exception:
            pass
    uid_assign = str(session.get("user_id") or "").strip()
    _attach_student_local_exam_statuses(uid_assign, normalized)
    return jsonify(
        {
            "code": 0,
            "message": "ok（含本地 assignment 兜底）" if normalized else "ok（无可用考试任务）",
            "data": {"assignments": normalized},
            "trace_id": uuid.uuid4().hex,
            "request": {"upstreamTried": tried},
        }
    ), 200


@bp.get("/api/exam-center/student/history")
@login_required
def api_exam_student_history():
    from .observer_view import (
        exam_activity_observer_fields,
        exam_student_view_mode,
        user_team_filter_options_for_exam,
    )
    from sqlalchemy import false as sql_false

    _ensure_exam_scope_data_repaired()
    view_mode = exam_student_view_mode()
    team_filter = str(request.args.get("teamId") or request.args.get("team_id") or "").strip()
    user_filter = str(request.args.get("userId") or request.args.get("user_id") or "").strip()
    try:
        lim = int(str(request.args.get("limit") or "100"))
    except (TypeError, ValueError):
        lim = 100
    lim = max(1, min(200, lim))
    try:
        offset = int(str(request.args.get("offset") or "0"))
    except (TypeError, ValueError):
        offset = 0
    offset = max(0, offset)
    pass_score = _exam_pass_score()

    if view_mode == "normal":
        uid = str(session.get("user_id") or "").strip()
        if not uid:
            return jsonify({"code": "UNAUTHORIZED", "message": "未登录", "data": None, "trace_id": uuid.uuid4().hex}), 401
        base_q = ExamCenterActivity.query.filter_by(user_id=uid)
        org_id = _current_exam_scope_organization_id()
        org_clause = _exam_org_id_sql_clause(org_id, ExamCenterActivity.organization_id)
        if org_clause is not None:
            base_q = base_q.filter(org_clause)
        total = base_q.count()
        rows = (
            base_q.order_by(ExamCenterActivity.created_at.desc())
            .offset(offset)
            .limit(lim)
            .all()
        )
        filter_opts = {"teams": [], "users": []}
    else:
        base_q = _scope_exam_activity_query(ExamCenterActivity.query)
        scoped_uids: set[str] = set()
        from .exam_display_labels import normalize_user_key

        for row in base_q.with_entities(ExamCenterActivity.user_id).distinct().all():
            uid = row[0] if isinstance(row, (tuple, list)) else row
            nk = normalize_user_key(str(uid or ""))
            if nk:
                scoped_uids.add(nk)
        filter_opts = user_team_filter_options_for_exam(
            scoped_uids,
            include_teams=(view_mode == "super_admin_readonly"),
        )
        if user_filter:
            from .exam_display_labels import exam_activity_user_id_match_keys

            match_keys = exam_activity_user_id_match_keys(user_filter)
            if match_keys:
                base_q = base_q.filter(ExamCenterActivity.user_id.in_(list(match_keys)))
            else:
                base_q = base_q.filter(sql_false())
        elif team_filter:
            scope_tid = _active_exam_team_id_for_observer()
            if scope_tid and team_filter != scope_tid:
                base_q = base_q.filter(sql_false())
            elif not scope_tid:
                from .models import UserTeamMembership

                team_uids = {
                    str(m.user_id).strip()
                    for m in UserTeamMembership.query.filter_by(team_id=team_filter).all()
                    if str(m.user_id).strip()
                }
                expanded_uids: set[str] = set()
                from .exam_display_labels import exam_activity_user_id_match_keys

                for tu in team_uids:
                    expanded_uids.update(exam_activity_user_id_match_keys(tu))
                if not expanded_uids:
                    base_q = base_q.filter(sql_false())
                else:
                    base_q = base_q.filter(ExamCenterActivity.user_id.in_(list(expanded_uids)))
        total = base_q.count()
        rows = (
            base_q.order_by(ExamCenterActivity.created_at.desc())
            .offset(offset)
            .limit(lim)
            .all()
        )

    ids = [str(r.id) for r in rows if getattr(r, "id", None)]
    det_map: dict[str, ExamCenterActivityDetail] = {}
    if ids:
        ds = ExamCenterActivityDetail.query.filter(ExamCenterActivityDetail.activity_id.in_(ids)).all()
        det_map = {str(d.activity_id): d for d in ds}
    records: list[dict[str, Any]] = []
    for a in rows:
        d = det_map.get(str(a.id))
        mode_l = (a.mode or "").strip().lower()
        mode_label = "练习" if mode_l == "practice" else ("考试" if mode_l == "exam" else (a.mode or "-"))
        tgt = (a.assignment_label or a.set_id or a.attempt_id or "-") or "-"
        obs = exam_activity_observer_fields(
            str(getattr(a, "user_id", "") or ""),
            preferred_team_id=_active_exam_team_id_for_observer(),
            activity_display=str(getattr(a, "display_name", None) or "").strip() or None,
            activity_username=str(getattr(a, "username", None) or "").strip() or None,
        )
        records.append(
            {
                "id": a.id,
                "created_at": a.created_at.isoformat() if a.created_at else "",
                "mode": a.mode,
                "mode_label": mode_label,
                "assignment_id": str(a.assignment_id).strip() if getattr(a, "assignment_id", None) else "",
                "set_id": str(getattr(a, "set_id", None) or "").strip() or None,
                "setId": str(getattr(a, "set_id", None) or "").strip() or None,
                "target_label": str(tgt),
                "result": (a.result_summary or "-")[:500],
                "score": d.score if d else None,
                "total_score": d.total_score if d else None,
                "correct_count": d.correct_count if d else None,
                "wrong_count": d.wrong_count if d else None,
                "pass_score": pass_score,
                "passed": (d.score is not None and float(d.score) >= pass_score) if d and d.score is not None else None,
                "weakness": d.weakness if d else None,
                "recommendation": d.recommendation if d else None,
                "teamId": obs.get("teamId") or None,
                "teamName": obs.get("teamName") or None,
                "userId": obs.get("userId") or None,
                "displayName": obs.get("displayName") or None,
            }
        )
    has_more = (offset + len(records)) < int(total)
    return jsonify(
        {
            "code": 0,
            "message": "ok",
            "data": {
                "records": records,
                "total": int(total),
                "has_more": has_more,
                "viewMode": view_mode,
                "observerMode": view_mode != "normal",
                "readOnly": view_mode == "super_admin_readonly",
                "filterOptions": filter_opts,
            },
            "trace_id": uuid.uuid4().hex,
        }
    ), 200


@bp.post("/api/exam-center/student/exams/start")
@login_required
def api_exam_student_start_exam():
    # 兼容旧前端路径：考试链路已迁移到 aiword 本地，不再调用上游 /quiz/exams/*
    return api_exam_student_start_exam_local()


@bp.post("/api/exam-center/student/exams/submit")
@login_required
def api_exam_student_submit_exam():
    # 兼容旧前端路径：考试链路已迁移到 aiword 本地，不再调用上游 /quiz/exams/*
    return api_exam_student_submit_exam_local()


def _assignment_visible_to_user(*, assignment_id: str, user_id: str) -> bool:
    """若该任务配置了受众，则仅受众可见；未配置受众的历史任务默认可见。"""
    aid = (assignment_id or "").strip()
    uid = (user_id or "").strip()
    if not aid or not uid:
        return False
    try:
        row = ExamCenterAssignment.query.filter_by(assignment_id=aid).first()
        if not row:
            return False
        org_id = _current_exam_scope_organization_id()
        from .authz import exam_row_organization_id_matches

        if not exam_row_organization_id_matches(getattr(row, "organization_id", None), org_id):
            return False
        any_rows = ExamCenterAssignmentAudience.query.filter_by(assignment_id=aid).limit(1).all()
        if not any_rows:
            return True
        ok = (
            ExamCenterAssignmentAudience.query.filter_by(assignment_id=aid, user_id=uid)
            .limit(1)
            .all()
        )
        return bool(ok)
    except Exception:
        # 安全兜底：数据库异常时不扩大可见范围
        return False


def _merge_local_exam_attempt_status(attempts: list[ExamAttempt]) -> str:
    """同一学生对同一 assignment 多条 attempt 时合并状态：graded > grading > in_progress > open。"""
    if not attempts:
        return "open"
    states = {str(getattr(x, "state", None) or "").strip().lower() for x in attempts}
    if "graded" in states:
        return "graded"
    if "grading" in states:
        return "grading"
    if "started" in states:
        return "in_progress"
    return "open"


def _attach_student_local_exam_statuses(uid_cur: str, normalized: list[dict[str, Any]]) -> None:
    """
    为学生任务列表附加本地考试进度（ExamAttempt），用于：
    - 学生端：已提交/已出分后不再显示「开始考试」；
    - 后续钉钉：与页面3一致可用字段区分「仍需参加/未提交」与「已交卷」。
    """
    for r in normalized or []:
        if isinstance(r, dict):
            r.setdefault("local_exam_status", "open")
            r.setdefault("exam_task_closed_for_student", False)
            r.setdefault("student_exam_needs_submit", True)
    if not uid_cur or not normalized:
        return
    aids = [str(r.get("id") or "").strip() for r in normalized if isinstance(r, dict) and str(r.get("id") or "").strip()]
    if not aids:
        return
    try:
        all_atts = ExamAttempt.query.filter(ExamAttempt.user_id == uid_cur, ExamAttempt.assignment_id.in_(aids)).all()
        by_aid: dict[str, list[ExamAttempt]] = {}
        for atx in all_atts:
            k = str(atx.assignment_id or "").strip()
            if k:
                by_aid.setdefault(k, []).append(atx)
        for r in normalized:
            if not isinstance(r, dict):
                continue
            aid_k = str(r.get("id") or "").strip()
            st = _merge_local_exam_attempt_status(by_aid.get(aid_k) or [])
            r["local_exam_status"] = st
            # 已交卷（含阅卷中）：前端不再展示「开始考试」；钉钉「未完成」可筛 not finalized 或 needs_submit
            closed = st in ("grading", "graded")
            r["exam_task_closed_for_student"] = closed
            r["student_exam_needs_submit"] = st in ("open", "in_progress")
            r["student_exam_finalized"] = st == "graded"
    except Exception:
        for r in normalized:
            if isinstance(r, dict):
                r["local_exam_status"] = "open"
                r["exam_task_closed_for_student"] = False
                r["student_exam_needs_submit"] = True
                r["student_exam_finalized"] = False


def _subjective_answer_text_from_item(it: ExamAttemptItem) -> str:
    ua_raw = None
    if isinstance(it.user_answer, dict):
        ua_raw = it.user_answer.get("value")
    ua_txt = str(ua_raw or "").strip() if ua_raw is not None else ""
    if not ua_txt and isinstance(it.user_answer, dict):
        ua_txt = " ".join(
            str(v).strip()
            for v in it.user_answer.values()
            if v is not None and str(v).strip()
        ).strip()
    return ua_txt


def _cap_subjective_score_by_answer_text(score: float, ua_text: str) -> float:
    """空答/过短/明显不全时封顶，避免 LLM 误给高分。"""
    sc = max(0.0, min(1.0, float(score)))
    txt = str(ua_text or "").strip()
    if not txt:
        return 0.0
    n = len(txt)
    if n < 8:
        return 0.0
    if n < 30:
        return min(sc, 0.15)
    if n < 80:
        return min(sc, 0.35)
    if n < 150:
        return min(sc, 0.65)
    return sc


def _attempt_items_total_score_100(rows: list[ExamAttemptItem]) -> tuple[float, int, int]:
    """总分按 0~100 百分制：每题满分 1，客观题 0/1，主观题使用 0~1 score。"""
    if not rows:
        return 0.0, 0, 0
    total = len(rows)
    score_sum = 0.0
    corr = wrong = 0
    for it in rows:
        if getattr(it, "subjective_needed", False):
            sc = getattr(it, "subjective_score", None)
            if sc is None:
                continue
            try:
                score_sum += max(0.0, min(1.0, float(sc)))
            except Exception:
                pass
        else:
            ic = getattr(it, "is_correct", None)
            if ic is True:
                score_sum += 1.0
                corr += 1
            elif ic is False:
                wrong += 1
    pct = (score_sum / float(total)) * 100.0 if total > 0 else 0.0
    if pct < 0.0:
        pct = 0.0
    if pct > 100.0:
        pct = 100.0
    return pct, corr, wrong


def _local_exam_attempt_items_payload(attempt_id_key: str, att: Optional[ExamAttempt] = None) -> Optional[list[dict[str, Any]]]:
    """
    aiword 本地 ExamAttempt 的题目行，字段形态对齐练习上游 quiz/attempts/*/answers 的 items，
    供「练习/考试详情」弹窗与活动快照降级展示（options / answer / user_answer）。
    """
    aid = (attempt_id_key or "").strip()
    if not aid:
        return None
    att_row = att if att is not None else ExamAttempt.query.filter_by(attempt_id=aid).first()
    if not att_row:
        return None
    state_l = str(getattr(att_row, "state", None) or "").strip().lower()
    rows = ExamAttemptItem.query.filter_by(attempt_id=aid).order_by(ExamAttemptItem.created_at.asc()).all()
    items: list[dict[str, Any]] = []
    for it in rows:
        ua = (it.user_answer or {}).get("value") if isinstance(it.user_answer, dict) else None
        ans = (it.answer_snapshot or {}).get("answer") if isinstance(it.answer_snapshot, dict) else None
        opts = it.options_snapshot if isinstance(it.options_snapshot, list) else []
        entry: dict[str, Any] = {
            "question_id": it.question_id,
            "question_type": it.question_type,
            "stem": it.stem_snapshot,
            "options": opts,
            "user_answer": ua,
            "answer": ans,
            "is_correct": it.is_correct,
            "subjective_needed": bool(it.subjective_needed),
            "subjective_score": it.subjective_score,
            "subjective_reason": it.subjective_reason,
            "subjective_recommendation": it.subjective_recommendation,
            "evidence_used": it.evidence_used if isinstance(it.evidence_used, list) else None,
        }
        if (
            bool(it.subjective_needed)
            and state_l == "grading"
            and getattr(it, "subjective_score", None) is None
        ):
            entry["teacher_comment"] = "pending_subjective_grading"
        items.append(entry)
    return items


def _local_exam_attempt_ids_with_pending_subjective(attempt_ids: list[str]) -> set[str]:
    """给定 attempt_id 列表，返回其中「存在主观题且尚未给出 subjective_score」的 attempt_id 集合。"""
    ids = [str(x or "").strip() for x in attempt_ids if str(x or "").strip()]
    if not ids:
        return set()
    q = (
        db.session.query(ExamAttemptItem.attempt_id)
        .filter(ExamAttemptItem.attempt_id.in_(ids))
        .filter(ExamAttemptItem.subjective_needed.is_(True))
        .filter(ExamAttemptItem.subjective_score.is_(None))
        .distinct()
    )
    return {str(r[0]) for r in q.all() if r and r[0]}


def _enqueue_local_exam_ai_grading(att: ExamAttempt, rows: Optional[list[ExamAttemptItem]] = None) -> dict[str, Any]:
    """
    调用上游 quiz/grading/paper-by-ai 并 upsert ExamGradingJob（不 commit）。
    用于提交后首次入队，或统计端「重新阅卷」。
    """
    attempt_id = str(att.attempt_id or "").strip()
    if not attempt_id:
        return {"ok": False, "reason": "missing_attempt_id"}
    if rows is None:
        rows = ExamAttemptItem.query.filter_by(attempt_id=attempt_id).all()
    subj_items: list[dict[str, Any]] = []
    for it in rows:
        if not getattr(it, "subjective_needed", False):
            continue
        subj_items.append(
            {
                "question_id": it.question_id,
                "question_type": it.question_type,
                "stem": it.stem_snapshot,
                "options": it.options_snapshot or [],
                "user_answer": (it.user_answer or {}).get("value") if isinstance(it.user_answer, dict) else None,
            }
        )
    if not subj_items:
        return {"ok": False, "reason": "no_subjective_items"}
    payload = {
        "attempt_id": attempt_id,
        "exam_track": att.exam_track,
        "assignment_id": att.assignment_id,
        "items": subj_items,
    }
    st_g, pl_g = _quiz_api_call("quiz/grading/paper-by-ai", method="POST", payload=payload)
    job_id = ""
    if 200 <= int(st_g) < 300 and isinstance(pl_g, dict):
        root = _unwrap_quiz_api_success_data(pl_g)
        if isinstance(root, dict):
            inner = root.get("data") if isinstance(root.get("data"), dict) else root
            job_id = str(inner.get("job_id") or inner.get("jobId") or inner.get("id") or "").strip()
    gj = ExamGradingJob.query.filter_by(attempt_id=attempt_id).first()
    if not gj:
        gj = ExamGradingJob(attempt_id=attempt_id)
    gj.upstream_job_id = job_id or gj.upstream_job_id
    gj.status = "running" if job_id else "failed"
    gj.last_upstream_http_status = int(st_g)
    gj.last_upstream_trace_id = str((pl_g or {}).get("trace_id") or "").strip() if isinstance(pl_g, dict) else None
    gj.last_upstream_payload = pl_g if isinstance(pl_g, dict) else {"data": pl_g}
    gj.last_message = str((pl_g or {}).get("message") or "")[:500] if isinstance(pl_g, dict) else None
    db.session.add(gj)
    return {
        "ok": True,
        "job_id": job_id,
        "http_status": int(st_g),
        "payload": pl_g if isinstance(pl_g, dict) else None,
        "job_status": gj.status,
    }


def _local_exam_admin_regrade_allowed(att: ExamAttempt, gj: Optional[ExamGradingJob]) -> tuple[bool, str]:
    """统计/老师端：仅允许对「仍有待判主观分」且（阅卷中 / 阅卷失败 / 异常 graded）的记录重新入队。"""
    aid = str(att.attempt_id or "").strip()
    if not aid:
        return False, "缺少 attempt"
    pending_set = _local_exam_attempt_ids_with_pending_subjective([aid])
    if aid not in pending_set:
        return False, "主观题均已判分或无主观题，无需重新阅卷"
    ast = str(att.state or "").strip().lower()
    gj_st = str(gj.status or "").strip().lower() if gj else ""
    if ast == "grading":
        return True, ""
    if gj_st == "failed":
        return True, ""
    if ast == "graded":
        # 数据不一致：已标记 graded 但主观分未回填，仍允许管理员重试入队
        return True, ""
    return False, "仅支持对阅卷中、阅卷失败或判分未完成的本地考试重新发起判分"


def _update_exam_center_activity_snapshot_for_local_attempt(attempt_id_key: str) -> bool:
    """
    将 ExamAttempt 当前分数/阅卷状态写回 exam_center_activities 与 exam_center_activity_details，
    供统计端列表与「练习/考试详情」弹窗与 submit-local 落库口径一致。
    """
    aid = (attempt_id_key or "").strip()
    if not aid:
        return False
    att = ExamAttempt.query.filter_by(attempt_id=aid).first()
    if not att:
        return False
    act = ExamCenterActivity.query.filter_by(attempt_id=aid).order_by(ExamCenterActivity.created_at.desc()).first()
    if not act:
        return False
    assign_row = None
    aid_assign = str(att.assignment_id or "").strip()
    if aid_assign:
        assign_row = ExamCenterAssignment.query.filter_by(assignment_id=aid_assign).first()
    set_id_log = str(getattr(assign_row, "set_id", None) or "").strip() or None
    items_snap = _local_exam_attempt_items_payload(aid, att=att) or []
    gs = "pending" if str(getattr(att, "state", None) or "").strip().lower() == "grading" else "graded"
    merged_log: dict[str, Any] = {
        "code": 0,
        "message": "local-grading-synced",
        "grading_status": gs,
        "score": getattr(att, "score", None),
        "total_score": getattr(att, "total_score", None),
        "correct_count": getattr(att, "correct_count", None),
        "wrong_count": getattr(att, "wrong_count", None),
        "data": {
            "score": getattr(att, "score", None),
            "total_score": getattr(att, "total_score", None),
            "correct_count": getattr(att, "correct_count", None),
            "wrong_count": getattr(att, "wrong_count", None),
            "grading_status": gs,
            "attempt_id": att.attempt_id,
            "assignment_id": att.assignment_id,
            "set_id": set_id_log,
        },
        "attempt_items": items_snap,
    }
    mx = _extract_result_metrics(merged_log)
    summ = _exam_activity_history_result_text("exam", mx, merged_log.get("message"))
    if summ:
        act.result_summary = str(summ)[:500]
    det = ExamCenterActivityDetail.query.filter_by(activity_id=str(act.id)).first()
    if det is None:
        det = ExamCenterActivityDetail(activity_id=str(act.id), mode=str(act.mode or "")[:16] or "exam")
    det.mode = str(act.mode or "")[:16] or det.mode
    det.score = getattr(att, "score", None)
    det.total_score = getattr(att, "total_score", None)
    det.correct_count = getattr(att, "correct_count", None)
    det.wrong_count = getattr(att, "wrong_count", None)
    wk = mx.get("weakness")
    det.weakness = str(wk)[:1000] if wk is not None and str(wk).strip() else None
    rc = mx.get("recommendation")
    det.recommendation = str(rc)[:1000] if rc is not None and str(rc).strip() else None
    det.upstream_payload = _json_sanitize_for_db(merged_log)
    db.session.add(act)
    db.session.add(det)
    return True


def _local_exam_pull_grading_job_and_persist(aid_raw: str) -> dict[str, Any]:
    """
    拉取上游 quiz/grading/jobs/{job_id}，写回 ExamAttemptItem/ExamAttempt/ExamGradingJob，
    若已出分则同步活动快照（内含 db.session.commit）。
    返回 {"error": bool, "http": int, "message": str, "data": dict|None, "trace_id": str}
    """
    aid = (aid_raw or "").strip()
    trace = uuid.uuid4().hex
    if not aid:
        return {"error": True, "http": 400, "message": "缺少 attempt_id", "data": None, "trace_id": trace}
    att = ExamAttempt.query.filter_by(attempt_id=aid).first()
    if not att:
        return {"error": True, "http": 404, "message": "attempt 不存在", "data": None, "trace_id": trace}
    gj = ExamGradingJob.query.filter_by(attempt_id=aid).first()
    if not gj or not str(gj.upstream_job_id or "").strip():
        return {"error": True, "http": 404, "message": "未找到判分任务 job_id", "data": None, "trace_id": trace}

    st, pl = _quiz_api_call(f"quiz/grading/jobs/{_urlquote(str(gj.upstream_job_id), safe='')}", method="GET")
    if isinstance(pl, dict):
        gj.last_upstream_http_status = int(st)
        gj.last_upstream_trace_id = str(pl.get("trace_id") or "").strip() or None
        gj.last_upstream_payload = pl
        gj.last_message = str(pl.get("message") or "")[:500] or None
    try:
        root = _unwrap_quiz_api_success_data(pl) if isinstance(pl, dict) else {}
        inner = root.get("data") if isinstance(root, dict) and isinstance(root.get("data"), dict) else root
        status = str((inner or {}).get("status") or (inner or {}).get("state") or "").strip().lower()
    except Exception:
        inner = {}
        status = ""

    activity_synced = False
    if status in ("done", "success", "completed", "graded", "complete"):
        try:
            items = (inner or {}).get("items") if isinstance(inner, dict) else None
            if not isinstance(items, list):
                items = []
            by_q: dict[str, dict[str, Any]] = {}
            for x in items:
                if not isinstance(x, dict):
                    continue
                qid = str(x.get("question_id") or x.get("questionId") or "").strip()
                if qid:
                    by_q[qid] = x
            rows = ExamAttemptItem.query.filter_by(attempt_id=aid).all()
            for it in rows:
                if not getattr(it, "subjective_needed", False):
                    continue
                qid = str(it.question_id or "").strip()
                x = by_q.get(qid) or {}
                sc = x.get("score")
                ua_txt = _subjective_answer_text_from_item(it)
                try:
                    if sc is not None:
                        sc_f = _cap_subjective_score_by_answer_text(float(sc), ua_txt)
                        if not ua_txt:
                            it.subjective_reason = (str(x.get("reason") or "") or "未作答或答案为空")[:2000]
                        it.subjective_score = sc_f
                    else:
                        it.subjective_score = None
                except Exception:
                    it.subjective_score = None
                if not str(it.subjective_reason or "").strip():
                    it.subjective_reason = str(x.get("reason") or "")[:2000] or None
                it.subjective_recommendation = str(x.get("recommendation") or "")[:2000] or None
                ev = x.get("evidence_used") or x.get("evidenceUsed") or []
                it.evidence_used = ev if isinstance(ev, list) else None
                it.updated_at = now_local()
            sc2, corr2, wrong2 = _attempt_items_total_score_100(rows)
            att.state = "graded"
            att.score = sc2
            att.total_score = 100.0
            att.correct_count = corr2
            att.wrong_count = wrong2
            gj.status = "done"
            _update_exam_center_activity_snapshot_for_local_attempt(aid)
            activity_synced = True
            db.session.add(att)
            db.session.add(gj)
            db.session.commit()
        except Exception as e:
            db.session.rollback()
            return {"error": True, "http": 500, "message": str(e), "data": None, "trace_id": trace}
    elif status in ("failed", "error"):
        gj.status = "failed"
        try:
            db.session.add(gj)
            db.session.commit()
        except Exception as e2:
            db.session.rollback()
            return {"error": True, "http": 500, "message": str(e2), "data": None, "trace_id": trace}
    else:
        try:
            db.session.add(gj)
            db.session.commit()
        except Exception as e3:
            db.session.rollback()
            return {"error": True, "http": 500, "message": str(e3), "data": None, "trace_id": trace}

    return {
        "error": False,
        "http": 200,
        "message": "ok",
        "data": {
            "attempt_id": aid,
            "job_status": gj.status,
            "state": att.state,
            "activity_synced": activity_synced,
        },
        "trace_id": trace,
    }


@bp.post("/api/exam-center/student/exams/start-local")
@login_required
def api_exam_student_start_exam_local():
    from .observer_view import exam_student_mutation_allowed, observer_mutation_blocked_response

    if not exam_student_mutation_allowed():
        return observer_mutation_blocked_response()
    body = _json_payload()
    assignment_id = (body.get("assignment_id") or body.get("assignmentId") or "").strip()
    if not assignment_id:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 assignment_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    uid = str(session.get("user_id") or "").strip()
    if not uid:
        return jsonify({"code": "UNAUTHORIZED", "message": "未登录", "data": None, "trace_id": uuid.uuid4().hex}), 401
    if not _assignment_visible_to_user(assignment_id=assignment_id, user_id=uid):
        return jsonify({"code": "FORBIDDEN", "message": "无权参加该考试任务", "data": None, "trace_id": uuid.uuid4().hex}), 403

    row = ExamCenterAssignment.query.filter_by(assignment_id=assignment_id).first()
    if not row:
        return jsonify({"code": "NOT_FOUND", "message": "考试任务不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404
    # 截止后仍允许开考；是否按时完成由统计端按提交时间与 deadline 比对（见 /stats/exam/*）
    set_id = str(getattr(row, "set_id", None) or "").strip()
    if not set_id:
        return jsonify({"code": "BAD_REQUEST", "message": "任务缺少 set_id，无法开考", "data": None, "trace_id": uuid.uuid4().hex}), 400

    st_set, pl_set = _quiz_api_call(
        f"quiz/sets/{_urlquote(set_id, safe='')}",
        method="GET",
        query={},
        organization_id=str(getattr(row, "organization_id", "") or "").strip(),
    )
    if not (200 <= int(st_set) < 300) or not isinstance(pl_set, dict):
        return jsonify(pl_set), st_set
    up = _unwrap_quiz_api_success_data(pl_set)
    items = _find_set_item_dicts(up)
    if not items:
        return jsonify({"code": "UPSTREAM_EMPTY_SET", "message": "上游套题明细为空", "data": {"set_id": set_id}, "trace_id": uuid.uuid4().hex}), 502

    attempt_id = uuid.uuid4().hex
    try:
        att = ExamAttempt(
            organization_id=(str(getattr(row, "organization_id", "") or "").strip() or None),
            attempt_id=attempt_id,
            assignment_id=assignment_id,
            user_id=uid,
            exam_track=str(getattr(row, "exam_track", None) or "").strip() or None,
            exam_category=_normalize_exam_category(getattr(row, "exam_category", None) or "daily"),
            state="started",
            started_at=now_local(),
        )
        db.session.add(att)
        out_items: list[dict[str, Any]] = []
        for ix, it in enumerate(items):
            if not isinstance(it, dict):
                continue
            qid = _question_id_from_set_item(it) or f"q-{ix}"
            qt = str(it.get("question_type") or it.get("type") or "").strip().lower() or None
            stem = str(it.get("stem") or it.get("title") or "").strip() or None
            opts = it.get("options") if isinstance(it.get("options"), list) else None
            ans = it.get("answer") if "answer" in it else (it.get("correct_answer") if "correct_answer" in it else None)
            # 案例分析题：不展示选项，作答为文本；标准答案不按单选处理（快照置空，避免上游误带 options/answer）
            if qt == "case_analysis":
                opts = None
                ans = None
            # 主观题：无标准答案可立即判分（按用户要求）
            subj = ans is None or (qt is not None and qt not in ("single_choice", "multiple_choice", "true_false"))
            db.session.add(
                ExamAttemptItem(
                    attempt_id=attempt_id,
                    question_id=str(qid),
                    question_type=qt,
                    stem_snapshot=stem,
                    options_snapshot=opts,
                    answer_snapshot={"answer": ans} if ans is not None else None,
                    user_answer=None,
                    is_correct=None,
                    score=None,
                    subjective_needed=bool(subj),
                )
            )
            out_items.append(
                {"question_id": str(qid), "question_type": qt, "stem": stem or "", "options": opts or []}
            )
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({"code": "DB_ERROR", "message": str(e), "data": None, "trace_id": uuid.uuid4().hex}), 500

    return jsonify({"code": 0, "message": "ok", "data": {"attempt_id": attempt_id, "assignment_id": assignment_id, "set_id": set_id, "items": out_items}, "trace_id": uuid.uuid4().hex}), 200


@bp.post("/api/exam-center/student/exams/submit-local")
@login_required
def api_exam_student_submit_exam_local():
    from .observer_view import exam_student_mutation_allowed, observer_mutation_blocked_response

    if not exam_student_mutation_allowed():
        return observer_mutation_blocked_response()
    body = _json_payload()
    attempt_id = str(body.get("attempt_id") or body.get("attemptId") or "").strip()
    if not attempt_id:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 attempt_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    uid = str(session.get("user_id") or "").strip()
    if not uid:
        return jsonify({"code": "UNAUTHORIZED", "message": "未登录", "data": None, "trace_id": uuid.uuid4().hex}), 401

    att = ExamAttempt.query.filter_by(attempt_id=attempt_id).first()
    if not att or str(att.user_id or "").strip() != uid:
        return jsonify({"code": "NOT_FOUND", "message": "attempt 不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404

    answers = body.get("answers") or body.get("items") or []
    ans_map: dict[str, Any] = {}
    if isinstance(answers, dict):
        ans_map = {str(k).strip(): v for k, v in answers.items() if str(k).strip()}
    elif isinstance(answers, list):
        for r in answers:
            if not isinstance(r, dict):
                continue
            qid = str(r.get("question_id") or r.get("questionId") or "").strip()
            if not qid:
                continue
            if "answer" in r:
                ans_map[qid] = r.get("answer")
            elif "user_answer" in r:
                ans_map[qid] = r.get("user_answer")

    rows = ExamAttemptItem.query.filter_by(attempt_id=attempt_id).all()
    if not rows:
        return jsonify({"code": "BAD_REQUEST", "message": "attempt 无题目明细", "data": None, "trace_id": uuid.uuid4().hex}), 400

    try:
        # 写入作答与客观题判分
        for it in rows:
            qid = str(it.question_id or "").strip()
            if not qid:
                continue
            ua = ans_map.get(qid)
            it.user_answer = {"value": ua} if ua is not None else None
            if not getattr(it, "subjective_needed", False):
                ca = None
                try:
                    ca = (it.answer_snapshot or {}).get("answer") if isinstance(it.answer_snapshot, dict) else None
                except Exception:
                    ca = None
                base_row = {
                    "question_type": it.question_type,
                    "type": it.question_type,
                    "options": it.options_snapshot,
                }
                ok = _objective_answers_equivalent_aiword(base_row, ua, ca)
                it.is_correct = bool(ok)
                it.score = 1.0 if ok else 0.0
            it.updated_at = now_local()

        # 状态与汇总
        att.submitted_at = now_local()
        has_subj = any(bool(getattr(x, "subjective_needed", False)) for x in rows)
        if has_subj:
            att.state = "grading"
        else:
            att.state = "graded"
            sc, corr, wrong = _attempt_items_total_score_100(rows)
            att.score = sc
            att.total_score = 100.0
            att.correct_count = corr
            att.wrong_count = wrong

        db.session.add(att)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({"code": "DB_ERROR", "message": str(e), "data": None, "trace_id": uuid.uuid4().hex}), 500

    # 触发上游整卷主观判分 job（异步）
    if att.state == "grading":
        try:
            enq = _enqueue_local_exam_ai_grading(att, rows)
            if enq.get("ok"):
                db.session.commit()
            else:
                db.session.rollback()
        except Exception:
            db.session.rollback()

    # 学生历史列表读 exam_center_activities；本地考试提交原先只更新 exam_attempts，导致「提交成功但无记录」
    try:
        assign_row = ExamCenterAssignment.query.filter_by(assignment_id=str(att.assignment_id or "").strip()).first()
        set_id_log = str(getattr(assign_row, "set_id", None) or "").strip() or None
        assignment_label_log = str(getattr(assign_row, "title", None) or "").strip() or None
        items_snap = _local_exam_attempt_items_payload(str(att.attempt_id), att=att) or []
        gs = "pending" if str(getattr(att, "state", None) or "").strip().lower() == "grading" else "graded"
        merged_log: dict[str, Any] = {
            "code": 0,
            "message": "submit-local",
            "grading_status": gs,
            "score": getattr(att, "score", None),
            "total_score": getattr(att, "total_score", None),
            "correct_count": getattr(att, "correct_count", None),
            "wrong_count": getattr(att, "wrong_count", None),
            "data": {
                "score": getattr(att, "score", None),
                "total_score": getattr(att, "total_score", None),
                "correct_count": getattr(att, "correct_count", None),
                "wrong_count": getattr(att, "wrong_count", None),
                "grading_status": gs,
                "attempt_id": att.attempt_id,
                "assignment_id": att.assignment_id,
                "set_id": set_id_log,
            },
            "attempt_items": items_snap,
        }
        mx = _extract_result_metrics(merged_log)
        summ = _exam_activity_history_result_text("exam", mx, merged_log.get("message"))
        _log_student_exam_center_activity(
            mode="exam",
            exam_track=str(getattr(att, "exam_track", None) or "").strip() or None,
            exam_category=str(getattr(att, "exam_category", None) or "").strip() or "daily",
            set_id=set_id_log,
            assignment_id=str(att.assignment_id or "").strip() or None,
            assignment_label=assignment_label_log,
            attempt_id=str(att.attempt_id or "").strip() or None,
            upstream_http_status=200,
            upstream_trace_id=None,
            result_summary=summ,
            upstream_result_payload=merged_log,
        )
    except Exception:
        try:
            current_app.logger.exception("exam_submit_local_activity_log_failed attempt_id=%s", attempt_id)
        except Exception:
            pass

    return jsonify({"code": 0, "message": "ok", "data": {"attempt_id": attempt_id, "state": att.state}, "trace_id": uuid.uuid4().hex}), 200


@bp.get("/api/exam-center/student/attempts/<attempt_id>/grading-status")
@login_required
def api_exam_student_attempt_grading_status_local(attempt_id: str):
    aid = (attempt_id or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 attempt_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    uid = str(session.get("user_id") or "").strip()
    att = ExamAttempt.query.filter_by(attempt_id=aid).first()
    if not att or str(att.user_id or "").strip() != uid:
        return jsonify({"code": "NOT_FOUND", "message": "attempt 不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404
    gj = ExamGradingJob.query.filter_by(attempt_id=aid).first()
    return jsonify(
        {
            "code": 0,
            "message": "ok",
            "data": {
                "attempt_id": aid,
                "state": att.state,
                "score": att.score,
                "total_score": att.total_score,
                "correct_count": att.correct_count,
                "wrong_count": att.wrong_count,
                "job": {
                    "upstream_job_id": gj.upstream_job_id if gj else None,
                    "status": gj.status if gj else None,
                    "last_message": gj.last_message if gj else None,
                }
                if gj
                else None,
            },
            "trace_id": uuid.uuid4().hex,
        }
    ), 200


@bp.post("/api/exam-center/student/attempts/<attempt_id>/sync-grading")
@login_required
def api_exam_student_attempt_sync_grading_local(attempt_id: str):
    aid = (attempt_id or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 attempt_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    uid = str(session.get("user_id") or "").strip()
    att = ExamAttempt.query.filter_by(attempt_id=aid).first()
    if not att or str(att.user_id or "").strip() != uid:
        return jsonify({"code": "NOT_FOUND", "message": "attempt 不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404
    out = _local_exam_pull_grading_job_and_persist(aid)
    if out.get("error"):
        code = "DB_ERROR" if int(out.get("http") or 500) == 500 else "NOT_FOUND"
        if int(out.get("http") or 400) == 400:
            code = "BAD_REQUEST"
        return (
            jsonify(
                {
                    "code": code,
                    "message": str(out.get("message") or "同步失败"),
                    "data": out.get("data"),
                    "trace_id": str(out.get("trace_id") or uuid.uuid4().hex),
                }
            ),
            int(out.get("http") or 500),
        )
    return jsonify({"code": 0, "message": "ok", "data": out.get("data"), "trace_id": out.get("trace_id")}), 200


@bp.post("/api/exam-center/teacher/local-exam/attempts/<attempt_id>/sync-grading")
@page13_access_required
def api_exam_teacher_sync_local_exam_grading(attempt_id: str):
    """统计/老师端（page13）：将上游主观阅卷结果拉回 aiword 并更新活动快照（学生端同逻辑，但不校验答卷归属）。"""
    aid = (attempt_id or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 attempt_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    out = _local_exam_pull_grading_job_and_persist(aid)
    if out.get("error"):
        code = "DB_ERROR" if int(out.get("http") or 500) == 500 else "NOT_FOUND"
        if int(out.get("http") or 400) == 400:
            code = "BAD_REQUEST"
        return (
            jsonify(
                {
                    "code": code,
                    "message": str(out.get("message") or "同步失败"),
                    "data": out.get("data"),
                    "trace_id": str(out.get("trace_id") or uuid.uuid4().hex),
                }
            ),
            int(out.get("http") or 500),
        )
    return jsonify({"code": 0, "message": "ok（本地阅卷已同步）", "data": out.get("data"), "trace_id": out.get("trace_id")}), 200


@bp.post("/api/exam-center/teacher/local-exam/attempts/<attempt_id>/retry-grading")
@page13_access_required
def api_exam_teacher_retry_local_exam_grading(attempt_id: str):
    """统计端/老师端（page13）：对阅卷中或阅卷失败的本地考试重新发起上游主观判分任务。"""
    aid = (attempt_id or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 attempt_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    att = ExamAttempt.query.filter_by(attempt_id=aid).first()
    if not att:
        return jsonify({"code": "NOT_FOUND", "message": "本地考试 attempt 不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404
    gj = ExamGradingJob.query.filter_by(attempt_id=aid).first()
    ok_elig, msg_elig = _local_exam_admin_regrade_allowed(att, gj)
    if not ok_elig:
        return jsonify({"code": "BAD_STATE", "message": msg_elig, "data": None, "trace_id": uuid.uuid4().hex}), 400
    ast0 = str(att.state or "").strip().lower()
    if ast0 != "grading":
        att.state = "grading"
        db.session.add(att)
    try:
        rows = ExamAttemptItem.query.filter_by(attempt_id=aid).all()
        enq = _enqueue_local_exam_ai_grading(att, rows)
        if not enq.get("ok"):
            db.session.rollback()
            return jsonify(
                {
                    "code": "BAD_REQUEST",
                    "message": str(enq.get("reason") or "无法入队"),
                    "data": None,
                    "trace_id": uuid.uuid4().hex,
                }
            ), 400
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({"code": "DB_ERROR", "message": str(e), "data": None, "trace_id": uuid.uuid4().hex}), 500
    return jsonify(
        {
            "code": 0,
            "message": "ok（已重新发起主观题判分）",
            "data": {
                "attempt_id": aid,
                "state": att.state,
                "upstream_job_id": str(enq.get("job_id") or "").strip() or None,
                "job_status": enq.get("job_status"),
                "http_status": enq.get("http_status"),
            },
            "trace_id": uuid.uuid4().hex,
        }
    ), 200


@bp.get("/api/exam-center/student/attempts/<attempt_id>")
@login_required
def api_exam_student_attempt_detail_local(attempt_id: str):
    aid = (attempt_id or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 attempt_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    uid = str(session.get("user_id") or "").strip()
    att = ExamAttempt.query.filter_by(attempt_id=aid).first()
    if not att or str(att.user_id or "").strip() != uid:
        return jsonify({"code": "NOT_FOUND", "message": "attempt 不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404
    items = _local_exam_attempt_items_payload(aid, att=att) or []
    return jsonify(
        {
            "code": 0,
            "message": "ok",
            "data": {
                "attempt": {
                    "attempt_id": aid,
                    "assignment_id": att.assignment_id,
                    "state": att.state,
                    "exam_track": att.exam_track,
                    "started_at": att.started_at.isoformat(timespec="seconds") if att.started_at else None,
                    "submitted_at": att.submitted_at.isoformat(timespec="seconds") if att.submitted_at else None,
                    "score": att.score,
                    "total_score": att.total_score,
                    "correct_count": att.correct_count,
                    "wrong_count": att.wrong_count,
                },
                "items": items,
            },
            "trace_id": uuid.uuid4().hex,
        }
    ), 200


@bp.get("/api/exam-center/activity/<activity_id>")
@_page13_or_login_required
def api_exam_activity_detail(activity_id: str):
    aid = (activity_id or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 activity_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    row = ExamCenterActivity.query.filter_by(id=aid).first()
    if not row:
        return jsonify({"code": "NOT_FOUND", "message": "记录不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404
    from .authz import is_exam_center_staff

    if not is_exam_center_staff():
        uid = str(session.get("user_id") or "").strip()
        if not uid or uid != str(row.user_id or ""):
            return jsonify({"code": "FORBIDDEN", "message": "无权查看该记录", "data": None, "trace_id": uuid.uuid4().hex}), 403
    elif not _exam_activity_in_staff_scope(row):
        return jsonify({"code": "NOT_FOUND", "message": "记录不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404
    det = ExamCenterActivityDetail.query.filter_by(activity_id=aid).first()
    pass_score = _exam_pass_score()
    return jsonify(
        {
            "code": 0,
            "message": "ok",
            "data": {
                "activity": {
                    "id": row.id,
                    "created_at": row.created_at.isoformat() if row.created_at else "",
                    "user_id": row.user_id,
                    "username": row.username,
                    "display_name": row.display_name,
                    "mode": row.mode,
                    "exam_track": row.exam_track,
                    "set_id": row.set_id,
                    "assignment_id": row.assignment_id,
                    "assignment_label": row.assignment_label,
                    "attempt_id": row.attempt_id,
                    "result_summary": row.result_summary,
                },
                "detail": {
                    "score": det.score if det else None,
                    "total_score": det.total_score if det else None,
                    "correct_count": det.correct_count if det else None,
                    "wrong_count": det.wrong_count if det else None,
                    "pass_score": pass_score,
                    "passed": (det.score is not None and float(det.score) >= pass_score) if det and det.score is not None else None,
                    "weakness": det.weakness if det else None,
                    "recommendation": det.recommendation if det else None,
                    "upstream_payload": det.upstream_payload if det else None,
                },
                # attempt_items 改为异步拉取：避免详情接口阻塞导致弹窗长时间“加载中…”
                "attempt_items": None,
            },
            "trace_id": uuid.uuid4().hex,
        }
    ), 200


@bp.get("/api/exam-center/activity/<activity_id>/attempt-items")
@_page13_or_login_required
def api_exam_activity_attempt_items(activity_id: str):
    """异步拉取答题明细：前端弹窗先展示基础信息，再调用本接口补全题目对错/答案等。"""
    aid = (activity_id or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 activity_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    row = ExamCenterActivity.query.filter_by(id=aid).first()
    if not row:
        return jsonify({"code": "NOT_FOUND", "message": "记录不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404
    from .authz import is_exam_center_staff

    if not is_exam_center_staff():
        uid = str(session.get("user_id") or "").strip()
        if not uid or uid != str(row.user_id or ""):
            return jsonify({"code": "FORBIDDEN", "message": "无权查看该记录", "data": None, "trace_id": uuid.uuid4().hex}), 403
    att_id = str(row.attempt_id or "").strip()
    if not att_id:
        return jsonify({"code": 0, "message": "ok（无 attempt_id）", "data": {"items": []}, "trace_id": uuid.uuid4().hex}), 200

    det_snap = ExamCenterActivityDetail.query.filter_by(activity_id=aid).first()
    local_items = _local_exam_attempt_items_payload(att_id)
    if local_items is not None:
        return jsonify(
            {
                "code": 0,
                "message": "ok（本地考试明细）",
                "data": {"items": local_items},
                "trace_id": uuid.uuid4().hex,
            }
        ), 200

    st, pl = _quiz_api_call(
        f"quiz/attempts/{att_id}/answers",
        method="GET",
        query=request.args.to_dict(),
        timeout_seconds=_quiz_attempt_answers_quick_timeout_seconds(),
    )
    if 200 <= int(st) < 300 and isinstance(pl, dict):
        root = pl.get("data") if isinstance(pl.get("data"), dict) else None
        if isinstance(root, dict) and isinstance(root.get("data"), dict):
            root = root.get("data")
        items = root.get("items") if isinstance(root, dict) else None
        if not isinstance(items, list):
            items = []
        return jsonify({"code": 0, "message": "ok", "data": {"items": items}, "trace_id": uuid.uuid4().hex}), 200

    fb = _items_from_activity_detail_snapshot(det_snap)
    if fb:
        return jsonify(
            {
                "code": 0,
                "message": "ok（上游明细不可用，已降级为本地提交快照）",
                "data": {"items": fb},
                "trace_id": uuid.uuid4().hex,
                "request": {"url": "", "method": "GET", "upstreamPath": f"quiz/attempts/{att_id}/answers → snapshot"},
            }
        ), 200
    return jsonify({"code": "UPSTREAM_ERROR", "message": "上游明细不可用", "data": pl, "trace_id": uuid.uuid4().hex}), 502


@bp.get("/api/exam-center/stats/mode")
@page13_access_required
def api_exam_stats_mode():
    """按考试/练习维度的统计（本地聚合），用于第三块看板。"""
    mode = str(request.args.get("mode") or "").strip().lower()
    if mode not in {"exam", "practice"}:
        return jsonify({"code": "BAD_REQUEST", "message": "mode 仅支持 exam/practice", "data": None, "trace_id": uuid.uuid4().hex}), 400
    pass_score = _exam_pass_score()
    rows = _scope_exam_activity_query(
        ExamCenterActivity.query.filter_by(mode=mode)
    ).order_by(ExamCenterActivity.created_at.desc()).all()
    ids = [str(r.id) for r in rows if getattr(r, "id", None)]
    det_map: dict[str, ExamCenterActivityDetail] = {}
    if ids:
        ds = ExamCenterActivityDetail.query.filter(ExamCenterActivityDetail.activity_id.in_(ids)).all()
        det_map = {str(d.activity_id): d for d in ds}
    pass_count = 0
    fail_count = 0
    graded = 0
    for r in rows:
        d = det_map.get(str(r.id))
        if not d or d.score is None:
            continue
        graded += 1
        if float(d.score) >= pass_score:
            pass_count += 1
        else:
            fail_count += 1
    pass_rate = round((pass_count * 100.0 / graded), 2) if graded > 0 else None
    return jsonify(
        {
            "code": 0,
            "message": "ok",
            "data": {
                "mode": mode,
                "total_count": len(rows),
                "graded_count": graded,
                "pass_score": pass_score,
                "pass_count": pass_count,
                "fail_count": fail_count,
                "pass_rate_percent": pass_rate,
            },
            "trace_id": uuid.uuid4().hex,
        }
    ), 200


@bp.delete("/api/exam-center/activity/<activity_id>")
@page13_access_required
def api_exam_activity_delete(activity_id: str):
    """老师端删除学生练习/考试记录（超级管理员 / 项目管理员）。"""
    from .authz import is_page13_super_admin, is_project_admin

    aid = (activity_id or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 activity_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    if not (is_page13_super_admin() or is_project_admin()):
        return jsonify(
            {"code": "FORBIDDEN", "message": "仅超级管理员或项目管理员可删除记录", "data": None, "trace_id": uuid.uuid4().hex}
        ), 403
    row = ExamCenterActivity.query.filter_by(id=aid).first()
    if not row:
        return jsonify({"code": "NOT_FOUND", "message": "记录不存在", "data": None, "trace_id": uuid.uuid4().hex}), 404
    if not _exam_activity_deletable_by_staff(row):
        msg = (
            "无权删除该记录（项目管理员仅可删除所属项目组学员的记录）"
            if is_project_admin() and not is_page13_super_admin()
            else "无权删除该记录（不在当前公司范围内）"
        )
        return jsonify({"code": "FORBIDDEN", "message": msg, "data": None, "trace_id": uuid.uuid4().hex}), 403
    try:
        attempt_id = str(getattr(row, "attempt_id", None) or "").strip()
        ExamCenterActivityDetail.query.filter_by(activity_id=aid).delete()
        ExamCenterActivity.query.filter_by(id=aid).delete()
        if attempt_id:
            from .models import ExamAttempt, ExamAttemptItem, ExamGradingJob

            ExamAttemptItem.query.filter_by(attempt_id=attempt_id).delete()
            ExamGradingJob.query.filter_by(attempt_id=attempt_id).delete()
            ExamAttempt.query.filter_by(attempt_id=attempt_id).delete()
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({"code": "DB_ERROR", "message": f"删除失败：{e}", "data": None, "trace_id": uuid.uuid4().hex}), 500
    return jsonify({"code": 0, "message": "ok", "data": {"deleted": 1, "activity_id": aid}, "trace_id": uuid.uuid4().hex}), 200


@bp.get("/api/exam-center/stats/overview")
@page13_access_required
def api_exam_stats_overview():
    local = _local_stats_overview_all()
    return jsonify(
        {
            "code": 0,
            "message": "ok（本地统计）",
            "data": local,
            "trace_id": uuid.uuid4().hex,
            "request": {"url": "", "method": "GET", "upstreamPath": "local aggregate"},
        }
    ), 200


@bp.get("/api/exam-center/stats/student/<student_id>")
@page13_access_required
def api_exam_stats_student(student_id: str):
    from .authz import user_in_exam_team_scope

    uid = str(student_id or "").strip()
    if not uid or not user_in_exam_team_scope(uid):
        return jsonify({"code": "NOT_FOUND", "message": "学生不存在或无权查看", "data": None, "trace_id": uuid.uuid4().hex}), 404
    local = _local_stats_for_student(uid)
    return jsonify(
        {
            "code": 0,
            "message": "ok（本地统计）",
            "data": local,
            "trace_id": uuid.uuid4().hex,
            "request": {"url": "", "method": "GET", "upstreamPath": "local aggregate"},
        }
    ), 200


@bp.get("/api/exam-center/stats/students")
@page13_access_required
def api_exam_stats_students():
    """按学生多维统计表格（本地聚合）。供统计看板一次展示全部学生。"""
    rows = _local_stats_all_students_rows()
    return jsonify(
        {
            "code": 0,
            "message": "ok",
            "data": {
                "rows": rows,
                "pass_score": _exam_pass_score(),
                "count": len(rows),
            },
            "trace_id": uuid.uuid4().hex,
            "request": {"url": "", "method": "GET", "upstreamPath": "local aggregate → exam_center_activities"},
        }
    ), 200


@bp.get("/api/exam-center/stats/students-by-mode")
@bp.get("/api/exam-center/stats/students_by_mode")
@page13_access_required
def api_exam_stats_students_by_mode():
    """按学生 × 考试/练习 分组统计（本地聚合）。"""
    rows = _local_stats_rows_student_by_mode()
    return jsonify(
        {
            "code": 0,
            "message": "ok",
            "data": {
                "rows": rows,
                "pass_score": _exam_pass_score(),
                "count": len(rows),
            },
            "trace_id": uuid.uuid4().hex,
            "request": {"url": "", "method": "GET", "upstreamPath": "local aggregate student×mode → exam_center_activities"},
        }
    ), 200


@bp.get("/api/exam-center/stats/exam/<assignment_id>")
@page13_access_required
def api_exam_stats_exam(assignment_id: str):
    aid = str(assignment_id or "").strip()
    if not aid:
        return jsonify({"code": "BAD_REQUEST", "message": "缺少 assignment_id", "data": None, "trace_id": uuid.uuid4().hex}), 400
    rows = _scope_exam_activity_query(
        ExamCenterActivity.query.filter_by(assignment_id=aid)
    ).order_by(ExamCenterActivity.created_at.desc()).all()
    ids = [str(r.id) for r in rows if getattr(r, "id", None)]
    det_map: dict[str, ExamCenterActivityDetail] = {}
    if ids:
        ds = ExamCenterActivityDetail.query.filter(ExamCenterActivityDetail.activity_id.in_(ids)).all()
        det_map = {str(d.activity_id): d for d in ds}
    pass_score = _exam_pass_score()
    pass_count = 0
    fail_count = 0
    submitted = 0
    by_student: dict[str, dict[str, Any]] = {}
    for r in rows:
        uid = str(r.user_id or "").strip()
        d = det_map.get(str(r.id))
        submitted += 1
        if d and d.score is not None:
            if float(d.score) >= pass_score:
                pass_count += 1
            else:
                fail_count += 1
        if uid and uid not in by_student:
            by_student[uid] = _local_stats_for_student(uid)
    graded = pass_count + fail_count
    pass_rate = round((pass_count * 100.0 / graded), 2) if graded > 0 else None
    assign_row = ExamCenterAssignment.query.filter_by(assignment_id=aid).first()
    due_end = getattr(assign_row, "due_at", None) if assign_row else None
    deadline_completion = _deadline_completion_from_exam_rows(rows, due_end)
    return jsonify(
        {
            "code": 0,
            "message": "ok（本地统计兜底）",
            "data": {
                "assignment_id": aid,
                "submitted_count": submitted,
                "graded_count": graded,
                "pass_score": pass_score,
                "pass_count": pass_count,
                "fail_count": fail_count,
                "pass_rate_percent": pass_rate,
                "students": list(by_student.values()),
                "deadline_completion": deadline_completion,
            },
            "trace_id": uuid.uuid4().hex,
            "request": {"url": "", "method": "GET", "upstreamPath": "local aggregate"},
        }
    ), 200


@bp.get("/api/exam-center/health")
def api_exam_center_health():
    """
    用于排查上游连通性：优先请求 /health；若上游没有该路径，会返回上游错误信息。
    """
    # 不要用 3s 硬超时：本机/Windows/反代首次握手或 API 冷启动时容易误判超时；
    # 与 QUIZ_API_TIMEOUT_SECONDS 对齐，但健康检查仍限制上限，避免按钮长时间无反馈。
    cfg = _quiz_api_timeout_seconds()
    health_timeout = min(15, max(8, cfg))
    status, payload = _quiz_api_call(
        "health",
        method="GET",
        query=request.args.to_dict(),
        timeout_seconds=health_timeout,
    )
    return jsonify(payload), status


# ---------- 用户管理 API（页面1管理账号） ----------


@bp.get("/api/task-author-candidates")
@page13_access_required
def api_task_author_candidates():
    """任务录入编写人员下拉：项目绑定项目组成员 + 当前登录用户。"""
    from .authz import is_page13_super_admin, project_in_scope, rbac_enforced
    from .models import Project
    from .user_access import (
        _resolve_project_for_author_pick,
        _serialize_task_author_user,
        list_task_author_candidates,
    )

    project_id = (
        request.args.get("projectId") or request.args.get("project_id") or ""
    ).strip()
    team_id = (request.args.get("teamId") or request.args.get("team_id") or "").strip()
    if project_id:
        proj = _resolve_project_for_author_pick(project_id)
        if proj is None:
            return jsonify({"message": "项目不存在"}), 404
        if rbac_enforced() and not is_page13_super_admin() and not project_in_scope(proj):
            return jsonify({"message": "无权访问该项目"}), 403
    users = list_task_author_candidates(project_id=project_id or None, team_id=team_id or None)
    return jsonify({"users": [_serialize_task_author_user(u) for u in users]})


@bp.get("/api/users/feature-permission-schema")
@page13_access_required
def api_users_feature_permission_schema():
    """账号新建/编辑与批量功能权限共用的字段分组（勿在 app.js 重复维护）。"""
    from .user_feature_permissions import feature_permission_schema_for_client

    return jsonify({"ok": True, **feature_permission_schema_for_client()})


@bp.get("/api/users")
@page13_access_required
def api_users_list():
    from .user_access import serialize_user_access

    users = User.query.order_by(User.created_at.desc()).all()
    return jsonify({
        "users": [
            {
                "id": u.id,
                "username": u.username,
                "displayName": u.display_name,
                "mobile": getattr(u, "mobile", None) or None,
                "adminRole": (getattr(u, "admin_role", None) or "none").strip() or "none",
                "createdAt": u.created_at.isoformat() if u.created_at else None,
                **serialize_user_access(u),
            }
            for u in users
        ]
    })


@bp.post("/api/users")
@super_admin_required
def api_users_create():
    data = request.get_json(force=True) or {}
    username = (data.get("username") or "").strip()
    password = (data.get("password") or "").strip()
    display_name = (data.get("displayName") or "").strip() or None
    from .notify_content import format_mobile_for_storage

    mobile = format_mobile_for_storage(data.get("mobile"))
    from .models import ADMIN_ROLES

    admin_role = (data.get("adminRole") or "none").strip()
    if admin_role not in ADMIN_ROLES:
        admin_role = "none"
    if not username or not password:
        return jsonify({"message": "用户名和密码不能为空"}), 400
    existing = User.query.filter_by(username=username).first()
    if existing:
        return jsonify({"message": "用户名已存在"}), 409
    user = User(
        username=username,
        display_name=display_name,
        mobile=mobile,
        is_admin=False,
        admin_role=admin_role,
    )
    user.set_password(password)
    db.session.add(user)
    db.session.flush()
    from .user_access import (
        ensure_role_access_requirements,
        apply_user_access_fields,
        serialize_user_access,
    )

    try:
        ensure_role_access_requirements(user, data)
        apply_user_access_fields(user, data)
    except ValueError as e:
        db.session.rollback()
        return jsonify({"message": str(e)}), 400
    db.session.add(user)
    db.session.commit()
    return jsonify({
        "message": "用户创建成功",
        "user": {
            "id": user.id,
            "username": user.username,
            "displayName": user.display_name,
            "mobile": user.mobile,
            "adminRole": getattr(user, "admin_role", None) or "none",
            **serialize_user_access(user),
        },
    })


@bp.patch("/api/users/<user_id>")
@super_admin_required
def api_users_update(user_id: str):
    """更新用户显示名称、手机号（钉钉 @ 用）、分级角色与权限"""
    user = User.query.get(user_id)
    if not user:
        return jsonify({"message": "用户不存在"}), 404
    data = request.get_json(force=True) or {}
    from .notify_content import format_mobile_for_storage

    if "displayName" in data:
        user.display_name = (data["displayName"] or "").strip() or None
    if "mobile" in data:
        user.mobile = format_mobile_for_storage(data.get("mobile"))
    if "adminRole" in data:
        from .models import ADMIN_ROLE_COMPANY, ADMIN_ROLES

        role = (data.get("adminRole") or "none").strip()
        user.admin_role = role if role in ADMIN_ROLES else "none"
        if role == ADMIN_ROLE_COMPANY:
            user.can_access_company_registry = True
        else:
            user.can_access_company_registry = False
    from .user_access import (
        ensure_role_access_requirements,
        apply_user_access_fields,
        serialize_user_access,
    )

    try:
        ensure_role_access_requirements(user, data)
        apply_user_access_fields(user, data)
    except ValueError as e:
        return jsonify({"message": str(e)}), 400
    db.session.add(user)
    db.session.commit()
    db.session.expire(user)
    if session.get("user_id") == user.id:
        from .authz import _refresh_session_from_user

        _refresh_session_from_user(user)
    else:
        session.pop("country_scopes", None)
    return jsonify({
        "message": "已更新",
        "user": {
            "id": user.id,
            "username": user.username,
            "displayName": user.display_name,
            "mobile": user.mobile,
            "adminRole": getattr(user, "admin_role", None) or "none",
            **serialize_user_access(user),
        },
    })


@bp.post("/api/users/batch-feature-permissions")
@super_admin_required
def api_users_batch_feature_permissions():
    """批量设置多个账号的功能权限（仅合并指定项，未选字段保持各账号原值）。"""
    data = request.get_json(force=True) or {}
    user_ids = data.get("userIds")
    if not isinstance(user_ids, list) or not user_ids:
        return jsonify({"message": "请至少选择一个账号"}), 400
    from .user_feature_permissions import (
        apply_feature_permission_patch,
        parse_batch_feature_permission_patch,
        read_user_feature_permissions,
        write_user_feature_permissions,
    )

    try:
        patch = parse_batch_feature_permission_patch(data)
    except ValueError as exc:
        return jsonify({"message": str(exc)}), 400
    if not patch:
        return jsonify({"message": "请至少选择一项要修改的功能权限"}), 400

    ids = [str(x).strip() for x in user_ids if str(x).strip()]
    users = User.query.filter(User.id.in_(ids)).all() if ids else []
    found = {u.id for u in users}
    missing = [i for i in ids if i not in found]
    if missing:
        return jsonify({"message": f"账号不存在：{missing[0]}"}), 404

    updated = 0
    for user in users:
        merged = apply_feature_permission_patch(read_user_feature_permissions(user), patch)
        write_user_feature_permissions(user, merged or None)
        db.session.add(user)
        updated += 1
    db.session.commit()

    cur_uid = session.get("user_id")
    if cur_uid and str(cur_uid) in found:
        from .authz import _refresh_session_from_user

        u = User.query.get(cur_uid)
        if u:
            _refresh_session_from_user(u)

    return jsonify({"message": f"已更新 {updated} 个账号的功能权限", "updated": updated})


@bp.delete("/api/users/<user_id>")
@super_admin_required
def api_users_delete(user_id: str):
    from .models import UserCountryScope, UserOrganizationMembership, UserTeamMembership

    user = User.query.get(user_id)
    if not user:
        return jsonify({"message": "用户不存在"}), 404
    UserTeamMembership.query.filter_by(user_id=user_id).delete(synchronize_session=False)
    UserOrganizationMembership.query.filter_by(user_id=user_id).delete(
        synchronize_session=False
    )
    UserCountryScope.query.filter_by(user_id=user_id).delete(synchronize_session=False)
    db.session.delete(user)
    db.session.commit()
    return jsonify({"message": "用户已删除"})


@bp.get("/api/project-teams")
@page13_access_required
def api_project_teams_list():
    """页面4（超级管理员）：项目组列表（不依赖公司总览功能开关）。"""
    from .models import ProjectTeam

    rows = ProjectTeam.query.order_by(ProjectTeam.sort_order.asc(), ProjectTeam.name.asc()).all()
    return jsonify(
        [
            {
                "id": t.id,
                "name": t.name,
                "sortOrder": t.sort_order,
                "isActive": bool(t.is_active),
            }
            for t in rows
        ]
    )


@bp.get("/api/registered-countries")
@page13_access_required
def api_registered_countries_suggest():
    """注册国家字典（只读候选，维护请至页面0）。"""
    from .registered_countries import list_registered_countries

    return jsonify({"countries": list_registered_countries()})


# ---------- 配置项 API ----------

def _normalize_task_type_category(raw) -> str:
    """归一化任务类型一级分类：仅 file/matter；未识别值回退为 file。"""
    from .models import TASK_TYPE_CATEGORIES, TASK_TYPE_CATEGORY_FILE

    v = (str(raw or "").strip().lower())
    if v in TASK_TYPE_CATEGORIES:
        return v
    return TASK_TYPE_CATEGORY_FILE


@bp.get("/api/configs/task-types")
@_page13_or_login_required
def api_task_types():
    """获取任务类型配置列表（含一级分类 category：file=文件型；matter=事项型）。"""
    items = TaskTypeConfig.query.filter_by(is_active=True).order_by(TaskTypeConfig.sort_order).all()
    return jsonify({
        "taskTypes": [
            {
                "id": t.id,
                "name": t.name,
                "category": _normalize_task_type_category(getattr(t, "category", None)),
            }
            for t in items
        ]
    })


@bp.post("/api/configs/task-types")
@super_admin_required
def api_task_types_create():
    """新增任务类型；可选 category 字段（file/matter，默认 file）。"""
    data = request.get_json(force=True) or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"message": "名称不能为空"}), 400
    category = _normalize_task_type_category(data.get("category"))
    existing = TaskTypeConfig.query.filter_by(name=name).first()
    if existing:
        return jsonify({"message": "该类型已存在"}), 409
    max_order = db.session.query(db.func.max(TaskTypeConfig.sort_order)).scalar() or 0
    item = TaskTypeConfig(name=name, sort_order=max_order + 1, category=category)
    db.session.add(item)
    db.session.commit()
    return jsonify({
        "message": "创建成功",
        "id": item.id,
        "name": item.name,
        "category": item.category,
    })


@bp.patch("/api/configs/task-types/<item_id>")
@super_admin_required
def api_task_types_update(item_id: str):
    """更新任务类型分类（用于把已有类型在「文件型/事项型」间切换）。"""
    item = TaskTypeConfig.query.get(item_id)
    if not item:
        return jsonify({"message": "不存在"}), 404
    data = request.get_json(force=True) or {}
    if "category" in data:
        item.category = _normalize_task_type_category(data.get("category"))
    db.session.add(item)
    db.session.commit()
    return jsonify({"message": "已更新", "id": item.id, "name": item.name, "category": item.category})


@bp.delete("/api/configs/task-types/<item_id>")
@super_admin_required
def api_task_types_delete(item_id: str):
    """删除任务类型"""
    item = TaskTypeConfig.query.get(item_id)
    if not item:
        return jsonify({"message": "不存在"}), 404
    db.session.delete(item)
    db.session.commit()
    return jsonify({"message": "已删除"})


@bp.get("/api/configs/completion-statuses")
@_page13_or_login_required
def api_completion_statuses():
    """获取完成状态配置列表"""
    items = CompletionStatusConfig.query.filter_by(is_active=True).order_by(CompletionStatusConfig.sort_order).all()
    return jsonify({
        "completionStatuses": [{"id": s.id, "name": s.name} for s in items]
    })


@bp.post("/api/configs/completion-statuses")
@super_admin_required
def api_completion_statuses_create():
    """新增完成状态"""
    data = request.get_json(force=True) or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"message": "名称不能为空"}), 400
    existing = CompletionStatusConfig.query.filter_by(name=name).first()
    if existing:
        return jsonify({"message": "该状态已存在"}), 409
    max_order = db.session.query(db.func.max(CompletionStatusConfig.sort_order)).scalar() or 0
    item = CompletionStatusConfig(name=name, sort_order=max_order + 1)
    db.session.add(item)
    db.session.commit()
    return jsonify({"message": "创建成功", "id": item.id, "name": item.name})


@bp.delete("/api/configs/completion-statuses/<item_id>")
@super_admin_required
def api_completion_statuses_delete(item_id: str):
    """删除完成状态"""
    item = CompletionStatusConfig.query.get(item_id)
    if not item:
        return jsonify({"message": "不存在"}), 404
    db.session.delete(item)
    db.session.commit()
    return jsonify({"message": "已删除"})


@bp.get("/api/configs/audit-statuses")
@_page13_or_login_required
def api_audit_statuses():
    """获取审核状态配置列表（页面1使用）"""
    items = AuditStatusConfig.query.filter_by(is_active=True).order_by(AuditStatusConfig.sort_order).all()
    return jsonify({
        "auditStatuses": [{"id": s.id, "name": s.name} for s in items]
    })


@bp.post("/api/configs/audit-statuses")
@super_admin_required
def api_audit_statuses_create():
    """新增审核状态"""
    data = request.get_json(force=True) or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"message": "名称不能为空"}), 400
    existing = AuditStatusConfig.query.filter_by(name=name).first()
    if existing:
        return jsonify({"message": "该状态已存在"}), 409
    max_order = db.session.query(db.func.max(AuditStatusConfig.sort_order)).scalar() or 0
    item = AuditStatusConfig(name=name, sort_order=max_order + 1)
    db.session.add(item)
    db.session.commit()
    return jsonify({"message": "创建成功", "id": item.id, "name": item.name})


@bp.delete("/api/configs/audit-statuses/<item_id>")
@super_admin_required
def api_audit_statuses_delete(item_id: str):
    """删除审核状态"""
    item = AuditStatusConfig.query.get(item_id)
    if not item:
        return jsonify({"message": "不存在"}), 404
    db.session.delete(item)
    db.session.commit()
    return jsonify({"message": "已删除"})


# ---------- 上传与任务管理 API ----------

@bp.post("/api/upload")
@page13_access_required
def api_upload():
    project_id = (request.form.get("projectId") or "").strip() or None
    project_name = request.form.get("projectName", "").strip()
    project_code = request.form.get("projectCode", "").strip() or None
    file_name = request.form.get("fileName", "").strip()
    task_type = request.form.get("taskType", "").strip() or None
    author = request.form.get("author", "").strip()
    notes = request.form.get("notes", "").strip() or None
    project_notes = request.form.get("projectNotes", "").strip() or None
    replace = str(request.form.get("replace", "")).strip().lower() in ("true", "1", "yes")
    file = request.files.get("file")
    template_links_raw = request.form.get("templateLinks")
    template_links = (
        (template_links_raw or "").strip() or None
        if template_links_raw is not None
        else None
    )
    if template_links:
        template_links = _normalize_template_links(template_links) or None
    assignee_name = request.form.get("assigneeName", "").strip() or None
    due_date_str = request.form.get("dueDate", "").strip() or None
    business_side = request.form.get("businessSide", "").strip() or None
    product = request.form.get("product", "").strip() or None
    country = request.form.get("country", "").strip() or None
    registered_product_name = request.form.get("registeredProductName", "").strip() or None
    model = request.form.get("model", "").strip() or None
    registration_version = request.form.get("registrationVersion", "").strip() or None
    file_version = request.form.get("fileVersion", "").strip() or None
    document_display_date_str = request.form.get("documentDisplayDate", "").strip() or None
    reviewer = request.form.get("reviewer", "").strip() or None
    approver = request.form.get("approver", "").strip() or None
    belonging_module = request.form.get("belongingModule", "").strip() or None
    displayed_author = request.form.get("displayedAuthor", "").strip() or None

    if not project_name or not file_name or not author:
        return (
            jsonify({"message": "项目名称、文件名称、编写人员为必填项。"}),
            400,
        )
    # 链接和文件均可为空保存，后续可在页面2补充链接

    # 确保项目元数据存在（默认：中优先级/进行中）
    if project_id:
        p = Project.query.get(project_id)
        if p:
            project_name = _project_display_label(p)
        else:
            project_id = None
    _ensure_project_row(project_name)
    from .project_teams import resolve_organization_id_for_project_upload

    upload_org_id = resolve_organization_id_for_project_upload(project_id=project_id)

    upload_record_id = (request.form.get("uploadRecordId") or request.form.get("uploadId") or "").strip()
    existing: UploadRecord | None = None
    # 页面1「编辑任务」按记录 ID 替换（可改名后仍更新同一条，并触发 FTP）
    if upload_record_id and replace:
        existing = UploadRecord.query.get(upload_record_id)
        if not existing:
            return jsonify({"message": "未找到要更新的记录"}), 404
    if existing is None:
        existing = UploadRecord.query.filter_by(
            project_name=project_name, file_name=file_name, task_type=task_type, author=author
        ).first()

    if existing and not replace:
        return (
            jsonify(
                {
                    "message": f"存在同名项目+文件+类型+编写人({task_type or '无'}/{author})，是否需要替换原有内容？",
                    "needsConfirmation": True,
                }
            ),
            409,
        )

    uploads_dir = Path(current_app.config["UPLOAD_FOLDER"])
    stored_file_name = None
    storage_path = None
    original_file_name = None
    placeholders = []
    template_file_blob = None

    if file and file.filename:
        stored_file_name, storage_path = _save_file(file, uploads_dir)
        original_file_name = file.filename
        try:
            template_file_blob, placeholders = resolve_task_template_from_saved_path(
                Path(storage_path), file_name_hint=file_name
            )
        except Exception as exc:
            Path(storage_path).unlink(missing_ok=True)
            return jsonify({"message": f"解析模板失败：{exc}"}), 400
        Path(storage_path).unlink(missing_ok=True)
        storage_path = None
    elif template_links:
        links = [line.strip() for line in template_links.split("\n") if line.strip()]
        if links:
            first_link = links[0].lower()
            is_direct_docx = first_link.endswith('.docx') or first_link.endswith('.doc')
            if is_direct_docx:
                try:
                    temp_path = uploads_dir / f"temp_{now_local().strftime('%Y%m%d%H%M%S%f')}.docx"
                    download_template_from_url(links[0], str(temp_path))
                    placeholders = extract_placeholders(str(temp_path))
                    temp_path.unlink(missing_ok=True)
                except Exception as exc:
                    temp_path.unlink(missing_ok=True) if 'temp_path' in dir() else None
                    return jsonify({"message": f"下载或解析模板链接失败：{exc}"}), 400

    due_date = None
    if due_date_str:
        try:
            due_date = datetime.strptime(due_date_str, "%Y-%m-%d").date()
        except ValueError:
            return jsonify({"message": "截止日期格式应为 YYYY-MM-DD"}), 400
    document_display_date = None
    if document_display_date_str:
        try:
            document_display_date = datetime.strptime(document_display_date_str, "%Y-%m-%d").date()
        except ValueError:
            document_display_date = None

    if existing and replace:
        fm = request.form
        # 批量保存第二次请求（replace=true）常为「只 append 非空字段」；未出现的键不得写成 None，否则会误清空库内已有数据。
        def _sent(k: str) -> bool:
            return k in fm

        replacing_template_file = bool(file and file.filename)
        if replacing_template_file:
            _unlink_task_template_cache_files(existing.id)
            if existing.storage_path:
                previous_path = Path(existing.storage_path)
                if previous_path.exists():
                    previous_path.unlink()

        if replacing_template_file:
            old_ftp = (getattr(existing, "ftp_path", None) or "").strip()
            if old_ftp:
                try:
                    from .ftp_store import delete_path

                    delete_path(old_ftp)
                except Exception:
                    pass
            existing.ftp_path = None
            existing.ftp_last_error = None
            existing.template_file_blob = template_file_blob
            existing.stored_file_name = stored_file_name
            existing.storage_path = None
            existing.original_file_name = original_file_name
            # 页面1/2：上传文件覆盖时来源改为文件，清空文档链接
            existing.template_links = None
        else:
            existing.stored_file_name = stored_file_name or existing.stored_file_name
            existing.storage_path = storage_path if storage_path else existing.storage_path
            existing.original_file_name = original_file_name or existing.original_file_name
            if _sent("templateLinks"):
                existing.template_links = template_links
        existing.author = author
        if _sent("taskType"):
            existing.task_type = task_type
        if _sent("notes"):
            existing.notes = notes
        if _sent("projectNotes"):
            existing.project_notes = project_notes
        existing.project_name = project_name
        existing.project_id = project_id
        existing.organization_id = upload_org_id
        existing.file_name = file_name
        if (file and file.filename) or placeholders:
            existing.placeholders = placeholders
        if _sent("assigneeName"):
            existing.assignee_name = assignee_name
        if _sent("dueDate"):
            existing.due_date = due_date
        # 仅在上传新模板文件时重置完成状态；页面1 仅改任务类型等元数据时保留原状态
        if replacing_template_file:
            existing.task_status = "pending"
            existing.completion_status = None
            existing.quick_completed = False
        if _sent("businessSide"):
            existing.business_side = business_side
        if _sent("product"):
            existing.product = product
        if _sent("country"):
            existing.country = country
        if _sent("projectCode"):
            existing.project_code = project_code
        if _sent("fileVersion"):
            existing.file_version = file_version
        if _sent("documentDisplayDate"):
            existing.document_display_date = document_display_date
        if _sent("reviewer"):
            existing.reviewer = reviewer
        if _sent("approver"):
            existing.approver = approver
        if _sent("belongingModule"):
            existing.belonging_module = belonging_module
        if _sent("displayedAuthor"):
            existing.displayed_author = displayed_author
        if _sent("registeredProductName"):
            existing.registered_product_name = registered_product_name
        if _sent("model"):
            existing.model = model
        if _sent("registrationVersion"):
            existing.registration_version = registration_version
        if _sent("auditStatus"):
            prev_completion = existing.completion_status
            prev_audit = existing.audit_status
            ast = (request.form.get("auditStatus") or "").strip() or None
            _maybe_bump_audit_reject_count(
                existing,
                previous_completion_status=prev_completion,
                previous_audit_status=prev_audit,
                target_audit_status=ast,
            )
            existing.audit_status = ast or None
            if ast == AUDIT_REJECT_PENDING_STATUS:
                existing.completion_status = None
                existing.task_status = "pending"
                existing.quick_completed = False
        other = UploadRecord.query.filter(
            UploadRecord.project_name == existing.project_name,
            UploadRecord.file_name == existing.file_name,
            UploadRecord.author == existing.author,
            UploadRecord.id != existing.id,
        ).first()
        if other and (existing.task_type or None) == (other.task_type or None):
            return jsonify({"message": "更新后会与另一条记录（项目+文件名称+任务类型+编写人员）重复"}), 409
        summary = _prepare_summary(existing)
        summary.project_name = project_name
        summary.project_id = project_id
        summary.file_name = file_name
        summary.author = author
        if replacing_template_file:
            summary.has_generated = False
            summary.total_generate_clicks = 0
            summary.last_generated_at = None
        db.session.add(existing)
        db.session.flush()
        if replacing_template_file and existing.template_file_blob:
            _apply_task_template_ftp_after_flush(
                existing, existing.original_file_name or file_name
            )
        db.session.commit()

        try:
            _send_task_notification(existing, due_date_str)
        except Exception as exc:
            current_app.logger.warning("替换任务后发送钉钉通知失败（数据已保存）: %s", exc)

        ph_out = existing.placeholders if isinstance(existing.placeholders, list) else (placeholders or [])
        msg = (
            "已替换现有记录，状态已重置为待办。"
            if replacing_template_file
            else "任务已更新。"
        )
        return jsonify(
            {
                "message": msg,
                "record": {
                    "id": existing.id,
                    "projectName": existing.project_name,
                    "fileName": existing.file_name,
                    "taskType": existing.task_type,
                    "author": existing.author,
                    "placeholders": ph_out,
                    "assigneeName": existing.assignee_name,
                    "dueDate": due_date_str,
                    "businessSide": existing.business_side,
                    "product": existing.product,
                    "country": existing.country,
                    "registeredProductName": getattr(existing, "registered_product_name", None),
                    "model": getattr(existing, "model", None),
                    "registrationVersion": getattr(existing, "registration_version", None),
                },
            }
        )

    upload = UploadRecord(
        project_id=project_id,
        organization_id=upload_org_id,
        project_name=project_name,
        project_code=project_code,
        file_name=file_name,
        task_type=task_type,
        author=author,
        stored_file_name=stored_file_name,
        storage_path=storage_path,
        template_file_blob=template_file_blob,
        original_file_name=original_file_name,
        template_links=template_links,
        notes=notes,
        project_notes=project_notes,
        placeholders=placeholders,
        assignee_name=assignee_name,
        due_date=due_date,
        business_side=business_side,
        product=product,
        country=country,
        registered_product_name=registered_product_name,
        model=model,
        registration_version=registration_version,
        file_version=file_version,
        document_display_date=document_display_date,
        reviewer=reviewer,
        approver=approver,
        belonging_module=belonging_module,
        displayed_author=displayed_author,
    )
    db.session.add(upload)
    db.session.flush()
    summary = _prepare_summary(upload)
    summary.project_id = project_id
    db.session.add(summary)
    if template_file_blob:
        _apply_task_template_ftp_after_flush(upload, original_file_name or file_name)
    db.session.commit()

    try:
        _send_task_notification(upload, due_date_str)
    except Exception as exc:
        current_app.logger.warning("新建任务后发送钉钉通知失败（数据已保存）: %s", exc)

    return jsonify(
        {
            "message": "上传信息已保存。",
            "record": {
                "id": upload.id,
                "projectName": upload.project_name,
                "fileName": upload.file_name,
                "taskType": upload.task_type,
                "author": upload.author,
                "placeholders": placeholders,
                "assigneeName": upload.assignee_name,
                "dueDate": due_date_str,
                "businessSide": upload.business_side,
                "product": upload.product,
                "country": upload.country,
                "registeredProductName": getattr(upload, "registered_product_name", None),
                "model": getattr(upload, "model", None),
                "registrationVersion": getattr(upload, "registration_version", None),
            },
        }
    )


def _send_task_notification(upload: UploadRecord, due_date_str: str):
    """如果设置了负责人，发送钉钉通知（按任务所属项目组选择 Webhook）。"""
    if not upload.assignee_name:
        return
    from .dingtalk_team import resolve_dingtalk_credentials, resolve_team_id_by_upload

    team_id = resolve_team_id_by_upload(upload)
    webhook, secret, source = resolve_dingtalk_credentials(team_id)
    if not webhook:
        current_app.logger.warning(
            "新建任务钉钉通知跳过：未配置 Webhook team_id=%s upload=%s",
            team_id,
            upload.id,
        )
        return
    template_source = (
        "已上传文件"
        if _upload_record_has_task_file(upload)
        else f"链接({len(upload.get_template_links_list())}个)"
    )
    sent = dingtalk_service.notify_task_assigned(
        upload.assignee_name,
        f"{upload.project_name} - {upload.file_name}",
        due_date_str or "未设置",
        template_source,
        webhook=webhook,
        secret=secret,
    )
    if sent:
        upload.dingtalk_notified_at = now_local()
        db.session.add(upload)
        db.session.commit()
        current_app.logger.info(
            "新建任务钉钉通知已发送 team_id=%s source=%s assignee=%s",
            team_id,
            source,
            upload.assignee_name,
        )


def _import_header_map():
    from .task_entry_schema import import_header_map

    return import_header_map()


def _decode_import_csv_bytes(raw_bytes: bytes) -> str:
    """
    Excel「另存为 CSV」在中文 Windows 上常为 GBK/GB18030；仅用 UTF-8 会导致乱码或列错位，
    进而「文件名称」等列解析错误甚至整行被跳过。
    """
    if not raw_bytes:
        return ""
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk", "cp936"):
        try:
            return raw_bytes.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw_bytes.decode("utf-8", errors="replace")


def _parse_import_file(file_storage) -> tuple[list[dict], str]:
    """解析上传的 CSV 或 Excel，返回 (rows, error)。rows 每项为字段名->值的字典。"""
    if not file_storage or not file_storage.filename:
        return [], "请选择文件"
    fn = (file_storage.filename or "").lower()
    file_storage.stream.seek(0)
    raw_bytes = file_storage.stream.read()
    if fn.endswith(".csv"):
        if isinstance(raw_bytes, bytes):
            raw = _decode_import_csv_bytes(raw_bytes)
        else:
            raw = str(raw_bytes)
        return _parse_import_csv(raw)
    if fn.endswith(".xlsx") or fn.endswith(".xls"):
        return _parse_import_excel(raw_bytes, fn)
    return [], "仅支持 .csv 或 .xlsx 文件"


def _parse_import_csv(raw: str) -> tuple[list[dict], str]:
    reader = csv.reader(io.StringIO(raw))
    rows = list(reader)
    if not rows:
        return [], "文件为空"
    headers = [h.strip() for h in rows[0]]
    out = []
    for i, row in enumerate(rows[1:], start=2):
        if not row or all(not (c or "").strip() for c in row):
            continue
        d = {}
        for j, h in enumerate(headers):
            if j < len(row):
                val = (row[j] or "").strip()
                key = _import_header_map().get(h) or _import_header_map().get(h.strip())
                if key:
                    d[key] = val
        out.append(d)
    return out, ""


def _parse_import_excel(raw: bytes, filename: str) -> tuple[list[dict], str]:
    try:
        import openpyxl
    except ImportError:
        return [], "请安装 openpyxl 以支持 Excel 导入：pip install openpyxl"
    if not isinstance(raw, bytes):
        raw = raw.encode("utf-8") if isinstance(raw, str) else b""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        if not ws:
            return [], "Excel 无有效工作表"
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception as e:
        return [], f"Excel 解析失败：{e}"
    if not rows:
        return [], "文件为空"
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for row in rows[1:]:
        if not row or all(v is None or (str(v).strip() == "") for v in row):
            continue
        d = {}
        for j, h in enumerate(headers):
            if j < len(row):
                val = row[j]
                val = (str(val).strip() if val is not None else "") or ""
                key = _import_header_map().get(h) or _import_header_map().get(h.strip())
                if key:
                    d[key] = val
        out.append(d)
    return out, ""


def _build_import_template_csv(include_sample: bool, project_name: Optional[str] = None) -> str:
    from .task_entry_schema import build_import_template_csv

    records_for_project = None
    fallback_sample = None
    if include_sample:
        if project_name and (project_name or "").strip():
            records_for_project = (
                UploadRecord.query.filter(UploadRecord.project_name == project_name.strip())
                .order_by(UploadRecord.sort_order.asc(), UploadRecord.created_at.asc())
                .all()
            )
        else:
            fallback_sample = UploadRecord.query.order_by(UploadRecord.created_at.desc()).first()
    return build_import_template_csv(
        include_sample,
        project_name,
        records_for_project=records_for_project,
        fallback_sample_record=fallback_sample,
    )


@bp.get("/api/uploads/project-names")
@page13_access_required
def api_uploads_project_names():
    """返回已有项目名称列表，用于示例模板选择项目。"""
    include_history = str(request.args.get("includeHistory") or "").strip() in ("1", "true", "True", "yes", "on")
    proj_meta = _project_meta_map(auto_create_from_uploads=True)
    ended = {n for n, m in proj_meta.items() if (m.get("status") or "").strip().lower() == Project.STATUS_ENDED}

    names = (
        db.session.query(UploadRecord.project_name)
        .filter(UploadRecord.project_name.isnot(None), UploadRecord.project_name != "")
        .distinct()
        .all()
    )
    arr = [n[0] for n in names if (n[0] or "").strip()]
    if (not include_history) and ended:
        arr = [n for n in arr if n not in ended]

    def _k(n: str):
        m = proj_meta.get(n) or {}
        st = (m.get("status") or Project.STATUS_ACTIVE).strip().lower()
        pr = int(m.get("priority") or Project.PRIORITY_MEDIUM)
        return (1 if st == Project.STATUS_ENDED else 0, -pr, n or "")

    arr = sorted(set(arr), key=_k)
    return jsonify({"projectNames": arr})


@bp.get("/api/task-entry/import-schema")
@page13_access_required
def api_task_entry_import_schema():
    """任务录入 / 下载模板 / 导入待办 共用字段定义（表头顺序与 key）。"""
    from .task_entry_schema import import_schema_for_client

    return jsonify({"ok": True, **import_schema_for_client()})


@bp.get("/api/uploads/import-template")
@page13_access_required
def api_uploads_import_template():
    """下载导入模板：type=empty 仅表头；type=with_sample 可带 project_name 填充该项目下所有任务。"""
    try:
        template_type = (request.args.get("type") or "empty").strip().lower()
        include_sample = template_type in ("with_sample", "sample", "1", "true")
        project_name = (request.args.get("project_name") or "").strip() or None
        content = _build_import_template_csv(include_sample=include_sample, project_name=project_name)
        if include_sample and project_name:
            safe_name = "".join(c for c in project_name[:20] if c.isalnum() or c in " _-")
            filename = f"待办导入模板_{safe_name}.csv"
        else:
            filename = "待办导入模板_含示例.csv" if include_sample else "待办导入模板_空.csv"
        body = content.encode("utf-8-sig")
        resp = make_response(body)
        resp.headers["Content-Type"] = "text/csv; charset=utf-8"
        resp.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
        return resp
    except Exception as e:
        return jsonify({"success": False, "message": f"生成模板失败：{e}"}), 500


@bp.post("/api/uploads/note-files")
@page13_access_required
def api_upload_note_file():
    """上传备注附件（PDF 等），存入数据库，返回可访问的下载 URL。"""
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"message": "请选择文件"}), 400
    fn_lower = (file.filename or "").lower()
    allowed = (
        ".pdf", ".doc", ".docx", ".xls", ".xlsx",
        ".png", ".jpg", ".jpeg",
        ".zip", ".tar", ".gz", ".tgz", ".rar",
    )
    from .upload_filename import normalized_upload_extension, preserved_secure_filename

    ext = normalized_upload_extension(file.filename or "")
    if ext not in allowed and not any(fn_lower.endswith(s) for s in allowed):
        return jsonify({"message": f"仅支持以下格式：{', '.join(allowed)}"}), 400
    raw = file.read()
    if len(raw) > current_app.config.get("MAX_CONTENT_LENGTH", 25 * 1024 * 1024):
        return jsonify({"message": "文件过大"}), 400
    stored_name = f"{now_local().strftime('%Y%m%d%H%M%S%f')}_{preserved_secure_filename(file.filename or '')}"
    db.session.add(
        NoteAttachmentFile(
            stored_name=stored_name,
            file_blob=raw,
            original_name=file.filename,
        )
    )
    db.session.commit()
    download_url = f"/api/uploads/note-files/{stored_name}"
    return jsonify({"success": True, "url": download_url, "fileName": file.filename, "storedName": stored_name})


@bp.get("/api/uploads/note-files/<path:filename>")
def api_download_note_file(filename: str):
    """下载备注附件（优先数据库，其次遗留磁盘文件）。"""
    row = NoteAttachmentFile.query.filter_by(stored_name=filename).first()
    if row:
        return send_file(
            io.BytesIO(row.file_blob),
            as_attachment=True,
            download_name=row.original_name or filename,
        )
    notes_dir = Path(current_app.config["UPLOAD_FOLDER"]) / "notes"
    return send_from_directory(str(notes_dir), filename, as_attachment=True)


@bp.post("/api/uploads/import")
@page13_access_required
def api_uploads_import():
    """导入另一工具审核后导出的待办任务点（CSV 或 Excel，首行为表头）。"""
    file_storage = request.files.get("file")
    rows, parse_err = _parse_import_file(file_storage)
    if parse_err:
        return jsonify({"success": False, "message": parse_err}), 400
    if not rows:
        return jsonify({"success": True, "created": 0, "updated": 0, "skipped": 0, "errors": [], "message": "无有效数据行"})

    created = 0
    updated = 0
    skipped = 0
    errors = []
    from .authz import _project_lookup_maps
    from .project_teams import resolve_organization_id_for_project_upload
    from .task_entry_schema import parse_import_date, validate_import_row_task_category

    for idx, d in enumerate(rows):
        row_no = idx + 2
        project_name = (d.get("project_name") or "").strip()
        file_name = (d.get("fileName") or "").strip()
        author = (d.get("author") or "").strip()
        if not project_name or not file_name or not author:
            skipped += 1
            errors.append({"row": row_no, "message": "缺少项目名称、文件名称或编写人员，已跳过"})
            continue

        cat_err = validate_import_row_task_category(d)
        if cat_err:
            skipped += 1
            errors.append({"row": row_no, "message": cat_err})
            continue

        task_type = (d.get("task_type") or "").strip() or None
        template_links = (d.get("template_links") or "").strip() or None
        if template_links:
            template_links = _normalize_template_links(template_links) or None
        due_date = parse_import_date(d.get("due_date") or "")
        document_display_date = parse_import_date(d.get("document_display_date") or "")

        existing = UploadRecord.query.filter_by(
            project_name=project_name,
            file_name=file_name,
            task_type=task_type,
            author=author,
        ).first()

        notes_raw = d.get("notes") or ""
        notes_val = "\n".join(ln.strip() for ln in notes_raw.replace(";", "\n").replace("；", "\n").split("\n") if ln.strip()) or None

        _ensure_project_row(project_name)
        _, by_label, by_name = _project_lookup_maps()
        proj = by_label.get(project_name) or by_name.get(project_name)
        import_org_id = resolve_organization_id_for_project_upload(project=proj)

        try:
            if existing:
                existing.organization_id = import_org_id
                existing.project_code = (d.get("project_code") or "").strip() or None
                existing.business_side = (d.get("business_side") or "").strip() or None
                existing.product = (d.get("product") or "").strip() or None
                existing.country = (d.get("country") or "").strip() or None
                existing.project_notes = (d.get("project_notes") or "").strip() or None
                existing.template_links = template_links
                existing.notes = notes_val
                existing.assignee_name = (d.get("assignee_name") or "").strip() or None
                existing.due_date = due_date
                existing.file_version = (d.get("file_version") or "").strip() or None
                existing.document_display_date = document_display_date
                existing.reviewer = (d.get("reviewer") or "").strip() or None
                existing.approver = (d.get("approver") or "").strip() or None
                existing.belonging_module = (d.get("belonging_module") or "").strip() or None
                existing.displayed_author = (d.get("displayed_author") or "").strip() or None
                existing.registered_product_name = (d.get("registered_product_name") or "").strip() or None
                existing.model = (d.get("model") or "").strip() or None
                existing.registration_version = (d.get("registration_version") or "").strip() or None
                existing.task_status = "pending"
                existing.completion_status = None
                db.session.add(existing)
                updated += 1
            else:
                upload = UploadRecord(
                    project_name=project_name,
                    organization_id=import_org_id,
                    project_code=(d.get("project_code") or "").strip() or None,
                    file_name=file_name,
                    task_type=task_type,
                    author=author,
                    template_links=template_links,
                    notes=notes_val,
                    project_notes=(d.get("project_notes") or "").strip() or None,
                    assignee_name=(d.get("assignee_name") or "").strip() or None,
                    due_date=due_date,
                    business_side=(d.get("business_side") or "").strip() or None,
                    product=(d.get("product") or "").strip() or None,
                    country=(d.get("country") or "").strip() or None,
                    file_version=(d.get("file_version") or "").strip() or None,
                    document_display_date=document_display_date,
                    reviewer=(d.get("reviewer") or "").strip() or None,
                    approver=(d.get("approver") or "").strip() or None,
                    belonging_module=(d.get("belonging_module") or "").strip() or None,
                    displayed_author=(d.get("displayed_author") or "").strip() or None,
                    registered_product_name=(d.get("registered_product_name") or "").strip() or None,
                    model=(d.get("model") or "").strip() or None,
                    registration_version=(d.get("registration_version") or "").strip() or None,
                )
                db.session.add(upload)
                summary = _prepare_summary(upload)
                db.session.add(summary)
                created += 1
        except Exception as e:
            errors.append({"row": row_no, "message": str(e)})

    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "message": f"保存失败：{e}"}), 500

    msg = f"导入完成：新增 {created} 条，更新 {updated} 条，跳过 {skipped} 条"
    if errors:
        msg += f"，{len(errors)} 条有误"
    return jsonify({
        "success": True,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors[:50],
        "message": msg,
    })


@bp.get("/api/template-options")
@bp.get("/api/upload-options")
@page13_access_required
def api_template_options():
    records = UploadRecord.query.order_by(
        UploadRecord.project_name, UploadRecord.file_name, UploadRecord.author
    ).all()
    return jsonify({"projects": _build_option_tree(records)})


@bp.get("/api/templates/<upload_id>")
@bp.get("/api/uploads/<upload_id>")
@page13_access_required
def api_template_detail(upload_id: str):
    upload = UploadRecord.query.get(upload_id)
    if not upload:
        return jsonify({"message": "未找到模板记录"}), 404
    placeholders = upload.placeholders or []
    if not placeholders and (
        upload.template_file_blob
        or upload.storage_path
        or (getattr(upload, "ftp_path", None) or "").strip()
        or upload.template_links
    ):
        try:
            if upload.template_file_blob:
                placeholders = extract_placeholders_from_bytes(upload.template_file_blob)
            else:
                template_path = _get_template_path_for_upload(upload)
                placeholders = extract_placeholders(template_path)
            upload.placeholders = placeholders
            db.session.add(upload)
            db.session.commit()
        except Exception as exc:
            return jsonify({"message": f"解析模板失败：{exc}"}), 400
    return jsonify(
        {
            "id": upload.id,
            "projectName": upload.project_name,
            "fileName": upload.file_name,
            "author": upload.author,
            "placeholders": placeholders,
            "hasLinks": bool(upload.template_links),
            "linksCount": len(upload.get_template_links_list()),
            "taskStatus": upload.task_status,
            "quickCompleted": upload.quick_completed,
            "assigneeName": upload.assignee_name,
            "dueDate": upload.due_date.strftime("%Y-%m-%d") if upload.due_date else None,
        }
    )


@bp.get("/api/uploads")
@page13_access_required
def api_uploads_list():
    """获取所有上传记录列表"""
    include_history = str(request.args.get("includeHistory") or "").strip() in ("1", "true", "True", "yes", "on")
    proj_meta = _project_meta_map(auto_create_from_uploads=True)
    ended = {n for n, m in proj_meta.items() if (m.get("status") or "").strip().lower() == Project.STATUS_ENDED}
    q = UploadRecord.query
    if (not include_history) and ended:
        q = q.filter(~UploadRecord.project_name.in_(list(ended)))
    records = _sort_upload_records_by_project_priority(
        q.order_by(UploadRecord.sort_order.asc(), UploadRecord.created_at.asc()).all(),
        proj_meta,
    )
    from .authz import filter_upload_records_in_scope, filter_uploads_by_organization

    records = filter_upload_records_in_scope(records)
    records = filter_uploads_by_organization(records)
    upload_ids = [str(r.id) for r in records if getattr(r, "id", None)]
    has_gen: set[str] = set()
    if upload_ids:
        rows = (
            db.session.query(GenerateRecord.upload_id)
            .filter(GenerateRecord.upload_id.in_(upload_ids), GenerateRecord.success == True)
            .distinct()
            .all()
        )
        has_gen = {str(x[0]) for x in rows if x and x[0]}
    return jsonify({
        "records": [
            {
                "seq": idx + 1,
                "id": r.id,
                "projectId": (str(r.project_id).strip() if getattr(r, "project_id", None) else None),
                "projectName": r.project_name,
                "projectPriority": int((proj_meta.get(r.project_name) or {}).get("priority") or Project.PRIORITY_MEDIUM),
                "projectPriorityLabel": _project_priority_label((proj_meta.get(r.project_name) or {}).get("priority")),
                "projectStatus": ((proj_meta.get(r.project_name) or {}).get("status") or Project.STATUS_ACTIVE),
                "projectStatusLabel": _project_status_label((proj_meta.get(r.project_name) or {}).get("status")),
                "fileName": r.file_name,
                "taskType": r.task_type,
                "author": r.author,
                "hasFile": _upload_record_has_task_file(r),
                "hasGenerated": str(r.id) in has_gen,
                "hasLinks": bool(r.template_links),
                "templateLinks": r.template_links,
                "linksCount": len(r.get_template_links_list()),
                "ftpPath": ((getattr(r, "ftp_path", None) or "").strip() or None),
                "ftpUploaded": bool((getattr(r, "ftp_path", None) or "").strip()),
                "ftpLastError": ((getattr(r, "ftp_last_error", None) or "").strip() or None),
                "assigneeName": r.assignee_name,
                "dueDate": r.due_date.strftime("%Y-%m-%d") if r.due_date else None,
                "taskStatus": r.task_status,
                "completionStatus": r.completion_status,
                "auditStatus": r.audit_status,
                "quickCompleted": r.quick_completed,
                "sortOrder": r.sort_order,
                "businessSide": r.business_side,
                "product": r.product,
                "country": r.country,
                "projectCode": getattr(r, "project_code", None),
                "fileVersion": getattr(r, "file_version", None),
                "documentDisplayDate": (lambda d: d.strftime("%Y-%m-%d") if d else None)(getattr(r, "document_display_date", None)),
                "reviewer": getattr(r, "reviewer", None),
                "approver": getattr(r, "approver", None),
                "belongingModule": getattr(r, "belonging_module", None),
                "displayedAuthor": getattr(r, "displayed_author", None),
                "createdAt": r.created_at.isoformat() if r.created_at else None,
                "notes": r.notes,
                "projectNotes": getattr(r, "project_notes", None),
                "executionNotes": getattr(r, "execution_notes", None),
                "registeredProductName": getattr(r, "registered_product_name", None),
                "model": getattr(r, "model", None),
                "registrationVersion": getattr(r, "registration_version", None),
            }
            for idx, r in enumerate(records)
        ]
    })


@bp.delete("/api/uploads/<upload_id>")
@page13_access_required
def api_upload_delete(upload_id: str):
    """删除任务记录"""
    upload = UploadRecord.query.get(upload_id)
    if not upload:
        return jsonify({"message": "未找到该记录"}), 404
    
    uploads_dir = Path(current_app.config["UPLOAD_FOLDER"])
    _unlink_task_template_cache_files(upload_id)
    old_ftp = (getattr(upload, "ftp_path", None) or "").strip()
    if old_ftp:
        try:
            from .ftp_store import delete_path

            delete_path(old_ftp)
        except Exception:
            pass
    if upload.storage_path:
        try:
            Path(upload.storage_path).unlink(missing_ok=True)
        except Exception:
            pass

    db.session.delete(upload)
    db.session.commit()
    return jsonify({"message": "已删除"})


@bp.patch("/api/uploads/<upload_id>")
@page13_access_required
def api_upload_update(upload_id: str):
    """更新任务记录（页面1编辑），仅更新可编辑字段，不涉及文件替换。"""
    upload = UploadRecord.query.get(upload_id)
    if not upload:
        return jsonify({"message": "未找到该记录"}), 404
    
    data = request.get_json(force=True) or {}
    
    project_name = (data.get("projectName") or "").strip() or None
    if project_name:
        _ensure_project_row(project_name)
    file_name = (data.get("fileName") or "").strip() or None
    task_type = (data.get("taskType") or "").strip() or None
    author = (data.get("author") or "").strip() or None
    due_date_str = (data.get("dueDate") or "").strip() or None
    business_side = (data.get("businessSide") or "").strip() or None
    product = (data.get("product") or "").strip() or None
    country = (data.get("country") or "").strip() or None
    assignee_name = (data.get("assigneeName") or "").strip() or None
    project_code = (data.get("projectCode") or "").strip() or None
    file_version = (data.get("fileVersion") or "").strip() or None
    document_display_date_str = (data.get("documentDisplayDate") or "").strip() or None
    reviewer = (data.get("reviewer") or "").strip() or None
    approver = (data.get("approver") or "").strip() or None
    project_notes = (data.get("projectNotes") or "").strip() or None
    has_project_notes = "projectNotes" in data
    belonging_module = (data.get("belongingModule") or "").strip() or None
    has_belonging_module = "belongingModule" in data
    displayed_author = (data.get("displayedAuthor") or "").strip() or None
    has_displayed_author = "displayedAuthor" in data
    registered_product_name = (data.get("registeredProductName") or "").strip() or None
    has_registered_product_name = "registeredProductName" in data
    model = (data.get("model") or "").strip() or None
    has_model = "model" in data
    registration_version = (data.get("registrationVersion") or "").strip() or None
    has_registration_version = "registrationVersion" in data

    if project_name is not None:
        upload.project_name = project_name
    if file_name is not None:
        upload.file_name = file_name
    if "taskType" in data:
        raw_tt = data.get("taskType")
        upload.task_type = (str(raw_tt).strip() if raw_tt is not None else "") or None
    if author is not None:
        if not author:
            return jsonify({"message": "编写人员不能为空"}), 400
        upload.author = author
    if business_side is not None:
        upload.business_side = business_side or None
    if product is not None:
        upload.product = product or None
    if country is not None:
        upload.country = country or None
    if "templateLinks" in data:
        val = data.get("templateLinks")
        s = (val if isinstance(val, str) else "").strip()
        upload.template_links = _normalize_template_links(s) or None
    if assignee_name is not None:
        upload.assignee_name = assignee_name or None
    if "notes" in data:
        val = data.get("notes")
        s = (val if isinstance(val, str) else "").strip()
        upload.notes = s or None
    if has_project_notes:
        upload.project_notes = project_notes
    if project_code is not None:
        upload.project_code = project_code
    if file_version is not None:
        upload.file_version = file_version
    if document_display_date_str is not None:
        if not document_display_date_str:
            upload.document_display_date = None
        else:
            try:
                upload.document_display_date = datetime.strptime(document_display_date_str, "%Y-%m-%d").date()
            except ValueError:
                pass
    if reviewer is not None:
        upload.reviewer = reviewer
    if approver is not None:
        upload.approver = approver
    if has_belonging_module:
        upload.belonging_module = belonging_module
    if has_displayed_author:
        upload.displayed_author = displayed_author
    if has_registered_product_name:
        upload.registered_product_name = registered_product_name
    if has_model:
        upload.model = model
    if has_registration_version:
        upload.registration_version = registration_version
    if "auditStatus" in data:
        raw_ast = data.get("auditStatus")
        audit_status = (str(raw_ast).strip() if raw_ast is not None else "") or None
        _maybe_bump_audit_reject_count(
            upload,
            previous_completion_status=upload.completion_status,
            previous_audit_status=upload.audit_status,
            target_audit_status=audit_status,
        )
        upload.audit_status = audit_status
        if audit_status == AUDIT_REJECT_PENDING_STATUS:
            upload.completion_status = None
            upload.task_status = "pending"
            upload.quick_completed = False
    if "completionStatus" in data:
        raw_cs = data.get("completionStatus")
        completion_status = (str(raw_cs).strip() if raw_cs is not None else "") or None
        prev_completion = upload.completion_status
        if completion_status:
            _maybe_bump_audit_reject_count(
                upload,
                previous_completion_status=prev_completion,
                previous_audit_status=upload.audit_status,
                target_completion_status=completion_status,
            )
        upload.completion_status = completion_status
        if upload.completion_status:
            upload.task_status = "completed"
        else:
            upload.task_status = "pending"
    
    if due_date_str is not None:
        if not due_date_str:
            upload.due_date = None
        else:
            try:
                upload.due_date = datetime.strptime(due_date_str, "%Y-%m-%d").date()
            except ValueError:
                return jsonify({"message": "截止日期格式应为 YYYY-MM-DD"}), 400
    
    if upload.project_name and upload.file_name and upload.author:
        other = UploadRecord.query.filter(
            UploadRecord.project_name == upload.project_name,
            UploadRecord.file_name == upload.file_name,
            UploadRecord.author == upload.author,
            UploadRecord.id != upload.id,
        ).first()
        if other and (upload.task_type or None) == (other.task_type or None):
            return jsonify({"message": "项目名称+文件名称+任务类型+编写人与已有记录重复"}), 409
    
    db.session.add(upload)
    if upload.summary:
        upload.summary.project_name = upload.project_name
        upload.summary.file_name = upload.file_name
        upload.summary.author = upload.author
        db.session.add(upload.summary)
    db.session.commit()
    # 页面1编辑不触发模块级联催办，仅在页面2标记完成时触发（见 api_update_completion_status）

    return jsonify({
        "message": "已更新",
        "record": {
            "id": upload.id,
            "projectName": upload.project_name,
            "fileName": upload.file_name,
            "taskType": upload.task_type,
            "author": upload.author,
            "dueDate": upload.due_date.strftime("%Y-%m-%d") if upload.due_date else None,
            "businessSide": upload.business_side,
            "product": upload.product,
            "country": upload.country,
            "assigneeName": upload.assignee_name,
        },
    })


@bp.get("/api/my-tasks")
@login_required
def api_my_tasks():
    """页面2：普通账号仅本人任务；超管/项管只读观察全部可见范围，支持项目组/人员筛选。"""
    from .observer_view import (
        build_observer_filter_options,
        page2_observer_mode,
        page2_query_upload_rows,
        page2_view_mode,
        prepare_page2_observer_rows,
        upload_record_mutable_by_current_user,
    )

    include_history = str(request.args.get("includeHistory") or "").strip() in ("1", "true", "True", "yes", "on")
    team_filter = str(request.args.get("teamId") or request.args.get("team_id") or "").strip()
    user_filter = str(request.args.get("userId") or request.args.get("user_id") or "").strip()
    proj_meta = _project_meta_map(auto_create_from_uploads=True)
    ended = {n for n, m in proj_meta.items() if (m.get("status") or "").strip().lower() == Project.STATUS_ENDED}

    view_mode = page2_view_mode()
    rows = page2_query_upload_rows(include_history=include_history, proj_meta=proj_meta, ended=ended)
    all_metas: list[dict[str, str]] = []
    if view_mode != "normal":
        _, all_metas = prepare_page2_observer_rows(rows)
        filter_opts = build_observer_filter_options(
            all_metas,
            include_teams=(view_mode == "super_admin_readonly"),
        )
        rows, all_metas = prepare_page2_observer_rows(
            rows,
            team_id=team_filter or None,
            user_id=user_filter or None,
        )
    else:
        filter_opts = {"teams": [], "users": []}

    records = _sort_upload_records_by_project_priority(rows, proj_meta)
    meta_by_id = {str(getattr(r, "id", "") or ""): m for r, m in zip(rows, all_metas)}

    out_records: list[dict[str, Any]] = []
    for idx, r in enumerate(records):
        meta = meta_by_id.get(str(r.id or ""), {})
        out_records.append(
            {
                "seq": idx + 1,
                "id": r.id,
                "projectId": (str(r.project_id).strip() if getattr(r, "project_id", None) else None),
                "projectName": r.project_name,
                "projectPriority": int((proj_meta.get(r.project_name) or {}).get("priority") or Project.PRIORITY_MEDIUM),
                "projectPriorityLabel": _project_priority_label((proj_meta.get(r.project_name) or {}).get("priority")),
                "projectStatus": ((proj_meta.get(r.project_name) or {}).get("status") or Project.STATUS_ACTIVE),
                "projectStatusLabel": _project_status_label((proj_meta.get(r.project_name) or {}).get("status")),
                "fileName": r.file_name,
                "taskType": r.task_type,
                "author": r.author,
                "hasFile": _upload_record_has_task_file(r),
                "hasLinks": bool(r.template_links),
                "templateLinks": r.template_links,
                "linksCount": len(r.get_template_links_list()),
                "ftpPath": ((getattr(r, "ftp_path", None) or "").strip() or None),
                "ftpUploaded": bool((getattr(r, "ftp_path", None) or "").strip()),
                "ftpLastError": ((getattr(r, "ftp_last_error", None) or "").strip() or None),
                "assigneeName": r.assignee_name,
                "dueDate": r.due_date.strftime("%Y-%m-%d") if r.due_date else None,
                "taskStatus": r.task_status,
                "completionStatus": r.completion_status,
                "quickCompleted": r.quick_completed,
                "placeholders": r.placeholders or [],
                "sortOrder": r.sort_order,
                "businessSide": r.business_side,
                "product": r.product,
                "country": r.country,
                "projectCode": getattr(r, "project_code", None),
                "fileVersion": getattr(r, "file_version", None),
                "documentDisplayDate": (lambda d: d.strftime("%Y-%m-%d") if d else None)(getattr(r, "document_display_date", None)),
                "reviewer": getattr(r, "reviewer", None),
                "approver": getattr(r, "approver", None),
                "belongingModule": getattr(r, "belonging_module", None),
                "displayedAuthor": getattr(r, "displayed_author", None),
                "createdAt": r.created_at.isoformat() if r.created_at else None,
                "notes": r.notes,
                "projectNotes": getattr(r, "project_notes", None),
                "executionNotes": getattr(r, "execution_notes", None),
                "registeredProductName": getattr(r, "registered_product_name", None),
                "model": getattr(r, "model", None),
                "registrationVersion": getattr(r, "registration_version", None),
                "teamId": meta.get("teamId") or None,
                "teamName": meta.get("teamName") or None,
                "assigneeUserId": meta.get("assigneeUserId") or None,
                "assigneeLabel": meta.get("assigneeLabel") or None,
                "canMutate": upload_record_mutable_by_current_user(r),
            }
        )

    return jsonify(
        {
            "viewMode": view_mode,
            "observerMode": page2_observer_mode(),
            "readOnly": view_mode == "super_admin_readonly",
            "filterOptions": filter_opts,
            "records": out_records,
        }
    )


@bp.patch("/api/uploads/<upload_id>/execution-notes")
@login_required
def api_update_execution_notes(upload_id: str):
    """更新执行任务备注（仅页面2可编辑）。"""
    from .observer_view import observer_mutation_blocked_response, upload_record_mutable_by_current_user

    upload = UploadRecord.query.get(upload_id)
    if not upload:
        return jsonify({"message": "未找到该记录"}), 404
    if not upload_record_mutable_by_current_user(upload):
        return observer_mutation_blocked_response(record_level=True)
    from .notify_content import (
        MATTER_COMPLETE_NOTES_INVALID_MSG,
        is_matter_task_upload,
        is_meaningful_matter_execution_notes,
    )

    data = request.get_json(force=True) or {}
    val = data.get("executionNotes")
    s = (val if isinstance(val, str) else "").strip() or None
    if s and is_matter_task_upload(upload) and not is_meaningful_matter_execution_notes(s):
        return jsonify({"message": MATTER_COMPLETE_NOTES_INVALID_MSG}), 400
    upload.execution_notes = s
    db.session.add(upload)
    db.session.commit()
    return jsonify({"message": "已更新", "executionNotes": upload.execution_notes})


@bp.get("/api/uploads/<upload_id>/template-file")
@_page13_or_login_required
def api_download_upload_template_file(upload_id: str):
    """下载或在线查看任务模板（BLOB / 本机路径 / FTP）。"""
    upload = UploadRecord.query.get(upload_id)
    if not upload:
        return jsonify({"message": "未找到该记录"}), 404
    if not _can_access_upload_template(upload):
        return jsonify({"message": "无权查看该任务模板"}), 403
    if not _upload_record_has_task_file(upload):
        return jsonify({"message": "该任务无模板文件（仅有链接或未上传）"}), 404

    download_name = _upload_template_download_filename(upload)
    inline = str(request.args.get("view") or request.args.get("inline") or "").strip().lower() in (
        "1",
        "true",
        "yes",
        "on",
    )
    mimetype = (
        "application/msword"
        if download_name.lower().endswith(".doc")
        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    if upload.template_file_blob:
        return send_file(
            io.BytesIO(upload.template_file_blob),
            mimetype=mimetype,
            as_attachment=not inline,
            download_name=download_name,
        )

    try:
        path = _get_template_path_for_upload(upload, 0)
    except Exception as exc:
        current_app.logger.warning(
            "下载任务模板失败 upload_id=%s: %s", upload_id, exc
        )
        return jsonify({"message": f"获取模板失败：{exc}"}), 500

    if not path or not Path(path).is_file():
        return jsonify({"message": "模板文件不存在或无法读取"}), 404

    return send_file(
        path,
        mimetype=mimetype,
        as_attachment=not inline,
        download_name=download_name,
    )


@bp.post("/api/uploads/<upload_id>/template-file")
@login_required
def api_upload_replace_template_file(upload_id: str):
    """页面2：上传模板文件覆盖该任务已有文件或链接，写入 FTP（每条任务仅保留一个文件）。"""
    from .observer_view import observer_mutation_blocked_response, upload_record_mutable_by_current_user

    upload = UploadRecord.query.get(upload_id)
    if not upload:
        return jsonify({"message": "未找到该记录"}), 404
    if not upload_record_mutable_by_current_user(upload):
        return observer_mutation_blocked_response(record_level=True)

    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"message": "请选择要上传的文件"}), 400

    uploads_dir = Path(current_app.config["UPLOAD_FOLDER"])
    stored_file_name, storage_path = _save_file(file, uploads_dir)
    original_file_name = file.filename
    try:
        template_file_blob, placeholders = resolve_task_template_from_saved_path(
            Path(storage_path), file_name_hint=upload.file_name
        )
    except Exception as exc:
        Path(storage_path).unlink(missing_ok=True)
        return jsonify({"message": f"解析模板失败：{exc}"}), 400
    Path(storage_path).unlink(missing_ok=True)

    had_file = _upload_record_has_task_file(upload)
    had_links = bool(upload.get_template_links_list())

    _unlink_task_template_cache_files(upload_id)
    if upload.storage_path:
        try:
            Path(upload.storage_path).unlink(missing_ok=True)
        except Exception:
            pass

    old_ftp = (getattr(upload, "ftp_path", None) or "").strip()
    if old_ftp:
        try:
            from .ftp_store import delete_path

            delete_path(old_ftp)
        except Exception:
            pass

    upload.template_links = None
    upload.template_file_blob = template_file_blob
    upload.original_file_name = original_file_name
    upload.stored_file_name = None
    upload.storage_path = None
    upload.ftp_path = None
    upload.ftp_last_error = None
    upload.placeholders = placeholders

    db.session.add(upload)
    db.session.flush()
    _apply_task_template_ftp_after_flush(
        upload, upload.original_file_name or upload.file_name
    )
    db.session.commit()

    msg = "模板文件已上传"
    if had_file or had_links:
        msg = "已覆盖该任务原有模板文件或文档链接，来源已更新为文件"
    return jsonify(
        {
            "message": msg,
            "uploadId": upload_id,
            "hasFile": _upload_record_has_task_file(upload),
            "hasLinks": bool(upload.template_links),
            "ftpPath": ((getattr(upload, "ftp_path", None) or "").strip() or None),
            "ftpUploaded": bool((getattr(upload, "ftp_path", None) or "").strip()),
            "ftpLastError": ((getattr(upload, "ftp_last_error", None) or "").strip() or None),
            "originalFileName": upload.original_file_name,
        }
    )


@bp.patch("/api/uploads/<upload_id>/completion-status")
@login_required
def api_update_completion_status(upload_id: str):
    """更新任务的完成状态（页面2使用）。可仅更新文档链接；文件型标记完成需文档；事项型需执行任务备注。"""
    from .observer_view import observer_mutation_blocked_response, upload_record_mutable_by_current_user

    upload = UploadRecord.query.get(upload_id)
    if not upload:
        return jsonify({"message": "未找到该记录"}), 404
    if not upload_record_mutable_by_current_user(upload):
        return observer_mutation_blocked_response(record_level=True)
    from .notify_content import (
        MATTER_COMPLETE_NOTES_INVALID_MSG,
        MATTER_COMPLETE_NOTES_MSG,
        is_matter_task_upload,
        is_meaningful_matter_execution_notes,
    )

    data = request.get_json(force=True) or {}
    completion_status = data.get("completionStatus")
    template_links = (data.get("templateLinks") or "").strip() or None
    
    if template_links is not None:
        if template_links and not _is_valid_doc_link(template_links):
            return jsonify({"message": "请填写有效的文档链接（需以 http:// 或 https:// 开头）"}), 400
        upload.template_links = _normalize_template_links(template_links) or None

    if "executionNotes" in data:
        upload.execution_notes = (str(data.get("executionNotes") or "").strip()) or None
    
    if completion_status is not None:
        prev_completion = upload.completion_status
        if completion_status:
            if is_matter_task_upload(upload):
                if not is_meaningful_matter_execution_notes(upload.execution_notes):
                    msg = (
                        MATTER_COMPLETE_NOTES_INVALID_MSG
                        if (upload.execution_notes or "").strip()
                        else MATTER_COMPLETE_NOTES_MSG
                    )
                    return jsonify({"message": msg}), 400
            elif not upload.has_template():
                return jsonify({"message": "请先填写文档链接后再标记完成状态"}), 400
            _maybe_bump_audit_reject_count(
                upload,
                previous_completion_status=prev_completion,
                previous_audit_status=upload.audit_status,
                target_completion_status=completion_status,
            )
            upload.completion_status = completion_status
            upload.task_status = "completed"
        else:
            upload.completion_status = None
            upload.task_status = "pending"
    
    db.session.add(upload)
    db.session.commit()
    if upload.completion_status and (upload.belonging_module or "").strip() in ("产品", "开发"):
        _maybe_enqueue_module_cascade(upload.project_name or "", (upload.belonging_module or "").strip())

    return jsonify({
        "message": "已更新完成状态",
        "uploadId": upload_id,
        "completionStatus": upload.completion_status,
        "taskStatus": upload.task_status,
    })


# ---------- 文档生成 API ----------

@bp.post("/api/generate")
@login_required
def api_generate():
    from .observer_view import observer_mutation_blocked_response, upload_record_mutable_by_current_user

    data = request.get_json(force=True) or {}
    upload_id = data.get("uploadId")
    triggered_by = data.get("triggeredBy") or session.get("username", "web")
    placeholder_values = data.get("values") or {}
    output_name = data.get("outputName")
    replace = str(data.get("replace")).lower() == "true"
    link_index = int(data.get("linkIndex", 0))

    if not upload_id:
        return jsonify({"message": "请提供 uploadId。"}), 400

    upload = UploadRecord.query.get(upload_id)
    if not upload:
        return jsonify({"message": "未找到对应的模板记录"}), 404
    if not upload_record_mutable_by_current_user(upload):
        return observer_mutation_blocked_response(record_level=True)

    template_bytes = None
    template_path = None
    try:
        if upload.template_file_blob:
            template_bytes = upload.template_file_blob
        else:
            template_path = _get_template_path_for_upload(upload, link_index)
    except Exception as e:
        return jsonify({"message": str(e)}), 400

    required_placeholders = upload.placeholders or []
    if not required_placeholders:
        try:
            if template_bytes is not None:
                required_placeholders = extract_placeholders_from_bytes(template_bytes)
            else:
                required_placeholders = extract_placeholders(template_path or "")
            upload.placeholders = required_placeholders
            db.session.add(upload)
            db.session.commit()
        except Exception as exc:
            return jsonify({"message": f"解析模板失败：{exc}"}), 400

    missing = [
        key for key in required_placeholders if not placeholder_values.get(key, "").strip()
    ]
    if missing:
        return jsonify({"message": f"以下占位符尚未填写：{', '.join(missing)}"}), 400

    existing_record = GenerateRecord.query.filter_by(upload_id=upload_id).first()
    if existing_record and not replace:
        return (
            jsonify(
                {
                    "message": "该项目与文件的生成记录已存在，是否替换原内容？",
                    "needsConfirmation": True,
                }
            ),
            409,
        )

    outputs_dir = Path(current_app.config["OUTPUT_FOLDER"])
    stem = Path(upload.file_name).stem if upload.file_name else "document"
    try:
        output_path = generate_document(
            output_dir=str(outputs_dir),
            data=placeholder_values,
            output_name=output_name,
            template_path=template_path,
            template_bytes=template_bytes,
            template_stem=stem,
        )
    except Exception as exc:
        return jsonify({"message": f"生成文档失败：{exc}"}), 500

    out_p = Path(output_path)
    try:
        output_blob = out_p.read_bytes()
    except Exception as exc:
        return jsonify({"message": f"读取生成文件失败：{exc}"}), 500
    out_p.unlink(missing_ok=True)

    if existing_record and replace:
        previous_path = existing_record.output_path
        if previous_path:
            Path(previous_path).unlink(missing_ok=True)
        record = existing_record
        record.placeholder_payload = placeholder_values
        record.status = "completed"
        record.success = True
        record.completed_at = now_local()
        record.output_file_name = out_p.name
        record.output_path = None
        record.output_file_blob = output_blob
        record.triggered_by = triggered_by
    else:
        record = GenerateRecord(
            upload=upload,
            triggered_by=triggered_by,
            status="completed",
            success=True,
            completed_at=now_local(),
            placeholder_payload=placeholder_values,
            output_file_name=out_p.name,
            output_path=None,
            output_file_blob=output_blob,
        )
    db.session.add(record)

    summary = _prepare_summary(upload)
    summary.total_generate_clicks += 1
    summary.has_generated = True
    summary.last_generated_at = now_local()
    db.session.add(summary)

    upload.task_status = "completed"
    db.session.add(upload)

    db.session.commit()

    return jsonify(
        {
            "message": "文档生成成功。",
            "recordId": record.id,
            "uploadId": upload_id,
            "downloadUrl": url_for("pages.api_download_generated", record_id=record.id),
        }
    )


@bp.get("/api/generate-records/<record_id>/download")
@login_required
def api_download_generated(record_id: str):
    """下载已生成 Word（内容存于数据库）。"""
    record = GenerateRecord.query.get(record_id)
    if not record or not record.upload:
        return jsonify({"message": "记录不存在"}), 404
    username = session.get("username")
    if record.upload.author != username:
        return jsonify({"message": "无权下载"}), 403
    blob = record.output_file_blob
    if not blob and record.output_path and Path(record.output_path).is_file():
        blob = Path(record.output_path).read_bytes()
    if not blob:
        return jsonify({"message": "生成文件不可用"}), 404
    name = record.output_file_name or "generated.docx"
    return send_file(
        io.BytesIO(blob),
        as_attachment=True,
        download_name=name,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


# ---------- 快速完成 API ----------

@bp.post("/api/quick-complete")
@login_required
def api_quick_complete():
    """
    快速完成文档：无需填写占位符，直接在链接中填写完文档后标记完成。
    适用于页面1填写了文档链接的情况。
    """
    data = request.get_json(force=True) or {}
    upload_id = data.get("uploadId")

    if not upload_id:
        return jsonify({"message": "请提供 uploadId。"}), 400

    upload = UploadRecord.query.get(upload_id)
    if not upload:
        return jsonify({"message": "未找到对应的记录"}), 404

    if not upload.template_links:
        return jsonify({"message": "该记录没有文档链接，无法使用快速完成功能。"}), 400

    upload.quick_completed = True
    upload.task_status = "completed"
    db.session.add(upload)

    summary = _prepare_summary(upload)
    summary.has_generated = True
    summary.total_generate_clicks += 1
    summary.last_generated_at = now_local()
    db.session.add(summary)

    db.session.commit()

    return jsonify({
        "message": "已快速完成。",
        "uploadId": upload_id,
        "taskStatus": upload.task_status,
    })


# ---------- 任务状态更新 API ----------

@bp.patch("/api/uploads/<upload_id>/status")
@page13_access_required
def api_upload_status_update(upload_id: str):
    upload = UploadRecord.query.get(upload_id)
    if not upload:
        return jsonify({"message": "未找到该记录。"}), 404
    data = request.get_json(force=True) or {}
    if data.get("status") == "completed":
        upload.task_status = "completed"
        db.session.add(upload)
        db.session.commit()
    return jsonify({
        "message": "已更新。",
        "uploadId": upload_id,
        "taskStatus": upload.task_status,
    })


@bp.get("/api/summary")
@page13_access_required
def api_summary():
    return jsonify(_summary_payload())


# ---------- 项目管理 API ----------

def _sync_company_registry_projects_to_page1() -> int:
    """从页面0 同步分配到当前账号所属项目组的公司总览项目到页面1（尚无关联时各建一条）。"""
    from .authz import company_project_in_scope, is_page13_super_admin, user_team_ids
    from .models import CompanyProject, REGISTRATION_SCOPE_LEGACY
    from .tenant_context import resolve_organization_context

    org_id, _ = resolve_organization_context()
    q = CompanyProject.query
    if org_id:
        q = q.filter(CompanyProject.organization_id == org_id)
    if not is_page13_super_admin():
        team_ids = [str(t).strip() for t in user_team_ids() if str(t).strip()]
        if not team_ids:
            return 0
        q = q.filter(CompanyProject.assigned_team_id.in_(team_ids))
    rows = [cp for cp in q.order_by(CompanyProject.name.asc()).all() if company_project_in_scope(cp)]
    n = 0
    uid = str(session.get("user_id") or "").strip() or None
    for cp in rows:
        cp_id = str(cp.id or "").strip()
        if not cp_id:
            continue
        if Project.query.filter(Project.company_project_id == cp_id).first():
            continue
        p = Project(
            organization_id=str(getattr(cp, "organization_id", "") or "").strip() or None,
            name=cp.name,
            registered_country=getattr(cp, "registered_country", None),
            registered_category=getattr(cp, "registered_category", None),
            product_type=getattr(cp, "product_type", None),
            assigned_team_id=getattr(cp, "assigned_team_id", None),
            expected_certification_date=getattr(cp, "expected_certification_date", None),
            expected_submission_date=getattr(cp, "expected_submission_date", None),
            progress_description=getattr(cp, "progress_description", None),
            progress_updated_at=getattr(cp, "progress_updated_at", None),
            priority=int(cp.priority or CompanyProject.PRIORITY_MEDIUM),
            status=cp.status or CompanyProject.STATUS_ACTIVE,
            company_project_id=cp_id,
            registration_scope=REGISTRATION_SCOPE_LEGACY,
            created_by_user_id=uid,
        )
        db.session.add(p)
        n += 1
    if n:
        db.session.commit()
    return n


def _project_api_item(
    p: Project,
    *,
    prod_hints: dict[str, str] | None = None,
    team_names: dict[str, str] | None = None,
    org_names: dict[str, str] | None = None,
) -> dict:
    from .project_teams import project_has_page1_upload_tasks

    tid = (getattr(p, "assigned_team_id", None) or "").strip()
    teams = team_names or {}
    org_id = str(getattr(p, "organization_id", "") or "").strip()
    org_map = org_names or {}
    org_locked = project_has_page1_upload_tasks(p.id)
    return {
        "id": p.id,
        "companyProjectId": (getattr(p, "company_project_id", None) or None),
        "organizationId": org_id or None,
        "organizationName": org_map.get(org_id) if org_id else None,
        "organizationIdLocked": org_locked,
        "name": p.name,
        "registeredCountry": getattr(p, "registered_country", None),
        "registeredCategory": getattr(p, "registered_category", None),
        "registeredProductName": (prod_hints or {}).get(p.id),
        "productType": getattr(p, "product_type", None),
        "assignedTeamId": tid or None,
        "assignedTeamName": teams.get(tid) if tid else None,
        "assignedTeamIdLocked": org_locked,
        "expectedCertificationDate": (
            p.expected_certification_date.strftime("%Y-%m-%d")
            if getattr(p, "expected_certification_date", None)
            else None
        ),
        "expectedSubmissionDate": (
            p.expected_submission_date.strftime("%Y-%m-%d")
            if getattr(p, "expected_submission_date", None)
            else None
        ),
        "progressDescription": getattr(p, "progress_description", None),
        "registrationScope": p.registration_scope_effective(),
        "projectKey": _project_display_label(p),
        "priority": int(p.priority or Project.PRIORITY_MEDIUM),
        "priorityLabel": _project_priority_label(p.priority),
        "status": p.status or Project.STATUS_ACTIVE,
        "statusLabel": _project_status_label(p.status),
        "updatedAt": p.updated_at.isoformat() if p.updated_at else None,
    }


@bp.get("/api/projects")
@page13_access_required
def api_projects_list():
    """列出项目（从 upload_records 自动补齐缺失项）。"""
    from .authz import filter_projects_by_organization, project_in_scope, rbac_enforced
    from .models import ProjectTeam

    _project_meta_map(auto_create_from_uploads=True)
    _backfill_project_ids()
    rows = Project.query.order_by(Project.priority.desc(), Project.name.asc()).all()
    if rbac_enforced():
        rows = [p for p in rows if project_in_scope(p)]
    rows = filter_projects_by_organization(rows)
    try:
        from .project_registry_sync import sync_company_to_page1

        cp_ids = {
            (getattr(p, "company_project_id", None) or "").strip()
            for p in rows
            if (getattr(p, "company_project_id", None) or "").strip()
        }
        if cp_ids:
            for cid in cp_ids:
                sync_company_to_page1(cid)
            db.session.commit()
            rows = Project.query.order_by(
                Project.priority.desc(), Project.name.asc()
            ).all()
            if rbac_enforced():
                rows = [p for p in rows if project_in_scope(p)]
            rows = filter_projects_by_organization(rows)
    except Exception:
        pass
    prod_hints = _project_registered_product_name_hints([p.id for p in rows])
    team_names = {t.id: t.name for t in ProjectTeam.query.all()}
    org_ids = {
        str(getattr(p, "organization_id", "") or "").strip()
        for p in rows
        if str(getattr(p, "organization_id", "") or "").strip()
    }
    org_names: dict[str, str] = {}
    if org_ids:
        from .models import Organization

        org_names = {
            str(o.id or "").strip(): str(o.name or "").strip() or str(o.id or "").strip()
            for o in Organization.query.filter(Organization.id.in_(list(org_ids))).all()
        }
    return jsonify(
        [_project_api_item(p, prod_hints=prod_hints, team_names=team_names, org_names=org_names) for p in rows]
    )


@bp.post("/api/projects/sync-from-company-registry")
@page13_access_required
def api_projects_sync_from_company_registry():
    """从页面0 同步「所属项目组」下的公司总览项目到页面1。"""
    from .authz import is_page13_super_admin, user_team_ids

    if not is_page13_super_admin() and not user_team_ids():
        return jsonify({"message": "请先在页面4账号管理中分配所属项目组"}), 403
    try:
        synced = _sync_company_registry_projects_to_page1()
    except Exception as e:
        db.session.rollback()
        return jsonify({"message": f"同步失败：{e}"}), 500
    return jsonify(
        {
            "message": f"已从页面0 同步 {synced} 个项目到页面1（仅新建尚未关联的页面1 项目）",
            "synced": synced,
        }
    )


@bp.get("/api/assignable-organizations")
@page13_access_required
def api_assignable_organizations():
    """页面1 等项目管理可选的所属公司列表（按账号绑定过滤）。"""
    from .tenant_context import integration_organizations_payload

    return jsonify({"organizations": integration_organizations_payload()})


@bp.post("/api/projects")
@page13_access_required
def api_projects_create_or_update():
    """按（三字段）创建/更新项目优先级与状态。"""
    data = request.get_json(force=True) or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"message": "项目名称不能为空"}), 400
    from .registered_countries import validate_registered_country_selection

    registered_country, country_err = validate_registered_country_selection(
        data.get("registeredCountry")
    )
    if country_err:
        return jsonify({"message": country_err}), 400
    registered_category = (data.get("registeredCategory") or "").strip() or None
    priority = data.get("priority")
    status = (data.get("status") or "").strip().lower() or Project.STATUS_ACTIVE
    if priority is None:
        priority = Project.PRIORITY_MEDIUM
    try:
        priority = int(priority)
    except Exception:
        priority = Project.PRIORITY_MEDIUM
    priority = max(Project.PRIORITY_LOW, min(Project.PRIORITY_HIGH, priority))
    if status not in (Project.STATUS_ACTIVE, Project.STATUS_ENDED):
        status = Project.STATUS_ACTIVE

    from .authz import (
        current_admin_role,
        is_company_admin,
        parse_optional_date,
        rbac_enforced,
        user_team_ids,
    )
    from .models import (
        ADMIN_ROLE_PROJECT,
        CompanyProject,
        REGISTRATION_SCOPE_LEGACY,
        REGISTRATION_SCOPE_TEAM_LOCAL,
    )

    scope = (data.get("registrationScope") or "").strip()
    if scope not in (
        REGISTRATION_SCOPE_LEGACY,
        REGISTRATION_SCOPE_TEAM_LOCAL,
        "company",
    ):
        scope = REGISTRATION_SCOPE_LEGACY
    if rbac_enforced():
        if is_company_admin():
            scope = REGISTRATION_SCOPE_LEGACY
        elif current_admin_role() == ADMIN_ROLE_PROJECT:
            inc = bool(data.get("includeInCompanyRegistry"))
            scope = REGISTRATION_SCOPE_LEGACY if inc else REGISTRATION_SCOPE_TEAM_LOCAL

    q = Project.query.filter(Project.name == name)
    q = _filter_nullable_eq(q, Project.registered_country, registered_country)
    q = _filter_nullable_eq(q, Project.registered_category, registered_category)
    row = q.first()
    if not row:
        row = Project(
            name=name,
            registered_country=registered_country,
            registered_category=registered_category,
            priority=priority,
            status=status,
            registration_scope=scope,
            created_by_user_id=session.get("user_id"),
        )
        from .tenant_context import resolve_organization_context

        try:
            oid, _ = resolve_organization_context()
        except ValueError as exc:
            return jsonify({"message": str(exc)}), 403
        row.organization_id = oid or None
    if "organizationId" in data or "organization_id" in data:
        from .company_routes import _resolve_assignable_organization_id

        raw = (
            data.get("organizationId")
            if "organizationId" in data
            else data.get("organization_id")
        )
        oid, oid_err = _resolve_assignable_organization_id(raw)
        if oid_err:
            return jsonify({"message": oid_err}), 400
        row.organization_id = oid
    row.priority = priority
    row.status = status
    row.registration_scope = scope
    tid = (data.get("assignedTeamId") or "").strip() or None
    if tid:
        row.assigned_team_id = tid
    elif rbac_enforced() and not row.assigned_team_id:
        tids = user_team_ids()
        if tids:
            row.assigned_team_id = tids[0]
    cp_link = (data.get("companyProjectId") or "").strip()
    if bool(data.get("includeInCompanyRegistry")):
        if cp_link:
            row.company_project_id = cp_link
        else:
            cp = CompanyProject(
                name=name,
                registered_country=registered_country,
                registered_category=registered_category,
                priority=priority,
                status=status,
                organization_id=getattr(row, "organization_id", None),
                assigned_team_id=getattr(row, "assigned_team_id", None),
                created_by_user_id=session.get("user_id"),
            )
            from .project_registry_sync import apply_payload_to_company, payload_from_api_data

            pl_cp = payload_from_api_data(data)
            if pl_cp:
                apply_payload_to_company(cp, pl_cp)
            db.session.add(cp)
            db.session.flush()
            row.company_project_id = cp.id
    db.session.add(row)
    db.session.flush()
    _apply_page1_registry_fields_and_sync(row, data)
    db.session.commit()
    return jsonify(
        {
            "message": "已保存",
            "project": _project_api_item(row),
        }
    )


def _apply_page1_registry_fields_and_sync(row: Project, data: dict) -> None:
    from .project_registry_sync import (
        apply_payload_to_page1,
        payload_from_api_data,
        sync_page1_to_company,
    )

    pl = payload_from_api_data(data)
    if not pl:
        return
    apply_payload_to_page1(row, pl)
    sync_page1_to_company(row, pl)


@bp.patch("/api/projects/<project_id>")
@page13_access_required
def api_projects_patch(project_id: str):
    row = Project.query.get(project_id)
    if not row:
        return jsonify({"message": "未找到该项目"}), 404
    data = request.get_json(force=True) or {}

    old_key = _project_display_label(row)
    new_name = row.name
    new_country = getattr(row, "registered_country", None)
    new_category = getattr(row, "registered_category", None)

    if "name" in data:
        n = (data.get("name") or "").strip()
        if n:
            new_name = n
    if "registeredCountry" in data:
        from .registered_countries import validate_registered_country_selection

        new_country, country_err = validate_registered_country_selection(
            data.get("registeredCountry")
        )
        if country_err:
            return jsonify({"message": country_err}), 400
    if "registeredCategory" in data:
        new_category = (data.get("registeredCategory") or "").strip() or None

    # 更新基础三字段时，检查是否与其它项目重复
    q = Project.query.filter(Project.id != row.id).filter(Project.name == new_name)
    q = _filter_nullable_eq(q, Project.registered_country, new_country)
    q = _filter_nullable_eq(q, Project.registered_category, new_category)
    other = q.first()
    if other:
        return (
            jsonify(
                {
                    "message": "该项目（三字段）与已有项目重复，请先调整后再保存",
                    "conflictProjectId": other.id,
                }
            ),
            409,
        )

    if "priority" in data:
        try:
            p = int(data.get("priority"))
            row.priority = max(Project.PRIORITY_LOW, min(Project.PRIORITY_HIGH, p))
        except Exception:
            pass
    if "status" in data:
        s = (data.get("status") or "").strip().lower()
        if s in (Project.STATUS_ACTIVE, Project.STATUS_ENDED):
            row.status = s

    from .authz import project_in_scope, rbac_enforced

    if rbac_enforced() and not project_in_scope(row):
        return jsonify({"message": "无权修改该项目"}), 403

    if "assignedTeamId" in data:
        from .project_teams import (
            apply_project_assigned_team_id,
            validate_project_assigned_team_change,
        )

        tid = (data.get("assignedTeamId") or "").strip() or None
        team_err = validate_project_assigned_team_change(row, tid)
        if team_err:
            return jsonify({"message": team_err}), 403
        old_tid = str(getattr(row, "assigned_team_id", "") or "").strip() or None
        if old_tid != (tid or None):
            apply_project_assigned_team_id(row, tid)
        else:
            row.assigned_team_id = tid

    if "organizationId" in data or "organization_id" in data:
        from .company_routes import _resolve_assignable_organization_id
        from .project_teams import (
            apply_project_organization_id,
            project_has_page1_upload_tasks,
            validate_organization_id_change,
        )

        raw = (
            data.get("organizationId")
            if "organizationId" in data
            else data.get("organization_id")
        )
        new_oid, oid_err = _resolve_assignable_organization_id(raw)
        if oid_err:
            return jsonify({"message": oid_err}), 400
        old_oid = str(getattr(row, "organization_id", "") or "").strip() or None
        locked = project_has_page1_upload_tasks(row.id)
        org_err = validate_organization_id_change(
            old_org_id=old_oid,
            new_org_id=new_oid,
            upload_tasks_locked=locked,
        )
        if org_err:
            return jsonify({"message": org_err}), 403
        if old_oid != new_oid:
            apply_project_organization_id(row, new_oid)

    _apply_page1_registry_fields_and_sync(row, data)

    # 真正应用三字段
    row.name = new_name
    row.registered_country = new_country
    row.registered_category = new_category

    new_key = _project_display_label(row)
    db.session.add(row)
    db.session.commit()

    # 若三字段变化导致展示键变化，要同步更新已存在的任务/模块级联/生成摘要
    if old_key != new_key:
        UploadRecord.query.filter_by(project_name=old_key).update(
            {"project_name": new_key}
        )
        ModuleCascadeReminder.query.filter_by(project_name=old_key).update(
            {"project_name": new_key}
        )
        # 生成摘要也用于展示/下载列表时的统计口径（尽量保持一致）
        try:
            GenerationSummary.query.filter_by(project_name=old_key).update(
                {"project_name": new_key}
            )
        except Exception:
            pass
        db.session.commit()
    return jsonify({"message": "已更新"})


@bp.route("/api/projects/batch", methods=["PUT", "POST"])
@page13_access_required
def api_projects_batch_update():
    """批量更新项目优先级/状态（用于页面1一键保存/批量编辑）。"""
    data = request.get_json(force=True) or {}
    items = data.get("projects") or []
    if not isinstance(items, list) or not items:
        return jsonify({"message": "projects 不能为空"}), 400

    from .models import UploadRecord, ModuleCascadeReminder, GenerationSummary

    # 防止“一键保存”把两个项目更新成同一（三字段）重复项目
    request_ids = []
    triples_to_ids: dict[tuple[str, Any, Any], list[str]] = {}
    for it in items:
        if not isinstance(it, dict):
            continue
        pid = (it.get("id") or "").strip()
        if not pid:
            continue
        request_ids.append(pid)

    request_id_set = set(request_ids)
    for it in items:
        if not isinstance(it, dict):
            continue
        pid = (it.get("id") or "").strip()
        if not pid:
            continue
        row = Project.query.get(pid)
        if not row:
            continue
        from .registered_countries import validate_registered_country_selection

        target_country, country_err = validate_registered_country_selection(
            it.get("registeredCountry")
        )
        if country_err:
            return jsonify({"message": country_err}), 400
        target_category = (it.get("registeredCategory") or "").strip() or None
        triple = (row.name, target_country, target_category)
        triples_to_ids.setdefault(triple, []).append(pid)
    for triple, ids in triples_to_ids.items():
        if len(set(ids)) > 1:
            return (
                jsonify(
                    {
                        "message": "批量保存失败：检测到（三字段）重复项目（已被请求批量设置为同一三字段组合）。请先避免重复或删除多余项目。",
                        "dupProjectIds": list(set(ids)),
                    }
                ),
                409,
            )

    # 若与数据库中非本次请求的其它项目重复，也拒绝
    for triple, ids in triples_to_ids.items():
        base_name, target_country, target_category = triple
        q = Project.query.filter(Project.name == base_name)
        q = _filter_nullable_eq(q, Project.registered_country, target_country)
        q = _filter_nullable_eq(q, Project.registered_category, target_category)
        q = q.filter(~Project.id.in_(list(request_id_set) or [""]))
        other = q.first()
        if other:
            return (
                jsonify(
                    {
                        "message": "批量保存失败：目标（三字段）与已有项目重复，请先调整后再保存",
                        "conflictProjectId": other.id,
                    }
                ),
                409,
            )

    updated = 0
    for it in items:
        if not isinstance(it, dict):
            continue
        pid = (it.get("id") or "").strip()
        if not pid:
            continue
        row = Project.query.get(pid)
        if not row:
            continue

        old_key = _project_display_label(row)

        if "priority" in it:
            try:
                p = int(it.get("priority"))
                row.priority = max(Project.PRIORITY_LOW, min(Project.PRIORITY_HIGH, p))
            except Exception:
                pass
        if "status" in it:
            s = (it.get("status") or "").strip().lower()
            if s in (Project.STATUS_ACTIVE, Project.STATUS_ENDED):
                row.status = s

        if "registeredCountry" in it:
            from .registered_countries import validate_registered_country_selection

            rc, country_err = validate_registered_country_selection(
                it.get("registeredCountry")
            )
            if country_err:
                return jsonify({"message": country_err}), 400
            row.registered_country = rc
        if "registeredCategory" in it:
            row.registered_category = (it.get("registeredCategory") or "").strip() or None

        if "assignedTeamId" in it:
            from .project_teams import (
                apply_project_assigned_team_id,
                validate_project_assigned_team_change,
            )

            tid = (it.get("assignedTeamId") or "").strip() or None
            team_err = validate_project_assigned_team_change(row, tid)
            if team_err:
                return jsonify({"message": team_err}), 403
            old_tid = str(getattr(row, "assigned_team_id", "") or "").strip() or None
            if old_tid != (tid or None):
                apply_project_assigned_team_id(row, tid)
            else:
                row.assigned_team_id = tid

        if "organizationId" in it or "organization_id" in it:
            from .authz import is_page13_super_admin
            from .company_routes import _resolve_assignable_organization_id
            from .project_teams import (
                ORGANIZATION_ID_LOCKED_MSG,
                apply_project_organization_id,
                project_has_page1_upload_tasks,
                validate_organization_id_change,
            )

            raw = (
                it.get("organizationId")
                if "organizationId" in it
                else it.get("organization_id")
            )
            new_oid, oid_err = _resolve_assignable_organization_id(raw)
            if oid_err:
                return jsonify({"message": oid_err}), 400
            old_oid = str(getattr(row, "organization_id", "") or "").strip() or None
            locked = project_has_page1_upload_tasks(row.id)
            org_err = validate_organization_id_change(
                old_org_id=old_oid,
                new_org_id=new_oid,
                upload_tasks_locked=locked,
            )
            if org_err:
                return jsonify({"message": org_err}), 403
            if old_oid != new_oid:
                apply_project_organization_id(row, new_oid)

        _apply_page1_registry_fields_and_sync(row, it)

        new_key = _project_display_label(row)
        if old_key != new_key:
            UploadRecord.query.filter_by(project_name=old_key).update(
                {"project_name": new_key}
            )
            ModuleCascadeReminder.query.filter_by(project_name=old_key).update(
                {"project_name": new_key}
            )
            try:
                GenerationSummary.query.filter_by(project_name=old_key).update(
                    {"project_name": new_key}
                )
            except Exception:
                pass

        db.session.add(row)
        updated += 1

    db.session.commit()
    return jsonify({"message": f"已保存 {updated} 项", "updated": updated})


@bp.delete("/api/projects/<project_id>")
@page13_access_required
def api_projects_delete(project_id: str):
    """删除项目：仅当项目未绑定任务/级联记录时允许。"""
    from .models import UploadRecord, ModuleCascadeReminder, GenerationSummary

    row = Project.query.get(project_id)
    if not row:
        return jsonify({"message": "未找到该项目"}), 404

    key = _project_display_label(row)

    upload_count = UploadRecord.query.filter(UploadRecord.project_id == row.id).count()
    cascade_count = ModuleCascadeReminder.query.filter(ModuleCascadeReminder.project_id == row.id).count()
    generation_count = GenerationSummary.query.filter(GenerationSummary.project_id == row.id).count()

    total = upload_count + cascade_count + generation_count
    if total > 0:
        return (
            jsonify(
                {
                    "message": "无法删除：该项目已绑定任务/级联/生成记录，请先清理后再删除",
                    "bound": {
                        "uploadCount": upload_count,
                        "cascadeCount": cascade_count,
                        "generationCount": generation_count,
                        "totalCount": total,
                    },
                }
            ),
            409,
        )

    db.session.delete(row)
    db.session.commit()
    return jsonify({"message": "项目已删除"})


@bp.get("/api/projects/<project_id>/bindings")
@page13_access_required
def api_projects_bindings_count(project_id: str):
    """获取项目绑定数量：用于删除前提示。"""
    from .models import UploadRecord, ModuleCascadeReminder, GenerationSummary

    row = Project.query.get(project_id)
    if not row:
        return jsonify({"message": "未找到该项目"}), 404

    key = _project_display_label(row)
    upload_count = UploadRecord.query.filter(UploadRecord.project_id == row.id).count()
    cascade_count = ModuleCascadeReminder.query.filter(ModuleCascadeReminder.project_id == row.id).count()
    generation_count = GenerationSummary.query.filter(GenerationSummary.project_id == row.id).count()
    total = upload_count + cascade_count + generation_count

    return jsonify(
        {
            "message": "ok",
            "projectKey": key,
            "bound": {
                "uploadCount": upload_count,
                "cascadeCount": cascade_count,
                "generationCount": generation_count,
                "totalCount": total,
            },
        }
    )


# ---------- 通知文案配置 API ----------

@bp.get("/api/configs/notify-templates")
@super_admin_required
def api_get_notify_templates():
    templates = NotifyTemplateConfig.query.order_by(NotifyTemplateConfig.template_key).all()
    return jsonify([
        {
            "id": t.id,
            "key": t.template_key,
            "name": t.template_name,
            "content": t.template_content,
            "isActive": t.is_active,
        }
        for t in templates
    ])


@bp.put("/api/configs/notify-templates/<template_id>")
@super_admin_required
def api_update_notify_template(template_id: str):
    template = NotifyTemplateConfig.query.get(template_id)
    if not template:
        return jsonify({"message": "未找到该通知模板"}), 404
    
    data = request.get_json(force=True) or {}
    if "content" in data:
        template.template_content = data["content"]
    if "name" in data:
        template.template_name = data["name"]
    if "isActive" in data:
        template.is_active = data["isActive"]
    
    db.session.add(template)
    db.session.commit()
    
    return jsonify({
        "message": "更新成功",
        "id": template.id,
    })


# ---------- 排序更新 API ----------

@bp.post("/api/uploads/reorder")
@login_required
def api_reorder_uploads():
    """更新任务排序"""
    from .observer_view import observer_mutation_blocked_response, upload_record_mutable_by_current_user

    data = request.get_json(force=True) or {}
    orders = data.get("orders", [])
    
    if not orders:
        return jsonify({"message": "排序数据为空"}), 400
    
    for item in orders:
        upload_id = item.get("id")
        sort_order = item.get("sortOrder", 0)
        if upload_id:
            upload = UploadRecord.query.get(upload_id)
            if not upload:
                continue
            if not upload_record_mutable_by_current_user(upload):
                return observer_mutation_blocked_response(record_level=True)
            upload.sort_order = sort_order
            db.session.add(upload)
    
    db.session.commit()
    return jsonify({"message": "排序更新成功"})


# ---------- 钉钉通知推送 API ----------

@bp.get("/chatbot-test")
@page4_access_required
def chatbot_test_page():
    """钉钉机器人本地联调页：输入问题→调用 aicheckword→展示生成结果（不发真实钉钉消息）。"""
    return render_template("chatbot_test.html")


@bp.get("/api/dingtalk/chatbot/callback-url")
@page4_access_required
def api_dingtalk_chatbot_callback_url():
    """由系统配置 BASE_URL 生成钉钉 HTTP 回调完整 URL，供页面复制到开放平台。"""
    from .app_settings import chatbot_callback_url_info, get_setting

    base_url = (get_setting("BASE_URL") or "").strip()
    info = chatbot_callback_url_info(base_url)
    return jsonify({"success": True, **info})


@bp.get("/api/dingtalk/chatbot/llm-settings")
@page13_access_required
def api_dingtalk_chatbot_llm_settings():
    """联调页 LLM 提供方选项（与初稿页白名单对齐；钉钉回调用系统配置 CHATBOT_LLM_PROVIDER）。"""
    from .draft_generation_routes import (
        AIWORD_DRAFT_LLM_PROVIDERS,
        _has_usable_stored_key_for_provider,
        _load_user_credential,
        _merged_allowed_providers_for_client,
        _refresh_upstream_interop_if_stale,
    )

    try:
        _refresh_upstream_interop_if_stale()
    except Exception:
        pass
    allowed: list[dict[str, Any]] = []
    for row in _merged_allowed_providers_for_client():
        pid = str(row.get("id") or "").strip().lower()
        if pid and pid != "cursor" and pid in _AIWORD_CHATBOT_LLM_PROVIDERS:
            allowed.append(row)
    if not allowed:
        allowed = [
            {"id": "deepseek", "label": "DeepSeek", "status": "ok"},
            {"id": "tongyi", "label": "通义千问", "status": "ok"},
            {"id": "ollama", "label": "Ollama（本机）", "status": "ok"},
        ]
    prov = _chatbot_llm_provider_from_settings()
    has_personal = False
    uid = session.get("user_id")
    if uid:
        row = _load_user_credential(str(uid))
        sk = str(current_app.config.get("SECRET_KEY") or "")
        for p in AIWORD_DRAFT_LLM_PROVIDERS:
            if p == "cursor":
                continue
            if _has_usable_stored_key_for_provider(row, p, sk=sk):
                has_personal = True
                break
    return jsonify(
        {
            "provider": prov,
            "systemProvider": prov,
            "allowedProviders": allowed,
            "hasPersonalKey": has_personal,
            "hint": (
                "联调页下拉优先于系统配置；钉钉群真实回调固定使用 CHATBOT_LLM_PROVIDER。"
                "若已登录页面2并在「文档初稿」保存个人 API Key，将透传 X-Client-Llm-* 头（deepseek/通义）。"
            ),
        }
    )


@bp.post("/api/dingtalk/chatbot/test")
@page13_access_required
def api_dingtalk_chatbot_test():
    """
    本地联调入口：模拟一次钉钉触发。
    Body：
      - text: 问题文本（必填）
      - group_id: 群 ID（可选）
      - send: true=同时发到群机器人 webhook；false=仅返回结果（默认 false）
      - session_webhook: 钉钉本会话回执 webhook（可选，优先级最高）
    """
    payload = request.get_json(silent=True) or {}
    if not isinstance(payload, dict):
        return jsonify({"success": False, "message": "payload 必须是 JSON 对象"}), 400

    text = str(payload.get("text") or "").strip()
    if not text:
        return jsonify({"success": False, "message": "text 不能为空"}), 400

    group_id = str(payload.get("group_id") or payload.get("groupId") or "").strip() or "__test__"
    send_real = bool(payload.get("send"))
    session_webhook = str(payload.get("session_webhook") or payload.get("sessionWebhook") or "").strip()
    req_provider = str(payload.get("provider") or payload.get("current_provider") or "").strip() or None
    eff_provider, provider_note = _chatbot_normalize_provider(req_provider)
    client_hdrs = _chatbot_client_headers(eff_provider)

    diagnostics = {
        "chatbot_enabled": _chatbot_enabled(),
        "trigger_keywords": _chatbot_keywords(),
        "enabled_groups": sorted(list(_chatbot_enabled_groups())),
        "confidence_threshold": _chatbot_confidence_threshold(),
        "cooldown_seconds": _chatbot_cooldown_seconds(),
        "api_base": _chatbot_api_base(),
        "api_base_configured": bool(_chatbot_api_base()),
        "chat_timeout_seconds": _chatbot_api_timeout_seconds(),
        "requested_provider": req_provider or "(系统配置)",
        "llm_provider": eff_provider,
        "client_llm_provider_header": client_hdrs.get("X-Client-Llm-Provider"),
        "has_personal_llm_key": bool(session.get("user_id")),
    }
    if provider_note:
        diagnostics["provider_note"] = provider_note

    trigger_type = "at_bot" if "@" in text else ("keyword" if _chatbot_should_trigger(text) else "manual")
    reply_data, err = _chatbot_call_aicheckword(
        query=text,
        group_id=group_id,
        message_id=uuid.uuid4().hex,
        trigger_type=trigger_type,
        recent_messages=[],
        provider=eff_provider,
    )
    if err:
        return jsonify({"success": False, "message": err, "diagnostics": diagnostics}), 200

    need_human = bool(reply_data.get("need_human"))
    answer_summary = (
        reply_data.get("answer_summary") or reply_data.get("answer") or ""
    ).strip()
    answer_detail = (reply_data.get("answer_detail") or answer_summary or "").strip()
    answer = answer_summary
    confidence = float(reply_data.get("confidence") or 0.0)
    conf_threshold = _chatbot_confidence_threshold()
    delivered_answer = answer_summary
    if need_human or confidence < conf_threshold or not answer_summary:
        delivered_answer = "这个问题我先帮你转人工确认，稍后给你准确答复。"
        answer_summary = ""
        answer_detail = ""

    sent_result: Optional[dict[str, Any]] = None
    if send_real:
        webhook = session_webhook or _chatbot_dingtalk_webhook_str()
        secret = None if session_webhook else _chatbot_dingtalk_secret_opt()
        if not webhook:
            sent_result = {
                "success": False,
                "error": "未配置体系记录机器人 Webhook（CHATBOT_DINGTALK_WEBHOOK）且未传 session_webhook",
            }
        else:
            res = dingtalk_service.send_text_message(delivered_answer, webhook=webhook, secret=secret)
            sent_result = res if isinstance(res, dict) else {"success": False, "error": str(res)}

    return jsonify(
        {
            "success": True,
            "trigger_type": trigger_type,
            "need_human": need_human,
            "confidence": confidence,
            "answer_raw": answer,
            "answer_summary": answer_summary,
            "answer_detail": answer_detail,
            "answer_delivered": delivered_answer,
            "references": reply_data.get("references") or [],
            "reason": reply_data.get("reason") or "",
            "model_used": reply_data.get("model_used") or "",
            "effective_provider": reply_data.get("effective_provider") or eff_provider,
            "llm_provider_used": reply_data.get("llm_provider_used") or reply_data.get("effective_provider") or eff_provider,
            "provider_note": reply_data.get("provider_note") or provider_note,
            "latency_ms": reply_data.get("latency_ms") or 0,
            "upstream_request_id": reply_data.get("request_id") or "",
            "sent_result": sent_result,
            "diagnostics": diagnostics,
        }
    )


def _chatbot_deliver_answer(payload: dict[str, Any], answer: str, *, need_human: bool, confidence: float) -> dict[str, Any]:
    """将答案发到 sessionWebhook 或群 webhook；返回可 JSON 化的结果。"""
    session_webhook = str(payload.get("sessionWebhook") or "").strip()
    if session_webhook:
        res = dingtalk_service.send_text_message(answer, webhook=session_webhook, secret=None)
        ok = bool(res and res.get("success"))
        return {
            "success": ok,
            "message": "reply sent" if ok else (res.get("error") or "reply failed"),
            "need_human": need_human,
            "confidence": confidence,
        }
    webhook = _chatbot_dingtalk_webhook_str()
    secret = _chatbot_dingtalk_secret_opt()
    if not webhook:
        return {
            "success": True,
            "message": "no sessionWebhook and no CHATBOT_DINGTALK_WEBHOOK",
            "reply": answer,
            "need_human": need_human,
            "confidence": confidence,
        }
    res = dingtalk_service.send_text_message(answer, webhook=webhook, secret=secret)
    ok = bool(res and res.get("success"))
    return {
        "success": ok,
        "message": "reply sent" if ok else (res.get("error") or "reply failed"),
        "need_human": need_human,
        "confidence": confidence,
    }


@bp.post("/api/dingtalk/chatbot/callback")
def api_dingtalk_chatbot_callback():
    """
    钉钉 HTTP 回调机器人入口（开放平台加密回调 + 明文联调兼容）。
    配置 DINGTALK_CALLBACK_* 后走加解密；否则按明文 JSON（本地调试）。
    """
    crypto = _dingtalk_callback_crypto_optional()
    if crypto:
        body = request.get_json(silent=True) or {}
        if not isinstance(body, dict):
            return jsonify({"success": False, "message": "body 须为 JSON"}), 400
        encrypt = str(body.get("encrypt") or "").strip()
        if not encrypt:
            return jsonify({"success": False, "message": "缺少 encrypt 字段"}), 400
        sig, ts, nonce = parse_callback_query_args(request.args)
        try:
            plain = crypto.verify_and_decrypt(sig, ts, nonce, encrypt)
        except DingTalkCallbackCryptoError as e:
            current_app.logger.warning("钉钉回调解密失败: %s", e)
            return jsonify({"success": False, "message": str(e)}), 403
        try:
            payload = json.loads(plain)
        except json.JSONDecodeError:
            payload = {"text": {"content": plain}}
        if not isinstance(payload, dict):
            payload = {}

        event_type = str(payload.get("EventType") or payload.get("eventType") or "").strip()
        if event_type in ("check_url", "check_create_suite_url"):
            return jsonify(crypto.encrypted_response_map("success"))

        result, err = _chatbot_process_incoming_payload(payload)
        if err:
            current_app.logger.warning("chatbot 调用 aicheckword 失败: %s", err)
            return jsonify(crypto.encrypted_response_map("success"))
        if result and result.get("ignored"):
            return jsonify(crypto.encrypted_response_map("success"))

        answer = str((result or {}).get("answer") or "").strip()
        need_human = bool((result or {}).get("need_human"))
        confidence = float((result or {}).get("confidence") or 0.0)
        session_webhook = str(payload.get("sessionWebhook") or "").strip()

        if session_webhook:
            _chatbot_deliver_answer(payload, answer, need_human=need_human, confidence=confidence)
            return jsonify(crypto.encrypted_response_map("success"))

        reply_plain = build_text_reply_json(answer)
        return jsonify(crypto.encrypted_response_map(reply_plain))

    payload = request.get_json(silent=True) or {}
    if not isinstance(payload, dict):
        return jsonify({"success": False, "message": "payload 必须是 JSON 对象"}), 400

    result, err = _chatbot_process_incoming_payload(payload)
    if err:
        current_app.logger.warning("chatbot 调用 aicheckword 失败: %s", err)
        return jsonify({"success": False, "message": err}), 200
    if result and result.get("ignored"):
        return jsonify({"success": True, "ignored": result["ignored"]})

    answer = str((result or {}).get("answer") or "").strip()
    need_human = bool((result or {}).get("need_human"))
    confidence = float((result or {}).get("confidence") or 0.0)
    out = _chatbot_deliver_answer(payload, answer, need_human=need_human, confidence=confidence)
    return jsonify(out), 200


def _get_notify_template(key: str) -> str:
    """获取通知文案模板"""
    template = NotifyTemplateConfig.query.filter_by(template_key=key, is_active=True).first()
    if template:
        return template.template_content
    return ""


def _resolve_mobiles_for_authors(author_names: list) -> list:
    """根据编写人员姓名解析钉钉 @ 用的手机号（从 User 表）。"""
    from .notify_content import resolve_mobiles_for_author_labels

    mobiles, _, _ = resolve_mobiles_for_author_labels(author_names)
    return mobiles


def _notify_at_result_suffix(
    unmatched: list,
    no_mobile: list,
    at_mobiles: Optional[list] = None,
) -> str:
    parts = []
    if at_mobiles:
        masked = []
        for m in at_mobiles:
            s = str(m)
            masked.append(f"{s[:3]}****{s[-4:]}" if len(s) >= 7 else s)
        parts.append(f"已传钉钉@手机号：{'、'.join(masked)}")
    if unmatched:
        parts.append(f"未匹配账号无法@：{'、'.join(unmatched)}")
    if no_mobile:
        parts.append(f"已匹配但无手机号无法@：{'、'.join(no_mobile)}")
    if not parts:
        return ""
    return "（" + "；".join(parts) + "）"


@bp.get("/api/notify/at-resolve")
@page13_access_required
def api_notify_at_resolve():
    """诊断：编写人/负责人姓名能否解析到钉钉 @ 用手机号（排查某人 @ 不到）。"""
    label = (request.args.get("author") or request.args.get("label") or "").strip()
    if not label:
        return jsonify({"message": "请提供 author 或 label 参数"}), 400
    from .notify_content import at_resolve_report

    return jsonify(at_resolve_report(label))


def _get_base_url() -> str:
    """获取对外访问的基础 URL（域名），用于催办等通知中的链接。优先使用配置的 BASE_URL，否则用当前请求的 host。"""
    from .app_settings import get_setting
    base = (get_setting("BASE_URL") or "").strip()
    if base:
        return base.rstrip("/")
    if request and request.host_url:
        return request.host_url.rstrip("/")
    return ""


def _maybe_enqueue_module_cascade(project_name: str, trigger_module: str) -> None:
    """
    当某项目下所属模块为 trigger_module（产品或开发）的最后一份文档被标记为完成后，入队一条延迟发送的模块级联催办。
    产品→开发；开发→测试。同一项目同一触发模块只保留一条待执行，新入队会覆盖旧待执行。
    """
    if not project_name or (trigger_module or "").strip() not in ("产品", "开发"):
        return
    pname = (project_name or "").strip()
    tmod = (trigger_module or "").strip()
    target_module = "开发" if tmod == "产品" else "测试"
    all_in_module = UploadRecord.query.filter(
        UploadRecord.project_name == pname,
        UploadRecord.belonging_module == tmod,
    ).all()
    if not all_in_module:
        return
    if any(getattr(u, "completion_status", None) is None for u in all_in_module):
        return
    row = AppConfig.query.filter_by(config_key="MODULE_CASCADE_DELAY_MINUTES").first()
    delay_minutes = 5
    if row and row.config_value:
        try:
            delay_minutes = max(1, min(1440, int(str(row.config_value).strip())))
        except ValueError:
            pass
    from datetime import timedelta
    run_at = now_local() + timedelta(minutes=delay_minutes)
    ModuleCascadeReminder.query.filter(
        ModuleCascadeReminder.project_name == pname,
        ModuleCascadeReminder.trigger_module == tmod,
        ModuleCascadeReminder.status == "pending",
    ).delete()
    db.session.add(ModuleCascadeReminder(
        project_name=pname,
        trigger_module=tmod,
        target_module=target_module,
        run_at=run_at,
        status="pending",
    ))
    db.session.commit()


def _send_notify_to_individuals(
    title: str,
    message_markdown: str,
    at_mobiles: Optional[list] = None,
    at_names: Optional[list] = None,
):
    """群消息发送成功后，再向被@人员的个人工作通知发一份（需配置 DINGTALK_APP_KEY/APP_SECRET/AGENT_ID）。"""
    mobiles = list(at_mobiles) if at_mobiles else []
    if not mobiles and at_names:
        mobiles = _resolve_mobiles_for_authors(list(at_names))
    if not mobiles:
        return
    dingtalk_service.send_work_notification_to_mobiles(title, message_markdown, mobiles)


def _notify_pending_task_list_md(pending_uploads: list) -> str:
    """将未完成任务列表格式化为钉钉 Markdown（按项目/业务方/产品分组）。"""
    if not pending_uploads:
        return ""

    def _group_key(u):
        return (u.project_name or "", u.business_side or "", u.product or "")

    groups: dict = {}
    for u in pending_uploads:
        groups.setdefault(_group_key(u), []).append(u)

    def _task_block_md(key, uploads_in_group):
        proj, bs, pr = key
        lines = []
        proj_code = getattr(uploads_in_group[0], "project_code", None) if uploads_in_group else None
        header = (
            f"**项目：{proj or '-'}  项目编号：{proj_code or '-'}  "
            f"影响业务方：{bs or '-'}  产品：{pr or '-'}**"
        )
        lines.append(header)
        from .notify_content import notify_doc_link_suffix_md

        for u in uploads_in_group:
            due = u.due_date.strftime("%Y-%m-%d") if u.due_date else "-"
            due_red = f'<font color="red">{due}</font>' if due != "-" else "-"
            file_label = u.file_name or "-"
            if u.task_type:
                file_label += f" ({u.task_type})"
            line = f" - 文件名称：{file_label}  截止日期：{due_red}"
            line += notify_doc_link_suffix_md(u)
            lines.append(line)
        return "\n".join(lines)

    return "\n\n".join(_task_block_md(k, grp) for k, grp in groups.items())


def _page2_url_for_notify() -> str:
    base_url = _get_base_url()
    page2_path = url_for("pages.generate_page", _external=False)
    if base_url:
        return f"{base_url}{page2_path}"
    return url_for("pages.generate_page", _external=True)


@bp.post("/api/notify/by-project")
@page13_access_required
def api_notify_by_project():
    """按项目推送钉钉通知"""
    data = request.get_json(force=True) or {}
    project_name = data.get("projectName")
    
    if not project_name:
        return jsonify({"message": "请提供项目名称"}), 400

    from .authz import filter_upload_records_in_scope, is_page13_super_admin, project_label_in_page3_scope

    if not project_label_in_page3_scope(project_name):
        return jsonify({"message": "无权催办所属项目组以外的项目"}), 403
    
    pending_uploads = UploadRecord.query.filter(
        UploadRecord.project_name == project_name,
        UploadRecord.completion_status.is_(None)
    ).all()
    if not is_page13_super_admin():
        pending_uploads = filter_upload_records_in_scope(pending_uploads)
    
    if not pending_uploads:
        return jsonify({"message": f"项目 {project_name} 没有未完成的任务"})
    
    from .notify_content import page2_my_tasks_link_md

    from .notify_content import (
        collect_notify_person_names_from_uploads,
        resolve_mobiles_for_author_labels,
    )
    team_id = _resolve_team_id_by_project_name(project_name)
    webhook, secret, _source = _resolve_dingtalk_for_team(team_id)
    if not webhook:
        return jsonify({"message": "未配置催办 Webhook，请在页面4 · 系统与钉钉「项目组钉钉」或系统配置「催办与定时通知」填写；未配置时将使用互联网产品部机器人"}), 400

    assignees = collect_notify_person_names_from_uploads(pending_uploads)
    task_list_with_links_md = _notify_pending_task_list_md(pending_uploads)
    page2_url = _page2_url_for_notify()
    at_mobiles, unmatched, no_mobile = resolve_mobiles_for_author_labels(assignees)
    message_plain = (
        "【项目任务催办】\n\n"
        f"项目：{project_name}\n\n"
        f"未完成任务数：{len(pending_uploads)}\n\n"
        f"请以下人员尽快完成：{'、'.join(assignees)}\n\n\n"
        "未完成列表：\n\n"
        f"{task_list_with_links_md}\n\n"
        f"{page2_my_tasks_link_md(page2_url)}\n\n"
        "### **编写完成后请在页面2中标记完成状态。**\n\n"
        "请抓紧处理！"
    )
    result = dingtalk_service.send_markdown_message(
        "项目任务催办",
        message_plain,
        at_all=False,
        at_mobiles=at_mobiles,
        at_names=assignees,
        webhook=webhook,
        secret=secret,
    )
    ok = result is not None and result.get("success") is True
    if ok:
        msg = "通知发送成功" + _notify_at_result_suffix(unmatched, no_mobile, at_mobiles)
        return jsonify({
            "success": True,
            "message": msg,
            "atNames": list(assignees),
            "atMobiles": at_mobiles,
            "unmatchedAuthors": unmatched,
            "authorsWithoutMobile": no_mobile,
        }), 200
    err = result.get("error", "未知错误") if result else "未知错误"
    if isinstance(err, dict):
        err = "未知错误"
    return jsonify({"success": False, "message": f"通知发送失败: {err}"}), 200


@bp.post("/api/notify/by-author")
@page13_access_required
def api_notify_by_author():
    """按编写人员推送钉钉通知"""
    data = request.get_json(force=True) or {}
    author = data.get("author")
    
    if not author:
        return jsonify({"message": "请提供编写人员"}), 400
    
    pending_uploads = UploadRecord.query.filter(
        UploadRecord.author == author,
        UploadRecord.completion_status.is_(None)
    ).all()
    from .authz import filter_upload_records_in_scope, is_page13_super_admin

    if not is_page13_super_admin():
        pending_uploads = filter_upload_records_in_scope(pending_uploads)
    
    if not pending_uploads:
        return jsonify({"message": f"{author} 没有未完成的任务"})
    
    def _group_key(u):
        return (u.project_name or "", u.business_side or "", u.product or "")
    
    groups = {}
    for u in pending_uploads:
        k = _group_key(u)
        groups.setdefault(k, []).append(u)
    
    def _task_block_md(key, uploads_in_group):
        proj, bs, pr = key
        lines = []
        proj_code = getattr(uploads_in_group[0], "project_code", None) if uploads_in_group else None
        header = f"**项目：{proj or '-'}  项目编号：{proj_code or '-'}  影响业务方：{bs or '-'}  产品：{pr or '-'}**"
        lines.append(header)
        from .notify_content import notify_doc_link_suffix_md

        for u in uploads_in_group:
            due = u.due_date.strftime("%Y-%m-%d") if u.due_date else "-"
            due_red = f'<font color="red">{due}</font>' if due != "-" else "-"
            file_label = u.file_name or "-"
            if u.task_type:
                file_label += f" ({u.task_type})"
            fv = getattr(u, "file_version", None) or "-"
            ddd = (u.document_display_date.strftime("%Y-%m-%d") if u.document_display_date else "-") if getattr(u, "document_display_date", None) else "-"
            rev = getattr(u, "reviewer", None) or "-"
            appr = getattr(u, "approver", None) or "-"
            line = f" - 文件名称：{file_label}  文件版本号：{fv}  截止日期：{due_red}  文档体现日期：{ddd}  审核：{rev}  批准：{appr}"
            line += notify_doc_link_suffix_md(u)
            lines.append(line)
        return "\n".join(lines)
    
    base_url = _get_base_url()
    page2_path = url_for("pages.generate_page", _external=False)
    page2_url = f"{base_url}{page2_path}" if base_url else url_for("pages.generate_page", _external=True)
    
    template = _get_notify_template("author_reminder")
    if not template:
        template = "【个人任务催办】\n致：{author}\n您有 {pending_count} 个任务待完成：\n{task_list}\n\n请抓紧处理！"
    
    from .notify_content import (
        normalize_person_name,
        page2_my_tasks_link_md,
        resolve_mobiles_for_author_labels,
    )

    author_key = normalize_person_name(author)

    at_mobiles, unmatched, no_mobile = resolve_mobiles_for_author_labels([author_key])
    from .dingtalk_team import group_uploads_by_team

    grouped = group_uploads_by_team(pending_uploads)
    if not grouped:
        return jsonify({"message": f"{author} 没有可发送的任务"})
    total_success = 0
    failed_errors: list[str] = []
    for team_id, uploads in grouped.items():
        if not uploads:
            continue
        webhook, secret, _source = _resolve_dingtalk_for_team(team_id)
        if not webhook:
            failed_errors.append("未配置钉钉 Webhook")
            continue
        team_groups = {}
        for u in uploads:
            team_groups.setdefault(_group_key(u), []).append(u)
        team_task_list = "\n\n".join(_task_block_md(k, grp) for k, grp in team_groups.items())
        team_message = template.format(
            author=author,
            pending_count=len(uploads),
            task_list=team_task_list,
            page2_url=page2_url,
        ).rstrip()
        if not team_message.endswith("请抓紧处理！"):
            team_message += "\n\n请抓紧处理！"
        team_message += f"\n\n{page2_my_tasks_link_md(page2_url)}\n\n### **编写完成后请在页面2中标记完成状态。**"
        result = dingtalk_service.send_markdown_message(
            "个人任务催办",
            team_message,
            at_all=False,
            at_mobiles=at_mobiles,
            at_names=[author_key],
            webhook=webhook,
            secret=secret,
        )
        if result is not None and result.get("success") is True:
            total_success += 1
        else:
            err = result.get("error", "未知错误") if result else "未知错误"
            if isinstance(err, dict):
                err = "未知错误"
            failed_errors.append(str(err))
    if total_success > 0:
        msg = "通知发送成功" + _notify_at_result_suffix(unmatched, no_mobile, at_mobiles)
        if failed_errors:
            msg += f"（部分项目组发送失败：{'; '.join(failed_errors[:2])}）"
        return jsonify({
            "success": True,
            "message": msg,
            "atNames": [author_key],
            "atMobiles": at_mobiles,
            "unmatchedAuthors": unmatched,
            "authorsWithoutMobile": no_mobile,
        }), 200
    return jsonify({"success": False, "message": f"通知发送失败: {'; '.join(failed_errors[:2]) or '未知错误'}"}), 200


@bp.post("/api/notify/by-project-author")
@page13_access_required
def api_notify_by_project_author():
    """按项目 + 编写人员推送个人任务催办（仅该项目下该人员的未完成任务）。"""
    data = request.get_json(force=True) or {}
    project_name = (data.get("projectName") or "").strip()
    author = (data.get("author") or "").strip()

    if not project_name or not author:
        return jsonify({"message": "请提供项目名称与编写人员"}), 400

    from .authz import filter_upload_records_in_scope, is_page13_super_admin, project_label_in_page3_scope

    if not project_label_in_page3_scope(project_name):
        return jsonify({"message": "无权催办所属项目组以外的项目"}), 403

    pending_uploads = UploadRecord.query.filter(
        UploadRecord.project_name == project_name,
        UploadRecord.author == author,
        UploadRecord.completion_status.is_(None),
    ).all()
    if not is_page13_super_admin():
        pending_uploads = filter_upload_records_in_scope(pending_uploads)

    if not pending_uploads:
        return jsonify({
            "message": f"项目「{project_name}」下 {author} 没有未完成的任务",
        })
    from .dingtalk_team import group_uploads_by_team
    team_groups = group_uploads_by_team(pending_uploads)
    success = 0
    failed_errors: list[str] = []

    page2_url = _page2_url_for_notify()

    template = _get_notify_template("project_author_reminder")
    if not template:
        template = (
            "【个人任务催办】\n致：{author}\n\n"
            "项目：{project_name}\n\n"
            "您在该项目下有 {pending_count} 个任务待完成：\n\n"
            "{task_list}\n\n"
            "请抓紧处理！"
        )

    from .notify_content import page2_my_tasks_link_md

    from .notify_content import resolve_mobiles_for_author_labels

    at_mobiles, unmatched, no_mobile = resolve_mobiles_for_author_labels([author])
    for team_id, uploads in team_groups.items():
        if not uploads:
            continue
        webhook, secret, _source = _resolve_dingtalk_for_team(team_id)
        if not webhook:
            failed_errors.append("未配置钉钉 Webhook")
            continue
        task_list = _notify_pending_task_list_md(uploads)
        message = template.format(
            author=author,
            project_name=project_name,
            pending_count=len(uploads),
            task_list=task_list,
            page2_url=page2_url,
        ).rstrip()
        if "页面2" not in message and page2_url:
            message += f"\n\n{page2_my_tasks_link_md(page2_url)}\n\n### **编写完成后请在页面2中标记完成状态。**"
        if not message.endswith("请抓紧处理！"):
            message += "\n\n请抓紧处理！"
        result = dingtalk_service.send_markdown_message(
            "个人任务催办",
            message,
            at_all=False,
            at_mobiles=at_mobiles,
            at_names=[author],
            webhook=webhook,
            secret=secret,
        )
        if result is not None and result.get("success") is True:
            success += 1
        else:
            err = result.get("error", "未知错误") if result else "未知错误"
            if isinstance(err, dict):
                err = "未知错误"
            failed_errors.append(str(err))
    if success > 0:
        msg = "通知发送成功" + _notify_at_result_suffix(unmatched, no_mobile, at_mobiles)
        if failed_errors:
            msg += f"（部分项目组发送失败：{'; '.join(failed_errors[:2])}）"
        return jsonify({
            "success": True,
            "message": msg,
            "atNames": [author],
            "atMobiles": at_mobiles,
            "unmatchedAuthors": unmatched,
            "authorsWithoutMobile": no_mobile,
            "projectName": project_name,
        }), 200
    return jsonify({"success": False, "message": f"通知发送失败: {'; '.join(failed_errors[:2]) or '未知错误'}"}), 200


@bp.post("/api/notify/single-task")
@page13_access_required
def api_notify_single_task():
    """单条任务推送钉钉通知"""
    data = request.get_json(force=True) or {}
    upload_id = data.get("uploadId")
    
    if not upload_id:
        return jsonify({"message": "请提供任务ID"}), 400
    
    upload = UploadRecord.query.get(upload_id)
    if not upload:
        return jsonify({"message": "未找到该任务"}), 404
    
    if upload.completion_status:
        return jsonify({"message": "该任务已完成，无需催办"})
    from .authz import is_page13_super_admin, upload_record_visible_to_user

    if not is_page13_super_admin() and not upload_record_visible_to_user(upload):
        return jsonify({"message": "无权催办所属项目组以外的任务"}), 403
    from .dingtalk_team import resolve_team_id_by_upload

    team_id = resolve_team_id_by_upload(upload)
    webhook, secret, _source = _resolve_dingtalk_for_team(team_id)
    if not webhook:
        return jsonify({"message": "未配置催办 Webhook，请在页面4 · 系统与钉钉「项目组钉钉」或系统配置「催办与定时通知」填写；未配置时将使用互联网产品部机器人"}), 400
    
    from .notify_content import notify_doc_link_md_for_template

    due_date_str = upload.due_date.strftime("%Y-%m-%d") if upload.due_date else "-"
    due_red = f'<font color="red">{due_date_str}</font>' if due_date_str != "-" else "-"
    business_side = upload.business_side or "-"
    product = upload.product or "-"
    country = upload.country or "-"
    project_code = getattr(upload, "project_code", None) or "-"
    file_version = getattr(upload, "file_version", None) or "-"
    doc_display_date = (upload.document_display_date.strftime("%Y-%m-%d") if upload.document_display_date else "-") if getattr(upload, "document_display_date", None) else "-"
    reviewer = getattr(upload, "reviewer", None) or "-"
    approver = getattr(upload, "approver", None) or "-"
    project_notes = getattr(upload, "project_notes", None) or "-"
    title = f"{upload.project_name}/{upload.file_name}" + (f" ({upload.task_type})" if upload.task_type else "")
    doc_link_md = notify_doc_link_md_for_template(upload)

    template = _get_notify_template("single_task_reminder")
    if not template:
        template = (
            "【任务催办】\n致：{author}\n\n"
            "- **{title}**\n"
            " - 截止日期：{due_date}\n"
            " - 影响业务方：{business_side}\n"
            " - 产品：{product}\n"
            " - 国家：{country}\n"
            " - 文档地址：{doc_link_md}\n\n"
            "请抓紧处理！"
        )

    base_url = _get_base_url()
    page2_path = url_for("pages.generate_page", _external=False)
    page2_url = f"{base_url}{page2_path}" if base_url else url_for("pages.generate_page", _external=True)

    message_plain = template.format(
        author=upload.author,
        project_name=upload.project_name,
        file_name=upload.file_name,
        task_type=upload.task_type or "-",
        title=title,
        due_date=due_red,
        business_side=business_side,
        product=product,
        country=country,
        project_code=project_code,
        project_notes=project_notes,
        file_version=file_version,
        document_display_date=doc_display_date,
        reviewer=reviewer,
        approver=approver,
        doc_link_md=doc_link_md,
        doc_link=doc_link_md,
        page2_url=page2_url,
    )
    message_plain = message_plain.rstrip()
    if not message_plain.endswith("请抓紧处理！"):
        message_plain += "\n\n请抓紧处理！"
    from .notify_content import page2_my_tasks_link_md, resolve_mobiles_for_author_labels

    message_plain += f"\n\n{page2_my_tasks_link_md(page2_url)}\n\n### **编写完成后请在页面2中标记完成状态。**"

    author_key = (upload.author or "").strip()
    at_mobiles, unmatched, no_mobile = resolve_mobiles_for_author_labels([author_key])
    result = dingtalk_service.send_markdown_message(
        "任务催办",
        message_plain,
        at_all=False,
        at_mobiles=at_mobiles,
        at_names=[author_key] if author_key else None,
        webhook=webhook,
        secret=secret,
    )
    ok = result is not None and result.get("success") is True
    if ok:
        msg = "通知发送成功" + _notify_at_result_suffix(unmatched, no_mobile, at_mobiles)
        return jsonify({
            "success": True,
            "message": msg,
            "atNames": [author_key] if author_key else [],
            "atMobiles": at_mobiles,
            "unmatchedAuthors": unmatched,
            "authorsWithoutMobile": no_mobile,
        }), 200
    err = result.get("error", "未知错误") if result else "未知错误"
    if isinstance(err, dict):
        err = "未知错误"
    return jsonify({"success": False, "message": f"通知发送失败: {err}"}), 200


@bp.get("/api/notify/next-schedule")
@page13_access_required
def api_notify_next_schedule():
    """获取下一次自动通知时间（使用统计页配置的定时）"""
    from . import scheduler_service
    from .scheduler import _get_schedule_config_from_db
    cfg = _get_schedule_config_from_db()
    next_times = scheduler_service.get_next_run_times(cfg)
    
    webhook = _dingtalk_webhook_str()
    team_webhook_exists = (
        ProjectTeam.query.filter(ProjectTeam.dingtalk_webhook.isnot(None), ProjectTeam.dingtalk_webhook != "").first()
        is not None
    )
    next_times["dingtalkConfigured"] = bool(webhook) or team_webhook_exists
    
    return jsonify(next_times)


@bp.get("/api/notify/schedule-config")
@page13_access_required
def api_get_schedule_config():
    """获取自动通知时间配置（统计页用）"""
    from .scheduler import _get_schedule_config_from_db
    cfg = _get_schedule_config_from_db()
    row = AppConfig.query.filter_by(config_key="MODULE_CASCADE_DELAY_MINUTES").first()
    delay_minutes = 5
    if row and row.config_value:
        try:
            delay_minutes = max(1, min(1440, int(str(row.config_value).strip())))
        except ValueError:
            pass
    return jsonify({
        "weekly": cfg["weekly"],
        "overdue": cfg["overdue"],
        "project": cfg["project"],
        "moduleCascadeDelayMinutes": delay_minutes,
    })


@bp.put("/api/notify/schedule-config")
@super_admin_required
def api_put_schedule_config():
    """保存自动通知时间配置并立即生效"""
    data = request.get_json() or {}
    weekly = (data.get("weekly") or "").strip() or "thu 16:00"
    overdue = (data.get("overdue") or "").strip() or "15:00"
    project = (data.get("project") or "").strip() or "mon,wed,fri 9:30"
    delay_minutes = 5
    if "moduleCascadeDelayMinutes" in data:
        try:
            delay_minutes = max(1, min(1440, int(data.get("moduleCascadeDelayMinutes", 5))))
        except (TypeError, ValueError):
            pass
    for key, value in [
        ("SCHEDULE_WEEKLY_REMINDER", weekly),
        ("SCHEDULE_OVERDUE_REMINDER", overdue),
        ("SCHEDULE_PROJECT_STATS", project),
        ("MODULE_CASCADE_DELAY_MINUTES", str(delay_minutes)),
    ]:
        row = AppConfig.query.filter_by(config_key=key).first()
        if row:
            row.config_value = value
        else:
            db.session.add(AppConfig(config_key=key, config_value=value))
    db.session.commit()
    try:
        from .scheduler import reschedule_jobs
        reschedule_jobs(current_app._get_current_object())
    except Exception:
        pass
    return jsonify({"success": True, "message": "已保存并已更新定时任务"})


@bp.get("/api/system-settings")
@super_admin_required
def api_get_system_settings():
    """系统配置：默认带出当前生效值（已通过页面4 超级管理员校验；数据库 URI 脱敏）。"""
    from .app_settings import (
        SYSTEM_CONFIG_KEYS,
        chatbot_callback_url_info,
        get_setting,
        persist_config_json_into_empty_db_keys,
        sync_authoritative_sources_into_db,
        system_config_sections_for_api,
        system_settings_for_api_get,
    )

    app = current_app._get_current_object()
    project_root = Path(app.root_path).resolve().parent
    sync_authoritative_sources_into_db(project_root, app)
    persist_config_json_into_empty_db_keys(project_root, app)
    keys_meta = [{"key": k, "label": lbl, "sensitive": sens} for k, lbl, sens in SYSTEM_CONFIG_KEYS]
    base_url = (get_setting("BASE_URL", default="", app=app) or "").strip()
    return jsonify(
        {
            "settings": system_settings_for_api_get(app, project_root),
            "keys": keys_meta,
            "sections": system_config_sections_for_api(),
            "derived": {
                "chatbotCallback": chatbot_callback_url_info(base_url),
            },
        }
    )


@bp.put("/api/system-settings")
@super_admin_required
def api_put_system_settings():
    """保存系统配置并刷新当前进程 app.config（数据库连接 URI 需重启后生效）。"""
    body = request.get_json(force=True) or {}
    project_root = Path(current_app.root_path).resolve().parent
    from .app_settings import apply_system_settings_to_flask, save_system_settings

    try:
        save_system_settings(
            {str(k): ("" if v is None else str(v)) for k, v in body.items()},
            project_root,
        )
        apply_system_settings_to_flask(current_app._get_current_object(), project_root)
    except Exception as e:
        db.session.rollback()
        current_app.logger.exception("保存系统配置失败")
        return jsonify({"message": f"保存失败：{e}"}), 500
    return jsonify({"success": True, "message": "已保存。若修改了数据库连接 URI，请重启服务后生效。"})


@bp.get("/api/system-settings/team-dingtalk")
@page13_access_required
def api_get_team_dingtalk_settings():
    from .authz import is_page13_super_admin, is_project_admin, user_team_ids
    from .project_teams import serialize_team_item

    q = ProjectTeam.query.order_by(ProjectTeam.sort_order.asc(), ProjectTeam.name.asc())
    rows = q.all()
    if not is_page13_super_admin() and is_project_admin():
        allow = set(user_team_ids())
        rows = [t for t in rows if t.id in allow]
    teams = [serialize_team_item(t) for t in rows]
    return jsonify({"teams": teams})


@bp.put("/api/system-settings/team-dingtalk/<team_id>")
@page13_access_required
def api_put_team_dingtalk_settings(team_id: str):
    from .authz import is_page13_super_admin, is_project_admin, user_team_ids

    t = ProjectTeam.query.get((team_id or "").strip())
    if not t:
        return jsonify({"message": "未找到该项目组"}), 404
    if not is_page13_super_admin() and is_project_admin():
        if t.id not in set(user_team_ids()):
            return jsonify({"message": "仅可配置所属项目组的钉钉 webhook"}), 403
    data = request.get_json(force=True) or {}
    from .project_teams import apply_team_dingtalk_settings, serialize_team_item

    flags = apply_team_dingtalk_settings(t, data)
    db.session.commit()
    msg = "已保存"
    if flags.get("webhookEchoesGlobal"):
        msg = "已保存：该 Webhook 与全局催办/体系机器人相同，将自动使用全局配置（无需重复填写）"
    return jsonify(
        {
            "success": True,
            "message": msg,
            "webhookEchoesGlobal": bool(flags.get("webhookEchoesGlobal")),
            "team": serialize_team_item(t),
        }
    )


@bp.get("/api/notify/module-cascade-status")
@page13_access_required
def api_module_cascade_status():
    """模块级联催办状态：可配置延迟分钟数、待执行列表、最近已执行列表。"""
    row = AppConfig.query.filter_by(config_key="MODULE_CASCADE_DELAY_MINUTES").first()
    delay_minutes = 5
    if row and row.config_value:
        try:
            delay_minutes = max(1, min(1440, int(str(row.config_value).strip())))
        except ValueError:
            pass
    pending = ModuleCascadeReminder.query.filter_by(status="pending").order_by(ModuleCascadeReminder.run_at).all()
    recent_sent = (
        ModuleCascadeReminder.query.filter_by(status="sent")
        .order_by(ModuleCascadeReminder.sent_at.desc())
        .limit(20)
        .all()
    )
    from .authz import is_page13_super_admin, project_label_in_page3_scope

    if not is_page13_super_admin():
        pending = [r for r in pending if project_label_in_page3_scope(r.project_name)]
        recent_sent = [r for r in recent_sent if project_label_in_page3_scope(r.project_name)]
    return jsonify({
        "delayMinutes": delay_minutes,
        "pending": [
            {
                "projectName": r.project_name,
                "triggerModule": r.trigger_module,
                "targetModule": r.target_module,
                "runAt": r.run_at.strftime("%Y-%m-%d %H:%M") if r.run_at else None,
            }
            for r in pending
        ],
        "recentSent": [
            {
                "projectName": r.project_name,
                "triggerModule": r.trigger_module,
                "targetModule": r.target_module,
                "sentAt": r.sent_at.strftime("%Y-%m-%d %H:%M") if r.sent_at else None,
            }
            for r in recent_sent
        ],
    })


@bp.post("/api/notify/module-cascade-manual")
@page13_access_required
def api_notify_module_cascade_manual():
    """手动模块级联催办：按项目检查，产品全部完成→催办开发；开发全部完成→催办测试。body 可传 projectName 仅处理该项目。"""
    data = request.get_json(silent=True) or {}
    project_name = (data.get("projectName") or "").strip() or None
    from .authz import is_page13_super_admin, project_label_in_page3_scope

    if project_name and not is_page13_super_admin() and not project_label_in_page3_scope(project_name):
        return jsonify({"success": False, "message": "无权对所属项目组以外的项目执行级联催办"}), 403
    try:
        from .scheduler import _run_module_cascade_manual
        _run_module_cascade_manual(project_name=project_name)
        if project_name:
            return jsonify({"success": True, "message": f"已执行「{project_name}」的模块级联催办"})
        return jsonify({"success": True, "message": "已执行模块级联催办（按项目：产品→开发、开发→测试）"})
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@bp.post("/api/notify/test-auto")
@page13_access_required
def api_notify_test_auto():
    """测试自动催办：按类型执行与定时任务完全相同的逻辑，仅时间提前到点击时（仅超级管理员）。"""
    from .authz import is_page13_super_admin

    if not is_page13_super_admin():
        return jsonify({"success": False, "message": "仅超级管理员可测试自动催办"}), 403
    payload = request.get_json(silent=True) or {}
    test_type = (payload.get("type") or "").strip().lower()

    if test_type == "thursday":
        from .scheduler import _run_thursday_reminder
        _run_thursday_reminder(skip_dedupe=True)
        return jsonify({
            "success": True,
            "webhook_configured": True,
            "message": "已按定时任务逻辑发送（每周任务完成提醒）",
            "type": test_type,
        })
    if test_type == "overdue":
        from .scheduler import _run_overdue_reminder
        res = _run_overdue_reminder(skip_dedupe=True)
        if res is None:
            return jsonify({
                "success": False,
                "webhook_configured": True,
                "message": "执行异常",
                "type": test_type,
            })
        if res.get("no_tasks"):
            return jsonify({
                "success": True,
                "webhook_configured": True,
                "message": "无明日截止任务，未发送",
                "type": test_type,
            })
        sent, failed = res.get("sent") or 0, res.get("failed") or 0
        if failed > 0 or sent == 0:
            err = res.get("last_error") or "钉钉未返回成功"
            return jsonify({
                "success": False,
                "webhook_configured": True,
                "message": f"发送失败：{err}" + (f"（成功 {sent} 条，失败 {failed} 条）" if sent else ""),
                "type": test_type,
            })
        return jsonify({
            "success": True,
            "webhook_configured": True,
            "message": f"已发送 {sent} 条逾期提醒",
            "type": test_type,
        })
    if test_type == "project_stats":
        from .scheduler import _run_project_stats
        _run_project_stats(skip_dedupe=True)
        return jsonify({
            "success": True,
            "webhook_configured": True,
            "message": "已按定时任务逻辑发送（每两天项目完成情况统计）",
            "type": test_type,
        })
    if test_type == "module_cascade":
        from .scheduler import _run_process_module_cascade_pending
        _run_process_module_cascade_pending()
        return jsonify({
            "success": True,
            "webhook_configured": True,
            "message": "已处理到期的模块级联催办（按项目：产品完成→催办开发；开发完成→催办测试）",
            "type": test_type,
        })

    from . import dingtalk_service
    webhook, secret, _source = _resolve_dingtalk_for_team(None)
    if not webhook:
        return jsonify({
            "success": False,
            "webhook_configured": False,
            "message": "未配置催办 Webhook，请在页面4 · 系统与钉钉「项目组钉钉配置」或系统配置「催办与定时通知」中填写",
            "type": test_type or "default",
        }), 400
    content = "【自动催办测试】通道正常。定时任务将按配置时间发送每周任务完成提醒、逾期前一日催告、每两天项目完成情况统计。"
    result = dingtalk_service.send_text_message(content, webhook=webhook, secret=secret)
    ok = result.get("success") is True
    errmsg = result.get("error") or ""
    return jsonify({
        "success": ok,
        "webhook_configured": True,
        "message": "测试消息已发送" if ok else (errmsg or "钉钉返回失败，请检查 Webhook/Secret 及网络"),
        "type": test_type or "default",
    })


def register_blueprint(app):
    app.register_blueprint(bp)
